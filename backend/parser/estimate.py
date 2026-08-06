"""Оркестратор разбора сметы ГП.

Точка входа парсера: XLSX → JSON-структура, которая ложится в
`estimate_raw_data.raw_data` и разбирается импортом фазы 4 (AGENTS.md §5, шаги 2–3).

Соответствует `app/parse.py` исходника, но без всего, что относилось к
тендерному конвейеру: регистрации на Go-сервере, генерации markdown, чанков и
архивации файлов. Здесь только разбор.

Модуль не знает ни о БД, ни о FastAPI — это условие фазы 3.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from typing import IO, Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .constants import CONTRACTOR_SCAN_ROW_START, JSON_KEY_EXECUTOR, JSON_KEY_LOTS, TABLE_PARSE_POSITION_COLUMN_HEADERS
from .errors import EstimateParseError
from .layout import check_estimate_layout
from .parse_contractor_row import SUPPORTED_CONTRACTOR_COLSPANS
from .postprocess import (
    normalize_lots_json_structure,
    replace_excel_errors_with_null,
    stringify_temporal_values,
)
from .read_contractors import read_contractors
from .read_executer_block import read_executer_block
from .read_headers import read_headers
from .read_lots_and_boundaries import find_lot_starts, read_lots_and_boundaries
from .sheet import normalized_cell_text

log = logging.getLogger(__name__)

# Версия формата разбора. Пишется в `estimate_raw_data.parser_version`
# (AGENTS.md §4): по ней видно, каким кодом получен сохранённый JSON.
#
# Это НЕ `norm_version` из §4: версия нормализации наименований живёт отдельно и
# вводится в фазе 4 вместе с матчингом, потому что её инкремент требует миграции
# перевыпуска ключей `matching_cache` (AGENTS.md §11).
#
# 1.1.0 (Ф2): в contractor_items появился ключ `additional_works`. Версия
# минорная — структура только дополнена, из `positions` ничего не убрано, поэтому
# существующий импортёр не ломается. Исключение агрегатной строки из `positions`
# (Ф4) будет ломающим и потребует следующего подъёма.
PARSER_VERSION = "1.1.0"


@dataclass(frozen=True)
class ParseResult:
    """Результат разбора одной сметы."""

    data: dict[str, Any]
    """Полная JSON-структура для `estimate_raw_data.raw_data`."""

    parser_version: str = PARSER_VERSION
    """Версия парсера, которой получен `data`."""

    warnings: list[str] = field(default_factory=list)
    """Некритичные расхождения; фаза 4 кладёт их в `import_jobs.warnings`."""


def _select_worksheet(wb: openpyxl.Workbook, warnings: list[str]) -> Worksheet:
    """Возвращает лист со сметой, предупреждая о лишних листах."""
    if not wb.sheetnames:
        raise EstimateParseError("В книге нет ни одного листа.")

    if len(wb.sheetnames) > 1:
        warnings.append(
            f"В книге {len(wb.sheetnames)} листов ({', '.join(wb.sheetnames)}); "
            f"разобран первый — «{wb.sheetnames[0]}»."
        )

    return wb[wb.sheetnames[0]]


def _validate_contractor_blocks(contractors: list[dict[str, Any]]) -> None:
    """Отвергает файлы, в которых смысл колонок подрядчика неизвестен.

    Парсер определяет смысл колонок по ширине объединённого блока подрядчика
    (`parse_contractor_row.get_column_keys`). Для ширин из
    `SUPPORTED_CONTRACTOR_COLSPANS` раскладка известна, и файл разбирается —
    несовпадение с ожидаемой для сметы ГП шириной 11 уходит предупреждением
    (`layout.check_estimate_layout`). Для любой другой ширины раскладки нет, и
    разбор был бы выдумкой: стоимости легли бы не в те поля молча.

    Это граница между «читается с оговорками» и «структурно непригодно»:
    предупреждение обещает импорт, поэтому его нельзя выдавать там, где импорт
    невозможен.

    Args:
        contractors: результат `read_contractors` целиком (нулевой элемент —
            ячейка-маркер, дальше подрядчики).

    Raises:
        EstimateParseError: заголовок подрядчика не объединён с колонками блока
            либо ширина блока не поддерживается.
    """
    expected = ", ".join(str(value) for value in SUPPORTED_CONTRACTOR_COLSPANS)

    for contractor in contractors[1:]:
        title = contractor.get("value")
        coordinate = contractor.get("coordinate")
        merged_shape = contractor.get("merged_shape")

        if not merged_shape:
            raise EstimateParseError(
                f"Заголовок подрядчика «{title}» ({coordinate}) не объединён с колонками "
                "своего блока, поэтому неизвестно, сколько их и что в них лежит. "
                "Смысл колонок подрядчика задаётся шириной объединённого блока."
            )

        colspan = merged_shape.get("colspan")
        if colspan not in SUPPORTED_CONTRACTOR_COLSPANS:
            raise EstimateParseError(
                f"Блок подрядчика «{title}» ({coordinate}) занимает {colspan} колонок; "
                f"парсер знает раскладку только для {expected}. Смысл колонок определяется "
                "их числом, поэтому блок неизвестной ширины разобрать нельзя — стоимости "
                "попали бы не в те поля."
            )


def _find_column_header_row(ws: Worksheet, search_start_row: int, search_end_row: int) -> int | None:
    """Ищет строку шапки таблицы позиций по маркеру в колонке A.

    Область поиска — строго между строкой заголовка контрагентов и первой
    строкой данных (маркером лота). Весь лист не сканируется: обе границы к
    моменту вызова уже известны, а за ними шапки заведомо нет.

    Строка 9 не зашита константой намеренно: файл со сдвинутой на строку шапкой
    имеет верную раскладку и обязан разбираться.

    Args:
        ws: лист Excel.
        search_start_row: строка заголовка контрагентов (не включается).
        search_end_row: первая строка данных (не включается).

    Returns:
        Номер строки шапки либо None, если маркер не найден.
    """
    expected = normalized_cell_text(TABLE_PARSE_POSITION_COLUMN_HEADERS[1]).casefold()
    for row in range(search_start_row + 1, search_end_row):
        if normalized_cell_text(ws.cell(row=row, column=1).value).casefold() == expected:
            return row
    return None


def _validate_column_headers(
    ws: Worksheet,
    contractors: list[dict[str, Any]],
    lot_starts: list[dict[str, Any]],
) -> None:
    """Отвергает файлы, у которых шапка общих колонок не та.

    Колонки A, B, C, D читаются по ФИКСИРОВАННЫМ позициям
    (`get_lot_positions`), поэтому чужая шапка означает, что недостоверен весь
    позиционный разбор: номер, раздел, статья и наименование могли бы прийти не
    из тех ячеек. Это отказ, а не предупреждение, — та же граница, что у
    `_validate_contractor_blocks`: предупреждение обещает импорт, а импортировать
    здесь нечего.

    Args:
        ws: лист Excel.
        contractors: результат `read_contractors` целиком.
        lot_starts: результат `find_lot_starts`.

    Raises:
        EstimateParseError: строка шапки не найдена либо хотя бы один заголовок
            не совпал с ожидаемым.
    """
    header_marker_row = contractors[0].get("row_start")
    first_lot_row = lot_starts[0]["start_row"]
    if header_marker_row is None:
        header_marker_row = CONTRACTOR_SCAN_ROW_START - 1

    header_row = _find_column_header_row(ws, header_marker_row, first_lot_row)
    if header_row is None:
        raise EstimateParseError(
            f"Не найдена шапка таблицы позиций: в колонке A строк "
            f"{header_marker_row + 1}–{first_lot_row - 1} нет ячейки «"
            f"{TABLE_PARSE_POSITION_COLUMN_HEADERS[1]}». Колонки A–D читаются по "
            "фиксированным позициям, и без шапки нечем подтвердить, что раскладка "
            "та самая."
        )

    for column, expected_title in TABLE_PARSE_POSITION_COLUMN_HEADERS.items():
        actual = normalized_cell_text(ws.cell(row=header_row, column=column).value)
        if actual.casefold() != normalized_cell_text(expected_title).casefold():
            letter = get_column_letter(column)
            raise EstimateParseError(
                f"Колонка {letter} шапки (строка {header_row}) озаглавлена "
                f"«{actual}», ожидалось «{expected_title}». Колонки A–D читаются "
                "по фиксированным позициям, поэтому при другой раскладке номер, "
                "раздел, статья и наименование пришли бы не из тех ячеек."
            )


def parse_worksheet(ws: Worksheet) -> ParseResult:
    """Разбирает уже открытый лист.

    Отделено от `parse_estimate`, чтобы тесты могли подавать синтетические листы
    без временных файлов.

    Args:
        ws: лист Excel со сметой.

    Returns:
        `ParseResult` с полной структурой и предупреждениями.

    Raises:
        EstimateParseError: не найдена строка заголовков контрагентов, нет
            подрядчиков, нет маркера лота либо ширина блока подрядчика такова,
            что смысл его колонок неизвестен (`_validate_contractor_blocks`),
            либо шапка общих колонок A–D не совпала с ожидаемой
            (`_validate_column_headers`).
    """
    warnings: list[str] = []

    contractors = read_contractors(ws)
    if not contractors:
        raise EstimateParseError(
            "Не найдена строка заголовков контрагентов: ни в одной из строк "
            "4–10 нет ячейки, начинающейся с «Наименование контрагента». "
            "Похоже, это не смета ГП в ожидаемом формате."
        )
    if len(contractors) < 2:
        raise EstimateParseError(
            "Строка заголовков контрагентов найдена, но самого подрядчика в ней нет — "
            "справа от ячейки «Наименование контрагента» пусто."
        )

    lot_starts = find_lot_starts(ws)
    if not lot_starts:
        raise EstimateParseError(
            "Не найден маркер лота: в колонке D нет ячейки, начинающейся с «Лот №». "
            "Без него не определить границы блока позиций."
        )

    _validate_contractor_blocks(contractors)
    _validate_column_headers(ws, contractors, lot_starts)

    warnings.extend(check_estimate_layout(ws, contractors, lot_starts))

    data: dict[str, Any] = {
        **read_headers(ws),
        JSON_KEY_EXECUTOR: read_executer_block(ws),
        JSON_KEY_LOTS: read_lots_and_boundaries(ws),
    }
    data = normalize_lots_json_structure(data)
    data = replace_excel_errors_with_null(data)
    data = stringify_temporal_values(data)

    return ParseResult(data=data, warnings=warnings)


def parse_estimate(source: str | IO[bytes]) -> ParseResult:
    """Разбирает XLSX-файл сметы ГП.

    Книга открывается с `data_only=True`: парсер читает кэшированные значения
    формул, самих формул он не видит.

    Args:
        source: путь к файлу либо открытый двоичный поток.

    Returns:
        `ParseResult` с полной структурой и предупреждениями.

    Raises:
        EstimateParseError: файл не разбирается как смета ГП.
    """
    wb = openpyxl.load_workbook(source, data_only=True)
    try:
        warnings: list[str] = []
        ws = _select_worksheet(wb, warnings)
        result = parse_worksheet(ws)
        return ParseResult(
            data=result.data,
            parser_version=result.parser_version,
            warnings=warnings + result.warnings,
        )
    finally:
        wb.close()
