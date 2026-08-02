"""Поиск строки заголовков контрагентов.

Перенос `app/excel_parser/read_contractors.py` из `parser_tender_xlsx@0e178c0`.

Отличие от исходника: ширина скана ограничена `MAX_HEADER_SCAN_COLUMN` вместо
`ws.max_column`. На реальной смете `ws.max_column` равен 16384 из-за одной
пустой стилизованной ячейки в XFD (docs/phase0-input-data.md), и скан семи строк
шапки материализовал ~115 тыс. ячеек.
"""

from __future__ import annotations

from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .build_merged_shape_map import build_merged_shape_map
from .constants import (
    CONTRACTOR_SCAN_ROW_END,
    CONTRACTOR_SCAN_ROW_START,
    MAX_HEADER_SCAN_COLUMN,
    TABLE_PARSE_CONTRACTOR_TITLE,
)


def read_contractors(ws: Worksheet) -> list[dict[str, Any]] | None:
    """Находит строку заголовков контрагентов и описывает её непустые ячейки.

    Строка ищется в диапазоне `CONTRACTOR_SCAN_ROW_START..END`: ею считается
    первая строка, где какая-либо ячейка (без учёта регистра и крайних пробелов)
    начинается с `TABLE_PARSE_CONTRACTOR_TITLE`.

    В смете ГП это строка 6: ячейка-маркер "Наименование контрагента" и рядом
    объединённый блок подрядчика на 11 колонок (J..T).

    Args:
        ws: лист Excel.

    Returns:
        Список словарей по одному на непустую ячейку найденной строки::

            {
                "value": Any,             # значение ячейки
                "coordinate": str,        # "J6"
                "column_start": int,      # 1-индексированная колонка
                "row_start": int,         # номер строки заголовков
                "merged_shape": {"rowspan": int, "colspan": int},  # если объединена
            }

        Первый элемент — сама ячейка-маркер; подрядчики идут за ней (этим
        пользуется `get_proposals`, начиная обход с индекса 1).
        None — если строка заголовков не найдена.
    """
    search_prefix_lower = TABLE_PARSE_CONTRACTOR_TITLE.lower()

    merged_cells_map = build_merged_shape_map(ws)

    for row_tuple in ws.iter_rows(
        min_row=CONTRACTOR_SCAN_ROW_START,
        max_row=CONTRACTOR_SCAN_ROW_END,
        max_col=MAX_HEADER_SCAN_COLUMN,
    ):
        for cell in row_tuple:
            cell_value = cell.value
            if isinstance(cell_value, str) and cell_value.strip().lower().startswith(search_prefix_lower):
                contractor_headers_list: list[dict[str, Any]] = []
                for header_cell in row_tuple:
                    if header_cell.value is not None:
                        cell_info: dict[str, Any] = {
                            "value": header_cell.value,
                            "coordinate": header_cell.coordinate,
                            "column_start": header_cell.column,
                            "row_start": header_cell.row,
                        }

                        if header_cell.coordinate in merged_cells_map:
                            cell_info["merged_shape"] = merged_cells_map[header_cell.coordinate]

                        contractor_headers_list.append(cell_info)

                if contractor_headers_list:
                    return contractor_headers_list

    return None
