"""Определение границ лотов и запуск сбора данных по каждому из них.

Перенос `app/excel_parser/read_lots_and_boundaries.py` из
`parser_tender_xlsx@0e178c0` без изменений логики.

В смете ГП лот один: маркер "Лот №1 …" стоит в колонке D первой строки данных,
и лот тянется до конца листа — реальные границы блока позиций дальше срежет
досрочный выход `get_lot_positions` по объединённой ячейке колонки A.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    JSON_KEY_LOT_INDEX,
    JSON_KEY_LOT_TITLE,
    JSON_KEY_PROPOSALS,
    PARSE_TABLE_LOT_NUMBER,
    START_INDEXING_LOT_ROW,
)
from .get_proposals import get_proposals


def find_lot_starts(ws: Worksheet) -> list[dict[str, Any]]:
    """Находит строки-маркеры начала лотов.

    От `START_INDEXING_LOT_ROW` ищутся ячейки колонки D, начинающиеся с
    `PARSE_TABLE_LOT_NUMBER` (регистронезависимо).

    Вынесено из `read_lots_and_boundaries` отдельной функцией, чтобы оркестратор
    (`estimate.parse_estimate`) мог узнать первую строку данных для проверки
    раскладки колонок, не разбирая лист второй раз.

    Args:
        ws: лист Excel.

    Returns:
        Список `[{"start_row": int, "title": str}, ...]` в порядке следования.
    """
    lot_starts: list[dict[str, Any]] = []

    for current_row_num in range(START_INDEXING_LOT_ROW, ws.max_row + 1):
        cell_value_col_d = ws.cell(row=current_row_num, column=4).value

        if isinstance(cell_value_col_d, str) and cell_value_col_d.strip().lower().startswith(
            PARSE_TABLE_LOT_NUMBER.lower()
        ):
            lot_starts.append({"start_row": current_row_num, "title": cell_value_col_d.strip()})

    return lot_starts


@dataclass(frozen=True)
class LotsResult:
    """Лоты и предупреждения, собранные при их разборе."""

    lots: dict[str, dict[str, Any]]
    warnings: list[str]


def read_lots_and_boundaries(ws: Worksheet) -> LotsResult:
    """Находит лоты, вычисляет их границы и собирает данные по каждому.

    Шаг 1 — `find_lot_starts`. Шаг 2 — для каждого лота конечной строкой служит
    строка перед следующим лотом (для последнего — последняя строка листа),
    после чего вызывается `get_proposals`.

    Args:
        ws: лист Excel.

    Returns:
        `LotsResult`: словарь `{"lot_1": {"lot_title": str, "proposals": {...}}}`
        (пустой, если маркеров лотов нет) и предупреждения, собранные при
        разборе предложений каждого лота.
    """
    max_sheet_row = ws.max_row
    lot_starts = find_lot_starts(ws)

    if not lot_starts:
        return LotsResult(lots={}, warnings=[])

    found_lots_data: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []

    for i, lot_info in enumerate(lot_starts):
        start_row = lot_info["start_row"]
        lot_title = lot_info["title"]

        # Лот кончается перед началом следующего; последний — на конце листа.
        end_row = lot_starts[i + 1]["start_row"] - 1 if i + 1 < len(lot_starts) else max_sheet_row

        lot = get_proposals(ws, start_row=start_row, end_row=end_row)
        warnings.extend(lot.warnings)

        lot_key = f"{JSON_KEY_LOT_INDEX}{i + 1}"
        found_lots_data[lot_key] = {
            JSON_KEY_LOT_TITLE: lot_title,
            JSON_KEY_PROPOSALS: lot.proposals,
        }

    return LotsResult(lots=found_lots_data, warnings=warnings)
