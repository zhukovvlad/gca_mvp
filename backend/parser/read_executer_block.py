"""Чтение блока исполнителя документа.

Перенос `app/excel_parser/read_executer_block.py` из
`parser_tender_xlsx@0e178c0` без изменений логики.

Здесь лежит источник `estimates.data_prepared_on_date` (AGENTS.md §4) — поле
`executor_date` из строки «Дата составления». В полученном образце сметы ГП
блока исполнителя нет: и в реальном файле, и в fixture все три поля пустые,
поэтому `data_prepared_on_date` будет NULL, а датой сравнения с нормативом
станет `contracts.signed_date` — штатный фолбэк §4. Логику поиска намеренно не
трогаем: расширять окно сканирования не на чем — примеров файлов с этим блоком
у нас пока нет.
"""

from __future__ import annotations

from contextlib import suppress

from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    JSON_KEY_EXECUTOR_DATE,
    JSON_KEY_EXECUTOR_NAME,
    JSON_KEY_EXECUTOR_PHONE,
    TABLE_PARSE_EXECUTOR,
    TABLE_PARSE_PREPARATION_DATE,
    TABLE_PARSE_TELEPHONE,
)


def read_executer_block(ws: Worksheet) -> dict[str, str | None]:
    """Извлекает имя исполнителя, телефон и дату составления.

    Сканируются строки `ws.max_row - 5`, `-4`, `-3`, колонка B; сравнение
    ключевых фраз регистронезависимое, по префиксу. Для исполнителя и телефона
    значение берётся после первого двоеточия; для даты — всё после ключевой
    фразы (двоеточие-разделитель снимается, двоеточия во времени сохраняются).

    Args:
        ws: лист Excel.

    Returns:
        Словарь `executor_name`, `executor_phone`, `executor_date`;
        ненайденные значения — None.
    """
    max_sheet_row = ws.max_row
    executor_info: dict[str, str | None] = {
        JSON_KEY_EXECUTOR_NAME: None,
        JSON_KEY_EXECUTOR_PHONE: None,
        JSON_KEY_EXECUTOR_DATE: None,
    }

    for row_to_scan in range(max_sheet_row - 5, max_sheet_row - 2):
        if row_to_scan < 1:
            continue

        cell_value_raw = ws.cell(row=row_to_scan, column=2).value

        if not isinstance(cell_value_raw, str):
            continue

        cell_value_lower = cell_value_raw.lower()

        if cell_value_lower.startswith(TABLE_PARSE_EXECUTOR.lower()):
            with suppress(IndexError):
                executor_info[JSON_KEY_EXECUTOR_NAME] = cell_value_raw.split(":", 1)[1].strip()
        elif cell_value_lower.startswith(TABLE_PARSE_TELEPHONE.lower()):
            with suppress(IndexError):
                executor_info[JSON_KEY_EXECUTOR_PHONE] = cell_value_raw.split(":", 1)[1].strip()
        elif cell_value_lower.startswith(TABLE_PARSE_PREPARATION_DATE.lower()):
            value_part = cell_value_raw[len(TABLE_PARSE_PREPARATION_DATE) :]

            # Формат "Ключ: Значение" — снимаем разделитель; иначе "Ключ Значение".
            # Двоеточия внутри времени ("18:49:35") при этом сохраняются.
            final_value = (
                value_part[value_part.find(":") + 1 :] if value_part.lstrip().startswith(":") else value_part
            )

            executor_info[JSON_KEY_EXECUTOR_DATE] = final_value.strip()

    return executor_info
