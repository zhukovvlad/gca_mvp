"""Чтение позиций одного лота для одного подрядчика.

Перенос `app/excel_parser/get_lot_positions.py` из `parser_tender_xlsx@0e178c0`.

Два отличия от исходника (оба — по факту прогона на реальной смете ГП,
docs/phase3-parser.md):

1. **Начало блока позиций.** Исходник брал `max(START_INDEXING_POSITION_ROW=13,
   lot_start_row)`. В смете ГП лот начинается со строки 11, и константа молча
   съедала две первые строки — включая строку-раздел лота. Авторитетна граница
   лота: она найдена по реальному маркеру "Лот №", а не угадана. Константа
   убрана.
2. **Стоимость проверок.** Признак «ячейка в колонке A объединена» брался
   линейным перебором всех объединённых диапазонов на каждой строке, а проверка
   пустоты строки разворачивала `ws[row]` до `ws.max_column`. На реальном файле
   это давало минуты на смету. Теперь: индекс объединённых строк строится один
   раз, пустота проверяется в пределах блока подрядчика.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .build_merged_shape_map import merged_rows_in_first_column
from .constants import (
    JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW,
    JSON_KEY_ARTICLE_SMR,
    JSON_KEY_CHAPTER_NUMBER,
    JSON_KEY_COMMENT_ORGANIZER,
    JSON_KEY_JOB_TITLE,
    JSON_KEY_JOB_TITLE_NORMALIZED,
    JSON_KEY_NUMBER,
    JSON_KEY_QUANTITY,
    JSON_KEY_UNIT,
    TABLE_PARSE_ADDITIONAL_WORKS_TITLE,
)
from .errors import EstimateParseError
from .get_items_dict import get_items_dict
from .parse_contractor_row import parse_contractor_row
from .sanitize_text import normalize_job_title_with_lemmatization
from .sheet import cell_text_is_blank, contractor_last_column, normalized_cell_text, row_is_empty

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class LotRows:
    """Строки лота: позиции и, если она есть, агрегатная строка допработ."""

    positions: dict[str, Any]
    """Позиции подрядчика; ключи — порядковые номера в виде строк."""

    additional_works: dict[str, Any] | None = None
    """Агрегатная строка «Дополнительные работы» либо None, если её в файле нет.

    В Ф2 та же строка ПРИСУТСТВУЕТ и в `positions` — переходное решение спеки
    §2.2; исключение делает Ф4.
    """


def get_lot_positions(ws: Worksheet, contractor: dict[str, Any], lot_start_row: int, lot_end_row: int) -> LotRows:
    """Извлекает позиции подрядчика строго в границах лота.

    Обход идёт от `lot_start_row` до `lot_end_row` включительно и **досрочно
    прекращается** на первой строке, где ячейка колонки A входит в объединённый
    диапазон: это признак конца блока позиций и начала итогов (AGENTS.md §11).
    Полностью пустые строки пропускаются.

    Общие колонки читаются по фиксированным позициям: A — номер, B — номер
    раздела, C — статья СМР, D — наименование, F — комментарий организатора,
    G — единица измерения, H — количество организатора. Колонки подрядчика
    читает `parse_contractor_row`.

    Строка с пустыми A и B — кандидат в агрегатную строку допработ (спека Ф2
    §2.2). Если её нормализованное название совпадает с «Дополнительные работы»,
    она копируется в `LotRows.additional_works` и ОСТАЁТСЯ в `positions`
    (переходное решение, снимается в Ф4). Любое другое название такой строки —
    структурный отказ, как и вторая подобная строка в пределах лота.

    Args:
        ws: лист Excel.
        contractor: словарь подрядчика; нужны `column_start` и `merged_shape`.
        lot_start_row: первая строка лота.
        lot_end_row: последняя строка лота (включительно).

    Returns:
        `LotRows` с позициями (ключи — порядковые номера в виде строк) и
        агрегатной строкой допработ либо `None`.

    Raises:
        EstimateParseError: строка-кандидат (пустые A и B) названа не так, как
            единственная известная агрегатная строка, либо такая строка
            встретилась в лоте второй раз.
    """
    positions: dict[str, Any] = {}
    item_index: int = 1
    additional_works: dict[str, Any] | None = None

    merged_first_column_rows = merged_rows_in_first_column(ws)
    last_column = contractor_last_column(contractor)

    for current_row_num in range(lot_start_row, lot_end_row + 1):
        # Объединённая ячейка в первой колонке — конец блока позиций.
        if current_row_num in merged_first_column_rows:
            log.debug(
                "get_lot_positions: строка %s содержит объединённую ячейку — конец блока позиций",
                current_row_num,
            )
            break

        if row_is_empty(ws, current_row_num, last_column):
            log.debug("get_lot_positions: строка %s пустая — пропускаем", current_row_num)
            continue

        item = get_items_dict(contractor["merged_shape"]["colspan"])
        item[JSON_KEY_NUMBER] = ws.cell(row=current_row_num, column=1).value
        item[JSON_KEY_CHAPTER_NUMBER] = ws.cell(row=current_row_num, column=2).value
        item[JSON_KEY_ARTICLE_SMR] = ws.cell(row=current_row_num, column=3).value
        original_job_title = ws.cell(row=current_row_num, column=4).value
        item[JSON_KEY_JOB_TITLE] = original_job_title
        item[JSON_KEY_JOB_TITLE_NORMALIZED] = normalize_job_title_with_lemmatization(original_job_title)
        item[JSON_KEY_COMMENT_ORGANIZER] = ws.cell(row=current_row_num, column=6).value
        item[JSON_KEY_UNIT] = ws.cell(row=current_row_num, column=7).value
        item[JSON_KEY_QUANTITY] = ws.cell(row=current_row_num, column=8).value

        contractor_specific_data = parse_contractor_row(ws, current_row_num, contractor)
        item.update(contractor_specific_data)

        if cell_text_is_blank(item[JSON_KEY_NUMBER]) and cell_text_is_blank(item[JSON_KEY_CHAPTER_NUMBER]):
            # Строка без номера и без раздела — не позиция и не раздел. В смете
            # ГП такая ровно одна: агрегатная строка допработ перед ИТОГО.
            title = normalized_cell_text(item[JSON_KEY_JOB_TITLE])
            if title.casefold() != TABLE_PARSE_ADDITIONAL_WORKS_TITLE.casefold():
                raise EstimateParseError(
                    f"Строка {current_row_num} не имеет ни номера, ни раздела, "
                    f"а называется «{title}». Единственная известная строка такого "
                    f"вида — «{TABLE_PARSE_ADDITIONAL_WORKS_TITLE}»; её деньги мы "
                    "умеем отнести, а эти — нет, и они молча легли бы в статью "
                    "последнего раздела."
                )
            if additional_works is not None:
                raise EstimateParseError(
                    f"Агрегатных строк «{TABLE_PARSE_ADDITIONAL_WORKS_TITLE}» "
                    f"больше одной: строки "
                    f"{additional_works[JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW]} и "
                    f"{current_row_num}. Ожидается не больше одной на лот."
                )
            additional_works = {
                JSON_KEY_JOB_TITLE: item[JSON_KEY_JOB_TITLE],
                JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW: current_row_num,
                **contractor_specific_data,
            }
            # Строка НЕ пропускается: в Ф2 она остаётся позицией (спека §2.2).

        positions[str(item_index)] = item
        item_index += 1

    return LotRows(positions=positions, additional_works=additional_works)
