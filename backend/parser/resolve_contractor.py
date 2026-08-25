"""Разрешение раскладки блока подрядчика по подписям двухъярусной шапки.

Смысл колонки — пара «тип группы + эффективная подпись» (спека §2.2); ширина
блока задаёт только физическую границу сканирования. Словарь пар ЗАКРЫТЫЙ:
нечётких эвристик нет — измерено, что подписи не плавают (спека §1.3), а
нечёткое правило вернуло бы ровно тот класс ошибки, от которого фича
избавляется.

Здесь же — единственное место семантического контракта (§2.6): неизвестная
пара, повтор ключа и пропуск обязательного ключа дают EstimateParseError с
именем подрядчика. Геометрию (наличие объединения) проверяет
estimate._validate_contractor_geometry ДО вызова; предупреждениями о НАБОРЕ
ключей занимается layout.check_estimate_layout ПОСЛЕ.
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    JSON_KEY_COMMENT_CONTRACTOR,
    JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
    JSON_KEY_INDIRECT_COSTS,
    JSON_KEY_MATERIALS,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_SUGGESTED_QUANTITY,
    JSON_KEY_TOTAL,
    JSON_KEY_TOTAL_COST,
    JSON_KEY_UNIT_COST,
    JSON_KEY_WORKS,
    TABLE_PARSE_COMMENT_CONTRACTOR_LABEL,
    TABLE_PARSE_DEVIATION_COLUMN_LABEL,
    TABLE_PARSE_ORGANIZER_QUANTITY_LABEL,
    TABLE_PARSE_SUGGESTED_QUANTITY,
    TABLE_PARSE_TOTAL_COST_GROUP_PREFIX,
    TABLE_PARSE_UNIT_COST_GROUP_PREFIX,
)
from .errors import EstimateParseError
from .sheet import normalized_cell_text

#: Подписи колонок внутри денежной группы → хвост составного ключа.
_COST_GROUP_SUBLABELS = {
    "материалы": JSON_KEY_MATERIALS,
    "смр": JSON_KEY_WORKS,
    "косвенные расходы": JSON_KEY_INDIRECT_COSTS,
    "всего": JSON_KEY_TOTAL,
}

#: Закрытый словарь §2.2: пара «тип группы + эффективная подпись (casefold)» →
#: ключ JSON. Подписи одиночных колонок берутся из constants (одна правда).
COLUMN_KEY_BY_PAIR: dict[tuple[str | None, str], str] = {
    **{
        (JSON_KEY_UNIT_COST, label): f"{JSON_KEY_UNIT_COST}.{part}"
        for label, part in _COST_GROUP_SUBLABELS.items()
    },
    **{
        (JSON_KEY_TOTAL_COST, label): f"{JSON_KEY_TOTAL_COST}.{part}"
        for label, part in _COST_GROUP_SUBLABELS.items()
    },
    (None, TABLE_PARSE_SUGGESTED_QUANTITY.casefold()): JSON_KEY_SUGGESTED_QUANTITY,
    (None, TABLE_PARSE_ORGANIZER_QUANTITY_LABEL.casefold()): JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    (None, TABLE_PARSE_COMMENT_CONTRACTOR_LABEL.casefold()): JSON_KEY_COMMENT_CONTRACTOR,
    (None, TABLE_PARSE_DEVIATION_COLUMN_LABEL.casefold()): JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
}

#: Восемь денежных ключей обязаны присутствовать ровно один раз (спека §2.2).
REQUIRED_COLUMN_KEYS = frozenset(
    key for key in COLUMN_KEY_BY_PAIR.values() if "." in key
)
#: Четыре остальных — опциональны, но не более одного раза каждый.
OPTIONAL_COLUMN_KEYS = frozenset(COLUMN_KEY_BY_PAIR.values()) - REQUIRED_COLUMN_KEYS


@dataclass(frozen=True)
class BlockLayout:
    """Раскладка блока: физический порядок ключей и якоря денежных групп.

    column_keys — tuple, а не list: frozen=True запрещает переприсваивание
    поля, но не мутацию списка внутри него (спека §2.4).

    unit_cost_offset и total_cost_offset — это АНКОРЫ денежных групп: смещение
    (относительно начала блока) левой границы горизонтального объединения
    группы, а не индекс какой-либо конкретной подписи внутри неё (в
    частности, НЕ индекс «.materials» — Материалы не обязана быть первой
    физической колонкой группы). Это обязательное условие, а не стиль: single
    consumer — get_proposals.py — читает по этому смещению
    `ws.cell(row=header_row, column=col_start + offset).value`, чтобы достать
    текст объединённого заголовка группы для build_vat_rate. openpyxl хранит
    значение только в ЛЕВОЙ ВЕРХНЕЙ ячейке объединённого диапазона; любая
    другая ячейка того же диапазона возвращает None. Смещение, указывающее не
    на якорь, тихо теряет заголовок группы — а с ним и ставку НДС.
    """

    column_keys: tuple[str, ...]
    unit_cost_offset: int
    total_cost_offset: int


@dataclass(frozen=True)
class ResolvedContractor:
    """Геометрия блока и его раскладка — один объект, а не два списка,
    связываемых индексом: тихий сдвиг индекса при нескольких подрядчиках
    означал бы деньги не в тех полях (спека §2.4)."""

    geometry: dict[str, Any]
    layout: BlockLayout


def _header_merge_map(
    ws: Worksheet, header_row: int
) -> dict[tuple[int, int], tuple[int, int, int, int]]:
    """Границы объединённых диапазонов, накрывающих два яруса шапки.

    Раскрытие объединения — строго в его реальном диапазоне (спека §2.3 п.1):
    карта строится по границам диапазонов, «последняя виденная» шапка вправо
    не тянется.
    """
    bounds_by_cell: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for merged_range in ws.merged_cells.ranges:
        if merged_range.max_row < header_row or merged_range.min_row > header_row + 1:
            continue
        bounds = (merged_range.min_row, merged_range.min_col, merged_range.max_row, merged_range.max_col)
        for row in (header_row, header_row + 1):
            if merged_range.min_row <= row <= merged_range.max_row:
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    bounds_by_cell[(row, col)] = bounds
    return bounds_by_cell


def resolve_contractor(ws: Worksheet, contractor: dict[str, Any], header_row: int) -> ResolvedContractor:
    """Разрешает раскладку одного блока по подписям шапки.

    Верхний ярус — строка header_row (та же, где шапка A–D), нижний —
    header_row+1. Тип группы приписывается колонке, только если она входит в
    горизонтально объединённую ячейку верхнего яруса шириной больше единицы;
    эффективная подпись — нижний ярус, если он непуст, иначе верхний (для
    колонки внутри группы — только нижний, §2.3 п.4). Значение объединённой
    ячейки читается из её якоря.

    Raises:
        EstimateParseError: неизвестная пара / повтор ключа / пропуск
            обязательного ключа — контракт §2.6, только здесь.
    """
    title = contractor.get("value")
    col_start: int = contractor["column_start"]
    colspan: int = contractor["merged_shape"]["colspan"]
    merges = _header_merge_map(ws, header_row)

    def raw_at(row: int, col: int) -> Any:
        bounds = merges.get((row, col))
        if bounds is not None:
            return ws.cell(row=bounds[0], column=bounds[1]).value
        return ws.cell(row=row, column=col).value

    keys: list[str] = []
    seen: dict[str, tuple[str, str]] = {}  # ключ -> (координата, показанная подпись)
    group_anchor_offset: dict[str, int] = {}  # тип группы -> смещение её якоря

    for offset in range(colspan):
        col = col_start + offset
        coordinate = f"{get_column_letter(col)}{header_row}"
        upper_bounds = merges.get((header_row, col))
        upper_raw = raw_at(header_row, col)
        lower_raw = raw_at(header_row + 1, col)
        upper_shown = normalized_cell_text(upper_raw)
        lower_shown = normalized_cell_text(lower_raw)

        in_group = upper_bounds is not None and upper_bounds[3] > upper_bounds[1]
        if in_group:
            head = upper_shown.casefold()
            if head.startswith(TABLE_PARSE_UNIT_COST_GROUP_PREFIX):
                group_type: str | None = JSON_KEY_UNIT_COST
            elif head.startswith(TABLE_PARSE_TOTAL_COST_GROUP_PREFIX):
                group_type = JSON_KEY_TOTAL_COST
            else:
                group_type = None
            key = COLUMN_KEY_BY_PAIR.get((group_type, lower_shown.casefold())) if group_type else None
            shown = lower_shown
            if group_type is not None and group_type not in group_anchor_offset:
                group_anchor_offset[group_type] = upper_bounds[1] - col_start
        else:
            shown = lower_shown or upper_shown
            key = COLUMN_KEY_BY_PAIR.get((None, shown.casefold()))

        if key is None:
            raise EstimateParseError(
                f"Не удалось опознать колонку {coordinate} блока подрядчика «{title}»: "
                f"заголовок группы «{upper_raw or ''}», подпись «{lower_raw or ''}». "
                "Смысл колонок определяется подписями шапки по закрытому словарю; "
                "незнакомая подпись — отказ, иначе стоимости легли бы не в те поля молча."
            )
        if key in seen:
            first_coordinate, first_shown = seen[key]
            raise EstimateParseError(
                f"Колонки {first_coordinate} и {coordinate} блока подрядчика «{title}» "
                f"дают один и тот же ключ «{key}» (подписи «{first_shown}» и «{shown}»). "
                "Каждый ключ допустим не более одного раза."
            )
        seen[key] = (coordinate, shown)
        keys.append(key)

    missing = sorted(REQUIRED_COLUMN_KEYS - set(keys))
    if missing:
        raise EstimateParseError(
            f"В блоке подрядчика «{title}» нет обязательных денежных колонок: "
            f"{', '.join(missing)}. Координаты у отсутствующей колонки не существует; "
            "без полной восьмёрки денег предложение не читается."
        )

    column_keys = tuple(keys)
    # Обе денежных группы обязательны (REQUIRED_COLUMN_KEYS выше уже проверил
    # полную восьмёрку), и тип группы колонка получает только внутри
    # горизонтального объединения шириной больше единицы (in_group выше) — то
    # есть к этой строке group_anchor_offset обязан содержать обе группы.
    return ResolvedContractor(
        geometry=contractor,
        layout=BlockLayout(
            column_keys=column_keys,
            unit_cost_offset=group_anchor_offset[JSON_KEY_UNIT_COST],
            total_cost_offset=group_anchor_offset[JSON_KEY_TOTAL_COST],
        ),
    )
