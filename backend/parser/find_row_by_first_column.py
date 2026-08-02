"""Поиск строки по содержимому первой колонки.

Перенос `app/excel_parser/find_row_by_first_column.py` из
`parser_tender_xlsx@0e178c0` без изменений логики.
"""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet


def find_row_by_first_column(
    ws: Worksheet, target_text: str, start_row: int = 1, end_row: int | None = None
) -> int | None:
    """Возвращает номер первой строки, где ячейка колонки A содержит `target_text`.

    Поиск подстрочный и регистрозависимый, в диапазоне `start_row..end_row`
    включительно.

    Args:
        ws: лист Excel.
        target_text: искомый текст.
        start_row: начало диапазона (1-индексация).
        end_row: конец диапазона включительно; по умолчанию — `ws.max_row`.

    Returns:
        Номер строки либо None, если совпадений нет.
    """
    if end_row is None:
        end_row = ws.max_row

    for row_index in range(start_row, end_row + 1):
        cell_value = ws.cell(row=row_index, column=1).value
        if isinstance(cell_value, str) and target_text in cell_value:
            return row_index

    return None
