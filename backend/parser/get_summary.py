"""Чтение блока итогов (summary).

Перенос `app/excel_parser/get_summary.py` из `parser_tender_xlsx@0e178c0`.

Отличия от исходника:

1. **Начало поиска.** Исходник искал блок итогов от константы
   `START_INDEXING_POSITION_ROW = 13`. Константа убрана вместе с ней же в
   `get_lot_positions`; теперь floor передаётся явно (`search_start_row`) —
   это строка начала позиций. Так поиск не может наткнуться на объединённую
   ячейку шапки таблицы (в смете ГП это A9:A10 «№ п/п»).
2. **Стоимость проверок** — индекс объединённых строк и ограниченная проверка
   пустоты, как в `get_lot_positions`.
3. **Правила вынесены в ядро.** Распознавание метки, инъективность ключа,
   сверка арифметики и предупреждения больше не живут здесь — они в
   `summary_block` (спека Ф4a §2.2–§2.7) и проверяются без файла. Здесь только
   обход листа.
"""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from .build_merged_shape_map import merged_rows_in_first_column
from .parse_contractor_row import parse_contractor_row
from .resolve_contractor import ResolvedContractor
from .sheet import contractor_last_column, row_is_empty
from .summary_block import SummaryBlock, SummaryRow, build_summary_block


def get_summary(ws: Worksheet, contractor: ResolvedContractor, search_start_row: int) -> SummaryBlock:
    """Извлекает итоговые строки подрядчика из блока итогов внизу таблицы.

    Здесь только обход листа: где блок начинается, где кончается и что стоит в
    каждой строке. Все правила — распознавание метки, инъективность ключа,
    сверка арифметики, предупреждения — живут в `summary_block` и проверяются
    без файла (спека Ф4a §2.2–§2.7).

    Args:
        ws: лист Excel.
        contractor: разрешённый блок подрядчика — геометрия в `.geometry`,
            раскладка в `.layout`.
        search_start_row: первая строка позиций, откуда начинается поиск блока.

    Returns:
        `SummaryBlock`: строки блока и parser warnings о нём.
    """
    merged_first_column_rows = merged_rows_in_first_column(ws)
    last_column = contractor_last_column(contractor.geometry)

    summary_start_row = -1
    for row_num in range(search_start_row, ws.max_row + 1):
        if row_num in merged_first_column_rows:
            summary_start_row = row_num
            break

    rows: list[SummaryRow] = []
    if summary_start_row != -1:
        for current_row_num in range(summary_start_row, ws.max_row + 1):
            if row_is_empty(ws, current_row_num, last_column):
                break  # Пустая строка означает конец блока
            rows.append(
                SummaryRow(
                    row=current_row_num,
                    label=ws.cell(row=current_row_num, column=1).value,
                    values=parse_contractor_row(ws, current_row_num, contractor),
                )
            )

    return build_summary_block(rows, search_start_row=search_start_row)
