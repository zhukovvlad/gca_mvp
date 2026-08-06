"""Границы сканирования листа.

Новый модуль фазы 3 (в исходнике аналога нет). Он существует ради одного
требования: **не вводить полностраничных обходов** (docs/phase0-input-data.md,
docs/phase3-start.md §6).

В реальных выгрузках `ws.max_column` равен 16384 из-за одной пустой
стилизованной ячейки в XFD, тогда как данных 24 колонки. Исходный парсер
опирался на `ws.max_column` и на `ws[row]` (который тоже разворачивается до
`max_column`), и openpyxl материализовал ~16 тыс. объектов ячеек на каждую
строку. Здесь все построчные проверки ограничены фактической правой границей
данных — правым краем блока подрядчика.
"""

from __future__ import annotations

from typing import Any

from openpyxl.worksheet.worksheet import Worksheet


def contractor_last_column(contractor: dict[str, Any]) -> int:
    """Правая граница данных подрядчика (1-индексация, включительно).

    Всё, что парсер читает по строке позиции, лежит левее: общие колонки
    (№, раздел, статья, наименование, комментарий, единица, количество) и блок
    самого подрядчика. Колонки правее к смете отношения не имеют.

    Args:
        contractor: словарь подрядчика от `read_contractors` — нужны
            `column_start` и `merged_shape.colspan`.

    Returns:
        Номер последней значащей колонки. Если ширина блока неизвестна,
        считается, что подрядчик занимает одну колонку.
    """
    column_start = contractor.get("column_start") or 1
    colspan = (contractor.get("merged_shape") or {}).get("colspan", 1)
    return column_start + colspan - 1


def row_is_empty(ws: Worksheet, row: int, max_col: int) -> bool:
    """Проверяет, пуста ли строка в пределах колонок 1..`max_col`.

    Замена исходной проверки `all(cell.value is None for cell in ws[row])`:
    та разворачивалась до `ws.max_column` и на раздутых листах стоила
    ~0,06 с на строку.
    """
    for row_cells in ws.iter_rows(min_row=row, max_row=row, max_col=max_col):
        return all(cell.value is None for cell in row_cells)
    return True


def normalized_cell_text(value: Any) -> str:
    """Текст ячейки в форме, пригодной для сверки с ожидаемым.

    Схлопывает любые пробельные последовательности (включая переносы строк и
    неразрывный пробел) в один пробел и обрезает края. `None` даёт пустую строку.
    Регистр НЕ трогает — за это отвечает вызывающий, чтобы не мешать сверку и
    показ значения человеку.
    """
    if value is None:
        return ""
    return " ".join(str(value).split())


def cell_text_is_blank(value: Any) -> bool:
    """Пуста ли ячейка по тексту.

    Пустотой считаются `None`, пустая строка, пробелы, табуляции, переносы и
    неразрывный пробел (все они схлопываются `normalized_cell_text`). Ноль
    пустотой НЕ считается: `0` — это значение.
    """
    return normalized_cell_text(value) == ""
