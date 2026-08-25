"""Сборка предложений подрядчиков по одному лоту.

Перенос `app/excel_parser/get_proposals.py` из `parser_tender_xlsx@0e178c0`.

Отличие от исходника: в `get_summary` передаётся начало блока позиций
(`start_row`) — раньше поиск итогов стартовал от константы 13, см. `get_summary`.

Про «один подрядчик» (AGENTS.md §4): отдельной ветки для сметы ГП не нужно.
Смета с одним блоком подрядчика — это `len(contractors) == 1`, поэтому обход
даёт ровно одно предложение `contractor_1`. Слой proposals сохраняется как
есть — это задел на возврат тендеров (§4).

`contractors` приходит сюда уже РАЗРЕШЁННЫМ из `parse_worksheet`
(`resolve_contractor` на каждый блок) — `read_contractors` отсюда убран: геометрия
не читается второй раз (спека §2.5). Якоря групповых шапок `unit_cost` и
`total_cost` — это смещения, уже посчитанные в `resolved.layout` из того же
перечня ключей, что и сами колонки (спека §2.4).

Ф4б: здесь же читаются две ячейки групповой шапки ценового блока и вызывается
`vat_rate.build_vat_rate` — это единственное место, где блок подрядчика и
результат `get_summary` уже рядом (спека Ф4б §1.4 факт 1, §2.8).
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    JSON_KEY_CONTRACTOR_ACCREDITATION,
    JSON_KEY_CONTRACTOR_ADDITIONAL_INFO,
    JSON_KEY_CONTRACTOR_ADDITIONAL_WORKS,
    JSON_KEY_CONTRACTOR_ADDRESS,
    JSON_KEY_CONTRACTOR_COORDINATE,
    JSON_KEY_CONTRACTOR_HEIGHT,
    JSON_KEY_CONTRACTOR_INDEX,
    JSON_KEY_CONTRACTOR_INN,
    JSON_KEY_CONTRACTOR_ITEMS,
    JSON_KEY_CONTRACTOR_POSITIONS,
    JSON_KEY_CONTRACTOR_SUMMARY,
    JSON_KEY_CONTRACTOR_TITLE,
    JSON_KEY_CONTRACTOR_WIDTH,
    JSON_KEY_VAT_RATE,
)
from .get_additional_info import get_additional_info
from .get_lot_positions import get_lot_positions
from .get_summary import get_summary
from .resolve_contractor import ResolvedContractor
from .vat_rate import build_vat_rate


@dataclass(frozen=True)
class LotProposals:
    """Предложения лота и предупреждения, собранные при их разборе."""

    proposals: dict[str, dict[str, Any]]
    warnings: list[str]


def get_proposals(
    ws: Worksheet, start_row: int, end_row: int, *, header_row: int, contractors: Sequence[ResolvedContractor]
) -> LotProposals:
    """Собирает предложения всех подрядчиков для одного лота.

    Позиции берутся строго в границах лота, итоги и дополнительная информация —
    общие по всему листу.

    Реквизиты подрядчика (ИНН, адрес, аккредитация) читаются из трёх строк под
    заголовком, и только если заголовок занимает одну строку (`rowspan == 1`).

    `header_row` обязателен и без значения по умолчанию: умолчание завело бы
    второй путь «найти строку шапки самому» рядом с уже вычисленным вызывающей
    стороной значением — то есть вторую правду о том, где шапка.

    `contractors` приходит РАЗРЕШЁННЫМ из `parse_worksheet`: каждый блок уже
    прошёл `resolve_contractor` выше по стеку, и `read_contractors` здесь
    больше не зовётся — геометрия не читается второй раз (спека §2.5). Пустой
    список — валидный вход, обход по нему пуст, отдельная ветка не нужна.

    Ставка НДС (`vat_rate`) читается из двух ячеек строки `header_row` —
    якорей групповых шапок `unit_cost` и `total_cost`. Смещения этих якорей
    относительно начала блока — `resolved.layout.unit_cost_offset` и
    `.total_cost_offset`, посчитанные `resolve_contractor` из того же перечня
    ключей, что и сами колонки (спека §2.4); захардкоженные смещения читали бы
    не те ячейки молча. Ключ кладётся в предложение ВСЕГДА, в том числе
    значением `None`.

    Args:
        ws: лист Excel.
        start_row: первая строка лота.
        end_row: последняя строка лота.
        header_row: номер строки шапки таблицы позиций, найденный выше по стеку
            (`estimate._validate_column_headers`).
        contractors: разрешённые блоки подрядчиков (`resolve_contractor` на
            каждый), в порядке следования по листу.

    Returns:
        `LotProposals`: словарь `{"contractor_1": {...}, ...}` (пустой, если
        подрядчики не найдены) и предупреждения, собранные при разборе блоков
        итогов подрядчиков и ставки НДС.
    """
    proposals: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []

    for i, resolved in enumerate(contractors, start=1):
        geometry = resolved.geometry
        contractor_name = geometry.get("value")
        contractor_row_start = geometry.get("row_start")
        contractor_col_start: int = geometry["column_start"]
        contractor_coordinate = geometry.get("coordinate")
        merged_shape: dict[str, int] = geometry.get("merged_shape", {})
        rowspan: int = merged_shape.get("rowspan", 1)
        colspan: int = merged_shape.get("colspan", 1)

        inn_val: Any = None
        address_val: Any = None
        accreditation_val: Any = None
        if rowspan == 1 and contractor_row_start is not None:
            inn_val = ws.cell(row=contractor_row_start + 1, column=contractor_col_start).value
            address_val = ws.cell(row=contractor_row_start + 2, column=contractor_col_start).value
            accreditation_val = ws.cell(row=contractor_row_start + 3, column=contractor_col_start).value

        lot_rows = get_lot_positions(ws, resolved, lot_start_row=start_row, lot_end_row=end_row)
        summary = get_summary(ws, resolved, search_start_row=start_row)
        warnings.extend(summary.warnings)

        # Якоря групповых шапок — из разрешённой раскладки: смещения живут в
        # BlockLayout и посчитаны из того же перечня ключей, что сами колонки.
        unit_cost_label = ws.cell(row=header_row, column=contractor_col_start + resolved.layout.unit_cost_offset).value
        total_cost_label = ws.cell(row=header_row, column=contractor_col_start + resolved.layout.total_cost_offset).value

        vat_rate_result = build_vat_rate(unit_cost_label, total_cost_label, summary.lines)
        warnings.extend(vat_rate_result.warnings)
        # Деньги и ставки в JSON — десятичные строки, не Decimal и не float
        # (AGENTS.md §3): `None` остаётся `None`.
        vat_rate_value = str(vat_rate_result.rate) if vat_rate_result.rate is not None else None

        contractor_items_data = {
            JSON_KEY_CONTRACTOR_POSITIONS: lot_rows.positions,
            JSON_KEY_CONTRACTOR_SUMMARY: summary.lines,
            JSON_KEY_CONTRACTOR_ADDITIONAL_WORKS: lot_rows.additional_works,
        }

        contractor_additional_info_data = get_additional_info(ws, geometry)

        proposal_key = f"{JSON_KEY_CONTRACTOR_INDEX}{i}"
        proposals[proposal_key] = {
            JSON_KEY_CONTRACTOR_TITLE: contractor_name,
            JSON_KEY_CONTRACTOR_INN: inn_val,
            JSON_KEY_CONTRACTOR_ADDRESS: address_val,
            JSON_KEY_CONTRACTOR_ACCREDITATION: accreditation_val,
            JSON_KEY_CONTRACTOR_COORDINATE: contractor_coordinate,
            JSON_KEY_CONTRACTOR_WIDTH: colspan,
            JSON_KEY_CONTRACTOR_HEIGHT: rowspan,
            JSON_KEY_VAT_RATE: vat_rate_value,
            JSON_KEY_CONTRACTOR_ITEMS: contractor_items_data,
            JSON_KEY_CONTRACTOR_ADDITIONAL_INFO: contractor_additional_info_data,
        }

    return LotProposals(proposals=proposals, warnings=warnings)
