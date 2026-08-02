"""Чтение блока «Дополнительная информация».

Перенос `app/excel_parser/get_additional_info.py` из
`parser_tender_xlsx@0e178c0`.

Отличие от исходника: проверка пустоты строки ограничена блоком подрядчика
вместо разворачивания `ws[row]` до `ws.max_column` (см. `sheet.row_is_empty`).
"""

from __future__ import annotations

from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from .constants import SEARCH_KEYWORD_ADDITIONAL_INFO
from .find_row_by_first_column import find_row_by_first_column
from .sheet import contractor_last_column, row_is_empty


def get_additional_info(ws: Worksheet, contractor: dict[str, Any]) -> dict[str, Any]:
    """Извлекает пары «ключ-значение» блока дополнительной информации.

    Блок ищется по маркеру `SEARCH_KEYWORD_ADDITIONAL_INFO` в колонке A. Данные
    начинаются со следующей строки: ключ — колонка B, значение — колонка
    подрядчика (`column_start`). Чтение прекращается на первой полностью пустой
    строке или на конце листа. Пустые значения превращаются в "".

    В смете ГП это последний блок листа — он лежит НИЖЕ итогов, и позиции до
    него не доходят: `get_lot_positions` останавливается на объединённой ячейке
    строки «ИТОГО». Это важно, потому что в колонке подрядчика здесь лежит текст
    («Представлено», «К обсуждению»), который при сломанном досрочном выходе был
    бы разобран как количество (docs/phase0-input-data.md).

    Args:
        ws: лист Excel.
        contractor: словарь подрядчика; нужен `column_start`.

    Returns:
        Словарь дополнительной информации; пустой, если блок не найден.
    """
    additional_info: dict[str, Any] = {}

    header_row_num = find_row_by_first_column(ws, SEARCH_KEYWORD_ADDITIONAL_INFO)

    if not header_row_num:
        return additional_info

    current_row_num = header_row_num + 1
    contractor_data_col = contractor["column_start"]
    last_column = max(contractor_last_column(contractor), 2)

    while True:
        if current_row_num > ws.max_row:
            break

        if row_is_empty(ws, current_row_num, last_column):
            break

        key_data = ws.cell(row=current_row_num, column=2).value
        value_data = ws.cell(row=current_row_num, column=contractor_data_col).value

        processed_key = None
        if key_data is not None:
            processed_key = str(key_data).strip() or None

        processed_value = str(value_data).strip() if value_data is not None else ""

        if processed_key:
            additional_info[processed_key] = processed_value

        current_row_num += 1

    return additional_info
