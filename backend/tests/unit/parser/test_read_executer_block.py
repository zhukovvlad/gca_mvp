"""Тесты чтения блока исполнителя.

Перенос `app/tests/excel_parser/test_read_executer_block.py` из
`parser_tender_xlsx@0e178c0` без содержательных изменений — логика при
адаптации не менялась.

Здесь берётся `executor_date` — источник `estimates.data_prepared_on_date`
(AGENTS.md §4). В полученном образце сметы ГП этого блока нет, см.
`test_gp_estimate_without_executor_block`.
"""
from __future__ import annotations

from openpyxl import Workbook

from parser.constants import (
    JSON_KEY_EXECUTOR_DATE,
    JSON_KEY_EXECUTOR_NAME,
    JSON_KEY_EXECUTOR_PHONE,
)
from parser.read_executer_block import read_executer_block


def set_max_row(ws, row_num):
    """Задаёт максимальную строку листа — от неё считается окно сканирования."""
    ws[f"A{row_num}"] = "dummy_data_to_set_max_row"


def test_happy_path_with_colon():
    """Идеальный сценарий: дата отделена двоеточием."""
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B15"] = "Исполнитель: Иванов И.И."
    ws["B16"] = "Телефон: +7 (999) 123-45-67"
    ws["B17"] = "Дата составления: 15.07.2025"

    result = read_executer_block(ws)

    assert result[JSON_KEY_EXECUTOR_NAME] == "Иванов И.И."
    assert result[JSON_KEY_EXECUTOR_PHONE] == "+7 (999) 123-45-67"
    assert result[JSON_KEY_EXECUTOR_DATE] == "15.07.2025"


def test_date_format_without_colon():
    """Дата без разделителя, но с двоеточиями во времени."""
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B17"] = "Дата составления 07.05.2025 18:49:35"

    result = read_executer_block(ws)

    assert result[JSON_KEY_EXECUTOR_NAME] is None
    assert result[JSON_KEY_EXECUTOR_DATE] == "07.05.2025 18:49:35"


def test_date_format_with_space_before_colon():
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B17"] = "Дата составления : 12.12.2025"

    assert read_executer_block(ws)[JSON_KEY_EXECUTOR_DATE] == "12.12.2025"


def test_is_case_insensitive():
    ws = Workbook().active
    set_max_row(ws, 10)  # сканируются строки 5, 6, 7
    ws["B5"] = "ИСПОЛНИТЕЛЬ: Петров П.П."
    ws["B6"] = "телефон: 88005553535"
    ws["B7"] = "ДАТА СОСТАВЛЕНИЯ: 01.01.2025"

    result = read_executer_block(ws)

    assert result[JSON_KEY_EXECUTOR_NAME] == "Петров П.П."
    assert result[JSON_KEY_EXECUTOR_PHONE] == "88005553535"
    assert result[JSON_KEY_EXECUTOR_DATE] == "01.01.2025"


def test_no_data_found():
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B15"] = "Просто какой-то текст"

    assert all(value is None for value in read_executer_block(ws).values())


def test_ignores_data_in_wrong_column():
    """Данные ищутся только в колонке B."""
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["A15"] = "Исполнитель: Иванов И.И."

    assert read_executer_block(ws)[JSON_KEY_EXECUTOR_NAME] is None


def test_handles_small_sheet_gracefully():
    """На коротком листе окно сканирования уходит в отрицательные строки."""
    ws = Workbook().active
    ws["A1"] = "data"
    ws["A2"] = "data"

    assert all(value is None for value in read_executer_block(ws).values())


def test_ignores_non_string_values():
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B15"] = 123456789

    assert all(value is None for value in read_executer_block(ws).values())


def test_handles_missing_colon_for_name():
    """Без двоеточия значение не извлекается, но и падения нет."""
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B15"] = "Исполнитель Иванов И.И."

    assert read_executer_block(ws)[JSON_KEY_EXECUTOR_NAME] is None


def test_handles_missing_colon_for_phone():
    ws = Workbook().active
    set_max_row(ws, 20)
    ws["B16"] = "Телефон 89991234567"

    assert read_executer_block(ws)[JSON_KEY_EXECUTOR_PHONE] is None


def test_ignores_data_in_wrong_rows():
    """Окно сканирования — ровно три строки: max_row-5, -4, -3."""
    ws = Workbook().active
    set_max_row(ws, 20)  # сканируются 15, 16, 17
    ws["B14"] = "Исполнитель: Неправильный"
    ws["B18"] = "Телефон: 555-55-55"

    assert all(value is None for value in read_executer_block(ws).values())


def test_gp_estimate_without_executor_block():
    """Смета ГП: блока исполнителя нет — все три поля пустые.

    Проверено и на реальном образце, и на fixture. Следствие для фазы 4:
    `estimates.data_prepared_on_date` будет NULL, и датой сравнения с
    нормативом станет `contracts.signed_date` — штатный фолбэк AGENTS.md §4.
    """
    ws = Workbook().active
    set_max_row(ws, 30)
    # Хвост сметы ГП — блок дополнительной информации, а не блок исполнителя
    ws["A25"] = "Дополнительная информация:"
    ws["B26"] = "График производства работ"
    ws["B27"] = "Согласие с проектом договора"

    result = read_executer_block(ws)

    assert result == {
        JSON_KEY_EXECUTOR_NAME: None,
        JSON_KEY_EXECUTOR_PHONE: None,
        JSON_KEY_EXECUTOR_DATE: None,
    }
