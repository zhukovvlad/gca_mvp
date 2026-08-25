"""Разрешение раскладки блока по подписям шапки — ядро фичи, без оркестратора.

Пять измеренных раскладок (план, «Замеры») — тестовые случаи, а не ветви кода
(спека §3). Отказы §2.6 проверяются на входах, нарушающих РОВНО ОДНО
ограничение каждый (docs/insights/verifying-guards.md, слой про выбор входа).
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from parser.errors import EstimateParseError
from parser.read_contractors import read_contractors
from parser.resolve_contractor import resolve_contractor

from .sheet_builders import (
    KEYS_8,
    KEYS_10,
    KEYS_12,
    KEYS_GP_11,
    KEYS_PERMUTED_11,
    KEYS_SUBLABELS_SWAPPED_11,
    KEYS_TENDER_11,
    gp_sheet,
)

HEADER_ROW = 9


def _contractor(ws):
    """Геометрия первого подрядчика листа — тем же кодом, что в продакшене."""
    return read_contractors(ws)[1]


class TestMeasuredLayouts:
    @pytest.mark.parametrize(
        ("columns", "offsets"),
        [
            (KEYS_8, (0, 4)),
            (KEYS_10, (1, 5)),
            (KEYS_GP_11, (1, 5)),
            (KEYS_TENDER_11, (1, 5)),
            (KEYS_12, (1, 5)),
        ],
    )
    def test_measured_layouts_resolve_to_measured_keys(self, columns, offsets):
        """Смещения записаны ЛИТЕРАЛАМИ — эталон не выводится из проверяемого."""
        ws = gp_sheet(columns)
        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert result.layout.column_keys == columns
        assert (result.layout.unit_cost_offset, result.layout.total_cost_offset) == offsets

    def test_permuted_columns_resolve_by_labels_not_by_position(self):
        ws = gp_sheet(KEYS_PERMUTED_11)
        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert result.layout.column_keys == KEYS_PERMUTED_11
        assert result.layout.unit_cost_offset == 6
        assert result.layout.total_cost_offset == 1

    def test_sublabels_swapped_inside_group_resolve_to_group_anchor(self):
        """Смещение — якорь группы (её левая физическая граница), а не индекс
        колонки Материалы. Здесь группы стоят на канонических местах группы
        GP-11 (1, 5), но подписи внутри каждой группы идут «СМР, Материалы,
        …» — Материалы сдвинута на одну позицию внутрь группы. Раз якорь
        зависит только от места группы, а не от места Материалы внутри неё,
        ожидаемые смещения совпадают с KEYS_GP_11: (1, 5), а не (2, 6)."""
        ws = gp_sheet(KEYS_SUBLABELS_SWAPPED_11)
        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert result.layout.column_keys == KEYS_SUBLABELS_SWAPPED_11
        assert (result.layout.unit_cost_offset, result.layout.total_cost_offset) == (1, 5)

    def test_real_gp_geometry_built_by_hand(self):
        """Геометрия fixture воспроизведена merge_cells-ами, не строителем —
        строитель не проверяет сам себя (замер плана: J..T, header_row 9)."""
        ws = Workbook().active
        ws["G6"] = "Наименование контрагента"
        ws.merge_cells("J6:T6")
        ws["J6"] = 'ООО "Тест"'
        ws["J10"] = "Предлагаемое количество"
        ws.merge_cells("K9:N9")
        ws["K9"] = "Цена за ед. изм., RUB, ОСН, с учетом НДС 20%"
        ws.merge_cells("O9:R9")
        ws["O9"] = "Стоимость всего, RUB, ОСН, с учетом НДС 20%"
        for col, label in ((11, "Материалы"), (12, "СМР"), (13, "Косвенные расходы"), (14, "Всего")):
            ws.cell(row=10, column=col, value=label)
            ws.cell(row=10, column=col + 4, value=label)
        ws.merge_cells("S9:S10")
        ws["S9"] = "Стоимость всего за объемы заказчика"
        ws["T9"] = "Комментарий участника"

        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)

        assert result.layout.column_keys == KEYS_GP_11
        assert (result.layout.unit_cost_offset, result.layout.total_cost_offset) == (1, 5)

    def test_group_header_without_rate_still_types_the_group(self):
        """`…, с учетом НДС` без ставки и вовсе без хвоста — тип группы тот же;
        ставкой занимается vat_rate, не словарь (спека §2.1)."""
        ws = gp_sheet(KEYS_GP_11, vat_suffix=None)
        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert result.layout.column_keys == KEYS_GP_11

    def test_geometry_is_passed_through_unchanged(self):
        ws = gp_sheet(KEYS_GP_11)
        contractor = _contractor(ws)
        assert resolve_contractor(ws, contractor, HEADER_ROW).geometry is contractor


class TestHeaderTiers:
    def test_labels_are_matched_after_normalization(self):
        """Регистр, двойные пробелы и неразрывный пробел не мешают паре."""
        ws = gp_sheet(KEYS_GP_11)
        ws.cell(row=HEADER_ROW, column=20, value="КОММЕНТАРИЙ  УЧАСТНИКА")
        result = resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert result.layout.column_keys == KEYS_GP_11

    def test_merged_group_header_does_not_bleed_right(self):
        """«Последняя виденная» шапка вправо не тянется (спека §2.3 п.1):
        колонка сразу за группой с пустыми обоими ярусами — отказ, а не
        продолжение группы."""
        ws = gp_sheet(KEYS_GP_11)
        # Сносим подпись колонки S (организатор): вертикальное объединение
        # остаётся, значение убирается — оба яруса пусты.
        # `ws.cell(..., value=None)` — не-операция в openpyxl (значение
        # присваивается только если value is not None); очистка — через
        # прямое присваивание атрибуту.
        ws.cell(row=HEADER_ROW, column=19).value = None
        with pytest.raises(EstimateParseError, match="S9"):
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)

    def test_column_inside_group_without_sublabel_is_a_refusal(self):
        """Пара (тип, «») — неопознана (спека §2.3 п.4)."""
        ws = gp_sheet(KEYS_GP_11)
        # `ws.cell(..., value=None)` не очищает существующее значение (см.
        # комментарий выше) — очистка через прямое присваивание атрибуту.
        ws.cell(row=HEADER_ROW + 1, column=12).value = None  # СМР группы цен
        with pytest.raises(EstimateParseError, match="L9"):
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)


class TestRefusals:
    def test_unknown_label_names_coordinate_and_both_raw_labels(self):
        ws = gp_sheet(KEYS_GP_11)
        ws.cell(row=HEADER_ROW + 1, column=12, value="Труд")  # вместо СМР
        with pytest.raises(EstimateParseError) as err:
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        message = str(err.value)
        assert "L9" in message
        assert "Труд" in message
        assert "Цена за ед. изм." in message      # верхняя подпись до нормализации
        assert 'ООО "Тест"' in message

    def test_merged_group_with_foreign_prefix_is_unknown(self):
        """Объединённая группа с чужим заголовком не получает тип — и колонка
        с легальной подписью «Материалы» внутри неё не опознаётся."""
        ws = gp_sheet(KEYS_GP_11)
        ws.cell(row=HEADER_ROW, column=11, value="Скидка, RUB")
        with pytest.raises(EstimateParseError, match="Скидка"):
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)

    def test_duplicate_key_names_both_coordinates(self):
        ws = gp_sheet(KEYS_GP_11)
        # Организатор S превращаем во второй комментарий: снять вертикальное
        # объединение нельзя, но подпись заменить можно — значение в S9.
        ws.cell(row=HEADER_ROW, column=19, value="Комментарий участника")
        with pytest.raises(EstimateParseError) as err:
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)
        assert "S9" in str(err.value) and "T9" in str(err.value)

    def test_missing_required_key_names_block_and_key_without_coordinate(self):
        """Ширина 7: восьмёрка денег без total_cost.total — группа стоимости из
        ТРЁХ колонок. Нарушено ровно одно ограничение — состав; все подписи
        легальны и уникальны. Лист рукописный: строитель групп из трёх колонок
        с легальной формой не описывает."""
        ws = Workbook().active
        ws["G6"] = "Наименование контрагента"
        ws.merge_cells("J6:P6")
        ws["J6"] = 'ООО "Тест"'
        ws.merge_cells("J9:M9")
        ws["J9"] = "Цена за ед. изм., RUB, ОСН, с учетом НДС 20%"
        ws.merge_cells("N9:P9")
        ws["N9"] = "Стоимость всего, RUB, ОСН, с учетом НДС 20%"
        for col, label in ((10, "Материалы"), (11, "СМР"), (12, "Косвенные расходы"), (13, "Всего")):
            ws.cell(row=10, column=col, value=label)
        for col, label in ((14, "Материалы"), (15, "СМР"), (16, "Косвенные расходы")):
            ws.cell(row=10, column=col, value=label)

        with pytest.raises(EstimateParseError) as err:
            resolve_contractor(ws, _contractor(ws), HEADER_ROW)

        message = str(err.value)
        assert "total_cost.total" in message
        assert 'ООО "Тест"' in message
