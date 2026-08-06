"""Тест производительности на листе с раздутой размерностью.

Хвост фазы 0 (`docs/phase0-input-data.md`, `docs/phase3-start.md` §6).

Суть дефекта: в реальной выгрузке есть одна пустая стилизованная ячейка в
колонке XFD, из-за которой `ws.max_column` равен 16384 при 24 колонках данных.
Всё, что разворачивается до `max_column` — явный `max_col=ws.max_column` и
неявный `ws[row]`, — начинает материализовать ~16 тыс. объектов ячеек на строку.

Обезличенный fixture этот дефект НЕ воспроизводит: при сохранении openpyxl
отбросил пустые стилизованные ячейки, и в обычном режиме открытия его
`max_column` равен 24. Поэтому тест строит синтетический лист в памяти — так
условие воспроизводится точно и без файла в репозитории.

Метод проверки — сравнение с контрольным листом без ячейки в XFD. Абсолютные
пороги плавают от машины к машине; отношение — нет. Если сканы ограничены
фактической шириной данных, лишняя колонка не стоит ничего и оба прогона идут
одинаково. Если кто-то вернёт полностраничный обход, раздутый лист станет
медленнее в сотни раз.
"""
from __future__ import annotations

import time

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from parser import parse_worksheet
from parser.read_contractors import read_contractors

# Последняя колонка Excel — та самая XFD.
LAST_EXCEL_COLUMN = 16384

# Строк позиций в синтетическом листе. Больше не нужно: разница между
# ограниченным и полностраничным сканом видна уже на сотнях строк.
POSITION_ROWS = 300

# Все позиции носят одно наименование: лемматизация кэшируется, и замер
# показывает стоимость обхода листа, а не работу spaCy.
JOB_TITLE = "Устройство монолитных стен"


def _build_estimate_sheet(*, inflate_to_xfd: bool) -> Worksheet:
    """Синтетическая смета ГП в ожидаемой раскладке.

    Args:
        inflate_to_xfd: создать ли пустую ячейку в XFD — она и раздувает
            `ws.max_column` до 16384.
    """
    ws = Workbook().active

    # Шапка (строки 2–4, как в смете ГП)
    ws["A2"] = "Предмет тендера: "
    ws["D2"] = "№001-ТУ Синтетика"
    ws["A3"] = "Объект: "
    ws["D3"] = "Синтетический объект"
    ws["A4"] = "Адрес объекта"
    ws["D4"] = "г. Тестоград"

    # Контрагент: маркер и блок на 11 колонок J..T
    ws["G6"] = "Наименование контрагента"
    ws.merge_cells("J6:T6")
    ws["J6"] = 'ООО "Синтетика"'
    ws["J7"] = "0000000000"
    ws["J8"] = "г. Тестоград, ул. Примерная"

    # Шапка таблицы: две строки, колонка A объединена по вертикали
    ws.merge_cells("A9:A10")
    ws["A9"] = "№ п/п"
    ws["B9"] = "№ раздела"
    ws["C9"] = "Статья СМР"
    ws["D9"] = "Наименование работ"
    ws["G9"] = "Ед. изм "
    ws["H9"] = "Общее кол-во"
    ws["J10"] = "Предлагаемое количество"

    # Строка-маркер лота
    ws["A11"] = "1"
    ws["B11"] = "1"
    ws["D11"] = "Лот №1 - Синтетика"

    # Позиции
    first_position_row = 12
    for offset in range(POSITION_ROWS):
        row = first_position_row + offset
        ws.cell(row=row, column=1, value=str(offset + 1))
        ws.cell(row=row, column=4, value=JOB_TITLE)
        ws.cell(row=row, column=7, value="м2")
        ws.cell(row=row, column=8, value=1)
        ws.cell(row=row, column=10, value=100)  # J — предлагаемое количество
        for col in range(11, 15):  # K..N — цена за единицу
            ws.cell(row=row, column=col, value=10)
        for col in range(15, 19):  # O..R — стоимость всего
            ws.cell(row=row, column=col, value=1000)
        ws.cell(row=row, column=19, value=10)  # S
        ws.cell(row=row, column=20, value="комментарий")  # T

    # Блок итогов: объединённая ячейка в колонке A — конец блока позиций
    total_row = first_position_row + POSITION_ROWS
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.cell(row=total_row, column=1, value="ИТОГО, руб. с учетом НДС")
    ws.cell(row=total_row, column=18, value=1000 * POSITION_ROWS)

    if inflate_to_xfd:
        # Пустая ячейка в крайней правой колонке — ровно то, что делает
        # стилизованная пустая ячейка в реальном файле.
        ws.cell(row=9, column=LAST_EXCEL_COLUMN)

    return ws


@pytest.fixture(scope="module")
def inflated_sheet() -> Worksheet:
    return _build_estimate_sheet(inflate_to_xfd=True)


@pytest.fixture(scope="module")
def control_sheet() -> Worksheet:
    return _build_estimate_sheet(inflate_to_xfd=False)


def _seconds(func) -> float:
    start = time.perf_counter()
    func()
    return time.perf_counter() - start


def test_synthetic_sheet_reproduces_the_defect(inflated_sheet, control_sheet):
    """Предусловие теста: раздутый лист действительно раздут.

    Без этой проверки замер ниже мог бы «проходить» на листе, где дефекта нет.
    """
    assert inflated_sheet.max_column == LAST_EXCEL_COLUMN
    assert control_sheet.max_column <= 24


def test_inflated_sheet_parses_correctly(inflated_sheet, control_sheet):
    """Раздутая размерность не меняет результат разбора."""
    inflated = parse_worksheet(inflated_sheet)
    control = parse_worksheet(control_sheet)

    assert inflated.warnings == []
    assert inflated.data == control.data

    positions = inflated.data["lots"]["lot_1"]["proposals"]["contractor_1"]["contractor_items"]["positions"]
    # Строка-маркер лота плюс сами позиции
    assert len(positions) == POSITION_ROWS + 1


def test_read_contractors_does_not_scan_full_width(inflated_sheet, control_sheet):
    """Поиск шапки контрагентов ограничен, а не идёт до `ws.max_column`.

    Это то самое место, которое фаза 0 замерила в 115 тыс. ячеек на реальном
    файле (`read_contractors.py:72` исходника).
    """
    control_time = _seconds(lambda: read_contractors(control_sheet))
    inflated_time = _seconds(lambda: read_contractors(inflated_sheet))

    assert inflated_time < control_time * 5 + 0.5, (
        f"поиск шапки на раздутом листе {inflated_time:.3f} c против "
        f"{control_time:.3f} c на контрольном — похоже, скан идёт по всей ширине"
    )


def test_full_parse_does_not_scan_full_width(inflated_sheet, control_sheet):
    """Разбор целиком не дорожает от лишних пустых колонок.

    Полностраничный обход строк — обращение к `ws[row]` вместо ограниченного
    `iter_rows(..., max_col=...)` — сделал бы раздутый лист медленнее в сотни
    раз (на реальном файле это давало минуты вместо секунд).
    """
    # Прогреваем кэш лемматизации: он не должен попасть в замер.
    parse_worksheet(control_sheet)

    control_time = _seconds(lambda: parse_worksheet(control_sheet))
    inflated_time = _seconds(lambda: parse_worksheet(inflated_sheet))

    assert inflated_time < control_time * 5 + 1.0, (
        f"разбор раздутого листа {inflated_time:.3f} c против {control_time:.3f} c "
        "на контрольном — где-то вернулся обход по всей ширине листа"
    )
