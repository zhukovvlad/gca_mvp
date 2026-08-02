"""Тесты поиска строки заголовков контрагентов.

Перенос `app/tests/excel_parser/test_read_contractors.py` из
`parser_tender_xlsx@0e178c0` без содержательных изменений: диапазон строк
поиска (4–10) при адаптации не менялся.
"""
from __future__ import annotations

import re

import pytest
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from parser.read_contractors import read_contractors


@pytest.fixture
def empty_worksheet():
    return Workbook().active


@pytest.fixture
def worksheet_with_contractors():
    """Лист с типичной строкой заголовков контрагентов (строка 5, 10 ячеек)."""
    ws = Workbook().active

    ws.cell(row=1, column=1, value="Тендерная документация")
    ws.cell(row=2, column=1, value="Объект строительства")
    ws.cell(row=3, column=1, value="Раздел работ")

    ws.cell(row=5, column=1, value="№ п/п")
    ws.cell(row=5, column=2, value="Наименование работ")
    ws.cell(row=5, column=3, value="Ед. изм.")
    ws.cell(row=5, column=4, value="Кол-во")
    ws.cell(row=5, column=5, value="Наименование контрагента №1")  # маркер
    ws.cell(row=5, column=6, value="Цена за единицу")
    ws.cell(row=5, column=7, value="Общая стоимость")
    ws.cell(row=5, column=8, value="Наименование контрагента №2")
    ws.cell(row=5, column=9, value="Цена за единицу")
    ws.cell(row=5, column=10, value="Общая стоимость")

    return ws


@pytest.fixture
def worksheet_with_merged_cells():
    """Лист с объединённой ячейкой в строке контрагентов (D6:F6)."""
    ws = Workbook().active

    ws.cell(row=6, column=1, value="№ п/п")
    ws.cell(row=6, column=2, value="Наименование работ")
    ws.cell(row=6, column=3, value="Наименование контрагента ООО 'Строй'")  # маркер

    ws.merge_cells("D6:F6")
    ws.cell(row=6, column=4, value="Подрядчик Альфа")

    ws.cell(row=6, column=7, value="Комментарии")

    return ws


def _clear(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


class TestReadContractorsBehavior:
    """Базовое поведение."""

    def test_returns_none_for_empty_worksheet(self, empty_worksheet):
        assert read_contractors(empty_worksheet) is None

    def test_returns_none_when_no_contractor_title_found(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=4, column=1, value="Обычные данные")
        ws.cell(row=5, column=1, value="Еще данные")
        ws.cell(row=6, column=1, value="Название работ")
        ws.cell(row=7, column=1, value="Прочая информация")

        assert read_contractors(ws) is None

    def test_returns_list_when_contractor_title_found(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        assert isinstance(result, list)
        assert len(result) > 0

    @pytest.mark.parametrize(
        "case_variant",
        [
            "НАИМЕНОВАНИЕ КОНТРАГЕНТА",
            "Наименование Контрагента",
            "наименование контрагента",
            "НаИмЕнОвАнИе КоНтРаГеНтА",
        ],
    )
    def test_finds_contractor_title_case_insensitive(self, empty_worksheet, case_variant):
        ws = empty_worksheet
        ws.cell(row=5, column=3, value=case_variant)
        ws.cell(row=5, column=4, value="Подрядчик 1")

        result = read_contractors(ws)

        assert result is not None, case_variant
        assert len(result) == 2

    def test_ignores_leading_trailing_whitespace_in_search(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=6, column=2, value="  наименование контрагента  ")
        ws.cell(row=6, column=3, value="ООО Строитель")

        result = read_contractors(ws)

        assert result is not None
        assert len(result) == 2


class TestReadContractorsSearchRange:
    """Диапазон поиска — строки 4–10."""

    @pytest.mark.parametrize("row_num", [1, 2, 3, 11, 15])
    def test_searches_only_in_rows_4_to_10(self, empty_worksheet, row_num):
        ws = empty_worksheet
        ws.cell(row=row_num, column=1, value="наименование контрагента")
        ws.cell(row=row_num, column=2, value="Подрядчик")

        assert read_contractors(ws) is None, f"строка {row_num} вне диапазона"

    @pytest.mark.parametrize("row_num", [4, 5, 6, 7, 8, 9, 10])
    def test_finds_in_each_valid_row(self, empty_worksheet, row_num):
        ws = empty_worksheet
        ws.cell(row=row_num, column=1, value="наименование контрагента")
        ws.cell(row=row_num, column=2, value="Подрядчик")

        result = read_contractors(ws)

        assert result is not None, f"строка {row_num}"
        assert len(result) == 2

    def test_returns_first_matching_row_only(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=5, column=1, value="наименование контрагента")
        ws.cell(row=5, column=2, value="Первый подрядчик")

        ws.cell(row=7, column=1, value="наименование контрагента")
        ws.cell(row=7, column=2, value="Второй подрядчик")
        ws.cell(row=7, column=3, value="Третий подрядчик")

        result = read_contractors(ws)

        assert len(result) == 2
        values = [c["value"] for c in result]
        assert "Первый подрядчик" in values
        assert "Второй подрядчик" not in values


class TestReadContractorsDataExtraction:
    """Извлечение данных из найденной строки."""

    def test_extracts_all_non_empty_cells_from_contractor_row(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        assert len(result) == 10
        assert all(c["value"] is not None for c in result)

    def test_skips_empty_cells_in_contractor_row(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=5, column=1, value="наименование контрагента")
        ws.cell(row=5, column=3, value="Подрядчик 1")
        ws.cell(row=5, column=6, value="Подрядчик 2")

        result = read_contractors(ws)

        assert len(result) == 3
        assert [c["value"] for c in result] == [
            "наименование контрагента",
            "Подрядчик 1",
            "Подрядчик 2",
        ]

    def test_cell_info_structure(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        for cell_info in result:
            for field in ("value", "coordinate", "column_start", "row_start"):
                assert field in cell_info, field

            assert isinstance(cell_info["coordinate"], str)
            assert isinstance(cell_info["column_start"], int)
            assert isinstance(cell_info["row_start"], int)
            assert cell_info["column_start"] >= 1
            assert cell_info["row_start"] >= 4

    def test_row_start_consistency(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        assert {c["row_start"] for c in result} == {5}

    def test_column_start_matches_coordinate(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        for cell_info in result:
            column_letter = "".join(filter(str.isalpha, cell_info["coordinate"]))
            assert cell_info["column_start"] == column_index_from_string(column_letter)


class TestReadContractorsMergedCells:
    """Обработка объединённых ячеек."""

    def test_adds_merged_shape_info_for_merged_cells(self, worksheet_with_merged_cells):
        result = read_contractors(worksheet_with_merged_cells)

        merged_cell_info = next((c for c in result if c["coordinate"] == "D6"), None)

        assert merged_cell_info is not None, "объединённая ячейка D6 должна быть найдена"
        assert merged_cell_info["merged_shape"] == {"rowspan": 1, "colspan": 3}

    def test_no_merged_shape_for_regular_cells(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        assert all("merged_shape" not in c for c in result)


class TestReadContractorsEdgeCases:
    """Граничные случаи."""

    def test_handles_partial_match_in_cell_value(self, empty_worksheet):
        """Маркер найдётся, если он в начале более длинной строки."""
        ws = empty_worksheet
        ws.cell(row=5, column=1, value="наименование контрагента и участников тендера")
        ws.cell(row=5, column=2, value="Подрядчик")

        assert len(read_contractors(ws)) == 2

    def test_handles_non_string_cell_values(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=5, column=1, value=123)
        ws.cell(row=5, column=2, value=None)
        ws.cell(row=5, column=3, value="наименование контрагента")
        ws.cell(row=5, column=4, value=45.67)
        ws.cell(row=5, column=5, value="Подрядчик")

        result = read_contractors(ws)

        assert len(result) == 4  # пустая ячейка пропущена
        assert [c["value"] for c in result] == [123, "наименование контрагента", 45.67, "Подрядчик"]

    def test_handles_worksheet_with_no_data_in_search_range(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=1, column=1, value="Заголовок")
        ws.cell(row=2, column=1, value="наименование контрагента")  # вне диапазона
        ws.cell(row=15, column=1, value="Конец документа")

        assert read_contractors(ws) is None

    def test_marker_text_must_be_at_start_of_cell_value(self, empty_worksheet):
        ws = empty_worksheet
        ws.cell(row=5, column=1, value="Это не наименование контрагента в начале")
        ws.cell(row=5, column=2, value="Подрядчик")

        assert read_contractors(ws) is None

        ws.cell(row=5, column=1, value="наименование контрагента - основная информация")

        assert read_contractors(ws) is not None


class TestReadContractorsDataIntegrity:
    """Качество возвращаемых данных."""

    def test_maintains_cell_order_in_result(self, empty_worksheet):
        ws = empty_worksheet
        for col, value in [
            (1, "наименование контрагента"),
            (3, "Подрядчик A"),
            (5, "Подрядчик B"),
            (7, "Подрядчик C"),
        ]:
            ws.cell(row=6, column=col, value=value)

        result = read_contractors(ws)

        assert [c["column_start"] for c in result] == [1, 3, 5, 7]
        assert [c["value"] for c in result] == [
            "наименование контрагента",
            "Подрядчик A",
            "Подрядчик B",
            "Подрядчик C",
        ]

    def test_coordinate_format_consistency(self, worksheet_with_contractors):
        result = read_contractors(worksheet_with_contractors)

        pattern = re.compile(r"^[A-Z]+\d+$")
        assert all(pattern.match(c["coordinate"]) for c in result)

    def test_returns_independent_data_structures(self, worksheet_with_contractors):
        result1 = read_contractors(worksheet_with_contractors)
        result2 = read_contractors(worksheet_with_contractors)

        assert result1 == result2
        assert result1 is not result2

        result1[0]["test_field"] = "modified"
        assert "test_field" not in result2[0]
