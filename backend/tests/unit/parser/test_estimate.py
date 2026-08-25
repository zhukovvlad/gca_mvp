"""Тесты оркестратора `parse_estimate` на реальной раскладке сметы ГП.

Два источника данных:

* `fixtures/gp_estimate_fixture.xlsx` — обезличенная копия реального образца
  (фаза 0). Коммитится, работает в CI, на нём держатся основные проверки.
* `samples/` — реальные оферты (может лежать несколько). В репозиторий не
  попадают (AGENTS.md §9), тесты на них пропускаются, если каталога нет.
  Проверяются инварианты формата и семантика §6 на КАЖДОЙ оферте, плюс
  присутствие первоисточника fixture — что обезличивание не изменило
  поведения парсера.

Числа в ожиданиях — не «что получилось», а замеры фазы 0
(`docs/phase0-input-data.md`): 2576 строк позиций, 1828 расценённых строк с
заполненным «Предлагаемым количеством». Они принадлежат первоисточнику
fixture, а не любому образцу.
"""
from __future__ import annotations

import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pytest

from parser import PARSER_VERSION, EstimateParseError, parse_estimate, parse_worksheet
from parser.constants import TABLE_PARSE_POSITION_COLUMN_HEADERS
from parser.postprocess import BASELINE_MISSING_TITLE

from .sheet_builders import (
    COLUMNS_BY_WIDTH,
    KEYS_8,
    KEYS_10,
    KEYS_12,
    KEYS_GP_11,
    KEYS_PERMUTED_11,
    KEYS_TENDER_11,
    add_contractor_block,
    gp_sheet,
)

REPO_ROOT = Path(__file__).resolve().parents[4]
FIXTURE_PATH = REPO_ROOT / "fixtures" / "gp_estimate_fixture.xlsx"
SAMPLES_DIR = REPO_ROOT / "samples"

# Замеры фазы 0 на реальном образце (docs/phase0-input-data.md).
EXPECTED_POSITIONS = 2576
EXPECTED_PRICED_ROWS = 1828

# Стоимости в файле округлены до копеек, поэтому произведение
# «цена за единицу × количество» сходится с точностью до копейки, а не побитово.
# На крупных суммах округление самих сомножителей даёт больше копейки, поэтому
# рядом стоит относительный допуск. Оба заведомо меньше, чем расхождение от
# перепутанных колонок: оно было бы на порядки.
KOPECK = Decimal("0.01")
RELATIVE_TOLERANCE = Decimal("1e-6")

# Денежные поля позиции — те, что AGENTS.md §3 требует держать строками.
MONEY_PATHS = (
    ("unit_cost", "materials"),
    ("unit_cost", "works"),
    ("unit_cost", "indirect_costs"),
    ("unit_cost", "total"),
    ("total_cost", "materials"),
    ("total_cost", "works"),
    ("total_cost", "indirect_costs"),
    ("total_cost", "total"),
    ("total_cost_for_organizer_quantity",),
)


# Все локальные оферты. Берутся только «Оферта_*»: в samples/ могут лежать и
# файлы других форматов, которые парсер текущей фазы намеренно отвергает.
SAMPLE_PATHS = sorted(SAMPLES_DIR.glob("Оферта*.xlsx")) if SAMPLES_DIR.is_dir() else []


@pytest.fixture(scope="module")
def fixture_result():
    """Разбор обезличенного fixture. Модульная область — разбор небыстрый."""
    if not FIXTURE_PATH.is_file():
        pytest.fail(f"Обезличенный fixture обязан лежать в репозитории: {FIXTURE_PATH}")
    return parse_estimate(str(FIXTURE_PATH))


@pytest.fixture(scope="module")
def sample_results():
    """Разбор всех локальных оферт; пропуск, если их нет.

    В сообщениях тестов образцы именуются по номеру, не по имени файла:
    в именах файлов фигурируют контрагенты, а вывод тестов может попасть
    в логи (AGENTS.md §9).
    """
    if not SAMPLE_PATHS:
        pytest.skip("Каталог samples/ пуст или отсутствует — реальные оферты не коммитятся (AGENTS.md §9)")
    return [(f"образец №{i}", parse_estimate(str(path))) for i, path in enumerate(SAMPLE_PATHS, start=1)]


class TestParseEstimateOnFixture:
    """Разбор обезличенного fixture."""

    def test_parses_without_warnings(self, fixture_result):
        """Раскладка совпадает с ожидаемой — предупреждений нет."""
        assert fixture_result.warnings == []

    def test_reports_parser_version(self, fixture_result):
        assert fixture_result.parser_version == PARSER_VERSION

    def test_extracts_header(self, fixture_result):
        """Шапка сметы ГП лежит в строке 2 — исходный диапазон 3–5 её терял."""
        data = fixture_result.data

        assert data["tender_id"] == "001-ТУ"
        assert data["tender_title"] == '"Тестовый ЖК_Генподряд"'
        assert data["tender_object"] == "Тестовый ЖК Корпус 1"
        assert data["tender_address"] == "г. Тестоград, ул. Примерная, вл 1"

    def test_single_lot_and_single_proposal(self, fixture_result):
        """Смета ГП: один лот, один подрядчик (AGENTS.md §4)."""
        lots = fixture_result.data["lots"]

        assert list(lots) == ["lot_1"]
        assert lots["lot_1"]["lot_title"].startswith("Лот №1")
        assert list(lots["lot_1"]["proposals"]) == ["contractor_1"]

    def test_contractor_block_is_eleven_columns(self, fixture_result):
        """J..T — раскладка, на которой держится смысл колонок стоимости."""
        proposal = fixture_result.data["lots"]["lot_1"]["proposals"]["contractor_1"]

        assert proposal["contractor_width"] == 11
        assert proposal["contractor_coordinate"] == "J6"
        assert proposal["title"] == 'ООО "ТЕСТПОДРЯД"'
        assert proposal["inn"] == "0000000000"

    def test_baseline_is_absent(self, fixture_result):
        """Колонки «Расчётная стоимость» в смете ГП нет — это норма, не ошибка."""
        lot = fixture_result.data["lots"]["lot_1"]

        assert lot["baseline_proposal"]["title"] == BASELINE_MISSING_TITLE

    def test_no_deviation_fields_on_positions(self, fixture_result):
        """Без baseline поля отклонения вычищены (AGENTS.md §4)."""
        positions = _positions(fixture_result)

        assert all("deviation_from_baseline_cost" not in p for p in positions.values())

    def test_position_count_matches_phase0_measurement(self, fixture_result):
        """Все строки блока позиций на месте — включая первые две.

        Исходная константа START_INDEXING_POSITION_ROW = 13 отрезала строки
        11 и 12, давая 2574 вместо 2576.
        """
        assert len(_positions(fixture_result)) == EXPECTED_POSITIONS

    def test_first_rows_are_not_dropped(self, fixture_result):
        """Первая позиция — строка-раздел лота, вторая — первый раздел сметы."""
        positions = _positions(fixture_result)

        assert positions["1"]["job_title"].startswith("Лот №1")
        assert positions["1"]["is_chapter"] is True
        assert positions["2"]["job_title"].startswith("Подготовительные работы")

    def test_priced_rows_match_phase0_measurement(self, fixture_result):
        """1828 строк с «Предлагаемым количеством» — цифра замера фазы 0.

        На ней держится вывод §6: вес позиции берётся из suggested_quantity.
        """
        positions = _positions(fixture_result)
        with_suggested = [p for p in positions.values() if p.get("suggested_quantity") is not None]

        assert len(with_suggested) == EXPECTED_PRICED_ROWS

    def test_suggested_quantity_drives_total_cost(self, fixture_result):
        """`Стоимость всего = цена за единицу × Предлагаемое количество`.

        Ключевая проверка семантики §6: именно это соотношение фаза 0 нашла на
        1828 из 1828 строк. Если раскладка колонок поедет, оно развалится.
        """
        positions = _positions(fixture_result)

        checked = 0
        for position in positions.values():
            unit_total = _money(position["unit_cost"]["total"])
            row_total = _money(position["total_cost"]["total"])
            weight = _amount(position.get("suggested_quantity"))

            if unit_total is None or row_total is None or not weight:
                continue

            _assert_close(row_total, unit_total * weight, position["job_title"])
            checked += 1

        assert checked > 1500, f"проверено слишком мало строк: {checked}"

    def test_organizer_quantity_total_uses_quantity(self, fixture_result):
        """Справочная колонка S считается по «Общему кол-ву», а не по J.

        Обратная сторона той же проверки: колонки не перепутаны местами.
        """
        positions = _positions(fixture_result)

        checked = 0
        for position in positions.values():
            unit_total = _money(position["unit_cost"]["total"])
            organizer_total = _money(position.get("total_cost_for_organizer_quantity"))
            quantity = _amount(position.get("quantity"))

            if unit_total is None or organizer_total is None or not quantity:
                continue

            _assert_close(organizer_total, unit_total * quantity, position["job_title"])
            checked += 1

        assert checked > 1500, f"проверено слишком мало строк: {checked}"

    def test_summary_block_gives_one_key_per_row(self, fixture_result):
        """Три строки блока в файле → три ключа. Прежде их было два: метки
        «с учетом НДС» и «без учета НДС» обе содержали «итого» и «ндс»."""
        summary = _proposal(fixture_result)["contractor_items"]["summary"]
        assert sorted(summary) == ["total_cost_excluding_vat", "total_cost_including_vat", "vat_amount"]
        assert summary["total_cost_including_vat"]["job_title"] == "ИТОГО, руб. с учетом НДС"
        assert summary["total_cost_excluding_vat"]["job_title"] == "ИТОГО, руб. без учета НДС"

    def test_additional_info_parsed(self, fixture_result):
        """Блок дополнительной информации — шесть пунктов."""
        additional = _proposal(fixture_result)["additional_info"]

        assert len(additional) == 6

    def test_positions_stop_before_additional_info_block(self, fixture_result):
        """Досрочный выход по merged-ячейке отработал (AGENTS.md §11).

        В блоке дополнительной информации колонка подрядчика содержит текст
        («Представлено», «К обсуждению»). Если бы позиции дотянулись до него,
        этот текст попал бы в suggested_quantity.
        """
        positions = _positions(fixture_result)

        for position in positions.values():
            weight = position.get("suggested_quantity")
            assert not isinstance(weight, str), f"текст в количестве: {weight!r}"

        titles = {str(p.get("job_title")) for p in positions.values()}
        assert not any("Дополнительная информация" in t for t in titles)

    def test_executor_block_is_empty(self, fixture_result):
        """В смете ГП блока исполнителя нет — data_prepared_on_date будет NULL."""
        assert fixture_result.data["executor"] == {
            "executor_name": None,
            "executor_phone": None,
            "executor_date": None,
        }

    def test_every_position_with_title_has_normalized_title(self, fixture_result):
        """Нормализованное наименование — вход матчинга фазы 4."""
        positions = _positions(fixture_result)

        for position in positions.values():
            if position.get("job_title"):
                assert position["job_title_normalized"], position["job_title"]


class TestParseEstimateOnRealSamples:
    """Инварианты формата на каждой локальной оферте + сверка первоисточника.

    Точные числа (2576/1828) принадлежат первоисточнику fixture; остальные
    оферты проверяются на инварианты формата и семантику §6 — так каждый новый
    образец в samples/ автоматически становится проверкой парсера.
    """

    def test_every_sample_has_only_expected_warnings(self, sample_results):
        """На каждой оферте либо предупреждений нет вовсе, либо ровно два —
        про пустую строку «В том числе НДС» (Ф4a) и про ненайденную ставку
        НДС в шапке ценового блока (Ф4б).

        Один реальный образец (449-ТУ) несёт блок итогов, где строка «В том
        числе НДС» физически присутствует, но все четыре её денежные ячейки
        пусты (Ф4a, спека §2.7), И групповые шапки его ценового блока не
        несут суффикса ставки НДС (Ф4б, спека §2.3) — те же два факта файла,
        что и раньше, только теперь их источников два, а не один. 449-ТУ
        прикрыт веткой некаскадирования дважды (спека Ф4б §2.7): у него нет
        ни заявленной ставки, ни ключа `total_cost_excluding_vat` в блоке
        итогов, поэтому сверка с блоком итогов не запускается вовсе и
        «сверка не проведена» здесь не добавляется третьим предупреждением —
        доказательством некаскадирования сам этот замер не служит, для этого
        есть синтетическая проверка (Task 8, спека §4.5 п.11).

        Утверждение — не «пусто ИЛИ любое число предупреждений»: другое их
        число на любой оферте, как и наличие предупреждения без ОБОИХ
        смысловых фрагментов ниже, обязаны покрасить тест. Проверяются именно
        фрагменты, а не текст целиком — он несёт номер строки и фактические
        метки шапки, которые меняются при обновлении файла оферты и являются
        диагностической деталью, а не частью проверяемого смысла.

        Предупреждение о блоке итогов не означает «НДС отсутствует», «НДС не
        заявлен» или «смета без НДС»: парсер установил только то, что строка
        СУЩЕСТВУЕТ и что её денежные ячейки пусты. Предупреждение о ставке не
        означает «в файле нет НДС»: оно означает только то, что шапка ценового
        блока не заявила ставку суффиксом известного вида.
        """
        for name, result in sample_results:
            if not result.warnings:
                continue
            assert len(result.warnings) == 2, (name, result.warnings)

            summary_warnings = [w for w in result.warnings if "В том числе НДС" in w]
            assert len(summary_warnings) == 1, (name, result.warnings)
            assert "суммы не указаны" in summary_warnings[0], name

            vat_rate_warnings = [w for w in result.warnings if "Ставка НДС не получена из шапки" in w]
            assert len(vat_rate_warnings) == 1, (name, result.warnings)

    def test_every_sample_has_gp_layout(self, sample_results):
        """J6, 11 колонок — раскладка, на которой держится смысл стоимостей."""
        for name, result in sample_results:
            proposal = _proposal(result)

            assert proposal["contractor_width"] == 11, name
            assert proposal["contractor_coordinate"] == "J6", name

    def test_every_sample_extracts_header_identifier(self, sample_results):
        """На реальных файлах исходный парсер оставлял оба поля пустыми."""
        for name, result in sample_results:
            assert result.data["tender_id"], name
            assert result.data["tender_title"], name

    def test_every_sample_has_no_baseline(self, sample_results):
        for name, result in sample_results:
            lot = result.data["lots"]["lot_1"]

            assert lot["baseline_proposal"]["title"] == BASELINE_MISSING_TITLE, name

    def test_weight_is_suggested_quantity_on_every_sample(self, sample_results):
        """`Стоимость всего = цена × Предлагаемое количество` (AGENTS.md §6).

        Фаза 0 доказала это на одном образце (1828/1828); каждая новая оферта
        в samples/ перепроверяет вывод. На втором образце (2026-08-03,
        экспорт из системы-источника): 1193/1193.
        """
        for name, result in sample_results:
            checked = 0
            for position in _positions(result).values():
                unit_total = _money(position["unit_cost"]["total"])
                row_total = _money(position["total_cost"]["total"])
                weight = _amount(position.get("suggested_quantity"))

                if unit_total is None or row_total is None or not weight:
                    continue

                _assert_close(row_total, unit_total * weight, f"{name}: {position['job_title']}")
                checked += 1

            assert checked > 500, f"{name}: проверено слишком мало строк: {checked}"

    def test_coalesce_weight_fallback_is_safe_on_every_sample(self, sample_results):
        """Нет строк, где `quantity` осмыслен (≠1) и отличается от suggested.

        Это условие, при котором `w = COALESCE(suggested_quantity, quantity)`
        из §6 не может подменить вес чужим числом (фаза 0, п. «Чем доказано»).
        """
        for name, result in sample_results:
            for position in _positions(result).values():
                quantity = _amount(position.get("quantity"))
                suggested = _amount(position.get("suggested_quantity"))

                if quantity is None or suggested is None or quantity == 1:
                    continue

                assert quantity == suggested, (
                    f"{name}: «{position['job_title']}»: quantity={quantity}, suggested={suggested}"
                )

    def test_fixture_source_is_among_samples(self, sample_results, fixture_result):
        """Первоисточник fixture лежит в samples/ и совпадает с ним по числам.

        Если ни один образец не дал числа fixture — либо первоисточник убрали
        из каталога, либо парсер разошёлся с fixture на реальном файле. И то
        и другое должно быть громким.
        """
        fixture_positions = _positions(fixture_result)

        def priced(positions):
            return sum(1 for p in positions.values() if p.get("suggested_quantity") is not None)

        matches = [
            name
            for name, result in sample_results
            if len(_positions(result)) == len(fixture_positions) == EXPECTED_POSITIONS
            and priced(_positions(result)) == priced(fixture_positions) == EXPECTED_PRICED_ROWS
        ]

        assert matches, (
            f"ни один из {len(sample_results)} образцов не совпал с fixture "
            f"({EXPECTED_POSITIONS} позиций / {EXPECTED_PRICED_ROWS} расценённых)"
        )


class TestMoneyContract:
    """Деньги в `raw_data` — десятичные строки, ни одного `float` (AGENTS.md §3, §11)."""

    def test_position_money_fields_are_decimal_strings(self, fixture_result):
        """Каждое денежное поле позиции — строка либо None (пустая стоимость → NULL)."""
        for position in _positions(fixture_result).values():
            for path in MONEY_PATHS:
                value = _dig(position, path)
                assert value is None or isinstance(value, str), (
                    f"{'.'.join(path)} = {value!r} ({type(value).__name__}) у «{position.get('job_title')}»"
                )

    def test_summary_money_fields_are_decimal_strings(self, fixture_result):
        """Итоги идут через тот же `parse_contractor_row` — контракт тот же."""
        summary = _proposal(fixture_result)["contractor_items"]["summary"]

        assert summary, "блок итогов пуст — проверять нечего"
        for key, line in summary.items():
            for path in MONEY_PATHS:
                value = _dig(line, path)
                assert value is None or isinstance(value, str), f"{key}.{'.'.join(path)} = {value!r}"

    def test_money_strings_parse_back_to_decimal(self, fixture_result):
        """Фаза 4 делает `Decimal(value)` — значит, строка обязана им приниматься."""
        checked = 0
        for position in _positions(fixture_result).values():
            for path in MONEY_PATHS:
                value = _dig(position, path)
                if value is None:
                    continue
                Decimal(value)  # InvalidOperation здесь и есть провал теста
                checked += 1

        assert checked > 1500, f"проверено слишком мало значений: {checked}"

    def test_data_survives_json_round_trip_without_encoder(self, fixture_result):
        """`ParseResult.data` кладётся в jsonb как есть — без своего энкодера.

        И обратно: после round-trip'а деньги остаются теми же строками, то есть
        сериализация ничего не переводит в число.
        """
        restored = json.loads(json.dumps(fixture_result.data, ensure_ascii=False))

        assert _find_money_floats(restored) == []
        assert restored["lots"] == fixture_result.data["lots"]

    def test_totals_are_not_floats_anywhere_in_data(self, fixture_result):
        """Сквозная проверка: под денежными ключами `float` не встречается нигде."""
        floats = _find_money_floats(fixture_result.data)

        assert floats == [], f"float в денежных полях: {floats[:5]}"

    def test_temporal_and_error_cells_become_json_safe(self):
        """Дата в ячейке и ошибка Excel не ломают контракт «data → jsonb как есть».

        openpyxl отдаёт date-форматированную ячейку объектом `datetime` —
        на fixture таких нет, поэтому проверка на синтетическом листе:
        дата → ISO-строка, `#N/A` → None, соседняя сумма не задета.
        """
        import datetime as dt

        ws = _minimal_sheet(contractor_colspan=11)
        # A и B заполнены не для красоты: строка без номера и без раздела —
        # кандидат в агрегатную строку допработ, и лист был бы отвергнут
        # (спека Ф2 §2.2). Реальные файлы несут здесь номер и раздел.
        ws.cell(row=12, column=1, value=1)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Работа с датой в комментарии")
        ws.cell(row=12, column=11, value="#N/A")  # unit_cost.materials
        ws.cell(row=12, column=14, value=60.5)  # unit_cost.total
        ws.cell(row=12, column=20, value=dt.datetime(2025, 2, 1))  # comment_contractor

        result = parse_worksheet(ws)

        json.dumps(result.data, ensure_ascii=False)  # TypeError здесь — провал теста

        position = _positions(result)["2"]
        assert position["comment_contractor"] == "2025-02-01T00:00:00"
        assert position["unit_cost"]["materials"] is None
        assert position["unit_cost"]["total"] == "60.5"


class TestAdditionalWorksInJson:
    """Поле доезжает до итоговой структуры рядом с positions и summary."""

    def test_additional_works_lands_next_to_positions(self):
        ws = _minimal_sheet(11)
        ws.cell(row=12, column=1, value=1)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Обычная работа")
        ws.cell(row=13, column=4, value="Дополнительные работы")

        items = _proposal(parse_worksheet(ws))["contractor_items"]

        assert set(items) >= {"positions", "summary", "additional_works"}
        assert items["additional_works"]["job_title"] == "Дополнительные работы"
        assert items["additional_works"]["source_row"] == 13

    def test_additional_works_is_none_when_row_absent(self):
        """42-ТУ и 449-ТУ: ключ есть, значение None — это валидное состояние."""
        ws = _minimal_sheet(11)
        ws.cell(row=12, column=1, value=1)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Обычная работа")

        items = _proposal(parse_worksheet(ws))["contractor_items"]

        assert items["additional_works"] is None


def test_parser_version_is_major_because_a_key_disappeared_for_width_ten_files():
    """4.0.0: мажор из-за того, что у файлов ширины 10 из позиций ИСЧЕЗ ключ
    `total_cost_for_organizer_quantity` (файл такой колонки не нёс — прежний
    разбор молча читал в него комментарий участника). Взамен появился
    `comment_contractor`, а на позициях впервые стал возможен
    `deviation_from_baseline_cost`. Для смет ГП ширины 11 JSON побайтно
    прежний (спека §7, регрессия по классам) — мажор пришёл не за них.
    """
    assert PARSER_VERSION == "4.0.0"


class TestWidthTenComment:
    """Класс 2 спеки §7: комментарий ширины 10 попадает в comment_contractor,
    ключа total_cost_for_organizer_quantity в JSON нет вовсе (§2.8)."""

    def test_comment_of_a_ten_wide_block_lands_in_comment_contractor(self):
        ws = gp_sheet(KEYS_10)
        ws.cell(row=12, column=1, value=2)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Работа")
        ws.cell(row=12, column=19, value="таймлайн уточним")  # 10-я колонка блока

        result = parse_worksheet(ws)

        position = _positions(result)["2"]
        assert position["comment_contractor"] == "таймлайн уточним"
        assert "total_cost_for_organizer_quantity" not in position


class TestOldWidthTenShapeWasNoisy:
    def test_the_old_shape_did_produce_the_false_warning(self):
        """Старая форма позиции (комментарий под денежным ключом) на том же
        пути значений даёт ровно то предупреждение, чьё исчезновение утверждает
        интеграционный тест. Без контроля пустой список мог бы означать
        «смотреть было нечем» (docs/insights/unobservable-in-the-runner.md).
        """
        from services.estimate_import import _money

        value_problems: list[str] = []
        result = _money("таймлайн уточним", value_problems, "позиция «2»")

        assert result is None
        assert value_problems == [
            "позиция «2»: значение «таймлайн уточним» не число, записано NULL"
        ]


class TestPermutedColumnsEndToEnd:
    """Спека §6: синтетическая фикстура перестановки, собранная кодом. Два
    яруса, горизонтальные и вертикальные объединения, прежний colspan 11 и
    РАЗЛИЧИМЫЙ маркер в каждой физической колонке — иначе тест не отличит
    правильный разбор от совпадения."""

    def test_every_marker_lands_under_its_own_key(self):
        ws = gp_sheet(KEYS_PERMUTED_11)
        ws.cell(row=12, column=1, value=2)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Работа")
        # Маркеры: физическая колонка 10+i несёт значение 100+i (деньги и
        # количество) либо строку-маркер (комментарий).
        markers = ["маркер-комментарий", 101, 102, 103, 104, 105, 106, 107, 108, 109, 110]
        for offset, marker in enumerate(markers):
            ws.cell(row=12, column=10 + offset, value=marker)

        result = parse_worksheet(ws)
        position = _positions(result)["2"]

        assert position["comment_contractor"] == "маркер-комментарий"
        assert position["total_cost"] == {
            "materials": "101", "works": "102", "indirect_costs": "103", "total": "104",
        }
        assert position["total_cost_for_organizer_quantity"] == "105"
        assert position["unit_cost"] == {
            "materials": "106", "works": "107", "indirect_costs": "108", "total": "109",
        }
        assert position["suggested_quantity"] == 110

    def test_vat_anchors_follow_the_permuted_layout(self):
        """Ставка найдена в якорях ПЕРЕСТАВЛЕННЫХ групп — доказательство, что
        смещения вычисляются раскладкой, а не формулой от ширины (замена
        прежнего test_vat_rate_not_found_when_suffix_uses_wrong_offset_formula)."""
        ws = gp_sheet(KEYS_PERMUTED_11)
        result = parse_worksheet(ws)
        assert _proposal(result)["vat_rate"] == "20"


class TestMultiContractorSheet:
    """Лист с более чем одним блоком подрядчика — не отбрасывается и не
    обрезается до первого: разбираются оба блока, каждый попадает в
    `proposals` под своим ключом, а превышение ожидаемого числа подрядчиков
    (AGENTS.md §4) фиксируется предупреждением, а не исключением."""

    def test_two_blocks_parse_with_a_count_warning(self):
        ws = gp_sheet(KEYS_GP_11)                                  # J..T
        add_contractor_block(ws, col_start=22, columns=KEYS_12, title='ООО "Тест-2"')
        result = parse_worksheet(ws)
        proposals = result.data["lots"]["lot_1"]["proposals"]
        assert list(proposals) == ["contractor_1", "contractor_2"]
        assert any("ожидается один подрядчик, найдено 2" in w for w in result.warnings)

    @staticmethod
    def _tender_sheet_with_baseline(*, baseline_total: float | None):
        """Тендерный лист: подрядчик KEYS_TENDER_11 (J..T) + базовый блок
        KEYS_8 (V..AC), позиция с «% от р/с» и блок итогов.

        baseline_total кладётся в колонку total_cost.total БАЗОВОГО блока
        (22+7=29): ненулевое значение делает базу валидной для
        postprocess._is_baseline_valid; None — база пуста.
        """
        ws = gp_sheet(KEYS_TENDER_11)
        add_contractor_block(ws, col_start=22, columns=KEYS_8,
                             title="Расчетная стоимость", vat_suffix=None)
        ws.cell(row=12, column=1, value=2)
        ws.cell(row=12, column=2, value="1")
        ws.cell(row=12, column=4, value="Работа")
        ws.cell(row=12, column=20, value=-0.05)   # % от р/с — 11-я колонка блока
        # Объединённая ячейка в колонке A — конец блока позиций (AGENTS.md §11).
        ws.merge_cells(start_row=14, start_column=1, end_row=14, end_column=5)
        ws.cell(row=14, column=1, value="ИТОГО, руб. с учетом НДС")
        if baseline_total is not None:
            ws.cell(row=14, column=29, value=baseline_total)
        return ws

    def test_deviation_reaches_json_when_the_baseline_is_valid(self):
        """% от р/с доезжает до позиции (§2.8) — при живой расчётной стоимости
        postprocess его не трогает."""
        ws = self._tender_sheet_with_baseline(baseline_total=100.0)

        result = parse_worksheet(ws)

        lot = result.data["lots"]["lot_1"]
        assert lot["baseline_proposal"]["title"] == "Расчетная стоимость"
        position = lot["proposals"]["contractor_1"]["contractor_items"]["positions"]["2"]
        assert position["deviation_from_baseline_cost"] == "-0.05"

    def test_empty_baseline_cleans_deviations_with_the_base(self):
        """DoD 8: postprocess.py не менялся, и его правило «нет валидной базы —
        нет осмысленного отклонения» покрыто на файле с ПУСТОЙ расчётной
        стоимостью: тот же лист, но без итога базового блока."""
        ws = self._tender_sheet_with_baseline(baseline_total=None)

        result = parse_worksheet(ws)

        lot = result.data["lots"]["lot_1"]
        assert lot["baseline_proposal"]["title"] == BASELINE_MISSING_TITLE
        position = lot["proposals"]["contractor_1"]["contractor_items"]["positions"]["2"]
        assert "deviation_from_baseline_cost" not in position


class TestResolutionFlow:
    """Геометрия блока разрешается РОВНО один раз на блок, и в
    `parse_contractor_row` для каждой строки приходит тот же объект
    `ResolvedContractor` (identity), а не пересчёт. Это ловит именно тот
    регресс, ради которого убран `read_contractors` из `get_proposals`:
    если модуль снова начнёт читать геометрию сам, тест увидит либо второй
    вызов `resolve_contractor`, либо чужой объект в строках."""

    def test_resolution_happens_once_per_block_and_the_same_object_reaches_rows(self, monkeypatch):
        """Два лота, один блок: resolve_contractor вызван РОВНО один раз, и в
        parse_contractor_row приходит ТОТ ЖЕ объект (identity), а не пересчёт.
        Регресс «get_proposals снова читает геометрию сам» этим и ловится."""
        import parser.estimate as estimate_module
        import parser.get_lot_positions as glp_module

        ws = _sheet_with_summary_rows(SUMMARY_TRIPLE, second_lot=True)

        resolved_seen = []
        real_resolve = estimate_module.resolve_contractor
        monkeypatch.setattr(
            estimate_module, "resolve_contractor",
            lambda ws_, c, hr: resolved_seen.append(real_resolve(ws_, c, hr)) or resolved_seen[-1],
        )
        row_contractors = []
        real_row = glp_module.parse_contractor_row
        monkeypatch.setattr(
            glp_module, "parse_contractor_row",
            lambda ws_, row, contractor: row_contractors.append(contractor) or real_row(ws_, row, contractor),
        )

        parse_worksheet(ws)

        assert len(resolved_seen) == 1
        assert row_contractors, "ни одна строка не прошла через parse_contractor_row"
        assert all(c is resolved_seen[0] for c in row_contractors)


class TestParseEstimateFailures:
    """Структурно непригодные файлы отвергаются с внятной причиной.

    Граница отказа — в двух местах: геометрия (объединение заголовка блока
    подрядчика) проверяет `_validate_contractor_geometry`, смысл его колонок —
    `resolve_contractor` (спека §2.6). Ни то, ни другое не дублируется.
    """

    def test_raises_without_contractor_header_row(self):
        from openpyxl import Workbook

        ws = Workbook().active
        ws["A1"] = "Просто таблица"

        with pytest.raises(EstimateParseError, match="Наименование контрагента"):
            parse_worksheet(ws)

    def test_raises_when_contractor_marker_has_no_contractor(self):
        from openpyxl import Workbook

        ws = Workbook().active
        ws["G6"] = "Наименование контрагента"

        with pytest.raises(EstimateParseError, match="самого подрядчика"):
            parse_worksheet(ws)

    def test_raises_without_lot_marker(self):
        from openpyxl import Workbook

        ws = Workbook().active
        ws["G6"] = "Наименование контрагента"
        ws["J6"] = 'ООО "Тест"'

        with pytest.raises(EstimateParseError, match="Лот №"):
            parse_worksheet(ws)

    def test_unknown_column_label_is_a_refusal_through_the_full_path(self):
        """Проводка контракта §2.6 через parse_worksheet — сами классы отказа
        проверены в test_resolve_contractor.py."""
        ws = _minimal_sheet(11)
        ws.cell(row=10, column=12, value="Труд")  # подпись СМР группы цен
        with pytest.raises(EstimateParseError, match="L9"):
            parse_worksheet(ws)

    def test_twelve_wide_block_parses_with_an_extra_key_warning(self):
        """Ширина 12 разбирается; лишний deviation_from_baseline_cost — отличие
        от эталона ГП и предупреждение поимённо (спека §2.5)."""
        ws = _minimal_sheet(12)
        result = parse_worksheet(ws)
        matching = [w for w in result.warnings if "не совпадает с ожидаемым" in w]
        assert len(matching) == 1
        assert "deviation_from_baseline_cost" in matching[0]

    def test_raises_when_contractor_header_is_not_merged(self):
        """Необъединённый заголовок раньше давал `KeyError` по `merged_shape`."""
        ws = _minimal_sheet(contractor_colspan=None)

        with pytest.raises(EstimateParseError, match="не объединён"):
            parse_worksheet(ws)

    def test_supported_but_non_gp_colspan_is_a_warning_not_an_error(self):
        """Ширина 8 разбирается: раскладка известна, просто это не смета ГП.

        Здесь предупреждение обещает импорт — и импорт действительно возможен.
        Утверждение — набор ключей поимённо (§2.5), не число колонок: прежняя
        мысль «блок занимает N колонок вместо 11» исчезла вместе с самой собой.
        """
        ws = _minimal_sheet(contractor_colspan=8)

        result = parse_worksheet(ws)

        matching = [w for w in result.warnings if "не совпадает с ожидаемым" in w]
        assert len(matching) == 1
        for key in ("suggested_quantity", "total_cost_for_organizer_quantity", "comment_contractor"):
            assert key in matching[0]
        assert not any("колонок вместо" in w for w in result.warnings)


# --- вспомогательное ---


def _minimal_sheet(contractor_colspan: int | None):
    """Лист с шапкой контрагентов, шапкой блока и маркером лота, без позиций.

    None — заголовок подрядчика не объединён (случай геометрии); ширины 8–12
    получают канонические измеренные раскладки sheet_builders.COLUMNS_BY_WIDTH.
    """
    if contractor_colspan is None:
        from openpyxl import Workbook
        ws = Workbook().active
        ws["G6"] = "Наименование контрагента"
        ws["J6"] = 'ООО "Тест"'
        ws["D11"] = "Лот №1 Тестовый"
        for column, title in TABLE_PARSE_POSITION_COLUMN_HEADERS.items():
            ws.cell(row=9, column=column, value=title)
        ws["A11"] = 1
        ws["B11"] = 1
        return ws
    return gp_sheet(COLUMNS_BY_WIDTH[contractor_colspan])


def _proposal(result):
    return result.data["lots"]["lot_1"]["proposals"]["contractor_1"]


def _positions(result):
    return _proposal(result)["contractor_items"]["positions"]


def _dig(container, path):
    """Достаёт значение по пути ключей; None, если по дороге нет словаря."""
    value = container
    for key in path:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def _money(value):
    """Денежное значение из raw_data → Decimal; None, если это не число-строка."""
    if not isinstance(value, str):
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def _amount(value):
    """Количество (оно осталось числом) → Decimal; None, если это не число."""
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return Decimal(str(value))


def _assert_close(actual: Decimal, expected: Decimal, label) -> None:
    """Сравнение денег в Decimal: копейка округления либо относительный допуск."""
    tolerance = max(KOPECK, abs(expected) * RELATIVE_TOLERANCE)

    assert abs(actual - expected) <= tolerance, f"{label}: {actual} != {expected} (допуск {tolerance})"


def _find_money_floats(node, path=()):
    """Возвращает пути до `float`, лежащих под денежными ключами."""
    found: list[str] = []
    if isinstance(node, dict):
        for key, value in node.items():
            found.extend(_find_money_floats(value, (*path, str(key))))
    elif isinstance(node, list):
        for index, value in enumerate(node):
            found.extend(_find_money_floats(value, (*path, str(index))))
    elif isinstance(node, float) and _is_money_path(path):
        found.append(".".join(path))
    return found


def _is_money_path(path) -> bool:
    """Оканчивается ли путь одним из денежных полей."""
    return any(path[-len(money_path) :] == money_path for money_path in MONEY_PATHS)


# Раскладка колонок замерена: 15..18 — total_cost.{materials,works,
# indirect_costs,total}. Тождество ниже сходится ТОЧНО: 120 = 100 + 20.
SUMMARY_TRIPLE = (
    ("ИТОГО, руб. с учетом НДС", {15: 120.0, 16: 120.0, 17: 120.0, 18: 120.0}),
    ("В том числе НДС", {15: 20.0, 16: 20.0, 17: 20.0, 18: 20.0}),
    ("ИТОГО, руб. без учета НДС", {15: 100.0, 16: 100.0, 17: 100.0, 18: 100.0}),
)

# Тот же блок с НУЛЕВЫМ налогом: тождество сходится (100 = 100 + 0), а
# отношение даёт ровно 0 — то есть сверка (§2.4) подтверждает заявленный ноль,
# а не мешает ему. Нужен тесту сквозного пути нулевой ставки.
SUMMARY_TRIPLE_ZERO_VAT = (
    ("ИТОГО, руб. с учетом НДС", {15: 100.0, 16: 100.0, 17: 100.0, 18: 100.0}),
    ("В том числе НДС", {15: 0.0, 16: 0.0, 17: 0.0, 18: 0.0}),
    ("ИТОГО, руб. без учета НДС", {15: 100.0, 16: 100.0, 17: 100.0, 18: 100.0}),
)


def _sheet_with_summary_rows(summary_rows, *, second_lot: bool = False):
    """Лист с блоком итогов под позициями; при `second_lot` — два лота, один блок.

    Блок итогов — факт уровня листа, а `get_summary` зовётся на каждое
    предложение каждого лота, поэтому при двух лотах один и тот же блок
    читается дважды. Строка сразу под блоком остаётся пустой — она терминатор,
    без неё обход прочитает то, что ниже, как ещё одну итоговую строку.

    Args:
        summary_rows: последовательность (метка колонки A, {номер колонки: значение}).
        second_lot: добавить второй маркер лота над блоком.
    """
    ws = _minimal_sheet(contractor_colspan=11)          # лот №1 в D11
    ws.cell(row=12, column=1, value=1)
    ws.cell(row=12, column=2, value="1")
    ws.cell(row=12, column=4, value="Работа первого лота")

    if second_lot:
        ws.cell(row=13, column=1, value=2)
        ws.cell(row=13, column=2, value="2")
        ws.cell(row=13, column=4, value="Лот №2 Второй")
        ws.cell(row=14, column=1, value=3)
        ws.cell(row=14, column=2, value="3")
        ws.cell(row=14, column=4, value="Работа второго лота")

    for offset, (label, money) in enumerate(summary_rows):
        row = 15 + offset
        ws.cell(row=row, column=1, value=label)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        for column, value in money.items():
            ws.cell(row=row, column=column, value=value)
    return ws


class TestSummaryWarningsReachTheTop:
    """Проводка предупреждений блока итогов наверх, а не только правило."""

    def test_warning_from_the_core_reaches_parse_result(self):
        """Один лот: предупреждение обязано доехать до ParseResult.warnings.

        Без этого теста обрыв проводки в get_proposals или
        read_lots_and_boundaries остался бы невидимым: тесты ядра зовут его
        напрямую.

        Блок — полная тройка плюс лишняя строка с меткой: если оставить
        ТОЛЬКО лишнюю строку, отсутствие валового ИТОГО даёт СВОЁ отдельное
        предупреждение («Валовое ИТОГО отсутствует»), и оно попутно
        перечисляет присутствующие метки — в такой урезанной сборке единственная
        присутствующая метка совпадает с лишней, и подстрока находит два разных
        предупреждения вместо одного, хотя дублирования нет ни на йоту. Полная
        тройка эту метку из перечня убирает, и подстрока остаётся однозначным
        следом только одного, проверяемого здесь предупреждения.
        """
        rows = [*SUMMARY_TRIPLE, ("Совершенно чужая метка", {})]
        ws = _sheet_with_summary_rows(rows)

        result = parse_worksheet(ws)

        matching = [w for w in result.warnings if "Совершенно чужая метка" in w]
        assert len(matching) == 1

    def test_the_same_warning_is_not_doubled_on_a_two_lot_sheet(self):
        """Два лота — один блок; наверху остаётся ОДИН экземпляр (спека §2.8)."""
        rows = [*SUMMARY_TRIPLE, ("Совершенно чужая метка", {})]
        ws = _sheet_with_summary_rows(rows, second_lot=True)

        result = parse_worksheet(ws)

        matching = [w for w in result.warnings if "Совершенно чужая метка" in w]
        assert len(matching) == 1, f"ожидался один экземпляр, получено {len(matching)}"

    def test_excel_error_in_a_summary_cell_is_named_in_the_warning_and_nulled_in_json(self):
        """Пара, которую спека §2.6 обязалась назвать вслух.

        Предупреждение говорит о том, ЧТО СТОЯЛО В ЯЧЕЙКЕ (`#REF!`), а в JSON
        на её месте `None` — штатная `replace_excel_errors_with_null`
        отрабатывает ПОСЛЕ разбора блока. Оба утверждения в одном тесте:
        порознь они выглядели бы противоречием.

        Блок ОБЯЗАН быть трёхстрочным: `check_arithmetic` выходит сразу, если
        нет хотя бы одной из трёх налоговых строк, — на однострочном блоке
        ветка с негодным значением не исполнилась бы вовсе, и тест краснел бы
        по чужой причине.
        """
        gross_label, gross_money = SUMMARY_TRIPLE[0]
        rows = [
            (gross_label, {**gross_money, 18: "#REF!"}),
            SUMMARY_TRIPLE[1],
            SUMMARY_TRIPLE[2],
        ]
        ws = _sheet_with_summary_rows(rows)

        result = parse_worksheet(ws)

        named = [w for w in result.warnings if "#REF!" in w]
        assert len(named) == 1, "негодное значение обязано быть названо ровно одним предупреждением"
        assert "не проверена" in named[0]
        assert not any("не сходится" in w for w in result.warnings), (
            "в трёх годных колонках тождество сходится точно — «не сходится» здесь быть не должно"
        )

        summary = _proposal(result)["contractor_items"]["summary"]
        total_cost = summary["total_cost_including_vat"]["total_cost"]
        assert total_cost["total"] is None, "#REF! обязан стать null штатной постобработкой"
        assert total_cost["materials"] == "120.0", "соседняя годная сумма не задета"


def test_parse_error_is_importable_from_the_package_root():
    """Публичный контракт: `from parser import EstimateParseError`.

    Класс переехал в `parser.errors` ради разрыва цикла импортов, но снаружи
    имя прежнее — на него завязан `services/import_pipeline`. Сверяется
    идентичность объекта, а не только импортируемость: два разных класса с одним
    именем ловились бы `except` мимо.
    """
    import parser as parser_package
    from parser.errors import EstimateParseError as FromErrors

    assert parser_package.EstimateParseError is FromErrors


class TestColumnHeaderGuard:
    """Чужая раскладка колонок A–D — структурный отказ (спека Ф2 §2.1).

    Не предупреждение: колонки A, B, C, D читаются по фиксированным позициям,
    поэтому при чужой шапке недостоверен весь позиционный разбор, а не только
    колонка статьи. Та же граница, что у `_validate_contractor_geometry`.
    """

    def test_correct_headers_parse(self):
        ws = _minimal_sheet(11)
        result = parse_worksheet(ws)
        assert result.data is not None

    @pytest.mark.parametrize(
        ("column", "letter", "wrong_value"),
        [
            (2, "B", "Глава"),
            (3, "C", "Артикул СМР"),
            (4, "D", "Наименование видов работ"),
        ],
    )
    def test_wrong_header_in_any_column_is_rejected(self, column, letter, wrong_value):
        """Колонка A сюда НЕ входит намеренно.

        Строка шапки ищется именно по маркеру в A, поэтому испорченный A даёт не
        «чужой заголовок», а «строка не найдена» — и то сообщение фактическое
        значение не называет (спека §2.1, пункт 4). Этот случай покрывает
        `test_missing_header_row_is_rejected_without_naming_a_row`.
        """
        ws = _minimal_sheet(11)
        ws.cell(row=9, column=column, value=wrong_value)

        with pytest.raises(EstimateParseError) as exc:
            parse_worksheet(ws)

        message = str(exc.value)
        assert letter in message
        assert wrong_value in message

    def test_missing_header_row_is_rejected_without_naming_a_row(self):
        """Маркер «№ п/п» испорчен — «ту самую» строку определить нельзя.

        Поэтому сообщение называет ожидаемый маркер и просмотренный диапазон,
        но НЕ фактическое значение: назвать его было бы выдумкой.
        """
        ws = _minimal_sheet(11)
        ws.cell(row=9, column=1, value="Порядковый номер")

        with pytest.raises(EstimateParseError) as exc:
            parse_worksheet(ws)

        message = str(exc.value)
        assert "№ п/п" in message
        # Диапазон: строка заголовка контрагентов у `_minimal_sheet` — 6, маркер
        # лота — 11, значит просмотрены строки 7–10.
        assert "7–10" in message, message
        # Ключевое утверждение docstring'а: фактического значения в сообщении нет.
        # Без этой строки тест был бы зелёным и у реализации, которая его называет
        # (найдено финальным ревью, проверено снятием защиты).
        assert "Порядковый номер" not in message, message

    def test_header_row_is_found_not_hardcoded(self):
        """Шапка сдвинута на строку — файл валиден и должен разбираться.

        Ради этого строка ищется по маркеру, а не берётся константой 9.

        Строитель кладёт и A–D, и двухъярусную шапку блока подрядчика в ОДНУ
        и ту же сдвинутую строку (`gp_sheet(header_row=8)`): раздельное
        затирание строки 9 и запись в 8 сдвинуло бы только общие колонки,
        оставив шапку блока подрядчика на месте — `resolve_contractor` искал
        бы группы там, где их больше нет, и тест падал бы по чужой причине
        (не «шапка не найдена», а «колонка не опознана»).
        """
        ws = gp_sheet(KEYS_GP_11, header_row=8)

        result = parse_worksheet(ws)
        assert result.data is not None

    def test_headers_are_compared_ignoring_case_and_extra_spaces(self):
        ws = _minimal_sheet(11)
        ws.cell(row=9, column=3, value="  статья  смр  ")

        result = parse_worksheet(ws)
        assert result.data is not None


class TestVatRateFullPath:
    """Полный путь Ф4б: XLSX → `ParseResult` (спека §4.2).

    Ядро (`vat_rate.py`) проверено без файла в `test_vat_rate.py`; здесь —
    только проводка: что лист со ставкой в шапке даёт `vat_rate` в
    `ParseResult.data`, а лист без неё — предупреждение в
    `ParseResult.warnings`, и что оба не путают лоты и строки шапки.
    """

    def test_vat_rate_reaches_parse_result_data_with_no_warnings(self):
        """Суффикс в обеих групповых шапках + сходящийся блок итогов
        (`SUMMARY_TRIPLE`) → ставка в `ParseResult.data`, предупреждений НЕТ
        ВООБЩЕ — не только о ставке: блок итогов у этого листа полный и
        годный, поэтому и сверка (§2.4) проходит тишиной, как на реальном
        fixture (`test_parses_without_warnings`). Шапку блока и суффикс
        `с учетом НДС 20%` даёт строитель (`_minimal_sheet(11)` →
        `gp_sheet(KEYS_GP_11)`) — ручных заплаток в строке 9 больше не нужно.
        """
        ws = _sheet_with_summary_rows(SUMMARY_TRIPLE)

        result = parse_worksheet(ws)

        assert _proposal(result)["vat_rate"] == "20"
        assert result.warnings == []

    def test_declared_zero_rate_reaches_parse_result_as_zero_not_as_absence(self):
        """Заявленный `0%` доезжает до `ParseResult.data` нулём, а не `None`.

        Найдено финальным ревью ветки: ноль проверялся только на самом
        внутреннем слое (`read_label_rate`), и весь путь наверх держался на том,
        что и ядро, и `get_proposals` сверяют `is None`, а не truthiness. Регресс
        вида `if not declared.rate` или `if not vat_rate_result.rate` прошёл бы
        весь набор зелёным, молча превратив единственный законный ноль в `NULL` —
        то самое неразличение «ноль» и «не заявлено», от которого спека §2.2
        отказалась явно.

        Блок итогов взят с нулевым НДС при ненулевой базе: `0 / 100 * 100 = 0`,
        то есть сверка не только не мешает, но и подтверждает ставку.

        Суффикс меняется поверх текста, положенного строителем: собственный
        текст без префикса типа группы (`TABLE_PARSE_UNIT_COST_GROUP_PREFIX` /
        `_TOTAL_COST_GROUP_PREFIX`) не опознался бы `resolve_contractor`'ом
        (спека §2.1) — `gp_sheet` не принимает суффикс отдельно от базового
        текста группы через `_sheet_with_summary_rows`, поэтому здесь пишется
        итоговый текст целиком, тот же, что даёт строитель, только с 0%.
        """
        ws = _sheet_with_summary_rows(SUMMARY_TRIPLE_ZERO_VAT)
        ws.cell(row=9, column=11).value = "Цена за ед. изм., RUB, ОСН, с учетом НДС 0%"
        ws.cell(row=9, column=15).value = "Стоимость всего, RUB, ОСН, с учетом НДС 0%"

        result = parse_worksheet(ws)

        assert _proposal(result)["vat_rate"] == "0"
        assert result.warnings == []

    def test_vat_rate_warning_reaches_parse_result_warnings(self):
        """Предупреждение ядра доезжает наверх проводкой `get_proposals` →
        `read_lots_and_boundaries` → `parse_worksheet`, а не гасится по дороге.

        `vat_suffix=None`: групповые шапки типизированы (замеренное написание,
        тип группы опознан), но без ставки — `read_label_rate` её не находит.
        """
        ws = gp_sheet(KEYS_GP_11, vat_suffix=None)

        result = parse_worksheet(ws)

        assert _proposal(result)["vat_rate"] is None
        assert any("Ставка НДС не получена из шапки" in w for w in result.warnings)

    def test_vat_rate_warning_is_not_doubled_on_a_two_lot_sheet(self):
        """Два лота — один блок подрядчика и одна шапка колонок; `dict.fromkeys`
        в `estimate.py` схлопывает повтор так же, как у предупреждений Ф4a
        (спека §2.8): тексты §2.7 не несут ничего, что различается между лотами.

        Суффикс снят с обеих групповых шапок (тот же довод, что у
        `test_vat_rate_warning_reaches_parse_result_warnings`): без этого
        `_minimal_sheet(11)` дал бы заявленную ставку по умолчанию, и
        предупреждение, которое стережёт тест, не возникло бы вовсе.
        """
        ws = _sheet_with_summary_rows(SUMMARY_TRIPLE, second_lot=True)
        ws.cell(row=9, column=11).value = "Цена за ед. изм., RUB, ОСН"
        ws.cell(row=9, column=15).value = "Стоимость всего, RUB, ОСН"

        result = parse_worksheet(ws)

        matching = [w for w in result.warnings if "Ставка НДС не получена из шапки" in w]
        assert len(matching) == 1, f"ожидался один экземпляр, получено {len(matching)}"

    def test_vat_rate_survives_a_header_row_shift(self):
        """Шапка сдвинута на строку выше 9-й — ставка всё равно найдена.

        `header_row` приходит из `_validate_column_headers`
        (`_find_column_header_row`), а не константы 9 (та же гарантия, что у
        `test_header_row_is_found_not_hardcoded`). Строитель кладёт и A–D, и
        двухъярусную шапку блока подрядчика в ОДНУ сдвинутую строку —
        отдельные затирания больше не нужны.
        """
        ws = gp_sheet(KEYS_GP_11, header_row=8)

        result = parse_worksheet(ws)

        assert _proposal(result)["vat_rate"] == "20"

    @pytest.mark.parametrize("colspan", sorted(COLUMNS_BY_WIDTH))
    def test_vat_rate_found_at_measured_offsets_for_every_supported_width(self, colspan):
        """Суффикс в шапке даёт ставку на каждой измеренной раскладке (спека
        §2.3, §4.2): якоря групповых шапок — смещения `BlockLayout`, посчитанные
        `resolve_contractor` из того же перечня ключей, что и сами колонки
        (спека §2.4) — строитель просто кладёт суффикс в анкер группы, где он
        физически стоит на каждой измеренной раскладке.
        """
        ws = gp_sheet(COLUMNS_BY_WIDTH[colspan])

        result = parse_worksheet(ws)

        assert _proposal(result)["vat_rate"] == "20"
