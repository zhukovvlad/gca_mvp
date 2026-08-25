"""Синтетические листы с настоящей двухъярусной шапкой блока подрядчика.

После фичи «колонки по заголовкам» ни один синтетический лист не разбирается
без полной шапки блока: resolve_contractor отказывает на неопознанной паре.
Геометрия — та же, что замерена на реальных файлах (план фичи, «Замеры»):
заголовок группы объединён по горизонтали в строке header_row, подписи группы —
в header_row+1; «Стоимость всего за объемы заказчика» объединена по вертикали;
«Комментарий участника» лежит в верхнем ярусе без объединения; «Предлагаемое
количество» и «% от р/с» — в нижнем ярусе.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from parser.constants import (
    JSON_KEY_COMMENT_CONTRACTOR,
    JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
    JSON_KEY_INDIRECT_COSTS,
    JSON_KEY_MATERIALS,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_SUGGESTED_QUANTITY,
    JSON_KEY_TOTAL,
    JSON_KEY_TOTAL_COST,
    JSON_KEY_UNIT_COST,
    JSON_KEY_WORKS,
    TABLE_PARSE_COMMENT_CONTRACTOR_LABEL,
    TABLE_PARSE_DEVIATION_COLUMN_LABEL,
    TABLE_PARSE_ORGANIZER_QUANTITY_LABEL,
    TABLE_PARSE_POSITION_COLUMN_HEADERS,
    TABLE_PARSE_SUGGESTED_QUANTITY,
)
from parser.resolve_contractor import BlockLayout, ResolvedContractor

_UC = JSON_KEY_UNIT_COST
_TC = JSON_KEY_TOTAL_COST
_MONEY_EIGHT = tuple(
    f"{group}.{part}"
    for group in (_UC, _TC)
    for part in (JSON_KEY_MATERIALS, JSON_KEY_WORKS, JSON_KEY_INDIRECT_COSTS, JSON_KEY_TOTAL)
)

KEYS_8 = _MONEY_EIGHT
KEYS_9 = (*_MONEY_EIGHT, JSON_KEY_COMMENT_CONTRACTOR)
KEYS_10 = (JSON_KEY_SUGGESTED_QUANTITY, *_MONEY_EIGHT, JSON_KEY_COMMENT_CONTRACTOR)
KEYS_GP_11 = (
    JSON_KEY_SUGGESTED_QUANTITY,
    *_MONEY_EIGHT,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_COMMENT_CONTRACTOR,
)
KEYS_TENDER_11 = (
    JSON_KEY_SUGGESTED_QUANTITY,
    *_MONEY_EIGHT,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
)
KEYS_12 = (
    JSON_KEY_SUGGESTED_QUANTITY,
    *_MONEY_EIGHT,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_COMMENT_CONTRACTOR,
    JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
)
#: Перестановка при ПРЕЖНЕМ colspan 11 (спека §6): группы целиком, но не на
#: канонических местах; одиночные колонки перемешаны.
KEYS_PERMUTED_11 = (
    JSON_KEY_COMMENT_CONTRACTOR,
    f"{_TC}.{JSON_KEY_MATERIALS}",
    f"{_TC}.{JSON_KEY_WORKS}",
    f"{_TC}.{JSON_KEY_INDIRECT_COSTS}",
    f"{_TC}.{JSON_KEY_TOTAL}",
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    f"{_UC}.{JSON_KEY_MATERIALS}",
    f"{_UC}.{JSON_KEY_WORKS}",
    f"{_UC}.{JSON_KEY_INDIRECT_COSTS}",
    f"{_UC}.{JSON_KEY_TOTAL}",
    JSON_KEY_SUGGESTED_QUANTITY,
)

#: Перестановка ВНУТРИ денежных групп при том же colspan 11 и тех же местах
#: самих групп (в отличие от KEYS_PERMUTED_11, где переставлены целиком
#: группы): подписи внутри каждой группы идут «СМР, Материалы, Косвенные
#: расходы, Всего» — Материалы не первая колонка группы. Ловит именно якорь
#: денежной группы: он обязан быть левой границей объединения (первой
#: физической колонкой группы), а не колонкой Материалы, которая здесь
#: сдвинута внутрь группы на одну позицию.
KEYS_SUBLABELS_SWAPPED_11 = (
    JSON_KEY_SUGGESTED_QUANTITY,
    f"{_UC}.{JSON_KEY_WORKS}", f"{_UC}.{JSON_KEY_MATERIALS}", f"{_UC}.{JSON_KEY_INDIRECT_COSTS}", f"{_UC}.{JSON_KEY_TOTAL}",
    f"{_TC}.{JSON_KEY_WORKS}", f"{_TC}.{JSON_KEY_MATERIALS}", f"{_TC}.{JSON_KEY_INDIRECT_COSTS}", f"{_TC}.{JSON_KEY_TOTAL}",
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_COMMENT_CONTRACTOR,
)

#: Канонические раскладки по ширине — ТОЛЬКО для удобства тестов: продакшен-код
#: ширину со смыслом не связывает.
COLUMNS_BY_WIDTH = {8: KEYS_8, 9: KEYS_9, 10: KEYS_10, 11: KEYS_GP_11, 12: KEYS_12}

_SUBLABEL_BY_PART = {
    JSON_KEY_MATERIALS: "Материалы",
    JSON_KEY_WORKS: "СМР",
    JSON_KEY_INDIRECT_COSTS: "Косвенные расходы",
    JSON_KEY_TOTAL: "Всего",
}
_GROUP_HEAD = {_UC: "Цена за ед. изм., RUB, ОСН", _TC: "Стоимость всего, RUB, ОСН"}


def add_contractor_block(
    ws: Worksheet,
    *,
    col_start: int,
    columns: tuple[str, ...],
    title: str = 'ООО "Тест"',
    contractor_row: int = 6,
    header_row: int = 9,
    vat_suffix: str | None = ", с учетом НДС 20%",
) -> Worksheet:
    """Пишет заголовок подрядчика и двухъярусную шапку его блока.

    Группы (подряд идущие ключи одного префикса) объединяются по горизонтали;
    группа из одной колонки — ошибка вызова: тип группы существует только у
    объединения шириной больше единицы (спека §2.3 п.1).
    """
    colspan = len(columns)
    ws.merge_cells(
        start_row=contractor_row, start_column=col_start,
        end_row=contractor_row, end_column=col_start + colspan - 1,
    )
    ws.cell(row=contractor_row, column=col_start, value=title)

    i = 0
    while i < colspan:
        key = columns[i]
        col = col_start + i
        group = key.split(".")[0] if "." in key else None
        if group in (_UC, _TC):
            run = 1
            while i + run < colspan and columns[i + run].startswith(f"{group}."):
                run += 1
            if run < 2:
                raise ValueError(f"группа {group} из одной колонки не имеет типа (спека §2.3 п.1)")
            ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row, end_column=col + run - 1)
            ws.cell(row=header_row, column=col, value=_GROUP_HEAD[group] + (vat_suffix or ""))
            for j in range(run):
                part = columns[i + j].split(".", 1)[1]
                ws.cell(row=header_row + 1, column=col + j, value=_SUBLABEL_BY_PART[part])
            i += run
            continue
        if key == JSON_KEY_SUGGESTED_QUANTITY:
            ws.cell(row=header_row + 1, column=col, value=TABLE_PARSE_SUGGESTED_QUANTITY)
        elif key == JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST:
            ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row + 1, end_column=col)
            ws.cell(row=header_row, column=col, value=TABLE_PARSE_ORGANIZER_QUANTITY_LABEL)
        elif key == JSON_KEY_COMMENT_CONTRACTOR:
            ws.cell(row=header_row, column=col, value=TABLE_PARSE_COMMENT_CONTRACTOR_LABEL)
        elif key == JSON_KEY_DEVIATION_FROM_CALCULATED_COST:
            ws.cell(row=header_row + 1, column=col, value=TABLE_PARSE_DEVIATION_COLUMN_LABEL)
        else:
            raise ValueError(f"строитель не знает ключа {key!r}")
        i += 1
    return ws


def gp_sheet(
    columns: tuple[str, ...] = KEYS_GP_11,
    *,
    col_start: int = 10,
    header_row: int = 9,
    **block_kwargs,
) -> Worksheet:
    """Минимальный лист сметы ГП: маркер контрагентов, шапка A–D, блок, лот.

    Наследник `_minimal_sheet` из test_estimate.py: та же строка маркера G6,
    маркер лота D11 и «позиция» A11/B11 — плюс настоящая шапка блока.
    """
    ws = Workbook().active
    ws["G6"] = "Наименование контрагента"
    add_contractor_block(ws, col_start=col_start, columns=columns, header_row=header_row, **block_kwargs)
    for column, title in TABLE_PARSE_POSITION_COLUMN_HEADERS.items():
        ws.cell(row=header_row, column=column, value=title)
    ws["D11"] = "Лот №1 Тестовый"
    ws["A11"] = 1
    ws["B11"] = 1
    return ws


def _group_anchor_offset(columns: tuple[str, ...], group: str) -> int:
    """Смещение ПЕРВОЙ физической колонки группы — якорь, а не «.materials».

    Совпадает с продакшен-контрактом BlockLayout: offset должен указывать на
    левую границу горизонтального объединения группы, потому что openpyxl
    отдаёт None для любой другой ячейки внутри объединения (см. докстринг
    BlockLayout в resolve_contractor.py).
    """
    return next(i for i, key in enumerate(columns) if key.startswith(f"{group}."))


def resolved(col_start: int, columns: tuple[str, ...], **geometry_extra) -> ResolvedContractor:
    """ResolvedContractor для юнит-тестов нижних слоёв — без листа и шапки.

    Смещения выводятся из columns как якоря групп (первая физическая колонка
    каждой группы), а не как индекс «.materials» — тесты, которые СТЕРЕГУТ
    смещения, пишут их литералами в test_resolve_contractor.py, а не берут
    отсюда.
    """
    geometry = {
        "column_start": col_start,
        "merged_shape": {"rowspan": 1, "colspan": len(columns)},
        **geometry_extra,
    }
    return ResolvedContractor(
        geometry=geometry,
        layout=BlockLayout(
            column_keys=tuple(columns),
            unit_cost_offset=_group_anchor_offset(columns, _UC),
            total_cost_offset=_group_anchor_offset(columns, _TC),
        ),
    )
