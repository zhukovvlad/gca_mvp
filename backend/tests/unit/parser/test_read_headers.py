"""Тесты чтения шапки документа.

Перенос `app/tests/excel_parser/test_read_headers.py` из
`parser_tender_xlsx@0e178c0`.

Изменения при переносе:

* фикстура на `pytest-mock` (`mocker.patch`) заменена на штатный `monkeypatch` —
  лишней зависимости в проекте нет;
* тест «данные вне диапазона сканирования игнорируются» переписан: диапазон
  расширен вверх до строки 2 (шапка сметы ГП), поэтому за границей теперь
  строки 1 и 6, а не 2 и 6;
* добавлен тест на реальную раскладку сметы ГП — ровно тот случай, который
  исходный диапазон 3–5 терял.
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from parser.constants import (
    JSON_KEY_TENDER_ADDRESS,
    JSON_KEY_TENDER_ID,
    JSON_KEY_TENDER_OBJECT,
    JSON_KEY_TENDER_TITLE,
    TABLE_PARSE_ADDRESS,
    TABLE_PARSE_OBJECT,
    TABLE_PARSE_TENDER_SUBJECT,
)
from parser.read_headers import read_headers


@pytest.fixture
def passthrough_sanitize(monkeypatch):
    """Делает sanitize_text прозрачным, чтобы тесты проверяли только read_headers."""
    import parser.read_headers as module

    monkeypatch.setattr(module, "sanitize_text", lambda text: text)


def test_read_headers_happy_path(passthrough_sanitize):
    """Идеальный сценарий: все данные на месте."""
    ws = Workbook().active
    ws["A3"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B3"] = "№ID-123 Закупка оборудования"
    ws["A4"] = TABLE_PARSE_OBJECT
    ws["C4"] = "Главный корпус"  # значение может быть в любой колонке
    ws["A5"] = TABLE_PARSE_ADDRESS
    ws["B5"] = "г. Тест, ул. Тестовая, д. 1"

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] == "ID-123"
    assert extracted[JSON_KEY_TENDER_TITLE] == "Закупка оборудования"
    assert extracted[JSON_KEY_TENDER_OBJECT] == "Главный корпус"
    assert extracted[JSON_KEY_TENDER_ADDRESS] == "г. Тест, ул. Тестовая, д. 1"


def test_read_headers_gp_estimate_layout(passthrough_sanitize):
    """Раскладка сметы ГП: шапка начинается со строки 2.

    Исходный диапазон 3–5 на этой раскладке молча терял tender_id и
    tender_title (замер на реальном файле — docs/phase3-parser.md).
    """
    ws = Workbook().active
    ws["A2"] = "Предмет тендера: "
    ws["D2"] = '№001-ТУ "Тестовый ЖК_Генподряд"'
    ws["A3"] = "Объект: "
    ws["D3"] = "Тестовый ЖК Корпус 1"
    ws["A4"] = "Адрес объекта"
    ws["D4"] = "г. Тестоград, ул. Примерная, вл 1"

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] == "001-ТУ"
    assert extracted[JSON_KEY_TENDER_TITLE] == '"Тестовый ЖК_Генподряд"'
    assert extracted[JSON_KEY_TENDER_OBJECT] == "Тестовый ЖК Корпус 1"
    assert extracted[JSON_KEY_TENDER_ADDRESS] == "г. Тестоград, ул. Примерная, вл 1"


def test_read_headers_handles_empty_sheet(passthrough_sanitize):
    ws = Workbook().active
    assert all(value is None for value in read_headers(ws).values())


def test_read_headers_subject_with_only_id(passthrough_sanitize):
    """Если после номера нет пробела и названия, номер идёт и в название."""
    ws = Workbook().active
    ws["A3"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B3"] = "№456789"

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] == "456789"
    assert extracted[JSON_KEY_TENDER_TITLE] == "456789"


def test_read_headers_ignores_data_outside_scan_range(passthrough_sanitize):
    """Строки 1 и 6 — за границами диапазона 2–5."""
    ws = Workbook().active
    ws["A1"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B1"] = "№ID-WRONG Закупка"
    ws["A6"] = TABLE_PARSE_OBJECT
    ws["B6"] = "Неправильный корпус"

    assert all(value is None for value in read_headers(ws).values())


def test_read_headers_handles_messy_data_and_extra_spaces(passthrough_sanitize):
    ws = Workbook().active
    ws["A4"] = "  Объект  "
    ws["C4"] = "   Здание АБК   "

    assert read_headers(ws)[JSON_KEY_TENDER_OBJECT] == "Здание АБК"


def test_read_headers_ignores_keyword_if_not_first_element(passthrough_sanitize):
    """Метка обязана быть первым непустым значением строки."""
    ws = Workbook().active
    ws["A4"] = "Дополнительно:"
    ws["B4"] = TABLE_PARSE_OBJECT
    ws["C4"] = "Какой-то корпус"

    assert read_headers(ws)[JSON_KEY_TENDER_OBJECT] is None


def test_read_headers_parses_subject_split_by_first_space(passthrough_sanitize):
    """Разбор предмета тендера идёт по первому пробелу — фиксируем как есть."""
    ws = Workbook().active
    ws["A3"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B3"] = '№777-ABC"Закупка ПО "Альфа""'

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] == '777-ABC"Закупка'
    assert extracted[JSON_KEY_TENDER_TITLE] == 'ПО "Альфа""'


def test_read_headers_with_partial_data(passthrough_sanitize):
    ws = Workbook().active
    ws["A4"] = TABLE_PARSE_OBJECT
    ws["B4"] = "Только объект"

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] is None
    assert extracted[JSON_KEY_TENDER_TITLE] is None
    assert extracted[JSON_KEY_TENDER_OBJECT] == "Только объект"
    assert extracted[JSON_KEY_TENDER_ADDRESS] is None


def test_read_headers_key_with_empty_or_whitespace_value(passthrough_sanitize):
    ws = Workbook().active
    ws["A4"] = TABLE_PARSE_OBJECT
    ws["B4"] = "    "
    ws["A5"] = TABLE_PARSE_ADDRESS  # ключ есть, значения дальше в строке нет

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_OBJECT] is None
    assert extracted[JSON_KEY_TENDER_ADDRESS] is None


def test_read_headers_duplicate_key_uses_last_one(passthrough_sanitize):
    """При дублировании ключа побеждает нижняя строка.

    Это свойство и делает расширение диапазона вверх безопасным для тендерных
    таблиц: их шапка ниже, она перекрывает то, что нашлось в строке 2.
    """
    ws = Workbook().active
    ws["A4"] = TABLE_PARSE_OBJECT
    ws["B4"] = "Старый объект"
    ws["A5"] = TABLE_PARSE_OBJECT
    ws["B5"] = "Новый объект"

    assert read_headers(ws)[JSON_KEY_TENDER_OBJECT] == "Новый объект"


def test_read_headers_handles_numeric_values(passthrough_sanitize):
    ws = Workbook().active
    ws["A3"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B3"] = 123456789

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] == "123456789"
    assert extracted[JSON_KEY_TENDER_TITLE] == "123456789"


def test_read_headers_subject_with_only_symbol(passthrough_sanitize):
    """Значение из одного «№» даёт пустой идентификатор, то есть None."""
    ws = Workbook().active
    ws["A3"] = TABLE_PARSE_TENDER_SUBJECT
    ws["B3"] = "№"

    extracted = read_headers(ws)

    assert extracted[JSON_KEY_TENDER_ID] is None
    assert extracted[JSON_KEY_TENDER_TITLE] is None
