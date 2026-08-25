"""Тесты сборки предложений подрядчиков.

Перенос `app/tests/excel_parser/test_get_proposals.py` из
`parser_tender_xlsx@0e178c0`.

Изменение при переносе: `get_summary` теперь получает третий аргумент —
начало блока позиций (`search_start_row`), см. `parser/get_summary.py`.
Тесты, проверявшие «в get_summary передаются ровно два аргумента»,
переписаны под новую сигнатуру.

Добавлен тест на смету ГП с одним подрядчиком (AGENTS.md §4).

Фича «колонки по заголовкам» (спека §2.5): `header_row` и `contractors`
приходят сверху уже готовыми — `header_row` от `estimate._validate_column_headers`,
`contractors` от `resolve_contractor`, вызванного на каждый блок ДО
`get_proposals`. `read_contractors` здесь больше не патчится — его нет в
модуле: геометрия не читается второй раз. Список `contractors` без
ячейки-маркера: обход `enumerate(contractors, start=1)` в продакшен-коде не
пропускает нулевой элемент, потому что пропускать больше нечего.
"""
from __future__ import annotations

from unittest.mock import patch

import pytest
from openpyxl import Workbook

from parser.constants import (
    JSON_KEY_CONTRACTOR_ACCREDITATION,
    JSON_KEY_CONTRACTOR_ADDITIONAL_INFO,
    JSON_KEY_CONTRACTOR_ADDRESS,
    JSON_KEY_CONTRACTOR_COORDINATE,
    JSON_KEY_CONTRACTOR_HEIGHT,
    JSON_KEY_CONTRACTOR_INN,
    JSON_KEY_CONTRACTOR_ITEMS,
    JSON_KEY_CONTRACTOR_POSITIONS,
    JSON_KEY_CONTRACTOR_SUMMARY,
    JSON_KEY_CONTRACTOR_TITLE,
    JSON_KEY_CONTRACTOR_WIDTH,
    JSON_KEY_VAT_RATE,
)
from parser.get_lot_positions import LotRows
from parser.get_proposals import get_proposals
from parser.summary_block import SummaryBlock

from .sheet_builders import KEYS_10, KEYS_GP_11, resolved

MODULE = "parser.get_proposals"

#: Номер строки шапки колонок, который `parse_worksheet` проводит вниз.
#: На синтетических листах этих тестов шапка не заполнена, поэтому значение
#: служит только для сигнатуры; 9 — та строка, где шапка стоит во всех четырёх
#: известных файлах.
COLUMN_HEADER_ROW = 9


@pytest.fixture
def empty_worksheet():
    return Workbook().active


@pytest.fixture
def sample_contractors_data():
    """Два разрешённых блока подрядчиков — БЕЗ ячейки-маркера.

    `read_contractors` в `get_proposals` больше не зовётся (спека §2.5):
    `contractors` приходит уже разрешённым списком, и маркер «Наименование
    контрагента» в нём никогда не было бы — его снимает `estimate.py` до
    `resolve_contractor`. `column_start` у двух подрядчиков — 2 и 4, их
    11-колоночные блоки при этом перекрываются; для чтения ячеек шапки это
    безразлично — лист пустой, обе метки дают `None`.
    """
    return [
        resolved(2, KEYS_GP_11, value="ООО Строитель", row_start=2, coordinate="B2"),
        resolved(4, KEYS_GP_11, value="АО Подрядчик", row_start=2, coordinate="D2"),
    ]


@pytest.fixture
def sample_positions_data():
    return {"1": {"job_title": "Позиция 1"}, "2": {"job_title": "Позиция 2"}}


@pytest.fixture
def sample_summary_data():
    return {"total_amount": 15000, "vat": 2700}


@pytest.fixture
def sample_additional_info():
    return {"comment": "Дополнительная информация", "deadline": "30 дней"}


def _patch_collaborators(positions=None, summary=None, additional=None):
    """Контекст с подменёнными соседями get_proposals.

    Без патча `read_contractors` — его в модуле больше нет: `contractors`
    приходит параметром вызова, а не читается с листа второй раз (спека §2.5).
    """
    summary_block = SummaryBlock(lines={} if summary is None else summary, warnings=[])
    return (
        patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={} if positions is None else positions)),
        patch(f"{MODULE}.get_summary", return_value=summary_block),
        patch(f"{MODULE}.get_additional_info", return_value={} if additional is None else additional),
    )


class TestGetProposalsBasicBehavior:
    """Базовое поведение."""

    def test_returns_empty_dict_when_no_contractors(self, empty_worksheet):
        """Пустой список — валидный вход: обход по нему пуст, отдельной ветки
        не нужно (спека §2.5)."""
        assert get_proposals(empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=[]).proposals == {}


class TestGetProposalsDataExtraction:
    """Извлечение реквизитов подрядчика из ячеек."""

    def test_extracts_basic_contractor_info(self, empty_worksheet, sample_contractors_data):
        empty_worksheet.cell(row=3, column=2, value="1234567890")
        empty_worksheet.cell(row=4, column=2, value="г. Москва, ул. Ленина, д.1")
        empty_worksheet.cell(row=5, column=2, value="Аккредитация есть")

        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        assert contractor_1[JSON_KEY_CONTRACTOR_TITLE] == "ООО Строитель"
        assert contractor_1[JSON_KEY_CONTRACTOR_INN] == "1234567890"
        assert contractor_1[JSON_KEY_CONTRACTOR_ADDRESS] == "г. Москва, ул. Ленина, д.1"
        assert contractor_1[JSON_KEY_CONTRACTOR_ACCREDITATION] == "Аккредитация есть"

    def test_extracts_coordinate_and_dimensions(self, empty_worksheet, sample_contractors_data):
        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        assert contractor_1[JSON_KEY_CONTRACTOR_COORDINATE] == "B2"
        assert contractor_1[JSON_KEY_CONTRACTOR_WIDTH] == 11
        assert contractor_1[JSON_KEY_CONTRACTOR_HEIGHT] == 1

    def test_handles_merged_cells(self, empty_worksheet):
        """При rowspan > 1 реквизиты под заголовком не читаются.

        `rowspan` = 2 — предмет самого теста, `merged_shape` переопределён
        поверх измеренной раскладки KEYS_10 через `geometry_extra`.
        """
        contractor_with_merged = resolved(
            2,
            KEYS_10,
            value="ООО Большая компания",
            row_start=2,
            coordinate="B2",
            merged_shape={"rowspan": 2, "colspan": 10},
        )

        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=[contractor_with_merged]
            ).proposals

        contractor_1 = result["contractor_1"]
        assert contractor_1[JSON_KEY_CONTRACTOR_WIDTH] == 10
        assert contractor_1[JSON_KEY_CONTRACTOR_HEIGHT] == 2
        assert contractor_1[JSON_KEY_CONTRACTOR_INN] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ADDRESS] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ACCREDITATION] is None


class TestGetProposalsModuleIntegration:
    """Взаимодействие с соседними модулями."""

    def test_calls_get_lot_positions_with_correct_parameters(self, empty_worksheet, sample_contractors_data):
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})) as mock_positions,
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])),
            patch(f"{MODULE}.get_additional_info", return_value={}),
        ):
            get_proposals(empty_worksheet, 15, 25, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        assert mock_positions.call_count == 2

        call_args = mock_positions.call_args_list[0]
        assert call_args[0][0] is empty_worksheet
        assert call_args[0][1] == sample_contractors_data[0]
        assert call_args[1]["lot_start_row"] == 15
        assert call_args[1]["lot_end_row"] == 25

    def test_calls_get_summary_for_each_contractor(self, empty_worksheet, sample_contractors_data):
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})),
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])) as mock_summary,
            patch(f"{MODULE}.get_additional_info", return_value={}),
        ):
            get_proposals(empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        assert mock_summary.call_count == 2

        for i, call_args in enumerate(mock_summary.call_args_list):
            assert call_args[0][0] is empty_worksheet
            assert call_args[0][1] == sample_contractors_data[i]

    def test_calls_get_additional_info_for_each_contractor(self, empty_worksheet, sample_contractors_data):
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})),
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])),
            patch(f"{MODULE}.get_additional_info", return_value={}) as mock_additional,
        ):
            get_proposals(empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        assert mock_additional.call_count == 2


class TestGetProposalsDataAggregation:
    """Сборка итоговой структуры."""

    def test_combines_positions_and_summary_in_items(
        self, empty_worksheet, sample_contractors_data, sample_positions_data, sample_summary_data
    ):
        p1, p2, p3 = _patch_collaborators(positions=sample_positions_data, summary=sample_summary_data)
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        items = result["contractor_1"][JSON_KEY_CONTRACTOR_ITEMS]
        assert items[JSON_KEY_CONTRACTOR_POSITIONS] == sample_positions_data
        assert items[JSON_KEY_CONTRACTOR_SUMMARY] == sample_summary_data

    def test_includes_additional_info_in_proposal(
        self, empty_worksheet, sample_contractors_data, sample_additional_info
    ):
        p1, p2, p3 = _patch_collaborators(additional=sample_additional_info)
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        assert result["contractor_1"][JSON_KEY_CONTRACTOR_ADDITIONAL_INFO] == sample_additional_info

    def test_creates_complete_proposal_structure(
        self,
        empty_worksheet,
        sample_contractors_data,
        sample_positions_data,
        sample_summary_data,
        sample_additional_info,
    ):
        empty_worksheet.cell(row=3, column=2, value="1234567890")
        empty_worksheet.cell(row=4, column=2, value="г. Москва")
        empty_worksheet.cell(row=5, column=2, value="Есть")

        p1, p2, p3 = _patch_collaborators(
            positions=sample_positions_data,
            summary=sample_summary_data,
            additional=sample_additional_info,
        )
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        for field in (
            JSON_KEY_CONTRACTOR_TITLE,
            JSON_KEY_CONTRACTOR_INN,
            JSON_KEY_CONTRACTOR_ADDRESS,
            JSON_KEY_CONTRACTOR_ACCREDITATION,
            JSON_KEY_CONTRACTOR_COORDINATE,
            JSON_KEY_CONTRACTOR_WIDTH,
            JSON_KEY_CONTRACTOR_HEIGHT,
            JSON_KEY_CONTRACTOR_ITEMS,
            JSON_KEY_CONTRACTOR_ADDITIONAL_INFO,
        ):
            assert field in contractor_1, field

        items = contractor_1[JSON_KEY_CONTRACTOR_ITEMS]
        assert JSON_KEY_CONTRACTOR_POSITIONS in items
        assert JSON_KEY_CONTRACTOR_SUMMARY in items


class TestGetProposalsMultipleContractors:
    """Несколько подрядчиков и нумерация ключей."""

    def test_processes_multiple_contractors(self, empty_worksheet, sample_contractors_data):
        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        assert result["contractor_1"][JSON_KEY_CONTRACTOR_TITLE] == "ООО Строитель"
        assert result["contractor_2"][JSON_KEY_CONTRACTOR_TITLE] == "АО Подрядчик"

    def test_contractor_keys_increment_correctly(self, empty_worksheet):
        """Ширина блока каждого подрядчика — 11 (измеренная раскладка ГП);
        список приходит уже разрешённым, без маркера — обход
        `enumerate(contractors, start=1)` нумерует все три подряд."""
        contractors_data = [
            resolved(col, KEYS_GP_11, value=name, row_start=2, coordinate=f"{chr(64 + col)}2")
            for name, col in [("Подрядчик 1", 2), ("Подрядчик 2", 4), ("Подрядчик 3", 6)]
        ]

        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=contractors_data
            ).proposals

        assert sorted(result) == ["contractor_1", "contractor_2", "contractor_3"]

    def test_gp_estimate_yields_exactly_one_proposal(self, empty_worksheet):
        """Смета ГП: один подрядчик — ровно одно предложение (AGENTS.md §4).

        Отдельной ветки «один подрядчик» в коде нет и не нужно: слой proposals
        сохранён как задел на возврат тендеров, а число предложений определяется
        длиной уже разрешённого списка `contractors` — маркер «Наименование
        контрагента» в нём не участвует, его снимает `estimate.py` раньше.
        """
        contractors_data = [resolved(10, KEYS_GP_11, value='ООО "ТЕСТПОДРЯД"', row_start=6, coordinate="J6")]

        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 11, 2600, header_row=COLUMN_HEADER_ROW, contractors=contractors_data
            ).proposals

        assert list(result) == ["contractor_1"]
        assert result["contractor_1"][JSON_KEY_CONTRACTOR_WIDTH] == 11


class TestGetProposalsVatRate:
    """Ключ vat_rate (спека Ф4б §2.1): контракт `get_proposals`, не полный путь."""

    def test_vat_rate_key_is_always_present_even_as_none(self, empty_worksheet, sample_contractors_data):
        """Синтетический лист пуст: обе метки шапки молчат, ставка не получена —
        но ключ обязан лечь в предложение, а не отсутствовать (спека §2.1)."""
        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        assert JSON_KEY_VAT_RATE in contractor_1
        assert contractor_1[JSON_KEY_VAT_RATE] is None

    def test_vat_rate_sits_next_to_width_not_inside_items(self, empty_worksheet, sample_contractors_data):
        """Ставка — свойство блока подрядчика, как contractor_width, а не
        содержимое contractor_items (спека §2.1)."""
        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        assert JSON_KEY_VAT_RATE not in contractor_1[JSON_KEY_CONTRACTOR_ITEMS]


class TestGetProposalsEdgeCases:
    """Граничные случаи."""

    def test_handles_missing_contractor_fields(self, empty_worksheet):
        """Геометрия без `row_start` — ИНН/адрес/аккредитация не читаются
        (`rowspan == 1 and contractor_row_start is not None` не выполняется),
        но предложение всё равно собирается: `column_start` присутствует —
        он обязателен у любого `ResolvedContractor` (`resolve_contractor`
        отказал бы раньше, не будь его)."""
        incomplete_contractors = [
            resolved(2, KEYS_GP_11, value="Неполные данные", coordinate="B2", merged_shape={"rowspan": 1, "colspan": 1}),
        ]

        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=incomplete_contractors
            ).proposals

        contractor_1 = result["contractor_1"]
        assert contractor_1[JSON_KEY_CONTRACTOR_TITLE] == "Неполные данные"
        assert contractor_1[JSON_KEY_CONTRACTOR_COORDINATE] == "B2"
        assert contractor_1[JSON_KEY_CONTRACTOR_WIDTH] == 1
        assert contractor_1[JSON_KEY_CONTRACTOR_HEIGHT] == 1
        assert contractor_1[JSON_KEY_CONTRACTOR_INN] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ADDRESS] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ACCREDITATION] is None

    def test_handles_none_values_in_cells(self, empty_worksheet, sample_contractors_data):
        p1, p2, p3 = _patch_collaborators()
        with p1, p2, p3:
            result = get_proposals(
                empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data
            ).proposals

        contractor_1 = result["contractor_1"]
        assert contractor_1[JSON_KEY_CONTRACTOR_INN] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ADDRESS] is None
        assert contractor_1[JSON_KEY_CONTRACTOR_ACCREDITATION] is None

    def test_propagates_exceptions_from_collaborators(self, empty_worksheet, sample_contractors_data):
        """Ошибка разбора не глотается: импорт фазы 4 должен упасть в error."""
        with (
            patch(f"{MODULE}.get_lot_positions", side_effect=RuntimeError("Ошибка позиций")),
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])),
            patch(f"{MODULE}.get_additional_info", return_value={}),
            pytest.raises(RuntimeError, match="Ошибка позиций"),
        ):
            get_proposals(empty_worksheet, 10, 20, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)


class TestGetProposalsLotBoundaries:
    """Передача границ лота вниз по стеку."""

    def test_passes_lot_boundaries_to_positions_function(self, empty_worksheet, sample_contractors_data):
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})) as mock_positions,
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])),
            patch(f"{MODULE}.get_additional_info", return_value={}),
        ):
            get_proposals(empty_worksheet, 100, 200, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        for call_args in mock_positions.call_args_list:
            assert call_args[1]["lot_start_row"] == 100
            assert call_args[1]["lot_end_row"] == 200

    def test_summary_gets_start_row_but_not_end_row(self, empty_worksheet, sample_contractors_data):
        """Итоги ищутся от начала позиций и не ограничены концом лота.

        Начало нужно, чтобы поиск не наткнулся на объединённую ячейку шапки
        таблицы; конец не нужен — блок итогов лежит ниже всех позиций.
        """
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})),
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])) as mock_summary,
            patch(f"{MODULE}.get_additional_info", return_value={}),
        ):
            get_proposals(empty_worksheet, 100, 200, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        for call_args in mock_summary.call_args_list:
            assert call_args.kwargs == {"search_start_row": 100}

    def test_additional_info_does_not_depend_on_lot_boundaries(self, empty_worksheet, sample_contractors_data):
        with (
            patch(f"{MODULE}.get_lot_positions", return_value=LotRows(positions={})),
            patch(f"{MODULE}.get_summary", return_value=SummaryBlock(lines={}, warnings=[])),
            patch(f"{MODULE}.get_additional_info", return_value={}) as mock_additional,
        ):
            get_proposals(empty_worksheet, 100, 200, header_row=COLUMN_HEADER_ROW, contractors=sample_contractors_data)

        for call_args in mock_additional.call_args_list:
            assert len(call_args.args) == 2
            assert call_args.kwargs == {}
