"""Тесты чтения позиций в границах лота.

Перенос `app/tests/excel_parser/test_get_lot_positions.py` из
`parser_tender_xlsx@0e178c0`.

Изменения при переносе:

* убран тест на несуществующий `app/tests/test_data/sample_tender.xlsx`
  (в исходнике он всегда пропускался); реальный файл проверяется отдельно —
  `test_estimate.py`;
* добавлен тест `test_starts_exactly_at_lot_start_row`: исходник брал
  `max(START_INDEXING_POSITION_ROW=13, lot_start_row)` и на смете ГП, где лот
  начинается со строки 11, молча терял две первые строки.
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from parser.constants import (
    JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW,
    JSON_KEY_JOB_TITLE,
    JSON_KEY_JOB_TITLE_NORMALIZED,
    JSON_KEY_NUMBER,
    JSON_KEY_QUANTITY,
    JSON_KEY_UNIT,
    TABLE_PARSE_ADDITIONAL_WORKS_TITLE,
)
from parser.errors import EstimateParseError
from parser.get_lot_positions import get_lot_positions
from parser.resolve_contractor import ResolvedContractor

from .sheet_builders import KEYS_8, resolved

CONTRACTOR = resolved(9, KEYS_8)


@pytest.fixture
def sample_worksheet():
    """Лист с тремя позициями в строках 13–15."""
    ws = Workbook().active

    headers = [
        "№ п/п",
        "Глава",
        "Артикул СМР",
        "Наименование видов работ",
        "Пропуск",
        "Комментарий",
        "Ед. изм.",
        "Кол-во",
        "Подрядчик 1 - Цена",
        "Подрядчик 1 - Стоимость",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    test_data = [
        ["1", "01", "01-01-003", "Земляные работы", "", "Основные работы", "м³", 100, 500, 50000],
        ["2", "02", "02-01-015", "Кирпичная кладка", "", "Каменные работы", "м³", 50, 1200, 60000],
        ["3", "03", "03-02-008", "Штукатурные работы", "", "Отделочные работы", "м²", 200, 300, 60000],
    ]
    for row_idx, row_data in enumerate(test_data, 13):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    return ws


class TestGetLotPositionsBehavior:
    """Основное поведение."""

    def test_extracts_positions_from_sample_data(self, sample_worksheet):
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=15).positions

        assert len(result) == 3

        for position_key in ("1", "2", "3"):
            position = result[position_key]
            for field in (JSON_KEY_NUMBER, JSON_KEY_JOB_TITLE, JSON_KEY_UNIT, JSON_KEY_QUANTITY):
                assert field in position, field
            if position.get(JSON_KEY_JOB_TITLE):
                assert JSON_KEY_JOB_TITLE_NORMALIZED in position

    def test_handles_empty_range_gracefully(self, sample_worksheet):
        """lot_start_row > lot_end_row — пустой результат, не падение."""
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=20, lot_end_row=15).positions

        assert result == {}

    def test_handles_empty_rows_correctly(self, sample_worksheet):
        """Пустые строки после данных пропускаются, а не обрывают обход."""
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=30).positions

        assert sorted(result) == ["1", "2", "3"]

    def test_respects_lot_boundaries(self, sample_worksheet):
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=14).positions

        assert sorted(result) == ["1", "2"]

        first_position = result["1"]
        assert first_position[JSON_KEY_JOB_TITLE] == "Земляные работы"
        assert first_position[JSON_KEY_UNIT] == "м³"
        assert first_position[JSON_KEY_QUANTITY] == 100

    def test_starts_exactly_at_lot_start_row(self):
        """Обход начинается с границы лота, без жёсткого нижнего порога.

        В смете ГП лот начинается со строки 11. Исходная константа
        START_INDEXING_POSITION_ROW = 13 съедала строки 11 и 12 — на реальном
        файле это были строка-раздел лота и первый раздел сметы.
        """
        ws = Workbook().active
        ws.cell(row=11, column=1, value="1")
        ws.cell(row=11, column=4, value="Лот №1 - Тестовый объект")
        ws.cell(row=12, column=1, value="1")
        ws.cell(row=12, column=4, value="Подготовительные работы")
        ws.cell(row=13, column=1, value="2")
        ws.cell(row=13, column=4, value="Обеспечение финансовых условий")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=11, lot_end_row=13).positions

        assert len(result) == 3
        assert result["1"][JSON_KEY_JOB_TITLE] == "Лот №1 - Тестовый объект"
        assert result["2"][JSON_KEY_JOB_TITLE] == "Подготовительные работы"
        assert result["3"][JSON_KEY_JOB_TITLE] == "Обеспечение финансовых условий"

    def test_processes_contractor_data_correctly(self, sample_worksheet):
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=13).positions

        position = result["1"]
        assert isinstance(position, dict)
        assert len(position) > 4  # больше, чем только общие поля

    def test_normalizes_job_titles(self, sample_worksheet):
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=15).positions

        for position in result.values():
            if position.get(JSON_KEY_JOB_TITLE):
                assert isinstance(position[JSON_KEY_JOB_TITLE_NORMALIZED], str)


class TestGetLotPositionsEdgeCases:
    """Граничные случаи."""

    def test_invalid_contractor_structure(self, sample_worksheet):
        """Без column_start разбор строки подрядчика обязан упасть, а не молчать."""
        invalid_contractor = ResolvedContractor(
            geometry={"merged_shape": {"colspan": 8}}, layout=CONTRACTOR.layout
        )

        with pytest.raises((KeyError, AttributeError, TypeError)):
            get_lot_positions(sample_worksheet, invalid_contractor, lot_start_row=13, lot_end_row=15)

    def test_extreme_row_ranges(self, sample_worksheet):
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=5000).positions

        assert len(result) == 3


class TestGetLotPositionsDataIntegrity:
    """Качество извлекаемых данных."""

    def test_preserves_data_types(self, sample_worksheet):
        position = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=13).positions["1"]

        assert isinstance(position[JSON_KEY_QUANTITY], int | float)
        assert isinstance(position[JSON_KEY_JOB_TITLE], str)
        assert isinstance(position[JSON_KEY_UNIT], str)

    def test_handles_missing_data_gracefully(self, sample_worksheet):
        ws = sample_worksheet
        ws.cell(row=16, column=1, value="4")
        ws.cell(row=16, column=4, value="Тест работа")
        # остальные ячейки строки пустые

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=16, lot_end_row=16).positions

        assert result["1"][JSON_KEY_JOB_TITLE] == "Тест работа"
        assert result["1"][JSON_KEY_UNIT] is None

    def test_stops_at_merged_cell_in_first_column(self, sample_worksheet):
        """Объединённая ячейка в колонке A — конец блока позиций.

        AGENTS.md §11: досрочный выход по merged-ячейке ломать нельзя. Без него
        парсер поехал бы в блок итогов и дополнительной информации, где в
        колонке подрядчика лежит текст, а не числа.
        """
        ws = sample_worksheet

        ws.cell(row=17, column=1, value="5")
        ws.cell(row=17, column=4, value="Обычная позиция")

        ws.merge_cells("A18:A19")
        ws.cell(row=18, column=1, value="ИТОГО")

        ws.cell(row=20, column=1, value="6")
        ws.cell(row=20, column=4, value="Не должна обрабатываться")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=17, lot_end_row=21).positions

        assert len(result) == 1
        assert result["1"][JSON_KEY_JOB_TITLE] == "Обычная позиция"

    def test_skips_completely_empty_rows(self, sample_worksheet):
        ws = sample_worksheet

        ws.cell(row=22, column=1, value="7")
        ws.cell(row=22, column=4, value="Первая позиция")
        # строка 23 пустая
        ws.cell(row=24, column=1, value="8")
        ws.cell(row=24, column=4, value="Вторая позиция")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=22, lot_end_row=24).positions

        assert len(result) == 2
        assert result["1"][JSON_KEY_JOB_TITLE] == "Первая позиция"
        assert result["2"][JSON_KEY_JOB_TITLE] == "Вторая позиция"

    def test_row_with_data_only_right_of_contractor_block_counts_as_empty(self):
        """Проверка пустоты ограничена блоком подрядчика.

        Это цена отказа от разворачивания `ws[row]` до `ws.max_column`: строка,
        где данные лежат правее блока подрядчика, считается пустой. Для сметы
        это верно — читать там нечего.
        """
        ws = Workbook().active
        ws.cell(row=13, column=1, value="1")
        ws.cell(row=13, column=4, value="Позиция")
        ws.cell(row=14, column=40, value="что-то далеко справа")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=14).positions

        assert len(result) == 1


def _floats_anywhere(value, path=""):
    """Пути до всех float внутри вложенной структуры. Пусто — значит их нет."""
    if isinstance(value, float):
        return [path or "<root>"]
    if isinstance(value, dict):
        found = []
        for key, nested in value.items():
            found.extend(_floats_anywhere(nested, f"{path}.{key}" if path else str(key)))
        return found
    return []


class TestAdditionalWorksRow:
    """Агрегатная строка «Дополнительные работы» (спека Ф2 §2.2).

    Признак — пустые A и B плюс точное название в D. Замер: строк с пустыми
    A и B во всех трёх реальных офертах ровно одна, ложных нет; наивное
    «D содержит „дополнительн“» дало бы 3–5 попаданий на файл, почти все —
    настоящие позиции.
    """

    def test_recognized_row_goes_to_additional_works(self, sample_worksheet):
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Дополнительные работы")
        ws.cell(row=16, column=10, value=12675964.53)

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        assert result.additional_works is not None
        assert result.additional_works["job_title"] == "Дополнительные работы"
        assert result.additional_works["source_row"] == 16

    def test_money_is_a_decimal_string_not_float(self, sample_worksheet):
        """Деньги идут через parse_contractor_row: строка, не float (AGENTS.md §3).

        Проверка адресная, а не «нет float среди values()»: деньги лежат ВЛОЖЕННО
        в `unit_cost` и `total_cost`, поэтому обход верхнего уровня их не видит и
        прошёл бы даже при float внутри. Раскладка замерена: у `CONTRACTOR`
        `column_start = 9` и `colspan = 8`, а ключи colspan-8 начинаются с
        `unit_cost.materials`, значит колонка 10 — это `unit_cost.works`.
        `money_to_json(12675964.53)` даёт ровно `"12675964.53"` (замерено).
        """
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Дополнительные работы")
        ws.cell(row=16, column=10, value=12675964.53)

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        value = result.additional_works["unit_cost"]["works"]
        assert value == "12675964.53"
        assert isinstance(value, str)
        assert _floats_anywhere(result.additional_works) == []

    def test_aggregate_row_is_not_a_position(self, sample_worksheet):
        """Ф4 (спека §2.1): строка не появляется ни под одним ключом `positions`.

        Дубль не снимается удалением — он не создаётся вовсе: `continue` стоит
        перед записью в `positions`, а `additional_works` собирается раньше и не
        затрагивается.
        """
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Дополнительные работы")
        ws.cell(row=16, column=10, value=12675964.53)

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        titles = [item[JSON_KEY_JOB_TITLE] for item in result.positions.values()]
        assert "Дополнительные работы" not in titles
        assert result.additional_works is not None

    def test_aggregate_row_in_the_middle_keeps_keys_contiguous(self, sample_worksheet):
        """Главный тест задачи (спека §2.1): ключи `positions` остаются `1..N`.

        В реальном файле агрегатная строка последняя, поэтому её ключ равен
        `len(positions)` и о непрерывности не говорит ничего. Здесь синтетический
        лист: после агрегатной строки идут ещё две обычные позиции. Резолвер Ф3
        отвергает неканоничный набор ключей — дыра в нумерации уронила бы весь
        импорт, а не испортила бы одну строку, поэтому свойство проверяется явно,
        а не выводится из механизма (`item_index` инкрементируется только вместе
        с записью в словарь, и это делает тест наблюдением за поведением, а не
        пересказом кода).
        """
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Дополнительные работы")
        ws.cell(row=16, column=10, value=12675964.53)
        ws.cell(row=17, column=1, value="4")
        ws.cell(row=17, column=4, value="Позиция сразу после агрегатной строки")
        ws.cell(row=18, column=1, value="5")
        ws.cell(row=18, column=4, value="Ещё одна позиция")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=18).positions

        n = len(result)
        assert set(result) == {str(i) for i in range(1, n + 1)}
        # Если бы агрегатная строка осталась в positions (Ф2), эта позиция
        # получила бы ключ "5"; без неё — ключом на единицу меньше, "4".
        assert result["4"][JSON_KEY_JOB_TITLE] == "Позиция сразу после агрегатной строки"

    def test_aggregate_row_lives_only_in_additional_works(self, sample_worksheet):
        """Контракт Ф4 (спека §2.1): строка живёт только в `additional_works`.

        Тест Ф2 `test_row_also_stays_in_positions` охранял переходное решение —
        строка ВРЕМЕННО остаётся ещё и в `positions` — и спека Ф2 §2.4 прямо
        поручила Ф4 перевернуть его на обратный контракт, когда исключение будет
        сделано.

        Усиление Ф2 не теряется, а меняет сторону сравнения. Тогда сверялось
        КАЖДОЕ поле копии в `positions` с копией в `additional_works`: обе
        стороны приходили из одного вызова `parse_contractor_row`, то есть
        сверка держалась на том, что копия не собрана из чужой строки. Копии
        больше нет — сверять не с чем, поэтому эталон здесь **независимый**: все
        восемь денежных колонок блока подрядчика заполнены РАЗНЫМИ значениями, и
        ожидание выписано литералом. Перепутанные местами колонки, потерянное
        поле и лишнее поле дают красный по отдельности, а не сливаются в одно
        «что-то не так».
        """
        ws = sample_worksheet
        ws.cell(row=16, column=4, value=TABLE_PARSE_ADDITIONAL_WORKS_TITLE)
        # Колонки 9–16 — весь блок подрядчика при colspan 8 (порядок задан
        # `parse_contractor_row.get_column_keys`). Значения различны: одинаковые
        # пропустили бы перестановку колонок молча.
        for column, value in enumerate([11.11, 22.22, 33.33, 44.44, 55.55, 66.66, 77.77, 88.88], 9):
            ws.cell(row=16, column=column, value=value)

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        copies = [
            item
            for item in result.positions.values()
            if item[JSON_KEY_JOB_TITLE] == TABLE_PARSE_ADDITIONAL_WORKS_TITLE
        ]
        assert copies == [], "агрегатная строка не должна попадать в positions ни разу"

        work = result.additional_works
        # Состав ключей целиком: пропажа поля и лишнее поле видны так же, как
        # неверное значение (тот же довод, что у DB_CHECKS/ORM_CHECKS в схеме).
        assert set(work) == {
            JSON_KEY_JOB_TITLE,
            JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW,
            "unit_cost",
            "total_cost",
        }
        assert work[JSON_KEY_JOB_TITLE] == TABLE_PARSE_ADDITIONAL_WORKS_TITLE
        assert work[JSON_KEY_ADDITIONAL_WORKS_SOURCE_ROW] == 16
        assert work["unit_cost"] == {
            "materials": "11.11",
            "works": "22.22",
            "indirect_costs": "33.33",
            "total": "44.44",
        }
        assert work["total_cost"] == {
            "materials": "55.55",
            "works": "66.66",
            "indirect_costs": "77.77",
            "total": "88.88",
        }
        assert _floats_anywhere(work) == [], "деньги обязаны остаться Decimal-строкой (AGENTS.md §3)"

    def test_absent_row_is_valid(self, sample_worksheet):
        """42-ТУ и 449-ТУ: строки нет вовсе, это не ошибка и не warning."""
        result = get_lot_positions(sample_worksheet, CONTRACTOR, lot_start_row=13, lot_end_row=15)

        assert result.additional_works is None
        assert result.positions != {}

    def test_candidate_with_other_title_is_rejected(self, sample_worksheet):
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Прочие затраты")

        with pytest.raises(EstimateParseError, match="Прочие затраты"):
            get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

    def test_second_candidate_is_rejected(self, sample_worksheet):
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="Дополнительные работы")
        ws.cell(row=17, column=4, value="Дополнительные работы")

        with pytest.raises(EstimateParseError, match="16"):
            get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=17)

    def test_title_is_compared_normalized(self, sample_worksheet):
        """Регистр и лишние пробелы в названии не мешают распознаванию.

        Без этого теста реализация с простым `==` тоже была бы зелёной, а спека
        §2.2 требует сверки нормализованного названия.
        """
        ws = sample_worksheet
        ws.cell(row=16, column=4, value="  ДОПОЛНИТЕЛЬНЫЕ   РАБОТЫ ")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        assert result.additional_works is not None
        assert result.additional_works["job_title"] == "  ДОПОЛНИТЕЛЬНЫЕ   РАБОТЫ "

    def test_blank_is_by_text_not_by_none(self, sample_worksheet):
        """Пробел и неразрывный пробел в A/B — тоже пустота (спека §2.2)."""
        ws = sample_worksheet
        ws.cell(row=16, column=1, value=" ")
        ws.cell(row=16, column=2, value=" ")
        ws.cell(row=16, column=4, value="Дополнительные работы")

        result = get_lot_positions(ws, CONTRACTOR, lot_start_row=13, lot_end_row=16)

        assert result.additional_works is not None
