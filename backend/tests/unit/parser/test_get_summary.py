"""Обход блока итогов по листу. Правила проверяются в test_summary_block.py."""
from __future__ import annotations

from openpyxl import Workbook

from parser.get_summary import get_summary

from .sheet_builders import KEYS_GP_11, resolved

# column_start=10, colspan=11 — те же координаты, что и раньше: литералы этого
# файла (TOTAL_COST_TOTAL_COLUMN=18 и все "20"/"120"/... ниже) сверены под
# КОНКРЕТНО эту раскладку, а не под resolved(9, KEYS_8) (та подошла бы
# test_get_lot_positions.py, но не этому файлу — здесь другая ширина и
# другой column_start).
CONTRACTOR = resolved(10, KEYS_GP_11)


# Блок подрядчика J..T (colspan 11) раскладывается так — замерено вызовом
# parse_contractor_row на листе, где в каждой ячейке лежит её номер колонки:
#   10 suggested_quantity | 11..14 unit_cost.{mat,wrk,ind,total}
#   15..18 total_cost.{mat,wrk,ind,total} | 19 организатор | 20 комментарий
# Итоговая стоимость — колонка 18, НЕ 17 (17 это indirect_costs).
TOTAL_COST_TOTAL_COLUMN = 18


def _sheet_with_summary(rows: list[tuple[str, str | None]]):
    """Лист, где с 20-й строки идёт блок итогов, а ниже — «Дополнительная информация».

    Args:
        rows: пары (метка колонки A, значение колонки `total_cost.total`);
            None означает строку без сумм.
    """
    ws = Workbook().active
    ws["A11"] = "первая строка позиций"
    for offset, (label, total) in enumerate(rows):
        row = 20 + offset
        ws.cell(row=row, column=1, value=label)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        if total is not None:
            ws.cell(row=row, column=TOTAL_COST_TOTAL_COLUMN, value=float(total))
    ws.cell(row=20 + len(rows) + 1, column=1, value="Дополнительная информация:")
    return ws


def test_total_is_read_from_the_column_the_helper_writes():
    """Предпосылка самого хелпера: колонка 18 — это `total_cost.total`.

    Проверяется внутри теста, а не «известна»: перепутанная колонка оставила бы
    все тесты блока зелёными, сверяя пустоту с пустотой
    (false-test-premises.md).
    """
    ws = _sheet_with_summary([("ИТОГО, руб. с учетом НДС", "120")])
    block = get_summary(ws, CONTRACTOR, search_start_row=11)
    line = block.lines["total_cost_including_vat"]
    assert line["total_cost"]["total"] == "120.0"


def test_walks_from_the_first_merged_row_to_the_first_empty_one():
    ws = _sheet_with_summary(
        [("ИТОГО, руб. с учетом НДС", "120"), ("В том числе НДС", "20"),
         ("ИТОГО, руб. без учета НДС", "100")]
    )
    block = get_summary(ws, CONTRACTOR, search_start_row=11)
    assert len(block.lines) == 3


def test_empty_row_stops_the_walk_before_additional_info():
    """Терминатор блока — пустая строка. Если её нет, обход прочитает
    «Дополнительную информацию» как ещё одну итоговую строку."""
    ws = _sheet_with_summary([("ИТОГО, руб. с учетом НДС", "120"), ("В том числе НДС", "20")])
    block = get_summary(ws, CONTRACTOR, search_start_row=11)
    assert "Дополнительная информация:" not in [
        line["job_title"] for line in block.lines.values()
    ]


def test_no_merged_row_means_block_not_found():
    ws = Workbook().active
    ws["A11"] = "первая строка позиций"
    block = get_summary(ws, CONTRACTOR, search_start_row=11)
    assert block.lines == {}
    assert any("не найден" in text for text in block.warnings)
