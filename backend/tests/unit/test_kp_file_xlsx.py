"""Книга xlsx файла подрядчику и общие подписи листов (спека 4б §2.8).

Книга читается обратно через `openpyxl.load_workbook` и сверяется со входом:
эталоны — литералы, а не значения, взятые у проверяемого кода.
"""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from services import excel, excel_changes_export, excel_comparison
from services.excel import FMT_MONEY, FMT_PCT, FMT_QTY, safe_str
from services.excel_kp_file import KP_FILE_COLUMNS, build_kp_file
from services.kp_file import FileHeader, FileLine, FileSection

YEARS = [
    {
        "year": 2024,
        "coefficient": Decimal("1.083"),
        "source": "Источник А",
        "is_forecast": False,
        "updated_at": "2026-01-12T10:00:00+00:00",
    },
    {
        "year": 2026,
        "coefficient": Decimal("1.04"),
        "source": "Источник Б",
        "is_forecast": True,
        "updated_at": None,
    },
]


def header(**over) -> FileHeader:
    base = FileHeader(
        tender="Тендер Т",
        object="Объект Т",
        round="Этап 2 от 05.03.2026",
        contractor="Подрядчик А",
        exported_on="28.10.2026",
        vat_label="Цены КП и наши расценки — с НДС 12 % (база НДС этого КП)",
        source_label="Цена источника — как записана в договоре, без НДС",
        level_label="номинальный, без приведения",
        in_progress_lines=0,
        inflation_years=None,
        inflation_series=None,
        inflation_target_month=None,
        filename="файл.xlsx",
    )
    return replace(base, **over)


def with_inflation(years=YEARS, **over) -> FileHeader:
    return header(
        inflation_years=years,
        inflation_series="Ряд А",
        inflation_target_month="2026-03",
        level_label="приведено по ИПЦ: Ряд А, к 03.2026; КП не приводится",
        **over,
    )


def line(**over) -> FileLine:
    base = FileLine(
        position_item_id=1,
        number="1.1",
        title="Устройство стяжки",
        unit="м2",
        quantity=Decimal("12.5"),
        kp_unit_price=Decimal("1500.5"),
        kp_total=Decimal("18756.25"),
        ours=Decimal("1400.25"),
        deviation_pct=Decimal("-6.7"),
        source="Объект Т, договор № 7, 2024 г.: 1 250 ₽ за м2 без НДС",
        inflation="Ряд А, к 03.2026: 1 300 ₽ за м2 без НДС",
        comment="Ниже цены КП",
        comment_kind="template",
        state="comparable",
    )
    return replace(base, **over)


def sheet_of(data: bytes):
    wb = load_workbook(BytesIO(data))
    assert len(wb.worksheets) == 1
    return wb.worksheets[0]


def column_header_row(ws) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "№":
            return row
    raise AssertionError("нет строки заголовков")


def col1_texts(ws) -> list:
    return [ws.cell(row=row, column=1).value for row in range(1, ws.max_row + 1)]


def labels(ws) -> list[str]:
    row = column_header_row(ws)
    return [ws.cell(row=row, column=c).value for c in range(1, ws.max_column + 1)]


# --- значения ячеек равны полям строк ---------------------------------------------------


def test_line_cells_equal_fields():
    ws = sheet_of(build_kp_file(header(), [line()]))
    first = column_header_row(ws) + 1
    values = [ws.cell(row=first, column=c).value for c in range(1, 11)]
    assert values == [
        "1.1",
        "Устройство стяжки",
        "м2",
        12.5,
        1500.5,
        18756.25,
        1400.25,
        -6.7,
        "Объект Т, договор № 7, 2024 г.: 1 250 ₽ за м2 без НДС",
        "Ниже цены КП",  # без ряда столбца ИПЦ нет, комментарий сдвигается влево
    ]
    assert ws.max_column == 10


def test_line_cells_with_inflation_column_and_comment():
    ws = sheet_of(build_kp_file(with_inflation(), [line()]))
    first = column_header_row(ws) + 1
    assert ws.cell(row=first, column=10).value == "Ряд А, к 03.2026: 1 300 ₽ за м2 без НДС"
    assert ws.cell(row=first, column=11).value == "Ниже цены КП"


def test_money_cells_are_numbers_with_formats_and_no_rounding():
    ours = Decimal("1400.123456789012345678")
    ws = sheet_of(build_kp_file(header(), [line(ours=ours, quantity=Decimal("0.0005"))]))
    row = column_header_row(ws) + 1
    for col in (5, 6, 7):
        assert isinstance(ws.cell(row=row, column=col).value, int | float)
    assert ws.cell(row=row, column=7).value == pytest.approx(1400.123456789012, abs=1e-9)
    # Объём — тоже значение поля без округления: 0,0005 не становится нулём.
    assert ws.cell(row=row, column=4).value == pytest.approx(0.0005, abs=1e-12)
    assert ws.cell(row=row, column=7).number_format == FMT_MONEY
    assert ws.cell(row=row, column=5).number_format == FMT_MONEY
    assert ws.cell(row=row, column=6).number_format == FMT_MONEY
    assert ws.cell(row=row, column=4).number_format == FMT_QTY
    assert ws.cell(row=row, column=8).number_format == FMT_PCT


def test_empty_fields_are_empty_cells():
    empty = line(
        number=None,
        quantity=None,
        kp_unit_price=None,
        kp_total=None,
        ours=None,
        deviation_pct=None,
        source=None,
        inflation=None,
        comment=None,
        comment_kind=None,
        state="no_comparison",
    )
    ws = sheet_of(build_kp_file(with_inflation(), [empty]))
    row = column_header_row(ws) + 1
    assert [ws.cell(row=row, column=c).value for c in (1, 4, 5, 6, 7, 8, 9, 10, 11)] == [None] * 9
    assert ws.cell(row=row, column=2).value == "Устройство стяжки"


def test_negative_and_positive_deviation_keep_sign():
    ws = sheet_of(
        build_kp_file(
            header(),
            [line(deviation_pct=Decimal("12.3")), line(position_item_id=2, deviation_pct=Decimal("-0.5"))],
        )
    )
    row = column_header_row(ws) + 1
    assert ws.cell(row=row, column=8).value == 12.3
    assert ws.cell(row=row + 1, column=8).value == -0.5


def test_rows_keep_order_and_sections_are_rows():
    rows = [
        FileSection(number="1", title="Глава А"),
        line(position_item_id=1, title="Позиция 1"),
        line(position_item_id=2, title="Позиция 2"),
        FileSection(number="2", title="Глава Б"),
        line(position_item_id=3, title="Позиция 3"),
    ]
    ws = sheet_of(build_kp_file(header(), rows))
    first = column_header_row(ws) + 1
    assert [ws.cell(row=first + i, column=2).value for i in range(5)] == [
        "Глава А",
        "Позиция 1",
        "Позиция 2",
        "Глава Б",
        "Позиция 3",
    ]
    assert ws.cell(row=first, column=1).value == "1"
    assert ws.max_row == first + 4


# --- шапка листа ------------------------------------------------------------------------


def test_header_texts_on_the_sheet():
    ws = sheet_of(build_kp_file(header(), [line()]))
    texts = [t for t in col1_texts(ws) if isinstance(t, str)]
    joined = "\n".join(texts)
    assert "Тендер Т" in texts[0] and "Объект Т" in texts[0]
    assert "Этап 2 от 05.03.2026" in joined
    assert "Подрядчик А" in joined
    assert "28.10.2026" in joined
    assert "Цены КП и наши расценки — с НДС 12 % (база НДС этого КП)" in texts
    assert "Цена источника — как записана в договоре, без НДС" in texts
    assert "Ценовой уровень: номинальный, без приведения" in texts


def test_header_level_label_with_inflation():
    ws = sheet_of(build_kp_file(with_inflation(), [line()]))
    assert "Ценовой уровень: приведено по ИПЦ: Ряд А, к 03.2026; КП не приводится" in col1_texts(ws)


def test_stage_caption_is_the_round_text_without_extra_label():
    ws = sheet_of(build_kp_file(header(), [line()]))
    texts = col1_texts(ws)
    assert "Этап 2 от 05.03.2026" in texts
    assert not any(isinstance(t, str) and t.startswith("Этап:") for t in texts)
    assert "Участник: Подрядчик А" in texts
    assert "Дата выгрузки: 28.10.2026" in texts


def test_header_texts_after_a_constant_prefix_have_no_stray_apostrophe():
    ws = sheet_of(build_kp_file(header(tender="=СУММ(A1)", object="-О", contractor="-Подрядчик"), [line()]))
    texts = col1_texts(ws)
    assert texts[0] == "Тендер: =СУММ(A1) · Объект: -О"
    assert "Участник: -Подрядчик" in texts
    assert not any(isinstance(t, str) and "'" in t for t in texts)
    assert ws.cell(row=1, column=1).data_type == "s"


def test_in_progress_count_is_not_printed():
    ws = sheet_of(build_kp_file(header(in_progress_lines=4242), [line()]))
    for row in ws.iter_rows():
        for cell in row:
            assert "4242" not in str(cell.value)


# --- блок приведения по ИПЦ --------------------------------------------------------------


def test_inflation_block_lists_each_year():
    ws = sheet_of(build_kp_file(with_inflation(), [line()]))
    texts = col1_texts(ws)
    lead = "Цены договоров приведены к: март 2026 · ряд: «Ряд А»"
    at = texts.index(lead)
    assert texts[at + 1] == "2024 · 1.083 · Источник А · факт · правлен 12.01.2026"
    assert texts[at + 2] == "2026 · 1.04 · Источник Б · прогноз"
    assert texts[at + 3] is None  # дальше пустая строка до таблицы


def test_inflation_block_with_no_years_says_so():
    ws = sheet_of(build_kp_file(with_inflation(years=[]), [line()]))
    texts = col1_texts(ws)
    at = texts.index("Цены договоров приведены к: март 2026 · ряд: «Ряд А»")
    assert texts[at + 1] == "Коэффициенты за годы не потребовались: цены уже в целевом уровне."


def test_no_inflation_block_without_series():
    ws = sheet_of(build_kp_file(header(), [line()]))
    joined = "\n".join(t for t in col1_texts(ws) if isinstance(t, str))
    assert "приведены к" not in joined
    assert "Коэффициенты за годы" not in joined


def test_inflation_column_only_with_series():
    plain = labels(sheet_of(build_kp_file(header(), [line()])))
    series = labels(sheet_of(build_kp_file(with_inflation(), [line()])))
    empty_years = labels(sheet_of(build_kp_file(with_inflation(years=[]), [line()])))
    assert "Приведение по ИПЦ" not in plain
    assert len(plain) == 10
    assert series[9] == "Приведение по ИПЦ"
    assert len(series) == 11
    assert empty_years == series
    assert series[:9] == [
        "№",
        "Наименование",
        "Ед.",
        "Кол-во",
        "Цена за ед.",
        "Стоимость",
        "Наша расценка за ед.",
        "Отклонение, %",
        "Источник",
    ]
    assert series[10] == "Комментарий"
    assert plain[9] == "Комментарий"


def test_columns_spec_shape():
    assert len(KP_FILE_COLUMNS) == 11
    assert all(len(spec) == 4 for spec in KP_FILE_COLUMNS)
    assert [spec[0] for spec in KP_FILE_COLUMNS][9] == "Приведение по ИПЦ"


def test_sheet_layout_title_widths_wrap_and_text_number_column():
    ws = sheet_of(build_kp_file(with_inflation(), [line()]))
    assert ws.title == "Наши расценки"
    for index, (_label, width, _fmt, _h) in enumerate(KP_FILE_COLUMNS, start=1):
        assert ws.column_dimensions[get_column_letter(index)].width == width
    row = column_header_row(ws) + 1
    # Тексты переносятся, числа — нет: «Источник» и комментарий длиннее ширины столбца.
    wrapped = {c for c in range(1, 12) if ws.cell(row=row, column=c).alignment.wrap_text}
    assert wrapped == {2, 9, 10, 11}
    # «№» — текст: «1.10» не превращается в число 1,1 при правке ячейки.
    assert ws.cell(row=row, column=1).number_format == "@"


# --- безопасные строки ------------------------------------------------------------------


def test_formula_like_comment_keeps_one_apostrophe_and_equals_field():
    for raw in ("=СУММ(A1)", "-скидка"):
        cell_text = safe_str(raw)  # то, что кладёт в строку файла собирающий её код
        ws = sheet_of(build_kp_file(header(), [line(comment=cell_text)]))
        value = ws.cell(row=column_header_row(ws) + 1, column=10).value
        assert value == "'" + raw
        assert value == cell_text


# --- заливка и разделы ------------------------------------------------------------------


def fill_of(ws, row, col):
    cell = ws.cell(row=row, column=col)
    return cell.fill.fgColor.rgb if cell.fill.fill_type == "solid" else None


def test_our_columns_are_filled_and_theirs_are_not():
    ws = sheet_of(build_kp_file(with_inflation(), [line()]))
    row = column_header_row(ws) + 1
    ours = {fill_of(ws, row, c) for c in range(7, 12)}
    assert len(ours) == 1 and None not in ours
    # Белая сплошная заливка на белом листе — не выделение.
    assert ours.pop()[-6:].upper() != "FFFFFF"
    assert all(fill_of(ws, row, c) is None for c in range(1, 7))


def test_our_columns_are_filled_even_when_empty():
    ws = sheet_of(build_kp_file(header(), [line(ours=None, comment=None)]))
    row = column_header_row(ws) + 1
    assert fill_of(ws, row, 7) is not None
    assert fill_of(ws, row, 10) is not None


@pytest.mark.parametrize("hdr, width", [(header(), 10), (with_inflation(), 11)])
def test_section_row_is_bold_and_filled_across_sheet(hdr, width):
    ws = sheet_of(build_kp_file(hdr, [FileSection(number="3", title="Глава В"), line()]))
    row = column_header_row(ws) + 1
    assert ws.cell(row=row, column=1).value == "3"
    assert ws.cell(row=row, column=2).value == "Глава В"
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        assert cell.font.b is True
        assert fill_of(ws, row, col) is not None
    assert ws.cell(row=row, column=width + 1).fill.fill_type is None
    line_cell = ws.cell(row=row + 1, column=2)
    assert not line_cell.font.b
    assert fill_of(ws, row + 1, 2) is None


# --- подъём общих подписей --------------------------------------------------------------


def test_lifted_names_are_public_in_common_module_and_gone_from_old_modules():
    for name in ("TAX_BASIS_CAPTION", "PRICE_LEVEL_NOMINAL_CAPTION", "write_caption", "write_inflation_block"):
        assert hasattr(excel, name)
    for name in ("month_label", "iso_date_to_ru"):
        assert hasattr(excel, name)
    for module, names in (
        (excel_changes_export, ("_TAX_BASIS_CAPTION", "_PRICE_LEVEL_CAPTION", "_write_caption")),
        (excel_comparison, ("_write_inflation_block", "_month_label", "_iso_date_to_ru", "_MONTHS_NOMINATIVE")),
    ):
        for name in names:
            assert not hasattr(module, name)


def test_lifted_captions_literal_texts():
    assert excel.TAX_BASIS_CAPTION == {
        "gross": "Деньги: валовые, как в файлах (одна известная ставка НДС на всех этапах)",
        "net": "Деньги: нетто (ставки НДС этапов различаются)",
        "none": "Деньги: сумм нет — ни на одном этапе не известна база НДС",
    }
    assert excel.PRICE_LEVEL_NOMINAL == "номинальный, без приведения"
    assert excel.PRICE_LEVEL_NOMINAL_CAPTION == "Ценовой уровень: номинальный, без приведения"


def test_changes_sheet_captions_keep_order_and_wrap_after_lift():
    """Подъём не меняет порядок и вид подписей листа «Изменений КП»: состав,
    уровень, затем «Что двигалось» с переносом."""
    from tests.unit.test_changes_export_sheet import book_of, group_row, sheet_of, stages

    wb = book_of(sheet_of(stages(2, rate="20"), groups=[group_row(1), group_row(2)], title="Участник А"))
    ws = wb[wb.sheetnames[0]]
    texts = col1_texts(ws)
    at = texts.index("Деньги: валовые, как в файлах (одна известная ставка НДС на всех этапах)")
    assert texts[at + 1] == "Ценовой уровень: номинальный, без приведения"
    assert texts[at + 2].startswith("Что двигалось между этапами (по классификатору): ")
    assert [bool(ws.cell(row=at + 1 + i, column=1).alignment.wrap_text) for i in range(3)] == [False, False, True]
    assert all(ws.cell(row=at + 1 + i, column=1).font.b for i in range(3))


def test_month_label_and_date():
    assert excel.month_label("2026-01") == "январь 2026"
    assert excel.month_label("2026-12") == "декабрь 2026"
    assert excel.iso_date_to_ru("2026-01-12T10:00:00+00:00") == "12.01.2026"
    assert excel.iso_date_to_ru("2026-11-02") == "02.11.2026"


def blank_sheet():
    from openpyxl import Workbook

    ws = Workbook().active
    return ws


def inflation_dict():
    return {"target_month": "2026-03", "series_name": "Ряд А", "used_years": YEARS}


def test_inflation_block_default_lead_is_unchanged():
    ws = blank_sheet()
    end = excel.write_inflation_block(ws, 1, inflation_dict())
    assert ws.cell(row=1, column=1).value == "Цены приведены к: март 2026 · ряд: «Ряд А»"
    assert ws.cell(row=2, column=1).value == "2024 · 1.083 · Источник А · факт · правлен 12.01.2026"
    assert ws.cell(row=3, column=1).value == "2026 · 1.04 · Источник Б · прогноз"
    assert end == 4


def test_inflation_block_named_lead():
    ws = blank_sheet()
    excel.write_inflation_block(ws, 5, inflation_dict(), lead="Цены договоров приведены к:")
    assert ws.cell(row=5, column=1).value == "Цены договоров приведены к: март 2026 · ряд: «Ряд А»"


def test_write_caption_wrap_and_safe():
    ws = blank_sheet()
    assert excel.write_caption(ws, 3, "=1+1", wrap=True) == 4
    cell = ws.cell(row=3, column=1)
    assert cell.value == "'=1+1"
    assert cell.alignment.wrap_text is True
    assert cell.font.b is True
    excel.write_caption(ws, 4, "текст")
    assert ws.cell(row=4, column=1).alignment.wrap_text is False
