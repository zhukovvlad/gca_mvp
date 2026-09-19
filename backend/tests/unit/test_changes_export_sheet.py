"""Билдер книги «Изменения КП» — лист Excel (спека
2026-09-16-tender-changes-export-design.md §2.3, §2.5, §2.9, §2.10; план,
Task 4, «Утверждения»).

Каждый тест носит в докстроке номер утверждения задачи 4 плана
(`docs/superpowers/plans/2026-09-17-tender-changes-export.md`, раздел Task 4)
в форме `A4.N` — только для трассировки отчёта задачи, не для цитирования в
самом коде продукта.

Границы DoD 13: тест утверждает СТРУКТУРНУЮ читаемость (`openpyxl.load_workbook`
разбирает `bytes` без исключения, денежные ячейки после чтения числовые) — не
«без окна восстановления Excel»: это ненаблюдаемо прогоном и проверяется
вручную на стенде (`docs/insights/unobservable-in-the-runner.md`).

Входы строятся через `services.changes_export.build_sheet` — тот же способ,
которым `services.excel_changes_export` реально получает `Sheet` на входе, а
не руками собранные датаклассы: так тесты этого модуля заодно проверяют, что
он верно потребляет РЕАЛЬНЫЙ выход задачи 2, а не придуманную форму.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, range_boundaries
from openpyxl.workbook import Workbook

from services import changes_export as ce
from services import excel_changes_export as ece
from services.excel import FMT_MONEY

D = Decimal

ARTICLE_A = ce.ArticleRef(code="7", title="Отделка", sort_order=5)
ARTICLE_B = ce.ArticleRef(code="2.5", title="Кровля", sort_order=2)


def stages(n: int, *, rate="20") -> list[ce.StageInput]:
    r = None if rate is None else D(rate)
    return [ce.StageInput(stage_no=i + 1, label=None, held_on=None, vat_rate_base=r) for i in range(n)]


def stages_rates(rates: list[str | None]) -> list[ce.StageInput]:
    return [
        ce.StageInput(stage_no=i + 1, label=None, held_on=None, vat_rate_base=None if r is None else D(r))
        for i, r in enumerate(rates)
    ]


def group_row(
    stage_no,
    *,
    article=ARTICLE_A,
    catalog_position_id=1,
    catalog_kind="POSITION",
    work_title="Работа 1",
    unit="м2",
    rows_all=1,
    rows_with_amount=1,
    rows_with_mix=1,
    rows_priced=1,
    amount="100",
    works="40",
    materials="50",
    indirect="10",
    quantity="2",
    price_num=None,
    price_den=None,
    unit_works=None,
    unit_materials=None,
    unit_indirect=None,
    numbers=("1",),
) -> ce.GroupRow:
    def d(v):
        return None if v is None else D(v)

    if price_num is None:
        price_num = amount
    if price_den is None:
        price_den = quantity
    if unit_works is None:
        unit_works = works
    if unit_materials is None:
        unit_materials = materials
    if unit_indirect is None:
        unit_indirect = indirect
    return ce.GroupRow(
        stage_no=stage_no,
        article=article,
        catalog_position_id=catalog_position_id,
        catalog_kind=catalog_kind,
        work_title=work_title,
        unit=unit,
        rows_all=rows_all,
        rows_with_amount=rows_with_amount,
        rows_with_mix=rows_with_mix,
        rows_priced=rows_priced,
        amount=d(amount),
        components=ce.Components(d(works), d(materials), d(indirect)),
        quantity=d(quantity),
        price_num=d(price_num),
        price_den=d(price_den),
        unit_components_num=ce.Components(d(unit_works), d(unit_materials), d(unit_indirect)),
        numbers=numbers,
    )


def additional_row(
    stage_no, *, article=ARTICLE_A, lot_key="ЛОТ-1", chapter_ref_raw="3.2.2",
    work_title="Допработа", amount="30", rows_all=1, rows_with_amount=1,
) -> ce.AdditionalRow:
    return ce.AdditionalRow(
        stage_no=stage_no, article=article, lot_key=lot_key, chapter_ref_raw=chapter_ref_raw,
        work_title=work_title, amount=None if amount is None else D(amount),
        rows_all=rows_all, rows_with_amount=rows_with_amount,
    )


def sheet_of(stage_list, groups=(), additional=(), *, title="Участник") -> ce.Sheet:
    return ce.build_sheet(ce.SheetInput(
        participant_title=title, stages=stage_list, groups=list(groups), additional=list(additional),
    ))


def book_of(*sheets: ce.Sheet, tender_header: dict | None = None) -> Workbook:
    data = ece.build_changes_export(list(sheets), tender_header=tender_header or {})
    return load_workbook(BytesIO(data))


def header_row(ws) -> int:
    """Строка меток колонок — первая строка, где колонка 1 несёт `LEFT_COLUMNS[0]`."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == ece.LEFT_COLUMNS[0]:
            return r
    raise AssertionError("header row (LEFT_COLUMNS[0]) not found")


def stage_first_col(idx: int) -> int:
    """Первая из четырёх колонок этапа с индексом `idx` (0-based)."""
    return len(ece.LEFT_COLUMNS) + 1 + idx * 4


def find_all_cells_text(ws) -> list[str]:
    return [
        c.value for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value
    ]


# --------------------------------------------------------------------------
# A4.1 — колонок ровно 4 + 4 × число этапов + 7.
# --------------------------------------------------------------------------


class TestColumnCount:
    def test_a4_1_column_count_is_4_plus_4_times_stages_plus_7(self):
        sheet = sheet_of(stages(3), groups=[group_row(1), group_row(2), group_row(3)])
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        assert ws.max_column == 4 + 4 * 3 + 7
        hr = header_row(ws)
        # Ревью задачи 4: `ws.max_column` сам по себе не различает «данные
        # физически заняли N колонок» от «заголовки колонок рассинхронились с
        # данными» — колонки данных пишутся фиксированной арифметикой
        # (`_write_data_row`), а заголовки строятся отдельно из `columns_spec`.
        # Пропавшая запись в `columns_spec` сдвинула бы подписи влево, и
        # последняя колонка данных осталась бы БЕЗ заголовка. Этот assert
        # предъявляет именно связку заголовок-колонка, а не только их число.
        assert ws.cell(hr, ws.max_column).value == ece.TAIL_COLUMNS[-1]


# --------------------------------------------------------------------------
# A4.2 — область данных ПЛОСКАЯ: ни одной строки-заголовка статьи внутри
# диапазона данных.
# --------------------------------------------------------------------------


class TestFlatDataArea:
    def test_a4_2_data_area_has_no_article_header_rows(self):
        sheet = sheet_of(
            stages(2),
            groups=[
                group_row(1, catalog_position_id=1, amount="100"),
                group_row(2, catalog_position_id=1, amount="100"),
                group_row(1, catalog_position_id=2, amount="50"),
                group_row(2, catalog_position_id=2, amount="50"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        data_rows = range(hr + 1, hr + 1 + len(sheet.rows))

        merged_row_starts = {rng.min_row for rng in ws.merged_cells.ranges if rng.min_col == 1}
        for r in data_rows:
            assert ws.cell(r, 1).value not in (None, "")
            assert r not in merged_row_starts, f"row {r} is a merged banner, not a flat data row"

        # Контраст: блок подытогов НИЖЕ данных — банером, который ДЕЙСТВИТЕЛЬНО
        # мержит колонку 1 на всю ширину, доказывая, что мерж-детектор рабочий.
        subtotal_banner_row = hr + len(sheet.rows) + 2
        assert subtotal_banner_row in merged_row_starts


# --------------------------------------------------------------------------
# A4.3 — автофильтр стоит на области данных и не покрывает блок подытогов.
# --------------------------------------------------------------------------


class TestAutoFilter:
    def test_a4_3_autofilter_excludes_subtotal_block(self):
        sheet = sheet_of(stages(2), groups=[group_row(1), group_row(2)])
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        last_data_row = hr + len(sheet.rows)

        ref = ws.auto_filter.ref
        assert ref is not None
        _min_col, min_row, _max_col, max_row = range_boundaries(ref)
        assert min_row == hr
        assert max_row == last_data_row

        subtotal_banner_row = last_data_row + 2
        assert subtotal_banner_row > max_row


# --------------------------------------------------------------------------
# A4.4 — окно закреплено по шапке и первым четырём колонкам.
# --------------------------------------------------------------------------


class TestFreezePanes:
    def test_a4_4_freeze_panes_locks_header_and_first_four_columns(self):
        sheet = sheet_of(stages(2), groups=[group_row(1), group_row(2)])
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        col, row = coordinate_from_string(ws.freeze_panes)
        assert col == "E"
        assert row == hr + 1


# --------------------------------------------------------------------------
# A4.5 — денежная ячейка ОДНОЯРУСНА: число с `FMT_MONEY`, суммируется как
# число, без переноса строки внутри ячейки.
# --------------------------------------------------------------------------


class TestMoneyCellSingleTier:
    def test_a4_5_money_cell_is_single_tier_and_summable(self):
        sheet = sheet_of(
            stages(1),
            groups=[
                group_row(1, catalog_position_id=1, amount="100", numbers=("1",)),
                group_row(1, catalog_position_id=2, amount="250", numbers=("2",)),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        amount_col = stage_first_col(0) + 3

        v1 = ws.cell(hr + 1, amount_col).value
        v2 = ws.cell(hr + 2, amount_col).value
        assert isinstance(v1, int | float)
        assert isinstance(v2, int | float)
        assert "\n" not in str(v1)
        assert round(float(v1) + float(v2), 2) == 350.0
        assert not ws.cell(hr + 1, amount_col).alignment.wrap_text


# --------------------------------------------------------------------------
# A4.6 — сумма различает три состояния: строки нет (прочерк), строка есть без
# конечной суммы (текст «суммы нет»), конечная сумма ноль (число 0).
# --------------------------------------------------------------------------


class TestThreeAmountStates:
    def test_a4_6_no_amount_zero_and_absent_are_three_distinct_renderings(self):
        sheet = sheet_of(
            stages(2),
            groups=[
                # Работа 10: этап 1 — строка есть, конечной суммы нет.
                group_row(1, catalog_position_id=10, work_title="Работа10", rows_with_amount=0,
                          amount=None, price_num="0", price_den="0"),
                # Работа 10: этап 2 — конечная сумма, действительно ноль.
                group_row(2, catalog_position_id=10, work_title="Работа10", amount="0",
                          works="0", materials="0", indirect="0", price_num="0", price_den="0"),
                # Работа 20: только на этапе 2 — на этапе 1 строки нет вовсе.
                group_row(2, catalog_position_id=20, work_title="Работа20", amount="80"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)

        def row_of(title: str) -> int:
            for r in range(hr + 1, hr + 1 + len(sheet.rows)):
                if ws.cell(r, 3).value == title:
                    return r
            raise AssertionError(f"row for {title!r} not found")

        r10 = row_of("Работа10")
        r20 = row_of("Работа20")
        amount_col_stage1 = stage_first_col(0) + 3
        amount_col_stage2 = stage_first_col(1) + 3

        assert ws.cell(r10, amount_col_stage1).value == ece.TEXT_NO_AMOUNT
        assert ws.cell(r10, amount_col_stage2).value == 0
        assert ws.cell(r20, amount_col_stage1).value == "—"


# --------------------------------------------------------------------------
# A4.7 — в блоке подытогов TEXT_NO_AMOUNT не появляется НИ РАЗУ: недоступность
# по оси печатает TEXT_UNKNOWN_VAT, недоступность по строке остаётся числом.
# --------------------------------------------------------------------------


class TestSubtotalNeverShowsNoAmount:
    def test_a4_7_row_without_final_amount_leaves_subtotal_numeric(self):
        # Статья с двумя работами на этапе 1: у одной сумма отсутствует, у
        # другой есть. Подытог статьи — ЧИСЛО (40), не «суммы нет».
        sheet = sheet_of(
            stages(2),
            groups=[
                group_row(1, catalog_position_id=1, amount=None, rows_with_amount=0,
                          price_num="0", price_den="0"),
                group_row(1, catalog_position_id=2, amount="40"),
                group_row(2, catalog_position_id=1, amount="10"),
                group_row(2, catalog_position_id=2, amount="40"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        block_header = hr + len(sheet.rows) + 3  # разделитель + баннер + строка меток
        subtotal_row = block_header + 1
        stage1_col = 4  # Статья, Наименование, Строк, затем этапы (этап 1 — первая из них)

        cell = ws.cell(subtotal_row, stage1_col)
        assert cell.value == 40
        assert cell.value != ece.TEXT_NO_AMOUNT

    def test_a4_7_unknown_vat_axis_shows_text_never_no_amount(self):
        sheet = sheet_of(
            stages_rates([None, None]),
            groups=[group_row(1, amount="100"), group_row(2, amount="100")],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        block_header = hr + len(sheet.rows) + 3
        subtotal_row = block_header + 1

        for stage_index in range(2):
            cell = ws.cell(subtotal_row, 4 + stage_index)
            assert cell.value == ece.TEXT_UNKNOWN_VAT
            assert cell.value != ece.TEXT_NO_AMOUNT


# --------------------------------------------------------------------------
# A4.8 — Money со значением печатается числом с FMT_MONEY, и в блоке
# подытогов тоже.
# --------------------------------------------------------------------------


class TestFmtMoneyEverywhere:
    def test_a4_8_fmt_money_applies_to_data_and_subtotal_cells(self):
        sheet = sheet_of(stages(2), groups=[group_row(1, amount="100"), group_row(2, amount="120")])
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        amount_col = stage_first_col(0) + 3
        assert ws.cell(hr + 1, amount_col).number_format == FMT_MONEY

        block_header = hr + len(sheet.rows) + 3
        subtotal_row = block_header + 1
        assert ws.cell(subtotal_row, 4).number_format == FMT_MONEY


# --------------------------------------------------------------------------
# A4.9 — подытог со значением и признаком `incomplete` остаётся ЧИСЛОМ, а
# неполнота — заливкой И соседней колонкой «Полнота» блока подытогов.
# --------------------------------------------------------------------------


class TestSubtotalIncompleteStaysNumeric:
    def test_a4_9_incomplete_subtotal_is_numeric_with_fill_and_neighbor_note(self):
        sheet = sheet_of(
            stages(2),
            groups=[
                group_row(1, catalog_position_id=1, amount=None, rows_with_amount=0,
                          price_num="0", price_den="0"),
                group_row(1, catalog_position_id=2, amount="40"),
                group_row(2, catalog_position_id=1, amount="10"),
                group_row(2, catalog_position_id=2, amount="40"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        block_header = hr + len(sheet.rows) + 3
        subtotal_row = block_header + 1
        n_stages = 2
        stage1_cell = ws.cell(subtotal_row, 4)
        note_col = 4 + n_stages + 1  # Статья,Наименование,Строк + этапы + Δ -> Полнота

        assert stage1_cell.value == 40
        assert isinstance(stage1_cell.value, int | float)
        assert stage1_cell.fill.fgColor.rgb not in (None, "00000000")
        note_text = ws.cell(subtotal_row, note_col).value
        assert note_text and "Э1" in note_text

    def test_review_d4_partial_cell_subtotal_gets_fill_and_note_stays_numeric(self):
        """Ревью финального (Д4): та же заливка и та же соседняя колонка
        «Полнота», что A4.9 уже проверяет для ячейки БЕЗ конечной суммы,
        обязаны появиться и для ЧАСТИЧНОЙ ячейки (`rows_with_amount < rows_all`,
        сумма при этом ЕСТЬ). До правки задачи сервисного слоя эта ячейка
        не поднимала `incomplete`, и билдер книги печатал число без заливки
        и без строки в «Полноте» — молчаливая частичная свёртка."""
        sheet = sheet_of(
            stages(2),
            groups=[
                group_row(1, catalog_position_id=1, amount="100", rows_all=2, rows_with_amount=1,
                          rows_with_mix=1),
                group_row(2, catalog_position_id=1, amount="130"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        block_header = hr + len(sheet.rows) + 3
        subtotal_row = block_header + 1
        n_stages = 2
        stage1_cell = ws.cell(subtotal_row, 4)
        note_col = 4 + n_stages + 1

        assert stage1_cell.value == 100
        assert isinstance(stage1_cell.value, int | float)
        assert stage1_cell.value != ece.TEXT_NO_AMOUNT
        assert stage1_cell.fill.fgColor.rgb not in (None, "00000000")
        note_text = ws.cell(subtotal_row, note_col).value
        assert note_text and "Э1" in note_text


# --------------------------------------------------------------------------
# A4.10 — ячейка без пригодной цены несёт текст TEXT_NO_PRICE.
# --------------------------------------------------------------------------


class TestNoPriceText:
    def test_a4_10_cell_without_priced_rows_shows_no_price_text(self):
        sheet = sheet_of(
            stages(1),
            groups=[group_row(1, amount="100", price_num="0", price_den="0")],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        price_col = stage_first_col(0) + 2
        assert ws.cell(hr + 1, price_col).value == ece.TEXT_NO_PRICE


# --------------------------------------------------------------------------
# A4.11 — у MONEY_ONLY_KINDS «№ строк», «Объём», «Цена за ед.» — прочерк, и
# три колонки Δ состава — прочерк, а не подпись состояния.
# --------------------------------------------------------------------------


class TestMoneyOnlyKindsDashes:
    def test_a4_11_money_only_kind_shows_dashes_not_state_labels(self):
        sheet = sheet_of(
            stages(2),
            groups=[],
            additional=[additional_row(1, amount="30"), additional_row(2, amount="30")],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        r = hr + 1
        first_col = stage_first_col(0)

        assert ws.cell(r, first_col).value == "—"
        assert ws.cell(r, first_col + 1).value == "—"
        assert ws.cell(r, first_col + 2).value == "—"
        assert ws.cell(r, first_col + 3).value == 30  # сумма — число

        # Δ работы / материалы / косвенные — три следующие за Δ сумма/Δ%, прочерк.
        n_stages = 2
        delta_amount_col = 4 + n_stages * 4 + 1
        works_col = delta_amount_col + 2
        for offset in range(3):
            value = ws.cell(r, works_col + offset).value
            assert value == "—"
            assert value != ece.TEXT_MIX_INCOMPLETE


# --------------------------------------------------------------------------
# A4.12 — наименование работы прогоняется через safe_str.
# --------------------------------------------------------------------------


class TestSafeStr:
    def test_a4_12_work_title_is_neutralised_via_safe_str(self):
        sheet = sheet_of(stages(1), groups=[group_row(1, work_title="=1+1", amount="100")])
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        assert ws.cell(hr + 1, 3).value == "'=1+1"


# --------------------------------------------------------------------------
# A4.13 — шапка листа печатает подпись налогового состава и подпись ценового
# уровня, обе, на каждом листе.
# --------------------------------------------------------------------------


class TestCaptions:
    def test_a4_13_tax_basis_and_price_level_captions_are_on_every_sheet(self):
        # Ревью задачи 4: заголовки участников НЕ должны содержать слова
        # «валов»/«нетто» — иначе баннер листа (который несёт `sheet.title`)
        # подтверждает assert одним лишь своим именем участника, а не подписью
        # налогового состава, которую тест якобы проверяет (ложная зелень по
        # выбору фикстуры, `docs/insights/verifying-guards.md`, слой 12).
        gross_sheet = sheet_of(stages(2, rate="20"), groups=[group_row(1), group_row(2)], title="Участник A")
        net_sheet = sheet_of(stages_rates(["20", "10"]), groups=[group_row(1), group_row(2)], title="Участник B")
        wb = book_of(gross_sheet, net_sheet)

        for name, expect_word in ((wb.sheetnames[0], "валов"), (wb.sheetnames[1], "нетто")):
            texts = find_all_cells_text(wb[name])
            joined = " ".join(texts).lower()
            assert expect_word in joined
            assert "номинальный" in joined
            assert "приведения" in joined


# --------------------------------------------------------------------------
# A4.14 — шапка листа печатает moves_note один раз.
# --------------------------------------------------------------------------


class TestMovesNoteOnce:
    def test_a4_14_moves_note_is_printed_exactly_once(self):
        sheet = sheet_of(
            stages(3),
            groups=[
                group_row(1, catalog_position_id=1, article=ARTICLE_A, amount="100"),
                group_row(2, catalog_position_id=1, article=ARTICLE_B, amount="100"),
                group_row(3, catalog_position_id=1, article=ARTICLE_B, amount="100"),
            ],
        )
        wb = book_of(sheet)
        ws = wb[wb.sheetnames[0]]
        texts = find_all_cells_text(ws)
        occurrences = sum(1 for t in texts if sheet.moves_note in t)
        assert occurrences == 1


# --------------------------------------------------------------------------
# A4.15 — имена листов книги — результат sheet_names, столько же элементов,
# порядок совпадает с порядком sheets.
# --------------------------------------------------------------------------


class TestSheetNamesMatchOrder:
    def test_a4_15_sheet_names_come_from_sheet_names_helper_in_order(self):
        s1 = sheet_of(stages(2), groups=[group_row(1), group_row(2)], title="Бета")
        s2 = sheet_of(stages(2), groups=[group_row(1), group_row(2)], title="Альфа")
        wb = book_of(s1, s2)
        expected = ce.sheet_names(["Бета", "Альфа"])
        assert wb.sheetnames == expected
        assert len(wb.sheetnames) == 2


# --------------------------------------------------------------------------
# A4.16 — книга СТРУКТУРНО читаема: `load_workbook` разбирает `bytes` без
# исключения, и каждая денежная ячейка после чтения — числового типа.
# --------------------------------------------------------------------------


class TestStructuralReadability:
    def test_a4_16_workbook_round_trips_and_money_cells_are_numeric(self):
        sheet = sheet_of(
            stages(2),
            groups=[
                group_row(1, catalog_position_id=1, amount="100"),
                group_row(2, catalog_position_id=1, amount=None, rows_with_amount=0,
                          price_num="0", price_den="0"),
            ],
        )
        data = ece.build_changes_export([sheet], tender_header={"tender_number": "12/2025", "tender_title": "Т"})
        wb = load_workbook(BytesIO(data))  # не должно бросать исключение
        ws = wb[wb.sheetnames[0]]
        hr = header_row(ws)
        amount_col_stage1 = stage_first_col(0) + 3
        amount_col_stage2 = stage_first_col(1) + 3

        assert isinstance(ws.cell(hr + 1, amount_col_stage1).value, int | float)
        assert ws.cell(hr + 1, amount_col_stage2).value == ece.TEXT_NO_AMOUNT
        assert isinstance(ws.cell(hr + 1, amount_col_stage2).value, str)


# --------------------------------------------------------------------------
# Контракт `tender_header` (ревью задачи 4): ключи `tender_number` и
# `tender_title` не просто существуют с мягким `.get()` — их печать на листе
# ПРЕДЪЯВЛЕНА тестом. Без этого теста опечатка в ключе на стороне задачи 5
# (другие имена в словаре) прошла бы молча: `.get()` не бросает исключение,
# шапка просто осталась бы словом «Тендер» без номера и названия.
# --------------------------------------------------------------------------


class TestTenderHeaderContract:
    def test_tender_number_and_title_are_printed_in_the_banner(self):
        sheet = sheet_of(stages(1), groups=[group_row(1, amount="100")])
        wb = book_of(sheet, tender_header={"tender_number": "12/2025", "tender_title": "Ремонт склада"})
        ws = wb[wb.sheetnames[0]]
        banner_text = ws.cell(1, 1).value
        assert "12/2025" in banner_text
        assert "Ремонт склада" in banner_text
