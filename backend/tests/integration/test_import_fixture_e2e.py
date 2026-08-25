"""Полный конвейер на обезличенном fixture: upload → … → done (DoD фазы 4).

Единственное место, где парсер работает по-настоящему, поэтому файл разбирается
ОДИН раз на модуль (~20 с, `docs/phase3-parser.md` §3), а оба сценария
пользуются одним `ParseResult`.

Числа fixture (2576 позиций, 746 разделов) — из `docs/phase3-parser.md`.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import pytest
import sqlalchemy as sa
from openpyxl import load_workbook

from models import (
    CatalogKind,
    CatalogPosition,
    Estimate,
    EstimateAdditionalWork,
    EstimateRawData,
    ImportJob,
    ImportJobStatus,
    Lot,
    PositionItem,
    Proposal,
    ProposalSummaryLine,
    WorkCategory,
)
from parser import parse_estimate
from parser.constants import TABLE_PARSE_ADDITIONAL_WORKS_TITLE
from parser.parse_contractor_row import parse_contractor_row
from parser.resolve_contractor import BlockLayout, ResolvedContractor
from services import import_pipeline
from services.review import set_kind
from services.unit_resolution import UnitResolver

pytestmark = pytest.mark.integration

FIXTURE = Path(__file__).resolve().parents[3] / "fixtures" / "gp_estimate_fixture.xlsx"

#: Замеры фазы 3 на этом же файле.
FIXTURE_POSITIONS = 2576
FIXTURE_CHAPTERS = 746
FIXTURE_WORK_ROWS = FIXTURE_POSITIONS - FIXTURE_CHAPTERS

# Замер спеки Ф3 §1.1 на этом файле. Сумма сходится с FIXTURE_CHAPTERS: 222+485+39=746.
FIXTURE_CHAPTERS_OWN = 222
FIXTURE_CHAPTERS_INHERITED = 485
FIXTURE_CHAPTERS_UNASSIGNED = 39
FIXTURE_POSITIONS_UNASSIGNED = 38
#: Разделов глубины 1 — единственные строки без родителя (строк вне структуры в этом
#: файле нет). Из гистограммы глубин пробника гейта 1: {1: 16, 2: 56, 3: 130, 4: 196,
#: 5: 337, 6: 11}; сумма 746 = FIXTURE_CHAPTERS.
FIXTURE_CHAPTERS_TOP_LEVEL = 16


def test_measured_counts_add_up_to_the_chapter_total():
    """Арифметика замера — до всякой БД: три класса разделов покрывают все разделы."""
    assert (
        FIXTURE_CHAPTERS_OWN + FIXTURE_CHAPTERS_INHERITED + FIXTURE_CHAPTERS_UNASSIGNED
        == FIXTURE_CHAPTERS
    )


@pytest.fixture(scope="module")
def parsed_fixture():
    """Разбор fixture — один раз на модуль: это самая дорогая операция в тестах."""
    if not FIXTURE.is_file():
        pytest.skip(f"Нет {FIXTURE}")
    return parse_estimate(str(FIXTURE))


@pytest.fixture
def stub_with_fixture(monkeypatch, parsed_fixture):
    """Пайплайн получает уже разобранный fixture вместо повторного чтения XLSX."""
    monkeypatch.setattr(import_pipeline, "parse_estimate", lambda _handle: parsed_fixture)
    return parsed_fixture


def upload(client, contract_id: int, content: bytes):
    response = client.post(
        "/api/v1/estimates/upload",
        files={"file": ("смета.xlsx", content, "application/octet-stream")},
        data={"contract_id": str(contract_id)},
    )
    assert response.status_code == 202, response.text
    return client.get(f"/api/v1/import-jobs/{response.json()['id']}").json()


def xlsx_stub(marker: str) -> bytes:
    """Байты-заглушка: настоящий разбор подменён, но эндпоинт проверяет ZIP-магию."""
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", marker)
    return buffer.getvalue()


@pytest.fixture
def imported_fixture_estimate(
    committing_client, committing_db, committing_factories, stub_with_fixture
) -> int:
    """Загруженная fixture-смета; отдаёт `estimate_id`."""
    contract = committing_factories.ContractFactory.create()
    committing_db.commit()
    job = upload(committing_client, contract.id, xlsx_stub("categories"))
    assert job["status"] == ImportJobStatus.done.value, job["error_text"]
    return job["estimate_id"]


@dataclass
class ImportedFixture:
    """То немногое из результата импорта, что нужно сквозным тестам Ф4a."""

    raw: dict
    proposal_id: int
    warnings: list[str]


@pytest.fixture
def imported_fixture(committing_db, imported_fixture_estimate) -> ImportedFixture:
    """`imported_fixture_estimate`, обёрнутый в сырой JSON, id единственного
    предложения и предупреждения задания — всё из БД, без повторного импорта."""
    estimate_id = imported_fixture_estimate
    raw = committing_db.get(EstimateRawData, estimate_id)
    proposal_id = committing_db.execute(
        sa.select(Proposal.id)
        .join(Lot, Lot.id == Proposal.lot_id)
        .where(Lot.estimate_id == estimate_id)
    ).scalar_one()
    estimate = committing_db.get(Estimate, estimate_id)
    job = committing_db.get(ImportJob, estimate.import_job_id)
    return ImportedFixture(
        raw=raw.raw_data, proposal_id=proposal_id, warnings=list(job.warnings)
    )


class TestFullPipelineOnFixture:
    def test_upload_to_done(
        self, committing_client, committing_db, committing_factories, stub_with_fixture
    ):
        contract = committing_factories.ContractFactory.create()
        committing_db.commit()

        job = upload(committing_client, contract.id, xlsx_stub("first"))

        assert job["status"] == ImportJobStatus.done.value
        assert job["error_text"] is None
        assert job["estimate_id"] is not None

        counters = job["counters"]
        assert counters["positions_total"] == FIXTURE_WORK_ROWS
        assert (
            counters["matched_cache"]
            + counters["matched_exact"]
            + counters["matched_nonposition"]
            + counters["to_review"]
            == FIXTURE_WORK_ROWS
        )
        # Каталог пуст — вся смета уходит в очередь Review.
        assert counters["to_review"] == FIXTURE_WORK_ROWS

        committing_db.expire_all()
        estimate_id = job["estimate_id"]
        assert (
            committing_db.execute(
                sa.select(sa.func.count())
                .select_from(PositionItem)
                .join(Proposal, Proposal.id == PositionItem.proposal_id)
                .join(Lot, Lot.id == Proposal.lot_id)
                .where(Lot.estimate_id == estimate_id)
            ).scalar_one()
            == FIXTURE_POSITIONS
        )
        assert (
            committing_db.execute(
                sa.select(sa.func.count())
                .select_from(PositionItem)
                .join(Proposal, Proposal.id == PositionItem.proposal_id)
                .join(Lot, Lot.id == Proposal.lot_id)
                .where(Lot.estimate_id == estimate_id, PositionItem.is_chapter.is_(True))
            ).scalar_one()
            == FIXTURE_CHAPTERS
        )

        raw = committing_db.get(EstimateRawData, estimate_id)
        assert raw.parser_version == stub_with_fixture.parser_version
        assert raw.raw_data["lots"]["lot_1"]["proposals"]["contractor_1"]["contractor_items"]

        # Разделы к матчингу не допускаются (§5, шаг 4).
        unmatched_chapters = committing_db.execute(
            sa.select(sa.func.count())
            .select_from(PositionItem)
            .join(Proposal, Proposal.id == PositionItem.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(
                Lot.estimate_id == estimate_id,
                PositionItem.is_chapter.is_(True),
                PositionItem.catalog_position_id.isnot(None),
            )
        ).scalar_one()
        assert unmatched_chapters == 0

    def test_second_upload_of_the_same_file_is_idempotent(
        self, committing_client, committing_db, committing_factories, stub_with_fixture
    ):
        contract = committing_factories.ContractFactory.create()
        committing_db.commit()
        content = xlsx_stub("same")
        first = upload(committing_client, contract.id, content)

        response = committing_client.post(
            "/api/v1/estimates/upload",
            files={"file": ("смета.xlsx", content, "application/octet-stream")},
            data={"contract_id": str(contract.id)},
        )

        assert response.status_code == 200
        assert response.json()["id"] == first["id"]
        committing_db.expire_all()
        assert (
            committing_db.execute(sa.select(sa.func.count()).select_from(Estimate)).scalar_one() == 1
        )


class TestMatchingMetric:
    """Метрика §10: после наполнения каталога те же строки матчатся ≥ 90 %.

    Загрузка идёт на НОВЫЙ договор — на старом сработала бы идемпотентность и
    вернулся бы прежний job (§10 требует именно контрольные договоры).
    """

    def test_metric_after_review_queue_is_processed(
        self, committing_client, committing_db, committing_factories, stub_with_fixture
    ):
        first_contract = committing_factories.ContractFactory.create()
        second_contract = committing_factories.ContractFactory.create()
        committing_db.commit()

        upload(committing_client, first_contract.id, xlsx_stub("first"))

        # Оператор разбирает очередь: всё утверждается как POSITION.
        committing_db.expire_all()
        resolver = UnitResolver(committing_db)
        queue = (
            committing_db.execute(
                sa.select(CatalogPosition.id).where(
                    CatalogPosition.kind == CatalogKind.TO_REVIEW.value
                )
            )
            .scalars()
            .all()
        )
        assert queue, "очередь Review не должна быть пустой после первой загрузки"
        for row_id in queue:
            set_kind(
                committing_db,
                to_review_id=row_id,
                kind=CatalogKind.POSITION.value,
                resolver=resolver,
            )
        committing_db.commit()

        job = upload(committing_client, second_contract.id, xlsx_stub("second"))

        counters = job["counters"]
        matched = counters["matched_cache"] + counters["matched_exact"]
        assert counters["positions_total"] == FIXTURE_WORK_ROWS
        assert matched / counters["positions_total"] >= 0.9
        # Ручные решения дают hit ветки 1, а не повторный точный поиск.
        assert counters["matched_cache"] == counters["positions_total"]
        assert counters["to_review"] == 0


class TestCategoryResolutionOnFixture:
    """Ф3 на закоммиченном файле: числа берутся из БД (спека §4.2)."""

    #: Общий хвост: только строки этой сметы.
    _SCOPE = (
        "from position_items p "
        "join proposals pr on pr.id = p.proposal_id "
        "join lots l on l.id = pr.lot_id "
        "where l.estimate_id = :eid"
    )

    def test_chapters_are_split_as_measured(self, committing_db, imported_fixture_estimate):
        counts = committing_db.execute(
            sa.text(
                "select "
                " count(*) filter (where p.is_chapter and p.smr_article_raw is not null "
                "                  and p.work_category_id is not null) as own, "
                " count(*) filter (where p.is_chapter and p.smr_article_raw is null "
                "                  and p.work_category_id is not null) as inherited, "
                " count(*) filter (where p.is_chapter "
                "                  and p.work_category_id is null) as unassigned, "
                " count(*) filter (where p.is_chapter and p.smr_article_raw is not null "
                "                  and p.work_category_id is null) as unreadable "
                + self._SCOPE
            ),
            {"eid": imported_fixture_estimate},
        ).one()
        assert counts.own == FIXTURE_CHAPTERS_OWN
        assert counts.inherited == FIXTURE_CHAPTERS_INHERITED
        assert counts.unassigned == FIXTURE_CHAPTERS_UNASSIGNED
        # Неизвестных кодов в файле нет — все 222 кода нашлись в справочнике.
        assert counts.unreadable == 0

    def test_every_work_row_is_attached_to_a_chapter(
        self, committing_db, imported_fixture_estimate
    ):
        """Строк вне структуры в этом файле нет, значит родитель есть у каждой позиции."""
        orphans = committing_db.execute(
            sa.text(
                "select count(*) " + self._SCOPE
                + " and not p.is_chapter and p.chapter_item_id is null"
            ),
            {"eid": imported_fixture_estimate},
        ).scalar_one()
        assert orphans == 0

    def test_unassigned_positions_match_the_measurement(
        self, committing_db, imported_fixture_estimate
    ):
        """Настоящий пробел: «14 SHELL & CORE» и «15 Рабочая документация» (спека §1.2)."""
        unassigned = committing_db.execute(
            sa.text(
                "select count(*) from position_items p "
                "join position_items c on c.id = p.chapter_item_id "
                "join proposals pr on pr.id = p.proposal_id "
                "join lots l on l.id = pr.lot_id "
                "where l.estimate_id = :eid and not p.is_chapter "
                "and c.work_category_id is null"
            ),
            {"eid": imported_fixture_estimate},
        ).scalar_one()
        assert unassigned == FIXTURE_POSITIONS_UNASSIGNED

    def test_parent_is_a_chapter_of_the_same_proposal_standing_earlier_in_the_file(
        self, committing_db, imported_fixture_estimate
    ):
        """Независимая проверка: эталон — порядок ключей файла, а не вывод резолвера.

        Слой 5 инсайта verifying-guards: если ожидание выводится из того же места,
        которое ломает контрпример, тест сравнивает величину с самой собой. Здесь
        ожидание берётся из `position_key_in_proposal` — его пишет импорт из ключей
        парсера, а не резолвер.
        """
        broken = committing_db.execute(
            sa.text(
                "select count(*) from position_items p "
                "join position_items c on c.id = p.chapter_item_id "
                "join proposals pr on pr.id = p.proposal_id "
                "join lots l on l.id = pr.lot_id "
                "where l.estimate_id = :eid and ("
                "  not c.is_chapter "
                "  or c.proposal_id <> p.proposal_id "
                "  or (c.position_key_in_proposal)::int >= (p.position_key_in_proposal)::int)"
            ),
            {"eid": imported_fixture_estimate},
        ).scalar_one()
        assert broken == 0
        # Ноль не должен быть вакуозным: ссылки в смете действительно есть.
        attached = committing_db.execute(
            sa.text("select count(*) " + self._SCOPE + " and p.chapter_item_id is not null"),
            {"eid": imported_fixture_estimate},
        ).scalar_one()
        assert attached == FIXTURE_POSITIONS - FIXTURE_CHAPTERS_TOP_LEVEL


# ---------------------------------------------------------------------------
# Ф4 (Task 6): агрегатная строка допработ, вписанная в fixture правкой этой
# задачи (docs/superpowers/specs/2026-08-07-additional-works-design.md §2.10).
# Скрипт правки в репозиторий не попадает (жил в scratchpad) — здесь только
# его наблюдаемый результат: контрольная сумма `T`, три строки "Сведений" и
# то, во что они обязаны резолвиться. Суммы синтетические и круглые.
# ---------------------------------------------------------------------------

#: Контрольная сумма агрегатной строки `T`; три строки "Сведений" в сумме дают
#: ровно её (полная расшивка, `P == T`, спека §2.6).
FIXTURE_ADDITIONAL_WORKS_TOTAL = Decimal("1500.00")

#: Строка 1: номер раздела, встречающийся среди строк-разделов РОВНО ОДИН РАЗ
#: и резолвящийся в единственную непустую статью (код "1.1").
FIXTURE_SVEDENIYA_REF_RESOLVED = "1.1"
FIXTURE_SVEDENIYA_REF_RESOLVED_AMOUNT = Decimal("1000.00")
FIXTURE_SVEDENIYA_REF_RESOLVED_TITLE = "Работы по разделу 1.1"
FIXTURE_SVEDENIYA_CATEGORY_CODE = "1.1"

#: Строка 2: номер "1" — в fixture он есть и у строки лота (без статьи), и у
#: первого настоящего раздела → кандидат без статьи → NULL + предупреждение.
#: Сумма строки — ноль (валидное состояние, спека §2.4).
FIXTURE_SVEDENIYA_REF_AMBIGUOUS = "1"
FIXTURE_SVEDENIYA_REF_AMBIGUOUS_AMOUNT = Decimal("0")

#: Строка 3: номер, которого в файле нет вовсе → ноль кандидатов → NULL +
#: предупреждение.
FIXTURE_SVEDENIYA_REF_ABSENT = "999.999"
FIXTURE_SVEDENIYA_REF_ABSENT_AMOUNT = Decimal("500.00")

#: Раскладка блока подрядчика fixture (J6, ширина 11) — тот же факт, что
#: проверяет `test_contractor_block_is_eleven_columns` в test_estimate.py:
#: column_start=10 (колонка J), colspan=11. Ключи и смещения — литералы, не
#: `sheet_builders.KEYS_GP_11`: эталон этого файла независим от строителя
#: (замер плана фичи, «Замеры»).
GP_11_COLUMN_KEYS = (
    "suggested_quantity",
    "unit_cost.materials",
    "unit_cost.works",
    "unit_cost.indirect_costs",
    "unit_cost.total",
    "total_cost.materials",
    "total_cost.works",
    "total_cost.indirect_costs",
    "total_cost.total",
    "total_cost_for_organizer_quantity",
    "comment_contractor",
)

FIXTURE_CONTRACTOR = ResolvedContractor(
    geometry={"column_start": 10, "merged_shape": {"colspan": 11}},
    layout=BlockLayout(column_keys=GP_11_COLUMN_KEYS, unit_cost_offset=1, total_cost_offset=5),
)

#: Точная метка строки ИТОГО (с учётом НДС) в колонке A листа fixture.
FIXTURE_TOTAL_WITH_VAT_LABEL = "ИТОГО, руб. с учетом НДС"


@pytest.fixture(scope="module")
def fixture_worksheet():
    """Лист fixture, открытый НАПРЯМУЮ через openpyxl — в обход парсера,
    `get_summary` и `proposal_summary_lines`. Нужен для независимых проверок
    §4.3 п.1 (агрегатная строка реально есть во входном листе) и п.8
    (независимое ИТОГО файла)."""
    if not FIXTURE.is_file():
        pytest.skip(f"Нет {FIXTURE}")
    wb = load_workbook(str(FIXTURE), data_only=True)
    try:
        yield wb.worksheets[0]
    finally:
        wb.close()


def _contractor_items(parsed) -> dict:
    return parsed.data["lots"]["lot_1"]["proposals"]["contractor_1"]["contractor_items"]


def _independent_total_with_vat(ws, contractor: ResolvedContractor) -> Decimal:
    """ИТОГО (с учётом НДС), прочитанное НАПРЯМУЮ из листа по точной метке в
    колонке A — в обход `get_summary`/`proposal_summary_lines`.

    Спека Ф4 §1.4: `get_summary` присваивает ключ словаря по метке
    (`"итого" in label and "ндс" in label`), и в трёхстрочной форме итогов
    "без учета НДС" молча перезаписывает "с учетом НДС" — три строки файла
    дают два ключа в JSON. Эталон п.8 обязан быть НЕЗАВИСИМ от этого кода:
    иначе проверка не отличила бы исправную сумму от дефекта (оба варианта
    сравнивались бы с одним и тем же испорченным числом).

    Денежный блок строки читается `parse_contractor_row` — той же функцией,
    что читает любую другую денежную строку листа, а не словарным поиском по
    ключу. Неизвестная метка — ГРОМКИЙ `pytest.fail`, а не молчаливый
    фолбэк: ноль совпадений означает, что эталона для сверки нет вовсе.
    """
    matches = [
        row
        for row in range(1, ws.max_row + 1)
        if str(ws.cell(row=row, column=1).value or "").strip() == FIXTURE_TOTAL_WITH_VAT_LABEL
    ]
    if not matches:
        pytest.fail(
            f"Метка ИТОГО «{FIXTURE_TOTAL_WITH_VAT_LABEL}» не найдена в колонке A листа "
            "fixture — независимый эталон для сверки недоступен."
        )
    if len(matches) > 1:
        pytest.fail(
            f"Метка ИТОГО «{FIXTURE_TOTAL_WITH_VAT_LABEL}» встретилась {len(matches)} раза "
            "в колонке A — неоднозначно, какую строку считать ИТОГО."
        )
    data = parse_contractor_row(ws, matches[0], contractor)
    value = data["total_cost"]["total"]
    if value is None:
        pytest.fail("Строка ИТОГО найдена, но total_cost.total пуст — сверка невозможна.")
    return Decimal(value)


SUMMARY_BLOCK_LABELS = (
    "ИТОГО, руб. с учетом НДС",
    "В том числе НДС",
    "ИТОГО, руб. без учета НДС",
)


def test_fixture_summary_block_has_three_filled_rows(fixture_worksheet):
    """Форма самого входа, прочитанная с листа, — не через парсер.

    Стережёт правку fixture: если блок снова станет двухстрочным, главный путь
    Ф4a останется без закоммиченного входа, и это должно быть видно сразу.
    """
    ws = fixture_worksheet
    rows = [2588, 2589, 2590]
    labels = [str(ws.cell(row=row, column=1).value or "").strip() for row in rows]
    assert labels == list(SUMMARY_BLOCK_LABELS)

    money = {
        row: [ws.cell(row=row, column=col).value for col in range(15, 19)]
        for row in rows
    }
    for row, values in money.items():
        assert all(value is not None for value in values), f"строка {row} заполнена не полностью"

    assert all(ws.cell(row=2591, column=col).value is None for col in range(1, 21)), (
        "строка 2591 обязана остаться пустой: она терминатор блока итогов"
    )

    for index in range(4):
        gross = Decimal(str(money[2588][index]))
        vat = Decimal(str(money[2589][index]))
        net = Decimal(str(money[2590][index]))
        assert gross == net + vat, f"колонка {15 + index}: тождество не сошлось"


class TestAggregateRowInSourceFile:
    """Спека §4.3, п.1-3: три факта о входе, каждый проверяется НЕЗАВИСИМО и
    ДО всякого импорта — п.1 читает workbook напрямую (не через парсер),
    п.2-3 читают уже разобранный JSON, но саму сумму/структуру не через БД."""

    def test_aggregate_row_is_present_in_the_input_sheet(self, fixture_worksheet):
        """П.1: строка реально есть в листе — чтение workbook, не парсера."""
        ws = fixture_worksheet
        matches = [
            row
            for row in range(11, ws.max_row + 1)
            if ws.cell(row=row, column=1).value is None
            and ws.cell(row=row, column=2).value is None
            and str(ws.cell(row=row, column=4).value or "").strip() == TABLE_PARSE_ADDITIONAL_WORKS_TITLE
        ]
        assert len(matches) == 1, f"ожидалась ровно одна агрегатная строка, найдено: {matches}"

    def test_aggregate_row_is_present_in_raw_additional_works(self, parsed_fixture):
        """П.2: строка присутствует в raw `additional_works`."""
        aggregate = _contractor_items(parsed_fixture)["additional_works"]
        assert aggregate is not None
        assert aggregate["job_title"] == TABLE_PARSE_ADDITIONAL_WORKS_TITLE
        assert Decimal(aggregate["total_cost"]["total"]) == FIXTURE_ADDITIONAL_WORKS_TOTAL

    def test_aggregate_row_is_not_in_raw_positions(self, parsed_fixture):
        """П.3: строки нет в raw `positions` — она не позиция (парсер 2.0.0, спека §2.1)."""
        positions = _contractor_items(parsed_fixture)["positions"]
        titles = {str(p.get("job_title")) for p in positions.values()}
        assert TABLE_PARSE_ADDITIONAL_WORKS_TITLE not in titles
        # Число позиций не выросло из-за агрегатной строки — она не заняла ключ.
        assert len(positions) == FIXTURE_POSITIONS


class TestAdditionalWorksInDatabase:
    """Спека §4.3, п.4-8: доменные записи после импорта. Каждый пункт — свой
    тест, читающий БД через `sa.select(Model.column)` (не `db_session.get()`
    после Core-операции — урок Ф3)."""

    def test_no_position_item_exists_for_the_aggregate_row(
        self, committing_db, imported_fixture_estimate
    ):
        """П.4: соответствующего `position_item` в БД нет."""
        count = committing_db.execute(
            sa.select(sa.func.count())
            .select_from(PositionItem)
            .join(Proposal, Proposal.id == PositionItem.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(
                Lot.estimate_id == imported_fixture_estimate,
                PositionItem.job_title_in_proposal == TABLE_PARSE_ADDITIONAL_WORKS_TITLE,
            )
        ).scalar_one()
        assert count == 0

    def test_three_additional_work_records_are_created(
        self, committing_db, imported_fixture_estimate
    ):
        """П.5: три записи `estimate_additional_works` (полная расшивка, P == T)."""
        ids = (
            committing_db.execute(
                sa.select(EstimateAdditionalWork.id)
                .join(Proposal, Proposal.id == EstimateAdditionalWork.proposal_id)
                .join(Lot, Lot.id == Proposal.lot_id)
                .where(Lot.estimate_id == imported_fixture_estimate)
            )
            .scalars()
            .all()
        )
        assert len(ids) == 3

    def test_additional_work_amounts_and_categories_match_the_three_lines(
        self, committing_db, imported_fixture_estimate
    ):
        """П.6: суммы и статьи ожидаемы — включая нулевую сумму и обе
        нераспределённые строки (кандидат без статьи; ноль кандидатов)."""
        rows = committing_db.execute(
            sa.select(
                EstimateAdditionalWork.ordinal,
                EstimateAdditionalWork.chapter_ref_raw,
                EstimateAdditionalWork.total_amount,
                EstimateAdditionalWork.work_category_id,
                EstimateAdditionalWork.title,
            )
            .join(Proposal, Proposal.id == EstimateAdditionalWork.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(Lot.estimate_id == imported_fixture_estimate)
            .order_by(EstimateAdditionalWork.ordinal)
        ).all()
        assert [r.ordinal for r in rows] == [1, 2, 3]
        resolved, ambiguous, absent = rows

        assert resolved.chapter_ref_raw == FIXTURE_SVEDENIYA_REF_RESOLVED
        assert resolved.total_amount == FIXTURE_SVEDENIYA_REF_RESOLVED_AMOUNT
        assert resolved.title == FIXTURE_SVEDENIYA_REF_RESOLVED_TITLE
        assert resolved.work_category_id is not None
        category_code = committing_db.execute(
            sa.select(WorkCategory.code).where(WorkCategory.id == resolved.work_category_id)
        ).scalar_one()
        assert category_code == FIXTURE_SVEDENIYA_CATEGORY_CODE

        assert ambiguous.chapter_ref_raw == FIXTURE_SVEDENIYA_REF_AMBIGUOUS
        assert ambiguous.total_amount == FIXTURE_SVEDENIYA_REF_AMBIGUOUS_AMOUNT  # ноль — валидная сумма
        assert ambiguous.work_category_id is None  # кандидат без статьи (лот + первый раздел)

        assert absent.chapter_ref_raw == FIXTURE_SVEDENIYA_REF_ABSENT
        assert absent.total_amount == FIXTURE_SVEDENIYA_REF_ABSENT_AMOUNT
        assert absent.work_category_id is None  # ноль кандидатов

    def test_records_sum_equals_the_aggregate_row_control_total(
        self, committing_db, imported_fixture_estimate
    ):
        """П.7: сумма расшивки равна контрольной сумме агрегатной строки `T`."""
        total = committing_db.execute(
            sa.select(sa.func.coalesce(sa.func.sum(EstimateAdditionalWork.total_amount), 0))
            .join(Proposal, Proposal.id == EstimateAdditionalWork.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(Lot.estimate_id == imported_fixture_estimate)
        ).scalar_one()
        assert Decimal(total) == FIXTURE_ADDITIONAL_WORKS_TOTAL

    def test_domain_total_matches_the_independent_grand_total_of_the_file(
        self, committing_db, imported_fixture_estimate, fixture_worksheet
    ):
        """П.8: `position_items` + допработы == независимое ИТОГО файла.

        Ради этого теста написана вся фича: инвариант "паспорт = позиции +
        допработы" (Ф4 §2.6) не ловит двойной счёт — обе стороны выросли бы
        одинаково. Ловит только сверка с ИТОГО, взятым НЕ из уже посчитанного
        паспорта, а прямо с листа (`_independent_total_with_vat`, минуя
        `get_summary`, спека §1.4).

        `is_chapter.is_(False)` в фильтре — не случайность и не вкусовщина: строки
        разделов в этом файле несут СВОЙ subtotal в `total_cost_total`
        (сумму своих детей), и суммирование ВСЕХ строк без фильтра считало бы
        каждый рубль по несколько раз — по разу на каждом уровне вложенности.
        Тот же фильтр использует паспорт (`crud/analytics.py:_passport_totals`)
        и каскад матчинга (`services/matching.py`).

        **Чего этот тест не ловит.** ИТОГО листа увеличено на `T` тем же
        скриптом правки, которым добавлена агрегатная строка (спека §2.10),
        поэтому неверное `T`, согласованно записанное в оба места, здесь
        прошло бы. А вот двойной счёт — то, ради чего тест написан, — так
        пройти не может: он дал бы `сумма + T + T`. Случай «`T` неверно с обеих сторон»
        закрывает сверка на настоящей оферте, где ИТОГО никто не правил
        (DoD спеки §6).
        """
        positions_total = committing_db.execute(
            sa.select(sa.func.coalesce(sa.func.sum(PositionItem.total_cost_total), 0))
            .join(Proposal, Proposal.id == PositionItem.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(Lot.estimate_id == imported_fixture_estimate, PositionItem.is_chapter.is_(False))
        ).scalar_one()
        additional_works_total = committing_db.execute(
            sa.select(sa.func.coalesce(sa.func.sum(EstimateAdditionalWork.total_amount), 0))
            .join(Proposal, Proposal.id == EstimateAdditionalWork.proposal_id)
            .join(Lot, Lot.id == Proposal.lot_id)
            .where(Lot.estimate_id == imported_fixture_estimate)
        ).scalar_one()

        domain_total = Decimal(positions_total) + Decimal(additional_works_total)
        independent_total = _independent_total_with_vat(fixture_worksheet, FIXTURE_CONTRACTOR)

        assert domain_total == independent_total


# ---------------------------------------------------------------------------
# Ф4a (Task 6): три строки файла доходят до трёх записей в БД. Дефект,
# которым эта фича доказывается: старый `get_summary` присваивал ключ по
# перекрывающимся подстрокам ("итого" и "ндс" есть и в "с учетом", и в "без
# учета") — три строки листа схлопывались в два ключа JSON, и уцелевшим
# оказывался НЕТТО, записанный под именем БРУТТО. Пять тестов ниже —
# независимые друг от друга проверки того, что теперь ключей и записей три,
# а брутто равен независимому эталону, прочитанному прямо с листа.
# ---------------------------------------------------------------------------


def test_three_sheet_rows_give_three_summary_keys_in_raw(imported_fixture):
    """Счёт, которым дефект был доказан: было три строки и два ключа."""
    summary = imported_fixture.raw["lots"]["lot_1"]["proposals"]["contractor_1"][
        "contractor_items"
    ]["summary"]
    assert len(summary) == 3
    assert sorted(summary) == [
        "total_cost_excluding_vat",
        "total_cost_including_vat",
        "vat_amount",
    ]


def test_three_summary_records_reach_the_database(committing_db, imported_fixture):
    lines = {
        line.summary_key: line
        for line in committing_db.scalars(
            sa.select(ProposalSummaryLine).where(
                ProposalSummaryLine.proposal_id == imported_fixture.proposal_id
            )
        )
    }
    assert sorted(lines) == [
        "total_cost_excluding_vat",
        "total_cost_including_vat",
        "vat_amount",
    ]
    assert lines["total_cost_including_vat"].job_title == "ИТОГО, руб. с учетом НДС"


def test_gross_total_matches_the_reference_read_from_the_sheet(
    committing_db, imported_fixture, fixture_worksheet
):
    """Главное доказательство починки: поле совпадает с эталоном, прочитанным
    с листа помимо парсера. До Ф4a в этом поле лежала сумма БЕЗ НДС."""
    expected = _independent_total_with_vat(fixture_worksheet, FIXTURE_CONTRACTOR)
    stored = committing_db.scalars(
        sa.select(ProposalSummaryLine).where(
            ProposalSummaryLine.proposal_id == imported_fixture.proposal_id,
            ProposalSummaryLine.summary_key == "total_cost_including_vat",
        )
    ).one()
    assert stored.total_cost == expected


@pytest.mark.parametrize(
    "field",
    ["materials_cost", "works_cost", "indirect_costs_cost", "total_cost"],
)
def test_identity_holds_on_the_stored_records(committing_db, imported_fixture, field):
    """Все ЧЕТЫРЕ денежные колонки, а не только итоговая (спека §2.6).

    Сверка одной колонки прошла бы и при разъехавшейся разбивке: три остальные
    поля доезжают до БД тем же путём и тем же `_money`, но проверялись бы
    ничем.
    """
    lines = {
        line.summary_key: line
        for line in committing_db.scalars(
            sa.select(ProposalSummaryLine).where(
                ProposalSummaryLine.proposal_id == imported_fixture.proposal_id
            )
        )
    }
    including = getattr(lines["total_cost_including_vat"], field)
    excluding = getattr(lines["total_cost_excluding_vat"], field)
    vat = getattr(lines["vat_amount"], field)
    # Каждое слагаемое проверяется ДО сложения: пустой компонент иначе даст
    # `TypeError: unsupported operand type(s)` вместо объясняющего падения, и
    # причина «в БД нет значения» осталась бы нечитаемой.
    for name, value in (("с НДС", including), ("без НДС", excluding), ("НДС", vat)):
        assert value is not None, f"{field}: значение «{name}» пусто — сверять нечего, тест был бы вакуозен"
    assert including == excluding + vat


SUMMARY_WARNING_MARKERS = (
    "Блок итогов не найден",
    "не распознано меток",
    "встретилась дважды",
    "суммы не указаны",
    "Валовое ИТОГО отсутствует",
    "Арифметика НДС не сходится",
    "Арифметика НДС не проверена из-за неполных, негодных или несравнимых точно данных",
)


def test_fixture_parses_without_any_summary_warning(imported_fixture):
    """Форма fixture полная и непротиворечивая — блок итогов молчит.

    Маркеры перечислены поимённо, а не отфильтрованы подстрокой «итог»: тексты
    про арифметику этого слова не содержат вовсе, и фильтр был бы вакуозен.
    """
    found = [
        text
        for text in imported_fixture.warnings
        for marker in SUMMARY_WARNING_MARKERS
        if marker in text
    ]
    assert found == []


# ---------------------------------------------------------------------------
# Ф4б (Task 2): ставка НДС заявлена суффиксом групповых шапок ценового блока.
# Правка вернула fixture в форму реальных файлов, где шапка и блок итогов
# согласованы (граница 2 Ф4a, спека Ф4б §2.10). Здесь — только форма ВХОДА;
# путь «лист → JSON → колонка БД» проверяется отдельно (Task 7).
# ---------------------------------------------------------------------------

#: Строка шапки колонок листа и якоря объединённых групповых шапок K9:N9 и O9:R9.
#: Номера замерены (спека §1.1) и одинаковы у всех четырёх известных файлов, но
#: продакшен их не зашивает: строку шапки он ищет `_validate_column_headers`.
FIXTURE_COLUMN_HEADER_ROW = 9
FIXTURE_UNIT_COST_HEADER_COLUMN = 11
FIXTURE_TOTAL_COST_HEADER_COLUMN = 15
FIXTURE_MONEY_GROUP_MERGES = ("K9:N9", "O9:R9")

#: Тексты обеих групповых шапок после правки Task 2. Это заголовки колонок, а не
#: коммерческие данные, — политика `samples/` их коммитить не запрещает.
FIXTURE_UNIT_COST_HEADER = "Цена за ед. изм., RUB, ОСН, с учетом НДС 20%"
FIXTURE_TOTAL_COST_HEADER = "Стоимость всего, RUB, ОСН, с учетом НДС 20%"


def test_fixture_column_headers_declare_the_vat_rate(fixture_worksheet):
    """Форма самого входа, прочитанная с листа, — не через парсер.

    Стережёт правку fixture ровно так же, как это делает
    `test_fixture_summary_block_has_three_filled_rows` для блока итогов: если
    суффикс из шапки пропадёт, главный положительный путь Ф4б останется без
    закоммиченного входа, и это должно быть видно сразу, а не через
    «предупреждений стало на одно больше».

    Целость объединений проверяется здесь же и не для красоты: ставка заявлена
    НАД группой из четырёх колонок, и снятое объединение оставило бы текст на
    месте при сломанной раскладке — то есть тест на один текст прошёл бы.
    """
    ws = fixture_worksheet
    unit_cost = ws.cell(
        row=FIXTURE_COLUMN_HEADER_ROW, column=FIXTURE_UNIT_COST_HEADER_COLUMN
    ).value
    total_cost = ws.cell(
        row=FIXTURE_COLUMN_HEADER_ROW, column=FIXTURE_TOTAL_COST_HEADER_COLUMN
    ).value
    assert unit_cost == FIXTURE_UNIT_COST_HEADER
    assert total_cost == FIXTURE_TOTAL_COST_HEADER

    merged = {str(rng) for rng in ws.merged_cells.ranges}
    for group in FIXTURE_MONEY_GROUP_MERGES:
        assert group in merged, f"объединение {group} потеряно"


# ---------------------------------------------------------------------------
# Ф4б (Task 7): ставка НДС от листа fixture до колонки `proposals.vat_rate`
# (спека §4.4, п.3-6). Пункты 1, 2, 7 и 8 закрыты существующими тестами этого
# файла (`test_fixture_column_headers_declare_the_vat_rate`, счётчики Ф3,
# `TestAdditionalWorksInDatabase`/`test_identity_holds_on_the_stored_records`
# для Ф4a) и здесь не дублируются. Утверждения ниже независимы друг от друга.
# ---------------------------------------------------------------------------

#: Тексты четырёх предупреждений Ф4б — подстроки взяты буквально из
#: `backend/parser/vat_rate.py` (`_rate_not_obtained_warning`,
#: `_rates_disagree_warning`, `build_vat_rate`). Перечислены поимённо, а не
#: одним словом «ставк»: такой фильтр не заметил бы, если тексты разойдутся.
VAT_RATE_WARNING_MARKERS = (
    "Ставка НДС не получена из шапки",
    "заявляют разные ставки НДС",
    "расходится с блоком итогов",
    "Сверка ставки НДС с блоком итогов не проведена",
)


def test_fixture_raw_vat_rate_is_a_decimal_string_on_the_contractor_level(imported_fixture):
    """П.3: в сыром JSON `vat_rate == "20"` — строка, не `Decimal` и не число.

    Ключ лежит на уровне подрядчика, рядом с `contractor_width` (спека §2.1),
    а не внутри `contractor_items` — второе утверждение проверяется здесь же,
    чтобы отличить «значения нет вовсе» от «значение лежит не там».
    """
    contractor = imported_fixture.raw["lots"]["lot_1"]["proposals"]["contractor_1"]
    assert contractor["vat_rate"] == "20"
    assert isinstance(contractor["vat_rate"], str)
    assert "contractor_width" in contractor  # соседство, на которое опирается решение §2.1
    assert "vat_rate" not in contractor["contractor_items"]


def test_fixture_database_vat_rate_is_a_decimal_twenty(committing_db, imported_fixture):
    """П.4: `proposals.vat_rate == Decimal("20")` — сравнение `Decimal` с `Decimal`."""
    stored = committing_db.execute(
        sa.select(Proposal.vat_rate).where(Proposal.id == imported_fixture.proposal_id)
    ).scalar_one()
    assert stored == Decimal("20")


@pytest.mark.parametrize(
    "column_index",
    [0, 1, 2, 3],
    ids=["materials", "works", "indirect_costs", "total"],
)
def test_fixture_summary_block_vat_ratio_is_within_tolerance(fixture_worksheet, column_index):
    """П.5: четыре отношения блока итогов проходят допуск.

    Эталон считается ЗДЕСЬ, из значений листа (строки 2589 «В том числе НДС» и
    2590 «ИТОГО без учета НДС», денежные колонки 15-18 — те же, что читает
    `test_fixture_summary_block_has_three_filled_rows`), а не берётся у
    парсера и не импортируется из
    `parser.vat_rate.VAT_RATE_TOLERANCE`: если эталон брать у проверяемого
    кода, обе стороны сравнения поедут вместе при поломке
    (docs/insights/verifying-guards.md, слой 5). Параметризация по всем
    четырём колонкам — чтобы разъехавшаяся разбивка была видна поколоночно, а
    не только по итогу.
    """
    ws = fixture_worksheet
    column = 15 + column_index
    vat = Decimal(str(ws.cell(row=2589, column=column).value))
    net = Decimal(str(ws.cell(row=2590, column=column).value))

    ratio = vat / net * 100
    tolerance = Decimal("0.01")  # литерал, не VAT_RATE_TOLERANCE — та же причина, что у эталона выше
    assert abs(ratio - Decimal("20")) <= tolerance


def test_fixture_parses_without_any_vat_rate_warning(imported_fixture):
    """П.6: предупреждений о ставке у fixture ноль.

    Форма — как у `test_fixture_parses_without_any_summary_warning`: маркеры
    перечислены поимённо (`VAT_RATE_WARNING_MARKERS`), а не отфильтрованы
    одной подстрокой вроде «ставк» — такой фильтр был бы вакуозен, если тексты
    предупреждений разойдутся.
    """
    found = [
        text
        for text in imported_fixture.warnings
        for marker in VAT_RATE_WARNING_MARKERS
        if marker in text
    ]
    assert found == []
