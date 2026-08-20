"""Лист выгрузки «Сравнение договоров» (спека 2026-08-17 §2.7, §2.8; план,
задача 6): числа листа — те же числа, что видит экран (DoD 19), три корзины
колонками с собственной суммой/₽ на м²/отклонением (DoD 21), подпись
налогового состава на листе (DoD 20), синтетические строки и три различимых
написания ячейки (прочерк / ноль / прочерк с причиной).

Ни один тест не пересчитывает арифметику независимо: ожидаемые числа берутся
из СЛОВАРЯ `crud.comparison.build_comparison`, а лист строится из того же
словаря — так и проверяется «один агрегат — два представления» (спека §2.7).
"""
from __future__ import annotations

import datetime as dt
from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import pytest
import sqlalchemy as sa
from openpyxl import load_workbook

from crud import comparison as cmp
from models import UserRole, WorkCategory
from money.inflation import YearMonth
from services.excel_comparison import _BUCKET_ORDER, _COLS_PER_CONTRACT, build_comparison_sheet
from tests import comparison_fixtures as fx

pytestmark = pytest.mark.integration


# ---------------------------------------------------------------------------
#  Помощники
# ---------------------------------------------------------------------------

def _category_title(db, code: str) -> str:
    return db.execute(
        sa.select(WorkCategory.title).where(WorkCategory.code == code)
    ).scalar_one()


def _sheet_from(data: dict):
    content = build_comparison_sheet(data, generated_at=dt.date(2026, 8, 17))
    return load_workbook(BytesIO(content)).active


def _cells(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def _text_of(ws) -> str:
    return "\n".join(
        " | ".join("" if v is None else str(v) for v in row) for row in _cells(ws)
    )


def _find_row_by_label(ws, label: str) -> int:
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return row[0].row
    raise AssertionError(f"строка с подписью {label!r} не найдена на листе")


def _col(column_index: int, bucket: str, metric: str) -> int:
    """Номер колонки метрики бакета договора по его позиции в `data["columns"]`.

    `column_index` — позиция в списке колонок агрегата, а НЕ id договора: обе
    стороны (лист и словарь `cells`) идут в одном и том же порядке, но порядок
    колонок — `signed_date DESC, id DESC` (спека §2.1, DoD 3), а не порядок
    вставки фикстур, поэтому тесты ищут позицию по `contract_id`, а не считают
    её вручную.
    """
    bucket_pos = _BUCKET_ORDER.index(bucket)
    metric_pos = {"sum": 0, "per_sqm": 1, "deviation": 2}[metric]
    return 2 + column_index * _COLS_PER_CONTRACT + bucket_pos * 3 + metric_pos


def _columns_index(data: dict) -> dict[int, int]:
    return {column["contract_id"]: index for index, column in enumerate(data["columns"])}


def _assert_bucket_cell_matches(ws, row_num: int, column_index: int, bucket: str, bucket_cell: dict) -> None:
    """Три метрики ОДНОЙ ячейки бакета — сверены с тем же словарём, что видит
    экран (DoD 19), включая различимость прочерк/ноль/прочерк-с-причиной
    (спека §2.1.2, §2.1.3) и то, что деньги — числа, а не строки (DoD 7)."""
    sum_value = ws.cell(row=row_num, column=_col(column_index, bucket, "sum")).value
    if bucket_cell["state"] == cmp.ABSENT:
        assert sum_value == "—", f"ABSENT обязан быть голым прочерком, получено {sum_value!r}"
    elif bucket_cell["shown"] is None:
        assert isinstance(sum_value, str) and sum_value.startswith("—"), (
            f"погашенная причинами ячейка обязана быть прочерком со словами, получено {sum_value!r}"
        )
        for code in bucket_cell["incomplete_reasons"]:
            label = {
                "unpriced_rows": "без цены",
                "not_finite_rows": "с ошибкой",
                "vat_base_unknown": "неизвестна база НДС",
                "display_rate_undefined": "ставка показа не определена",
            }[code]
            assert label in sum_value, f"причина {code!r} обязана быть названа словами на листе"
    else:
        assert isinstance(sum_value, int | float | Decimal), (
            f"деньги обязаны быть числом (СУММ по тексту даёт 0), получено {sum_value!r}"
        )
        assert Decimal(str(sum_value)) == bucket_cell["shown"]

    per_sqm_value = ws.cell(row=row_num, column=_col(column_index, bucket, "per_sqm")).value
    if bucket_cell["shown_per_sqm"] is None:
        assert per_sqm_value == "—"
    else:
        assert isinstance(per_sqm_value, int | float | Decimal)
        assert Decimal(str(per_sqm_value)) == bucket_cell["shown_per_sqm"]

    deviation_value = ws.cell(row=row_num, column=_col(column_index, bucket, "deviation")).value
    if bucket_cell["deviation_pct"] is None:
        assert deviation_value == "—"
    else:
        assert isinstance(deviation_value, int | float | Decimal)
        assert Decimal(str(deviation_value)) == bucket_cell["deviation_pct"]


def _assert_row_matches(ws, row_num: int, cells: list[dict], columns_index: dict[int, int]) -> None:
    for cell_entry in cells:
        column_index = columns_index[cell_entry["contract_id"]]
        for bucket in _BUCKET_ORDER:
            _assert_bucket_cell_matches(ws, row_num, column_index, bucket, cell_entry[bucket])


# ---------------------------------------------------------------------------
#  DoD 19 — числа листа совпадают с числами агрегата (нагрузочный тест)
# ---------------------------------------------------------------------------

def test_sheet_numbers_match_the_aggregate_dict(db_session, factories):
    """Лист строится из ТОГО ЖЕ словаря, что видит экран — числа не пересчитаны.

    Три договора с разной ценой статьи "1" и одинаковой площадью (100 000 м²)
    дают три РАЗНЫХ ₽/м² (10, 5, 20) — медиана 10, отклонения -50 % / 0 % /
    +100 %, поэтому тест ловит и перепутанный порядок колонок, и перепутанные
    местами сумма/удельная/отклонение.
    """
    c1 = fx.contract_with_area(db_session, factories, {"1": ["1200000.00"]}, vat_rate=fx.VAT_20)
    c2 = fx.contract_with_area(db_session, factories, {"1": ["600000.00"]}, vat_rate=fx.VAT_20)
    c3 = fx.contract_with_area(db_session, factories, {"1": ["2400000.00"]}, vat_rate=fx.VAT_20)
    db_session.flush()

    data = cmp.build_comparison(db_session, [c1, c2, c3], vat_mode=cmp.VAT_MODE_NET)
    ws = _sheet_from(data)
    columns_index = _columns_index(data)

    row_1 = next(r for r in data["rows"] if r["code"] == "1")
    row_num = _find_row_by_label(ws, _category_title(db_session, "1"))
    _assert_row_matches(ws, row_num, row_1["cells"], columns_index)

    totals_row_num = _find_row_by_label(ws, "ИТОГО ПО ДОГОВОРУ")
    _assert_row_matches(ws, totals_row_num, data["totals"], columns_index)

    # Замер, что тест действительно различает договоры, а не совпадает случайно:
    # три РАЗНЫХ отклонения на строке "1" в корзине "Итого".
    deviations = {
        cell["contract_id"]: cell[cmp.BUCKET_TOTAL]["deviation_pct"] for cell in row_1["cells"]
    }
    assert len({str(v) for v in deviations.values()}) == 3, "фикстура обязана давать три разных числа"


@pytest.mark.parametrize(
    ("vat_mode", "single_rate"),
    [(cmp.VAT_MODE_OWN, None), (cmp.VAT_MODE_SINGLE, Decimal("22"))],
)
def test_sheet_prints_the_shown_sum_not_the_netto_axis(db_session, factories, vat_mode, single_rate):
    """DoD 19 в режимах, где ПОКАЗ отличается от нетто-оси.

    Дыра, найденная снятием защиты: в режиме «нетто» `shown` равен `net`, поэтому
    набор, сверяющий числа только там, остаётся ЗЕЛЁНЫМ, даже если лист печатает
    нетто вместо показываемой суммы. Тогда экран и лист при одном режиме
    показывали бы разные числа — то самое расхождение представлений, ради
    исключения которого спека §2.7 требует один агрегат.

    Предпосылка утверждается в самом тесте: `shown != net`. Без неё тест снова
    стал бы вакуозным при первой же ставке 0 %.
    """
    contract_id = fx.contract_with_area(
        db_session, factories, {"1": ["1200000.00"]}, vat_rate=fx.VAT_20
    )
    db_session.flush()

    data = cmp.build_comparison(
        db_session, [contract_id], vat_mode=vat_mode, single_rate=single_rate
    )
    row_1 = next(r for r in data["rows"] if r["code"] == "1")
    cell = row_1["cells"][0][cmp.BUCKET_TOTAL]

    assert cell["shown"] != cell["net"], (
        "предпосылка: в этом режиме показываемая сумма ОБЯЗАНА отличаться от нетто, "
        "иначе тест не различает лист, печатающий нетто-ось"
    )

    ws = _sheet_from(data)
    row_num = _find_row_by_label(ws, _category_title(db_session, "1"))
    printed = ws.cell(
        row=row_num, column=_col(0, cmp.BUCKET_TOTAL, "sum")
    ).value

    assert Decimal(str(printed)) == cell["shown"], (
        "лист обязан печатать ПОКАЗЫВАЕМУЮ сумму — ту же, что видит экран в этом режиме"
    )

    printed_per_sqm = ws.cell(row=row_num, column=_col(0, cmp.BUCKET_TOTAL, "per_sqm")).value
    assert Decimal(str(printed_per_sqm)) == cell["shown_per_sqm"], (
        "₽/м² на листе — тоже показываемая величина, не вход медианы"
    )


# ---------------------------------------------------------------------------
#  DoD 21 — три корзины колонками, у каждой своя сумма/₽ на м²/отклонение
# ---------------------------------------------------------------------------

def test_each_bucket_has_its_own_sum_per_sqm_and_deviation(db_session, factories):
    """ДГП и ДС дают РАЗНЫЕ медианы -> отклонения одного договора в разных
    корзинах различаются (иначе тест прошёл бы и при одном отклонении на все
    три колонки, что и запрещает спека §2.8).

    Три договора, категория "1", ставка НДС 0 % (нетто = валовое, чтобы не
    примешивать пересчёт НДС к арифметике медианы):
        A: ДГП 1 000 000, ДС 5 000 000 -> ₽/м² 10 / 50 / 60
        B: ДГП 2 000 000, ДС 4 000 000 -> ₽/м² 20 / 40 / 60
        C: ДГП 3 000 000, ДС 3 000 000 -> ₽/м² 30 / 30 / 60
    Медианы: ДГП 20, ДС 40, Итого 60. У договора A отклонения: ДГП -50 %,
    ДС +25 %, Итого 0 % — три РАЗНЫХ числа на одной строке одного договора.
    """
    zero = Decimal("0")
    a = fx.contract_with_amendment(
        db_session, factories, base_rate=zero, amd_rate=zero, base="1000000.00", amd="5000000.00"
    )
    b = fx.contract_with_amendment(
        db_session, factories, base_rate=zero, amd_rate=zero, base="2000000.00", amd="4000000.00"
    )
    c = fx.contract_with_amendment(
        db_session, factories, base_rate=zero, amd_rate=zero, base="3000000.00", amd="3000000.00"
    )
    db_session.flush()

    data = cmp.build_comparison(db_session, [a.id, b.id, c.id], vat_mode=cmp.VAT_MODE_NET)
    ws = _sheet_from(data)
    columns_index = _columns_index(data)

    row_1 = next(r for r in data["rows"] if r["code"] == "1")
    row_num = _find_row_by_label(ws, _category_title(db_session, "1"))

    a_cell = next(cell for cell in row_1["cells"] if cell["contract_id"] == a.id)
    dev_base = a_cell[cmp.BUCKET_BASE]["deviation_pct"]
    dev_amd = a_cell[cmp.BUCKET_AMENDMENTS]["deviation_pct"]
    dev_total = a_cell[cmp.BUCKET_TOTAL]["deviation_pct"]
    assert dev_base is not None and dev_amd is not None and dev_total is not None
    assert len({dev_base, dev_amd, dev_total}) == 3, "три корзины обязаны нести три РАЗНЫХ отклонения"

    idx_a = columns_index[a.id]
    assert Decimal(str(ws.cell(row=row_num, column=_col(idx_a, cmp.BUCKET_BASE, "deviation")).value)) == dev_base
    assert Decimal(str(ws.cell(row=row_num, column=_col(idx_a, cmp.BUCKET_AMENDMENTS, "deviation")).value)) == dev_amd
    assert Decimal(str(ws.cell(row=row_num, column=_col(idx_a, cmp.BUCKET_TOTAL, "deviation")).value)) == dev_total

    # Заголовки колонок называют все три корзины и три метрики — DoD 21.
    text = _text_of(ws)
    for label in ("ДГП · Сумма", "ДГП · ₽/м²", "ДГП · Откл., %",
                  "ДС · Сумма", "ДС · ₽/м²", "ДС · Откл., %",
                  "Итого · Сумма", "Итого · ₽/м²", "Итого · Откл., %"):
        assert label in text, f"в заголовках нет {label!r}"


# ---------------------------------------------------------------------------
#  DoD 20 — подпись налогового состава на листе, включая состав по договору
# ---------------------------------------------------------------------------

def test_vat_mode_caption_and_per_contract_composition_are_on_the_sheet(db_session, factories):
    """Спека §2.3.1/§2.3.2, AGENTS.md §10 v6.8: без этой подписи лист позволил
    бы читателю сложить валовые суммы разных ставок, не заметив разницы."""
    c20 = fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
    c22 = fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_22)
    db_session.flush()

    data = cmp.build_comparison(db_session, [c20, c22], vat_mode=cmp.VAT_MODE_OWN)
    ws = _sheet_from(data)
    text = _text_of(ws)

    assert data["caption"] in text
    assert "Каждый договор показан в своей действующей ставке НДС" in text

    columns_by_id = {column["contract_id"]: column for column in data["columns"]}
    assert columns_by_id[c20]["composition_caption"] in text
    assert columns_by_id[c22]["composition_caption"] in text
    assert "20 %" in text
    assert "22 %" in text


def test_net_mode_caption_does_not_claim_per_contract_composition(db_session, factories):
    """Режим «нетто» не несёт состава по договору — состав нужен только там,
    где показ смешивает разные ставки (§2.3.2); в нетто ось и так одна."""
    c = fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
    db_session.flush()

    data = cmp.build_comparison(db_session, [c], vat_mode=cmp.VAT_MODE_NET)
    ws = _sheet_from(data)
    text = _text_of(ws)

    assert data["caption"] in text
    assert "Суммы показаны без НДС" in text


# ---------------------------------------------------------------------------
#  Синтетические строки — те же правила, что на экране
# ---------------------------------------------------------------------------

def test_synthetic_rows_appear_under_the_same_rules_as_the_screen(db_session, factories):
    """«Без подстатьи» — статья с прямыми деньгами И детьми; «Нераспределённое»
    — остаток вне классификатора (спека §2.1, §2.1.4)."""
    estimate = factories.EstimateFactory.create()
    proposal = fx.make_proposal(factories, estimate=estimate)
    fx.seed_chapter_with_positions(db_session, factories, proposal=proposal, code="3", amounts=["500.00"])
    fx.seed_chapter_with_positions(db_session, factories, proposal=proposal, code="3.1", amounts=["100.00"])
    fx.seed_additional_work(db_session, factories, proposal=proposal, code="3", amount="200.00")
    fx.seed_unallocated_positions(db_session, factories, proposal=proposal, amounts=["50.00"])
    db_session.flush()

    data = cmp.build_comparison(db_session, [estimate.contract_id], vat_mode=cmp.VAT_MODE_NET)
    ws = _sheet_from(data)

    # Строки листа обязаны существовать под теми же заголовками, что уже
    # проверены `crud.comparison.build_rows` (задача 3) — здесь проверяется
    # только то, что лист их не потерял при печати.
    assert any(r["kind"] == "own" for r in data["rows"]), "предпосылка: строка «Без подстатьи» есть в агрегате"
    assert any(r["kind"] == "unallocated" for r in data["rows"]), "предпосылка: «Нераспределённое» есть в агрегате"

    _find_row_by_label(ws, cmp.OWN_ROW_TITLE)  # не бросает — строка на листе есть
    _find_row_by_label(ws, cmp.UNALLOCATED_ROW_TITLE)


# ---------------------------------------------------------------------------
#  Три различимых написания ячейки: прочерк / ноль / прочерк с причиной
# ---------------------------------------------------------------------------

def test_dash_zero_and_reasoned_blank_are_three_distinct_renderings(db_session, factories):
    """Спека §2.1.2, §2.1.3: статьи нет — прочерк; статья есть и стоит ноль —
    настоящий ноль; статья есть, но погашена причиной — прочерк со словами."""
    contract_with_zero = fx.contract_with(db_session, factories, {"1": ["0.00"], "2": ["100.00"]})

    contract_without_one = fx.contract_with(db_session, factories, {"2": ["50.00"]})

    estimate_unknown_base = factories.EstimateFactory.create()
    proposal_unknown_base = fx.make_proposal(factories, estimate=estimate_unknown_base, vat_rate=None)
    fx.seed_chapter_with_positions(
        db_session, factories, proposal=proposal_unknown_base, code="1", amounts=["100.00"]
    )
    db_session.flush()

    ids = [contract_with_zero, contract_without_one, estimate_unknown_base.contract_id]
    data = cmp.build_comparison(db_session, ids, vat_mode=cmp.VAT_MODE_NET)
    ws = _sheet_from(data)
    columns_index = _columns_index(data)

    row_1 = next(r for r in data["rows"] if r["code"] == "1")
    row_num = _find_row_by_label(ws, _category_title(db_session, "1"))

    zero_cell = next(c for c in row_1["cells"] if c["contract_id"] == contract_with_zero)
    absent_cell = next(c for c in row_1["cells"] if c["contract_id"] == contract_without_one)
    blanked_cell = next(c for c in row_1["cells"] if c["contract_id"] == estimate_unknown_base.contract_id)

    assert zero_cell[cmp.BUCKET_TOTAL]["state"] == cmp.ZERO
    assert absent_cell[cmp.BUCKET_TOTAL]["state"] == cmp.ABSENT
    assert blanked_cell[cmp.BUCKET_TOTAL]["incomplete_reasons"] == ["vat_base_unknown"]

    zero_value = ws.cell(row=row_num, column=_col(columns_index[contract_with_zero], cmp.BUCKET_TOTAL, "sum")).value
    absent_value = ws.cell(
        row=row_num, column=_col(columns_index[contract_without_one], cmp.BUCKET_TOTAL, "sum")
    ).value
    blanked_value = ws.cell(
        row=row_num, column=_col(columns_index[estimate_unknown_base.contract_id], cmp.BUCKET_TOTAL, "sum")
    ).value

    assert zero_value == Decimal("0")
    assert isinstance(zero_value, int | float | Decimal), "настоящий ноль обязан быть числом, не текстом"
    assert absent_value == "—"
    assert blanked_value == "— (неизвестна база НДС)"
    # Три написания различимы попарно.
    assert len({str(zero_value), str(absent_value), str(blanked_value)}) == 3


# ---------------------------------------------------------------------------
#  Эндпоинт: медиатип, кириллица в имени файла, доступ member
# ---------------------------------------------------------------------------

class TestComparisonReportEndpoint:
    def test_endpoint_returns_xlsx_media_type_and_filename(self, client, factories, db_session):
        contract_id = fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
        db_session.flush()

        response = client.get(
            "/api/v1/reports/comparison",
            params={"ids": str(contract_id), "vat_mode": "net"},
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        disposition = unquote(response.headers["content-disposition"])
        assert "filename*=UTF-8''" in response.headers["content-disposition"]
        assert "Сравнение договоров" in disposition

    def test_filename_names_the_count_instead_of_gluing_every_number(
        self, client, factories, db_session
    ):
        """Предела на число колонок у сравнения нет (спека §3 отвергает «жёсткий
        предел 5-6 колонок»), поэтому склейка всех номеров упирается в предел
        длины пути, а не во вкус: имя из нескольких сотен символов не сохранится.
        Дальше порога имя называет ЧИСЛО договоров — обрезанный по середине
        список выглядел бы как полный.
        """
        ids = [
            fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
            for _ in range(4)
        ]
        db_session.flush()

        response = client.get(
            "/api/v1/reports/comparison",
            params={"ids": ",".join(str(i) for i in ids), "vat_mode": "net"},
        )

        assert response.status_code == 200, response.text
        disposition = unquote(response.headers["content-disposition"])
        assert "Сравнение договоров (4).xlsx" in disposition

    def test_empty_selection_does_not_produce_a_nonsense_filename(
        self, client, factories, db_session
    ):
        """Пустая выборка — законный ответ, и имя обязано это признавать, а не
        подставлять умолчание `_safe_filename_part` («…договоров договор.xlsx»)."""
        fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
        db_session.flush()

        response = client.get(
            "/api/v1/reports/comparison",
            params={"all": "1", "q": "такого объекта нет", "vat_mode": "net"},
        )

        assert response.status_code == 200, response.text
        disposition = unquote(response.headers["content-disposition"])
        assert "Сравнение договоров.xlsx" in disposition

    @pytest.mark.parametrize(
        "params",
        [
            pytest.param({"ids": "1,abc,3"}, id="broken-ids"),
            pytest.param({"ids": "1", "rate_class_id": "2,abc"}, id="broken-rate-class-id"),
            # Форма `all=1`, а не `ids`: у `ids` раньше сработала бы проверка
            # существования (404 «Договоры не найдены»), потому что порядок
            # отказов в `resolve_selection` — двусмысленность, потом форма, и
            # только потом сужение. Разбор адреса при этом опережает всё: он
            # живёт в роутере, до резолвера, — поэтому случай `2,abc` выше
            # спокойно обходится несуществующим id.
            pytest.param({"all": "1", "rate_class_id": ""}, id="empty-rate-class-id"),
        ],
    )
    def test_both_endpoints_answer_a_broken_param_the_same_way(self, client, params):
        """Один контракт выборки на два эндпоинта (спека §2.6, §2.7).

        Разбор `ids` был реализован ДВАЖДЫ — в роутере экрана и в роутере
        выгрузки, с разными сообщениями; один и тот же адрес получал два разных
        ответа в зависимости от того, куда его послали. Теперь разбор один
        (`crud.comparison.parse_ids_param`), и тест это стережёт: разойдись они
        снова — статусы или текст перестанут совпадать.

        **`rate_class_id` добавлен в тот же сторож** (задача 7): это ВТОРОЙ
        разбираемый параметр адреса с тем же режимом отказа, и у него ровно тот
        же способ разъехаться. Без него отказные пути маршрута ЛИСТА не были
        покрыты вовсе — он проверялся только счастливым путём DoD 6.

        Сравнивается `detail` ЦЕЛИКОМ, а не подстрокой: это заодно различает, из
        какого места пришёл отказ, тогда как подстрока «abc» есть в сообщениях и
        разбора адреса, и страховки фильтра.
        """
        screen = client.get("/api/v1/analytics/comparison", params=params)
        sheet = client.get("/api/v1/reports/comparison", params=params)

        assert screen.status_code == 400, screen.text
        assert sheet.status_code == 400, sheet.text
        assert screen.json()["detail"] == sheet.json()["detail"]

    def test_sheet_columns_match_the_screen_with_rate_class_id_filter(
        self, client, factories, db_session
    ):
        """DoD 6 (задача 7): лист на ТЕХ ЖЕ параметрах, включая многозначный
        `rate_class_id`, несёт тот же состав договоров, что экран (спека §2.7,
        «один агрегат — два представления»).

        Сравнивается состав колонок листа с составом колонок ответа
        `/analytics/comparison` на одинаковом query-string, а не арифметика
        внутри них — она уже покрыта DoD 19 выше.
        """
        class_a = factories.RateClassFactory.create()
        class_b = factories.RateClassFactory.create()
        other_class = factories.RateClassFactory.create()
        kept_a = factories.ContractFactory.create(rate_class=class_a)
        kept_b = factories.ContractFactory.create(rate_class=class_b)
        excluded = factories.ContractFactory.create(rate_class=other_class)
        db_session.commit()

        query = (
            f"ids={kept_a.id},{kept_b.id},{excluded.id}"
            f"&rate_class_id={class_a.id},{class_b.id}"
        )

        screen = client.get(f"/api/v1/analytics/comparison?{query}")
        sheet = client.get(f"/api/v1/reports/comparison?{query}")

        assert screen.status_code == 200, screen.text
        assert sheet.status_code == 200, sheet.text

        screen_ids = {column["contract_id"] for column in screen.json()["columns"]}
        assert screen_ids == {kept_a.id, kept_b.id}

        ws = load_workbook(BytesIO(sheet.content)).active
        text = _text_of(ws)
        assert kept_a.contract_number in text
        assert kept_b.contract_number in text
        assert excluded.contract_number not in text

    def test_member_can_download_the_report(self, client, factories, db_session):
        contract_id = fx.contract_with_area(db_session, factories, {"1": ["100.00"]}, vat_rate=fx.VAT_20)
        db_session.flush()

        client.auth_state["role"] = UserRole.member
        response = client.get(
            "/api/v1/reports/comparison",
            params={"ids": str(contract_id), "vat_mode": "net"},
        )
        assert response.status_code == 200, response.text


# ---------------------------------------------------------------------------
#  Лист выгрузки печатает ряд (спека инфляции §2.10, §2.12; план, задача 10)
# ---------------------------------------------------------------------------

def _inflation_sheet(db, factories, *, note=None, forecast=(2026,)):
    """Лист по приведённому агрегату и сам агрегат — парой, чтобы тест сравнивал
    напечатанное с тем, что пришло в словаре, а не с пересчитанным заново."""
    series_id = fx.series_with_years(
        db, "Росстат, ИПЦ, декабрь к декабрю",
        {2024: "1.0750", 2025: "1.0830", 2026: "1.0600"},
        note=note, forecast=forecast, source="бюллетень 01.2026",
    )
    ids = [
        fx.contract_with_dates(
            db, factories, {"1": ["1200000.00"]}, signed_date=signed_date
        )
        for signed_date in (dt.date(2024, 5, 28), dt.date(2026, 7, 6))
    ]
    data = cmp.build_comparison(
        db, ids, vat_mode=cmp.VAT_MODE_OWN,
        inflation_series_id=series_id, target_month=YearMonth(2026, 8),
    )
    return data, _sheet_from(data)


def test_sheet_prints_the_series_and_the_target_month(db_session, factories):
    """Ряд и целевой месяц — на самом листе (DoD 13)."""
    _data, ws = _inflation_sheet(db_session, factories)
    text = _text_of(ws)

    assert "Цены приведены к: август 2026" in text
    assert "Росстат, ИПЦ, декабрь к декабрю" in text


def test_sheet_prints_five_facts_for_every_used_year(db_session, factories):
    """Пять фактов на КАЖДЫЙ использованный год (DoD 13).

    Версионности у ряда нет (§2.10): воспроизводимость лежит на файле, и год без
    источника либо без даты правки эту роль не исполнил бы.
    """
    data, ws = _inflation_sheet(db_session, factories)
    text = _text_of(ws)

    assert [item["year"] for item in data["inflation"]["used_years"]] == [2024, 2025, 2026]
    for item in data["inflation"]["used_years"]:
        line = next(
            (row for row in _text_of(ws).splitlines() if row.startswith(f"{item['year']} · ")),
            None,
        )
        assert line is not None, f"строки года {item['year']} на листе нет"
        assert format(item["coefficient"], "f") in line
        assert item["source"] in line
        assert ("прогноз" if item["is_forecast"] else "факт") in line
        assert "правлен " in line

    # Прогнозный год назван прогнозом, фактический — фактом, и это РАЗНЫЕ слова
    # в одном и том же месте строки (DoD 14).
    assert "2026 · 1.0600 · бюллетень 01.2026 · прогноз" in text
    assert "2024 · 1.0750 · бюллетень 01.2026 · факт" in text


def test_sheet_prints_the_year_source_which_the_screen_does_not(db_session, factories):
    """Источник ГОДА есть на листе (DoD 32, половина про лист).

    Парная половина — «на экране сравнения его нет» — живёт в тесте экрана
    (задача 14): здесь её проверить нечем, а утверждать обе стороны в одном месте
    значило бы утверждать про экран из теста про файл.
    """
    data, ws = _inflation_sheet(db_session, factories)

    assert "бюллетень 01.2026" in _text_of(ws)
    # Примечание ряда — наоборот, дело экрана: на листе его нет.
    assert data["inflation"]["series_note"] is None


def test_sheet_prints_the_series_note_nowhere_but_keeps_the_screen_field(
    db_session, factories
):
    """Примечание ряда приходит в агрегате (его показывает полоса уровней), а на
    лист не печатается: состав листа и экрана РАЗНЫЙ, и это разделение, а не дубль."""
    data, ws = _inflation_sheet(db_session, factories, note="официальная публикация, по РФ")
    text = _text_of(ws)

    assert data["inflation"]["series_note"] == "официальная публикация, по РФ"
    assert "официальная публикация" not in text
    # Блок приведения при этом НАПЕЧАТАН, и признак берётся из самого блока:
    # название ряда попадает на лист ещё и через подпись оси, поэтому по нему
    # различить «блок есть» и «блока нет» нельзя — замерено снятием 2026-08-19,
    # тест оставался зелёным. «правлен ДД.ММ.ГГГГ» печатает только блок.
    assert "правлен " in text


def test_sheet_without_inflation_is_unchanged_cell_by_cell(db_session, factories):
    """Без блока `inflation` лист не меняется НИ НА ОДНУ ячейку (DoD 1).

    Сравниваются координаты и значения всех непустых ячеек, а не байты: XLSX —
    ZIP-контейнер, его байты расходятся метаданными архива при идентичном
    содержимом. Тот же приём, что у golden-снимка задачи 1; здесь он повторён на
    той же выборке до и после включения приведения, чтобы отличить «лист не
    изменился» от «снимок не заведён».
    """
    ids = [
        fx.contract_with_dates(
            db_session, factories, {"1": ["1200000.00"]}, signed_date=signed_date
        )
        for signed_date in (dt.date(2024, 5, 28), dt.date(2026, 7, 6))
    ]
    nominal = cmp.build_comparison(db_session, ids, vat_mode=cmp.VAT_MODE_OWN)

    first = _cells(_sheet_from(nominal))
    second = _cells(_sheet_from(nominal))
    assert first == second

    series_id = fx.series_with_years(
        db_session, "Ряд для сравнения листов", {2024: "1.075", 2025: "1.083", 2026: "1.060"}
    )
    adjusted = cmp.build_comparison(
        db_session, ids, vat_mode=cmp.VAT_MODE_OWN,
        inflation_series_id=series_id, target_month=YearMonth(2026, 8),
    )
    # Лист с приведением обязан отличаться, и отличаться ИМЕННО блоком: числа в
    # нём другие сами по себе, поэтому одного `!=` не хватает — оно осталось бы
    # зелёным и при неработающем блоке (проверено снятием 2026-08-19).
    adjusted_text = _text_of(_sheet_from(adjusted))
    assert _cells(_sheet_from(adjusted)) != first
    assert "Ряд для сравнения листов" in adjusted_text
    assert "Цены приведены к: август 2026" in adjusted_text


def test_sheet_says_so_when_no_year_was_needed(db_session, factories):
    """Цель совпала с месяцем сметы: годов нет, но лист об этом ГОВОРИТ (DoD 5).

    Молчание читалось бы как «ряд не доехал до листа», хотя приведение включено и
    сосчитано.
    """
    series_id = fx.series_with_years(db_session, "Ряд без нужных годов", {})
    contract_id = fx.contract_with_dates(
        db_session, factories, {"1": ["1200000.00"]}, signed_date=dt.date(2026, 8, 15)
    )
    data = cmp.build_comparison(
        db_session, [contract_id], vat_mode=cmp.VAT_MODE_OWN,
        inflation_series_id=series_id, target_month=YearMonth(2026, 8),
    )

    text = _text_of(_sheet_from(data))
    assert "Цены приведены к: август 2026" in text
    assert "не потребовались" in text
