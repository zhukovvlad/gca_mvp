"""Предпросмотр, выгрузка и отправленные файлы подрядчику: четыре маршрута
`…/our-price/offers/{offer_id}/{export-preview | export | sent-files | sent-files/{id}/file}`
(спека 4б §2.8-§2.11, §2.15).

Сцены строятся помощниками экрана этапа; названия и числа выдуманы, эталоны — числа и
строки, посчитанные руками. Каждое правило проверено с обеих сторон: вход, где оно
срабатывает, и соседний вход, нарушающий ровно одно условие. Хранилище — каталог теста;
запись сцены закрепляется перед каждым запросом: откат роутера не должен уносить саму сцену.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import unquote

import pytest
import sqlalchemy as sa
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from crud import comparison as cmp
from crud import kp_sent_files
from crud.our_price_inflation import plan_history_inflation
from models import (
    Estimate,
    InflationIndexValue,
    KpSentFile,
    PositionItem,
    SemanticJob,
    UnitOfMeasure,
)
from money.inflation import YearMonth
from services.excel_kp_file import KP_FILE_COLUMNS
from tests import comparison_fixtures as fx
from tests.integration.test_kp_file_scene import _agree, _comment, _user
from tests.integration.test_our_price_load import World
from tests.integration.test_our_price_offer_api import (
    FACTOR_FEB_2025,
    FULL_YEARS,
    SERIES_NAME,
    _inflation_scene,
)
from tests.integration.test_our_price_offer_api import _system_scene as _offer_system_scene
from tests.integration.test_our_price_round_api import Stage, _dated_history, _json, count_statements
from tests.integration.test_our_price_signal import _job

pytestmark = pytest.mark.integration

D = Decimal

HEADER_KEYS = {
    "tender",
    "object",
    "round",
    "contractor",
    "vat_label",
    "source_label",
    "level_label",
    "in_progress_lines",
    "inflation_years",
    "filename",
    "exported_on",
    "inflation_series",
    "inflation_target_month",
}
SECTION_KEYS = {"kind", "number", "title"}
LINE_KEYS = {
    "kind",
    "position_item_id",
    "number",
    "title",
    "unit",
    "quantity",
    "kp_unit_price",
    "kp_total",
    "ours",
    "deviation_pct",
    "source",
    "inflation",
    "comment",
    "comment_kind",
    "state",
}
COLUMN_KEYS = [
    "number",
    "title",
    "unit",
    "quantity",
    "kp_unit_price",
    "kp_total",
    "ours",
    "deviation_pct",
    "source",
    "inflation",
    "comment",
]
SENT_KEYS = {"id", "filename", "created_at", "created_by", "vat_rate", "level_label", "size_bytes"}


@pytest.fixture
def world(db_session, factories, admin_user):
    return World(db_session, factories, admin_user)


@pytest.fixture
def stage(world, db_session, factories):
    return Stage(world, db_session, factories)


@pytest.fixture
def member_client(member_client, db_session, tmp_storage):
    """Клиент `member`, который перед каждым запросом закрепляет сцену и пишет в каталог теста:
    `rollback` роутера откатывает транзакцию сессии, а в проде сцена к запросу уже закоммичена."""
    from main import app
    from storage import get_storage

    send = member_client.request

    def request(*args, **kwargs):
        Session.commit(db_session)
        return send(*args, **kwargs)

    member_client.request = request
    app.dependency_overrides[get_storage] = lambda: tmp_storage
    return member_client


@pytest.fixture
def saved_keys(tmp_storage, monkeypatch):
    """Ключи, под которыми хранилище сохранило файлы за время теста."""
    keys: list[str] = []
    save = tmp_storage.save

    def spy(data):
        key = save(data)
        keys.append(key)
        return key

    monkeypatch.setattr(tmp_storage, "save", spy)
    return keys


# ---------------------------------------------------------------------------
#  Помощники
# ---------------------------------------------------------------------------


def _base(stage, rnd, part) -> str:
    return f"/api/v1/tenders/{stage.tender.id}/rounds/{rnd.id}/our-price/offers/{part.offer.id}"


def _preview_get(client, stage, rnd, part, **params):
    return client.get(f"{_base(stage, rnd, part)}/export-preview", params=params)


def _preview(client, stage, rnd, part, **params) -> dict:
    response = _preview_get(client, stage, rnd, part, **params)
    assert response.status_code == 200, response.text
    return _json(response)


def _export(client, stage, rnd, part, **params):
    return client.post(f"{_base(stage, rnd, part)}/export", params=params)


def _export_ok(client, stage, rnd, part, **params):
    response = _export(client, stage, rnd, part, **params)
    assert response.status_code == 200, response.text
    return response


def _sent(client, stage, rnd, part):
    return client.get(f"{_base(stage, rnd, part)}/sent-files")


def _download(client, stage, rnd, part, file_id):
    return client.get(f"{_base(stage, rnd, part)}/sent-files/{file_id}/file")


def _code(response) -> str:
    detail = response.json()["detail"]
    assert isinstance(detail, dict), detail
    return detail["code"]


def _records(db):
    return db.execute(
        sa.select(
            KpSentFile.id,
            KpSentFile.offer_id,
            KpSentFile.file_key,
            KpSentFile.filename,
            KpSentFile.file_sha256,
            KpSentFile.size_bytes,
            KpSentFile.vat_rate,
            KpSentFile.level_label,
            KpSentFile.created_by,
        ).order_by(KpSentFile.id)
    ).all()


def _stored(storage) -> list[str]:
    return sorted(path.name for path in storage.root.glob("*")) if storage.root.exists() else []


def _disposition_name(response) -> str:
    header = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in header, header
    return unquote(header.split("filename*=UTF-8''", 1)[1])


def _dec(value) -> Decimal | None:
    return None if value is None else D(value)


def _unit_name(db, world) -> str:
    return db.get(UnitOfMeasure, world.unit_id).name


def _chapter(stage, proposal, number, title):
    return stage.f.PositionItemFactory.create(
        proposal=proposal,
        is_chapter=True,
        job_title_in_proposal=title,
        chapter_number_in_proposal=number,
        unit_cost_total=None,
        total_cost_total=None,
        quantity=None,
        suggested_quantity=None,
    )


def _kp_line(stage, proposal, chapter, price, *, number, qty=D("10"), variant=True):
    """Строка КП в главе с номером; `variant=False` — без семьи и варианта."""
    item = stage.world.line(
        proposal,
        chapter=chapter,
        unit_cost=price,
        total=price * qty,
        suggested=qty,
        quantity=qty,
        family=stage.family if variant else None,
        variant=stage.variant if variant else None,
        number=number,
    )
    stage.db.execute(sa.delete(SemanticJob).where(SemanticJob.context_id == item.context_id))
    return item


def _file_scene(stage, db_session, world, factories):
    """Смета КП (НДС 20 %) из двух лотов. История класса — 100 и 130 за единицу; авто-расценка —
    минимум 100 нетто, то есть 120 с НДС. Позиции лота 2 заведены РАНЬШЕ позиций лота 1, поэтому
    порядок `(лот, id)` отличается от порядка по одному `id`.

    Лот 2: глава 2; «2.1» 120 с НДС (нетто 100, отклонение 0,0); «2.2» без семьи.
    Лот 1: глава 1; «1.1» 180 (нетто 150, 50,0); «1.2» 240 (нетто 200) с «Договорились» 80 нетто
    (расценка 96, 150,0); «1.3» 210 (нетто 175, 75,0) со своим комментарием «=1+1».
    """
    c1, _, _ = stage.history(D("100"))
    stage.history(D("130"))
    rnd = stage.round()
    part = stage.participant(rnd, "Подрядчик А", vat=D("20"))
    lot2 = factories.LotFactory.create(estimate=part.estimate)
    proposal2 = factories.ProposalFactory.create(lot=lot2, contractor=part.contractor, vat_rate=D("20"))
    chapter2 = _chapter(stage, proposal2, "2", "Раздел 2")
    l3 = _kp_line(stage, proposal2, chapter2, D("120"), number="2.1")
    l4 = _kp_line(stage, proposal2, chapter2, D("50"), number="2.2", variant=False)
    chapter1 = _chapter(stage, part.proposal, "1", "Раздел 1")
    l1 = _kp_line(stage, part.proposal, chapter1, D("180"), number="1.1")
    l2 = _kp_line(stage, part.proposal, chapter1, D("240"), number="1.2", qty=D("5"))
    l5 = _kp_line(stage, part.proposal, chapter1, D("210"), number="1.3", qty=D("4"))
    user = _user(factories)
    _agree(db_session, l2, D("80"), user)
    _comment(db_session, l5, "=1+1", user)
    db_session.flush()
    return SimpleNamespace(
        rnd=rnd,
        part=part,
        contract=c1,
        chapters=(chapter1, chapter2),
        l1=l1,
        l2=l2,
        l3=l3,
        l4=l4,
        l5=l5,
        unit=_unit_name(db_session, world),
    )


def _by_item(body: dict, item) -> dict:
    (found,) = (row for row in body["rows"] if row.get("position_item_id") == item.id)
    return found


def _table(content: bytes):
    """Книга: заголовки таблицы и её строки (значения ячеек), без шапки листа."""
    ws = load_workbook(io.BytesIO(content)).active
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    start = next(index for index, row in enumerate(rows) if row[0] == "№")
    return rows[start], rows[start + 1 :], ws


def _num(cell) -> Decimal | None:
    return None if cell is None else D(str(cell))


# ---------------------------------------------------------------------------
#  Маршруты: отказы и доступ
# ---------------------------------------------------------------------------

ROUTES = ("preview", "export", "list", "download")


def _call(route, client, stage, rnd, part, **params):
    if route == "preview":
        return _preview_get(client, stage, rnd, part, **params)
    if route == "export":
        return _export(client, stage, rnd, part, **params)
    if route == "list":
        return _sent(client, stage, rnd, part)
    return _download(client, stage, rnd, part, 987654321)


class TestRouteGuards:
    @pytest.mark.parametrize("route", ROUTES)
    def test_unknown_tender_is_404(self, member_client, stage, route):
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А")
        url = f"/api/v1/tenders/987654321/rounds/{rnd.id}/our-price/offers/{part.offer.id}"
        url += {
            "preview": "/export-preview",
            "export": "/export",
            "list": "/sent-files",
            "download": "/sent-files/1/file",
        }[route]

        response = member_client.post(url) if route == "export" else member_client.get(url)

        assert response.status_code == 404

    @pytest.mark.parametrize("route", ROUTES)
    def test_round_of_another_tender_is_404(self, member_client, stage, factories, route):
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А")
        foreign = factories.TenderRoundFactory.create()

        assert _call(route, member_client, stage, foreign, part).status_code == 404

    @pytest.mark.parametrize("route", ROUTES)
    def test_offer_of_another_round_of_the_same_tender_is_404(self, member_client, stage, route):
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А")
        other = stage.round()

        assert _call(route, member_client, stage, other, part).status_code == 404

    @pytest.mark.parametrize("route", ROUTES)
    def test_offer_of_another_tender_is_404(self, member_client, stage, factories, route):
        rnd = stage.round()
        foreign = factories.OfferFactory.create()

        response = _call(route, member_client, stage, rnd, SimpleNamespace(offer=foreign))

        assert response.status_code == 404

    @pytest.mark.parametrize("route", ["preview", "export"])
    def test_participant_without_an_estimate_is_404(self, member_client, stage, route):
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", estimate=False)

        assert _call(route, member_client, stage, rnd, part).status_code == 404

    def test_participant_without_an_estimate_still_lists_an_empty_set(self, member_client, stage):
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", estimate=False)

        response = _sent(member_client, stage, rnd, part)

        assert (response.status_code, response.json()) == (200, [])

    @pytest.mark.parametrize("route", ["preview", "export"])
    def test_unknown_vat_base_is_422_without_a_record_or_a_file(
        self, member_client, stage, db_session, tmp_storage, saved_keys, route
    ):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=None)
        stage.kp(part, D("120"))

        response = _call(route, member_client, stage, rnd, part)

        assert (response.status_code, _code(response)) == (422, "unknown_vat_base")
        assert _records(db_session) == []
        assert saved_keys == []
        assert _stored(tmp_storage) == []

    def test_known_base_next_to_the_unknown_one_goes_through(self, member_client, stage):
        stage.history(D("100"))
        rnd = stage.round()
        unknown = stage.participant(rnd, "Подрядчик Б", vat=None)
        known = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(unknown, D("120"))
        stage.kp(known, D("120"))

        assert _preview_get(member_client, stage, rnd, known).status_code == 200

    @pytest.mark.parametrize("route", ROUTES)
    def test_member_reads_and_exports(self, member_client, stage, route):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        file_id = None
        if route == "download":
            _export_ok(member_client, stage, rnd, part)
            file_id = _sent(member_client, stage, rnd, part).json()[0]["id"]

        response = (
            _download(member_client, stage, rnd, part, file_id)
            if route == "download"
            else _call(route, member_client, stage, rnd, part)
        )

        assert response.status_code == 200


# ---------------------------------------------------------------------------
#  Чтение не пишет
# ---------------------------------------------------------------------------


class TestReadsWriteNothing:
    def test_preview_list_and_download_issue_no_insert_update_or_delete(self, member_client, stage, db_session):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        _export_ok(member_client, stage, rnd, part)
        file_id = _sent(member_client, stage, rnd, part).json()[0]["id"]
        Session.commit(db_session)

        with count_statements(db_session) as seen:
            assert _preview_get(member_client, stage, rnd, part).status_code == 200
            assert _sent(member_client, stage, rnd, part).status_code == 200
            assert _download(member_client, stage, rnd, part, file_id).status_code == 200

        writes = [s for s in seen if s.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))]
        assert writes == []
        assert any(s.lstrip().upper().startswith("SELECT") for s in seen)

    def test_the_export_does_write_a_record(self, member_client, stage, db_session):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        Session.commit(db_session)

        with count_statements(db_session) as seen:
            _export_ok(member_client, stage, rnd, part)

        assert any(s.lstrip().upper().startswith("INSERT INTO KP_SENT_FILES") for s in seen)


# ---------------------------------------------------------------------------
#  Предпросмотр: форма
# ---------------------------------------------------------------------------


class TestPreviewShape:
    def test_the_header_columns_and_rows_are_the_contract_of_the_spec(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        body = _preview(member_client, stage, scene.rnd, scene.part)

        assert set(body) == {"header", "columns", "rows"}
        assert set(body["header"]) == HEADER_KEYS
        for row in body["rows"]:
            assert set(row) == (SECTION_KEYS if row["kind"] == "section" else LINE_KEYS)

    def test_the_columns_are_key_and_title_in_the_order_of_the_sheet_without_a_series(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        titles = [spec[0] for spec in KP_FILE_COLUMNS if spec[0] != "Приведение по ИПЦ"]
        keys = [key for key in COLUMN_KEYS if key != "inflation"]

        body = _preview(member_client, stage, scene.rnd, scene.part)

        assert body["columns"] == [{"key": k, "title": t} for k, t in zip(keys, titles, strict=True)]

    def test_the_inflation_column_exists_only_with_a_series(self, member_client, stage, db_session):
        rnd, part, _, series_id = _inflation_scene(stage, db_session, "ok")

        with_series = _preview(member_client, stage, rnd, part, inflation_series_id=series_id)
        without = _preview(member_client, stage, rnd, part)

        assert [c["key"] for c in with_series["columns"]] == COLUMN_KEYS
        assert [c["title"] for c in with_series["columns"]] == [spec[0] for spec in KP_FILE_COLUMNS]
        assert [c["key"] for c in without["columns"]] == [k for k in COLUMN_KEYS if k != "inflation"]

    def test_the_response_carries_no_float_literal(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)

        # `_json` роняет тест на дробном литерале JSON
        _preview(member_client, stage, scene.rnd, scene.part)

    def test_the_header_values(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)
        today = dt.date.today()

        header = _preview(member_client, stage, scene.rnd, scene.part)["header"]

        assert header["tender"] == stage.tender.title
        assert header["object"] == stage.tender.object.title
        assert header["round"] == "Этап 1 от 10.08.2026"
        assert header["contractor"] == "Подрядчик А"
        assert header["exported_on"] == today.strftime("%d.%m.%Y")
        assert header["vat_label"] == "Цены КП и наши расценки — с НДС 20 % (база НДС этого КП)"
        assert header["source_label"] == "Цена источника — как записана в договоре, без НДС"
        assert header["level_label"] == "номинальный, без приведения"
        assert header["in_progress_lines"] == 0
        assert (header["inflation_years"], header["inflation_series"], header["inflation_target_month"]) == (
            None,
            None,
            None,
        )
        assert header["filename"] == f"КП Подрядчик А этап 1 — наши расценки {today:%Y-%m-%d}.xlsx"

    def test_a_zero_base_is_named_in_the_header(self, member_client, stage):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))

        header = _preview(member_client, stage, rnd, part)["header"]

        assert header["vat_label"] == "Цены КП и наши расценки — без НДС (база НДС этого КП — 0 %)"

    def test_the_stage_without_a_date_has_no_date_in_the_header(self, member_client, stage):
        stage.history(D("100"))
        rnd = stage.round(held_on=None)
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))

        assert _preview(member_client, stage, rnd, part)["header"]["round"] == "Этап 1"

    def test_the_contractor_title_is_made_safe_in_the_file_name_only(self, member_client, stage):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "ООО Альфа/Бета", vat=D("0"))
        stage.kp(part, D("120"))

        header = _preview(member_client, stage, rnd, part)["header"]

        assert header["contractor"] == "ООО Альфа/Бета"
        assert header["filename"].startswith("КП ООО Альфа-Бета этап 1 — ")

    def test_a_line_in_progress_is_counted_and_has_no_numbers_of_ours(self, member_client, stage, db_session):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        waiting = stage.kp(part, D("144"), variant=False)
        stage.kp(part, D("120"))
        _job(db_session, "family_suggestion", "pending", context_id=waiting.context_id)

        body = _preview(member_client, stage, rnd, part)

        row = _by_item(body, waiting)
        assert (row["state"], row["ours"], row["deviation_pct"], row["source"]) == ("in_progress", None, None, None)
        assert body["header"]["in_progress_lines"] == 1


# ---------------------------------------------------------------------------
#  Предпросмотр: строки
# ---------------------------------------------------------------------------


class TestPreviewRows:
    def test_sections_and_lines_follow_the_lot_then_the_id_with_the_chapters_in_place(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        body = _preview(member_client, stage, scene.rnd, scene.part)

        assert [(r["kind"], r["number"], r["title"] if r["kind"] == "section" else None) for r in body["rows"]] == [
            ("section", "1", "Раздел 1"),
            ("line", "1.1", None),
            ("line", "1.2", None),
            ("line", "1.3", None),
            ("section", "2", "Раздел 2"),
            ("line", "2.1", None),
            ("line", "2.2", None),
        ]
        assert [r["position_item_id"] for r in body["rows"] if r["kind"] == "line"] == [
            scene.l1.id,
            scene.l2.id,
            scene.l5.id,
            scene.l3.id,
            scene.l4.id,
        ]
        # по одному id порядок был бы другим: позиции лота 2 заведены раньше
        assert scene.l3.id < scene.l1.id

    def test_the_kp_columns_are_as_in_the_kp(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)

        body = _preview(member_client, stage, scene.rnd, scene.part)

        row = _by_item(body, scene.l1)
        assert (row["number"], row["unit"]) == ("1.1", scene.unit)
        assert (_dec(row["kp_unit_price"]), _dec(row["kp_total"])) == (D("180"), D("1800"))
        assert _dec(row["quantity"]) == D("10")
        assert row["title"] == scene.l1.job_title_in_proposal

    def test_our_rate_and_the_deviation_are_in_the_base_of_the_kp(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        body = _preview(member_client, stage, scene.rnd, scene.part)

        assert (_dec(_by_item(body, scene.l1)["ours"]), _by_item(body, scene.l1)["deviation_pct"]) == (D("120"), "50.0")
        assert (_dec(_by_item(body, scene.l5)["ours"]), _by_item(body, scene.l5)["deviation_pct"]) == (D("120"), "75.0")

    def test_a_kp_price_equal_to_the_target_is_zero_not_empty(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l3)

        assert (_dec(row["ours"]), row["deviation_pct"], row["state"]) == (D("120"), "0.0", "comparable")

    def test_a_line_without_a_comparison_has_no_numbers_of_ours(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l4)

        assert row["state"] != "comparable"
        assert (row["ours"], row["deviation_pct"], row["source"], row["inflation"]) == (None, None, None, None)

    def test_the_source_names_the_contract_and_the_nominal_rate_without_vat(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        contract = scene.contract

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)

        assert row["source"] == (
            f"{contract.object.title}, договор № {contract.contract_number}, 2025 г.: 100,00 ₽ за {scene.unit} без НДС"
        )
        assert row["inflation"] is None

    def test_the_agreement_replaces_the_source_and_moves_both_numbers(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l2)

        assert row["source"] == "Договорились 01.10.2026"
        assert _dec(row["ours"]) == D("96")
        assert row["deviation_pct"] == "150.0"
        assert row["inflation"] is None

    def test_a_template_comment_is_marked_template_and_a_line_without_a_case_has_none(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        body = _preview(member_client, stage, scene.rnd, scene.part)

        over = _by_item(body, scene.l1)
        assert (over["comment_kind"], bool(over["comment"])) == ("template", True)
        equal = _by_item(body, scene.l3)
        assert (equal["comment_kind"], equal["comment"]) == (None, None)

    def test_a_custom_comment_is_the_cell_value_with_the_formula_guard(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l5)

        assert (row["comment"], row["comment_kind"]) == ("'=1+1", "custom")

    def test_an_empty_custom_comment_wins_over_the_template_and_is_empty(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        _comment(db_session, scene.l1, "", _user(factories))

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)

        assert (row["comment"], row["comment_kind"]) == (None, "custom")

    def test_a_title_with_a_formula_start_is_the_cell_value(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)
        scene.l1.job_title_in_proposal = "-500 работ"
        db_session.flush()

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)

        assert row["title"] == "'-500 работ"

    def test_the_range_max_and_median_of_the_fork_are_nowhere_in_the_preview(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)  # вилка 100 – 130, медиана 115

        body = _preview(member_client, stage, scene.rnd, scene.part)

        forbidden = {D("130"), D("156"), D("115"), D("138")}
        for row in body["rows"]:
            if row["kind"] == "section":
                continue
            for key in ("quantity", "kp_unit_price", "kp_total", "ours", "deviation_pct"):
                assert _dec(row[key]) not in forbidden, (key, row[key])
            for key in ("source", "inflation", "comment"):
                assert "130" not in (row[key] or "") and "115" not in (row[key] or ""), (key, row[key])

    @pytest.mark.parametrize("bad", ["NaN", "Infinity"])
    @pytest.mark.parametrize(
        ("column", "field"),
        [
            ("unit_cost_total", "kp_unit_price"),
            ("total_cost_total", "kp_total"),
            ("suggested_quantity", "quantity"),
        ],
    )
    def test_a_non_finite_kp_number_is_null_in_the_preview_and_an_empty_cell_in_the_file(
        self, member_client, stage, db_session, world, factories, column, field, bad
    ):
        scene = _file_scene(stage, db_session, world, factories)
        setattr(scene.l1, column, D(bad))
        db_session.flush()

        body = _preview(member_client, stage, scene.rnd, scene.part)
        content = _export_ok(member_client, stage, scene.rnd, scene.part).content

        row = _by_item(body, scene.l1)
        assert row[field] is None
        other = {"kp_unit_price", "kp_total", "quantity"} - {field}
        assert all(row[key] is not None for key in other)
        titles, table, _ = _table(content)
        keys = [k for k in COLUMN_KEYS if k != "inflation"]
        cells = dict(zip(keys, table[1], strict=True))  # вторая строка таблицы — «1.1»
        assert cells["number"] == "1.1"
        assert cells[field] is None

    @pytest.mark.parametrize(
        ("column", "field", "value"),
        [
            ("total_cost_total", "kp_total", D("-1800")),
            ("unit_cost_total", "kp_unit_price", D("0")),
        ],
    )
    def test_a_negative_or_zero_kp_number_is_kept_in_the_preview_and_in_the_file(
        self, member_client, stage, db_session, world, factories, column, field, value
    ):
        """Правило предпросмотра — конечность, а не предикат цены: отрицательная стоимость и нулевая
        цена КП — числа КП, они идут в ячейку и в предпросмотр как есть."""
        scene = _file_scene(stage, db_session, world, factories)
        setattr(scene.l1, column, value)
        db_session.flush()

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)
        content = _export_ok(member_client, stage, scene.rnd, scene.part).content

        assert row[field] is not None and _dec(row[field]) == value
        _, table, _ = _table(content)
        keys = [k for k in COLUMN_KEYS if k != "inflation"]
        cells = dict(zip(keys, table[1], strict=True))  # вторая строка таблицы — «1.1»
        assert cells["number"] == "1.1"
        assert _num(cells[field]) == value

    def test_a_line_without_a_number_has_an_empty_number_not_the_key_of_the_position(
        self, member_client, stage, db_session
    ):
        """Номер — сырой `item_number_in_proposal`: у строки без номера он пуст, ключ позиции
        (`position_key_in_proposal`, подпись разбора) его не подменяет."""
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        item = stage.kp(part, D("120"))
        assert item.item_number_in_proposal is None and item.position_key_in_proposal

        body = _preview(member_client, stage, rnd, part)
        content = _export_ok(member_client, stage, rnd, part).content

        index = next(i for i, row in enumerate(body["rows"]) if row.get("position_item_id") == item.id)
        assert body["rows"][index]["number"] is None
        _, table, _ = _table(content)
        assert table[index][0] is None
        assert table[index][1] == body["rows"][index]["title"]

    def test_the_positions_are_those_of_the_scene_not_a_second_read(
        self, member_client, stage, db_session, world, factories, monkeypatch
    ):
        """Позиции файла — строки сцены, а не второе прочтение базы (Ruling R28). Сборка читает
        сцену, затем главы; правка строки КП, пришедшая между этими чтениями (здесь — внутри чтения
        глав), в файл не попадает: строка файла несёт цену КП той сцены, которую видел разбор.
        Перечитай сборка позиции после глав — цена КП была бы 999 при отклонении от неё же."""
        from crud import our_price as crud_our_price

        scene = _file_scene(stage, db_session, world, factories)
        skeleton = crud_our_price.load_file_skeleton

        def late_edit(db, estimate_id):
            db.execute(sa.update(PositionItem).where(PositionItem.id == scene.l1.id).values(unit_cost_total=D("999")))
            return skeleton(db, estimate_id)

        monkeypatch.setattr(crud_our_price, "load_file_skeleton", late_edit)

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)

        # 180 с НДС 20 % — нетто 150 против цели 100: 50,0
        assert (_dec(row["kp_unit_price"]), row["deviation_pct"]) == (D("180"), "50.0")
        assert db_session.get(PositionItem, scene.l1.id, populate_existing=True).unit_cost_total == D("999")

    def test_the_comment_names_a_contract_of_this_contractor_as_with_you(self, member_client, stage):
        """Договор-источник с самим участником назван в шаблоне «с вами»: участник тендера доходит до
        строк файла."""
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.history(D("100"), contractor=part.contractor)
        item = stage.kp(part, D("150"))

        row = _by_item(_preview(member_client, stage, rnd, part), item)

        assert row["comment_kind"] == "template"
        assert "В нашем договоре с вами по объекту" in row["comment"], row["comment"]

    def test_a_system_line_takes_the_area_of_the_tender_object(self, member_client, stage):
        """Цель системной строки — ₽/м² × площадь объекта тендера × доля строки: площадь
        приходит в строки файла из шапки сцены, без неё у системы нет ни отклонения, ни расценки."""
        rnd, part, items = _offer_system_scene(stage)

        body = _preview(member_client, stage, rnd, part)

        for item in items:
            row = _by_item(body, item)
            assert row["state"] == "comparable", row
            assert row["deviation_pct"] is not None and row["ours"] is not None, row


# ---------------------------------------------------------------------------
#  Предпросмотр равен файлу
# ---------------------------------------------------------------------------


class TestPreviewEqualsTheFile:
    @pytest.mark.parametrize("with_series", [False, True])
    def test_every_cell_of_the_book_is_the_field_of_the_preview_row(
        self, member_client, stage, db_session, world, factories, with_series
    ):
        if with_series:
            rnd, part, item, series_id = _inflation_scene(stage, db_session, "ok")
            _comment(db_session, item, "=1+1", _user(factories))
            params = {"inflation_series_id": series_id}
        else:
            scene = _file_scene(stage, db_session, world, factories)
            rnd, part, params = scene.rnd, scene.part, {}

        preview = _preview(member_client, stage, rnd, part, **params)
        content = _export_ok(member_client, stage, rnd, part, **params).content

        titles, table, ws = _table(content)
        assert titles == [c["title"] for c in preview["columns"]]
        assert len(table) == len(preview["rows"])
        keys = [c["key"] for c in preview["columns"]]
        for cells, row in zip(table, preview["rows"], strict=True):
            if row["kind"] == "section":
                assert cells[0] == row["number"] and cells[1] == row["title"]
                assert all(value is None for value in cells[2:])
                continue
            for key, cell in zip(keys, cells, strict=True):
                if key in ("quantity", "kp_unit_price", "kp_total", "ours", "deviation_pct"):
                    # книга хранит число как double: до ~15 значащих цифр без потерь
                    if row[key] is None:
                        assert cell is None, (key, cell)
                    else:
                        assert abs(_num(cell) - _dec(row[key])) <= abs(_dec(row[key])) * D("1e-14"), (
                            key,
                            cell,
                            row[key],
                        )
                else:
                    assert cell == row[key], (key, cell, row[key])

    def test_the_header_lines_of_the_sheet_are_the_header_of_the_preview(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        header = _preview(member_client, stage, scene.rnd, scene.part)["header"]
        content = _export_ok(member_client, stage, scene.rnd, scene.part).content

        _, _, ws = _table(content)
        texts = [row[0].value for row in ws.iter_rows(max_row=8) if row[0].value]
        assert f"Тендер: {header['tender']} · Объект: {header['object']}" in texts
        assert header["round"] in texts
        assert f"Участник: {header['contractor']}" in texts
        assert f"Дата выгрузки: {header['exported_on']}" in texts
        assert header["vat_label"] in texts
        assert header["source_label"] in texts
        assert f"Ценовой уровень: {header['level_label']}" in texts

    def test_the_cells_of_a_money_line_keep_the_long_digits_to_the_limit_of_the_format(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        scene.l1.unit_cost_total = D("180.123456789")
        scene.l1.total_cost_total = D("1801.23456789")
        db_session.flush()

        row = _by_item(_preview(member_client, stage, scene.rnd, scene.part), scene.l1)
        content = _export_ok(member_client, stage, scene.rnd, scene.part).content

        assert row["kp_unit_price"] == "180.123456789"
        _, table, _ = _table(content)
        assert _num(table[1][4]) == D("180.123456789")


# ---------------------------------------------------------------------------
#  ИПЦ
# ---------------------------------------------------------------------------


class TestInflation:
    def test_the_block_names_the_series_the_month_and_the_years_of_the_source_contract(
        self, member_client, stage, db_session
    ):
        rnd, part, item, series_id = _inflation_scene(stage, db_session, "ok")

        header = _preview(member_client, stage, rnd, part, inflation_series_id=series_id)["header"]

        assert header["inflation_series"] == SERIES_NAME
        assert header["inflation_target_month"] == "2026-08"
        assert header["level_label"] == f"приведено по ИПЦ: {SERIES_NAME}, к 08.2026; КП не приводится"
        years = header["inflation_years"]
        assert [y["year"] for y in years] == [2025, 2026]
        assert [(_dec(y["coefficient"]), y["is_forecast"]) for y in years] == [
            (D("1.083"), False),
            (D("1.060"), False),
        ]
        assert all(isinstance(y["source"], str) and isinstance(y["updated_at"], str) for y in years)
        assert {k for y in years for k in y} == {"year", "coefficient", "source", "is_forecast", "updated_at"}

    def test_the_years_are_those_of_the_source_contracts_only(self, member_client, stage, db_session):
        """Договор 2024 года в вилке, но ни одна цена источника не его: года 2024 в блоке нет."""
        _dated_history(stage, D("100"))
        contract, estimate, _ = stage.history(D("105"))
        contract.signed_date = dt.date(2024, 5, 20)
        estimate.data_prepared_on_date = None
        series_id = fx.series_with_years(db_session, SERIES_NAME, FULL_YEARS)
        rnd = stage.round(held_on=dt.date(2026, 8, 10))
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        db_session.flush()

        header = _preview(member_client, stage, rnd, part, inflation_series_id=series_id)["header"]

        assert [y["year"] for y in header["inflation_years"]] == [2025, 2026]

    def test_the_line_names_the_target_in_the_series_and_the_row_without_a_factor_has_none(
        self, member_client, stage, db_session, world
    ):
        rnd, part, item, series_id = _inflation_scene(stage, db_session, "ok")

        row = _by_item(_preview(member_client, stage, rnd, part, inflation_series_id=series_id), item)

        # цель 100 * 1,111034702… = 111,10 без НДС; база КП 0 %, расценка равна цели
        assert row["inflation"] == f"{SERIES_NAME}, к 08.2026: 111,10 ₽ за {_unit_name(db_session, world)} без НДС"
        assert abs(_dec(row["ours"]) - D("111.1034702376")) < D("0.0000000001")
        assert abs(FACTOR_FEB_2025 * 100 - _dec(row["ours"])) < D("0.0000000001")
        assert row["source"].endswith(f"2025 г.: 100,00 ₽ за {_unit_name(db_session, world)} без НДС")

    def test_the_same_month_needs_no_years_and_says_so_in_the_header_fields(self, member_client, stage, db_session):
        contract, estimate, _ = stage.history(D("100"))
        contract.signed_date = dt.date(2026, 8, 15)
        estimate.data_prepared_on_date = None
        series_id = fx.series_with_years(db_session, SERIES_NAME, {})
        rnd = stage.round(held_on=dt.date(2026, 8, 10))
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        item = stage.kp(part, D("120"))
        db_session.flush()

        body = _preview(member_client, stage, rnd, part, inflation_series_id=series_id)

        assert body["header"]["inflation_years"] == []
        assert _by_item(body, item)["inflation"] is None

    @pytest.mark.parametrize(
        ("kind", "code"),
        [
            ("target", "inflation_target_missing"),
            ("amendment", "amendment_date_missing"),
            ("years", "missing_inflation_years"),
        ],
    )
    @pytest.mark.parametrize("route", ["preview", "export"])
    def test_the_three_refusals_are_the_codes_of_the_breakdown_and_leave_no_file_and_no_record(
        self, member_client, stage, db_session, tmp_storage, saved_keys, kind, code, route
    ):
        rnd, part, _, series_id = _inflation_scene(stage, db_session, kind)

        response = _call(route, member_client, stage, rnd, part, inflation_series_id=series_id)

        assert (response.status_code, _code(response)) == (422, code)
        assert _records(db_session) == []
        assert saved_keys == [] and _stored(tmp_storage) == []

    def test_an_unknown_series_is_404(self, member_client, stage, db_session):
        rnd, part, _, _ = _inflation_scene(stage, db_session, "ok")

        assert _preview_get(member_client, stage, rnd, part, inflation_series_id=987654321).status_code == 404

    def test_a_series_that_got_a_hole_after_the_preview_refuses_the_export(
        self, member_client, stage, db_session, tmp_storage, saved_keys
    ):
        rnd, part, _, series_id = _inflation_scene(stage, db_session, "ok")
        assert _preview_get(member_client, stage, rnd, part, inflation_series_id=series_id).status_code == 200
        db_session.execute(
            sa.delete(InflationIndexValue).where(
                InflationIndexValue.series_id == series_id, InflationIndexValue.year == 2025
            )
        )
        db_session.flush()

        response = _export(member_client, stage, rnd, part, inflation_series_id=series_id)

        assert (response.status_code, _code(response)) == (422, "missing_inflation_years")
        assert _records(db_session) == []
        assert saved_keys == [] and _stored(tmp_storage) == []

    def test_the_export_with_a_series_names_the_level_in_the_record(self, member_client, stage, db_session):
        rnd, part, _, series_id = _inflation_scene(stage, db_session, "ok")

        _export_ok(member_client, stage, rnd, part, inflation_series_id=series_id)

        (record,) = _records(db_session)
        assert record.level_label == f"приведено по ИПЦ: {SERIES_NAME}, к 08.2026; КП не приводится"


# ---------------------------------------------------------------------------
#  Выгрузка
# ---------------------------------------------------------------------------


class TestExport:
    def test_the_response_is_the_book_and_the_record_is_the_file_of_the_preview(
        self, member_client, stage, db_session, world, factories, tmp_storage
    ):
        scene = _file_scene(stage, db_session, world, factories)
        header = _preview(member_client, stage, scene.rnd, scene.part)["header"]

        response = _export_ok(member_client, stage, scene.rnd, scene.part)

        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert _disposition_name(response) == header["filename"]
        (record,) = _records(db_session)
        assert record.filename == header["filename"]
        assert record.offer_id == scene.part.offer.id
        assert record.file_sha256 == hashlib.sha256(response.content).hexdigest()
        assert record.size_bytes == len(response.content)
        assert record.vat_rate == D("20")
        assert record.level_label == header["level_label"]
        assert record.created_by == member_client.user.id
        with tmp_storage.get(record.file_key) as handle:
            assert handle.read() == response.content

    def test_the_vat_rate_of_the_record_is_the_base_of_the_kp(self, member_client, stage, db_session):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))

        _export_ok(member_client, stage, rnd, part)

        (record,) = _records(db_session)
        assert record.vat_rate == D("0")

    def test_the_file_key_is_opaque_and_unique_for_the_same_file_name(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        _export_ok(member_client, stage, scene.rnd, scene.part)
        _export_ok(member_client, stage, scene.rnd, scene.part)

        first, second = _records(db_session)
        assert first.filename == second.filename
        assert first.file_key != second.file_key
        for record in (first, second):
            assert len(record.file_key) == 32 and set(record.file_key) <= set("0123456789abcdef")

    def test_a_failure_of_the_insert_after_the_save_removes_the_file_and_leaves_no_record(
        self, member_client, stage, db_session, world, factories, tmp_storage, saved_keys, monkeypatch
    ):
        scene = _file_scene(stage, db_session, world, factories)

        class Boom:
            def __init__(self, **_):
                raise RuntimeError("insert failed")

        monkeypatch.setattr(kp_sent_files, "KpSentFile", Boom)

        with pytest.raises(RuntimeError, match="insert failed"):
            _export(member_client, stage, scene.rnd, scene.part)

        assert len(saved_keys) == 1  # файл до вставки был сохранён
        assert not tmp_storage.exists(saved_keys[0])
        assert _stored(tmp_storage) == []
        assert _records(db_session) == []

    def test_a_failure_of_the_commit_removes_the_file_and_leaves_no_record(
        self, member_client, stage, db_session, world, factories, tmp_storage, saved_keys, monkeypatch
    ):
        scene = _file_scene(stage, db_session, world, factories)
        Session.commit(db_session)

        def boom():
            raise RuntimeError("commit failed")

        monkeypatch.setattr(db_session, "commit", boom)

        with pytest.raises(RuntimeError, match="commit failed"):
            _export(member_client, stage, scene.rnd, scene.part)

        assert len(saved_keys) == 1
        assert not tmp_storage.exists(saved_keys[0])
        assert _stored(tmp_storage) == []
        monkeypatch.undo()
        assert _records(db_session) == []

    def test_a_database_refusal_at_the_commit_removes_the_file_and_leaves_no_record(
        self, stage, db_session, world, factories, tmp_storage, saved_keys
    ):
        """Настоящий отказ базы при `commit` (внешний ключ автора записи), без подмены кода
        выгрузки: транзакция откатывается, файл по ключу удалён, записи нет."""
        scene = _file_scene(stage, db_session, world, factories)
        Session.commit(db_session)

        with pytest.raises(sa.exc.IntegrityError):
            kp_sent_files.export_file(
                db_session,
                tmp_storage,
                stage.tender.id,
                scene.rnd.id,
                scene.part.offer.id,
                inflation_series_id=None,
                user_id=987654321,
            )

        assert len(saved_keys) == 1
        assert not tmp_storage.exists(saved_keys[0])
        assert _stored(tmp_storage) == []
        assert _records(db_session) == []

    def test_a_refusal_before_the_save_never_touches_the_storage(
        self, member_client, stage, db_session, tmp_storage, saved_keys
    ):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=None)
        stage.kp(part, D("120"))

        response = _export(member_client, stage, rnd, part)

        assert response.status_code == 422
        assert saved_keys == []

    def test_two_exports_of_one_state_have_the_same_cells(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)

        first = _table(_export_ok(member_client, stage, scene.rnd, scene.part).content)[:2]
        second = _table(_export_ok(member_client, stage, scene.rnd, scene.part).content)[:2]

        assert first == second

    def test_the_range_max_and_median_of_the_fork_are_in_no_cell_of_the_file(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)

        content = _export_ok(member_client, stage, scene.rnd, scene.part).content

        _, table, ws = _table(content)
        forbidden = {D("130"), D("156"), D("115"), D("138")}
        for cells in table:
            for cell in cells:
                if isinstance(cell, int | float):
                    assert _num(cell) not in forbidden, cell
                elif isinstance(cell, str):
                    assert "130" not in cell and "115" not in cell, cell


# ---------------------------------------------------------------------------
#  Отправленные файлы
# ---------------------------------------------------------------------------


class TestSentFiles:
    def test_the_list_is_newest_first_with_the_author_and_the_keys_of_the_spec(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        first_author = member_client.user
        _export_ok(member_client, stage, scene.rnd, scene.part)
        second_author = _user(factories)
        member_client.set_user(second_author)
        _export_ok(member_client, stage, scene.rnd, scene.part)
        first, second = _records(db_session)

        body = _sent(member_client, stage, scene.rnd, scene.part).json()

        assert [item["id"] for item in body] == [second.id, first.id]
        assert all(set(item) == SENT_KEYS for item in body)
        assert [item["created_by"] for item in body] == [
            {"id": second_author.id, "email": second_author.email},
            {"id": first_author.id, "email": first_author.email},
        ]
        newest = body[0]
        assert newest["filename"] == second.filename
        assert newest["size_bytes"] == second.size_bytes
        assert newest["level_label"] == second.level_label
        assert _dec(newest["vat_rate"]) == D("20")
        assert isinstance(newest["created_at"], str) and "T" in newest["created_at"]

    def test_the_list_holds_only_the_files_of_this_offer(self, member_client, stage, db_session, world, factories):
        scene = _file_scene(stage, db_session, world, factories)
        other = stage.participant(scene.rnd, "Подрядчик Б", vat=D("0"))
        stage.kp(other, D("120"))
        _export_ok(member_client, stage, scene.rnd, scene.part)
        _export_ok(member_client, stage, scene.rnd, other)

        mine = _sent(member_client, stage, scene.rnd, scene.part).json()
        theirs = _sent(member_client, stage, scene.rnd, other).json()

        assert len(mine) == len(theirs) == 1
        assert mine[0]["id"] != theirs[0]["id"]

    def test_the_download_is_the_same_bytes_as_the_post_under_the_recorded_name(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        posted = _export_ok(member_client, stage, scene.rnd, scene.part)
        (item,) = _sent(member_client, stage, scene.rnd, scene.part).json()

        response = _download(member_client, stage, scene.rnd, scene.part, item["id"])

        assert response.status_code == 200
        assert response.content == posted.content
        assert _disposition_name(response) == item["filename"]
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def test_a_file_removed_from_the_disk_is_410_and_the_record_stays(
        self, member_client, stage, db_session, world, factories, tmp_storage
    ):
        scene = _file_scene(stage, db_session, world, factories)
        _export_ok(member_client, stage, scene.rnd, scene.part)
        (record,) = _records(db_session)
        assert tmp_storage.delete(record.file_key)

        response = _download(member_client, stage, scene.rnd, scene.part, record.id)

        assert response.status_code == 410
        assert len(_records(db_session)) == 1

    def test_a_file_of_another_offer_is_404_and_an_unknown_id_is_404(
        self, member_client, stage, db_session, world, factories
    ):
        scene = _file_scene(stage, db_session, world, factories)
        other = stage.participant(scene.rnd, "Подрядчик Б", vat=D("0"))
        stage.kp(other, D("120"))
        _export_ok(member_client, stage, scene.rnd, other)
        (foreign,) = _records(db_session)

        assert _download(member_client, stage, scene.rnd, scene.part, foreign.id).status_code == 404
        assert _download(member_client, stage, scene.rnd, scene.part, 987654321).status_code == 404
        assert _download(member_client, stage, scene.rnd, other, foreign.id).status_code == 200

    def test_a_file_of_another_stage_is_not_served_under_this_one(self, member_client, stage, db_session):
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        _export_ok(member_client, stage, rnd, part)
        (record,) = _records(db_session)
        other_round = stage.round()
        other = stage.participant(other_round, "Подрядчик А", vat=D("0"), package=part.offer.package)

        assert _download(member_client, stage, other_round, other, record.id).status_code == 404

    def test_the_files_survive_the_estimate_of_the_offer(self, member_client, stage, db_session, world, factories):
        """Список и скачивание держатся за оффер: смета ушла заменой раунда — файлы остались."""
        scene = _file_scene(stage, db_session, world, factories)
        _export_ok(member_client, stage, scene.rnd, scene.part)
        db_session.execute(sa.delete(Estimate).where(Estimate.id == scene.part.estimate.id))
        db_session.flush()

        (item,) = _sent(member_client, stage, scene.rnd, scene.part).json()

        assert _download(member_client, stage, scene.rnd, scene.part, item["id"]).status_code == 200

    @pytest.mark.parametrize("chain", ["tender", "round"])
    def test_an_existing_file_under_a_foreign_tender_or_round_is_404(
        self, member_client, stage, db_session, factories, chain
    ):
        """Файл существует и принадлежит офферу, но цепочка URL чужая: тендер другой либо этап
        другой при том же тендере — 404, как у списка."""
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        _export_ok(member_client, stage, rnd, part)
        (record,) = _records(db_session)
        if chain == "tender":
            tender_id, round_id = factories.TenderFactory.create().id, rnd.id
        else:
            tender_id, round_id = stage.tender.id, stage.round().id
        url = f"/api/v1/tenders/{tender_id}/rounds/{round_id}/our-price/offers/{part.offer.id}"

        assert member_client.get(f"{url}/sent-files/{record.id}/file").status_code == 404
        assert _download(member_client, stage, rnd, part, record.id).status_code == 200

    def test_a_broken_file_key_in_the_record_is_500(self, member_client, stage, db_session):
        """Испорченный ключ хранилища в записи — ошибка сервера с ответом 500, как у файла задания
        импорта, а не непойманное исключение."""
        stage.history(D("100"))
        rnd = stage.round()
        part = stage.participant(rnd, "Подрядчик А", vat=D("0"))
        stage.kp(part, D("120"))
        _export_ok(member_client, stage, rnd, part)
        (record,) = _records(db_session)
        db_session.execute(sa.update(KpSentFile).where(KpSentFile.id == record.id).values(file_key="../outside"))

        assert _download(member_client, stage, rnd, part, record.id).status_code == 500


# ---------------------------------------------------------------------------
#  Периоды смет договоров в плане приведения
# ---------------------------------------------------------------------------


class TestPeriodByEstimate:
    def test_the_history_plan_carries_the_period_of_each_estimate(self, stage, db_session):
        _, estimate, _ = _dated_history(stage)
        series_id = fx.series_with_years(db_session, SERIES_NAME, FULL_YEARS)

        plan = plan_history_inflation(db_session, [estimate.id], series_id=series_id, target_month=YearMonth(2026, 8))

        assert plan.period_by_estimate == {estimate.id: YearMonth(2025, 2)}
        assert set(plan.period_by_estimate) == set(plan.factor_by_estimate)

    def test_the_comparison_plan_carries_the_period_of_each_estimate(self, db_session, factories):
        series_id = fx.series_with_years(db_session, SERIES_NAME, FULL_YEARS)
        early = fx.contract_with_dates(db_session, factories, {"1": ["1200000.00"]}, signed_date=dt.date(2024, 5, 28))
        late = fx.contract_with_dates(db_session, factories, {"1": ["1200000.00"]}, signed_date=dt.date(2026, 7, 6))

        plan = cmp.resolve_inflation(db_session, [early, late], series_id=series_id, target_month=YearMonth(2026, 8))

        assert set(plan.period_by_estimate) == set(plan.factor_by_estimate)
        assert sorted(plan.period_by_estimate.values(), key=lambda p: (p.year, p.month)) == [
            YearMonth(2024, 5),
            YearMonth(2026, 7),
        ]

    def test_the_period_does_not_leak_into_the_comparison_response_keys(self, member_client, stage, db_session):
        rnd, part, _, series_id = _inflation_scene(stage, db_session, "ok")
        url = f"/api/v1/tenders/{stage.tender.id}/rounds/{rnd.id}/our-price"

        body = member_client.get(url, params={"inflation_series_id": series_id}).json()

        assert set(body["inflation"]) == {
            "series_id",
            "series_name",
            "series_note",
            "series_updated_at",
            "target_month",
            "has_forecast",
            "used_years",
            "sentence",
        }

    def test_the_period_does_not_leak_into_the_contract_comparison(self, db_session, factories):
        """Сравнение договоров отдаёт блок приведения явным перечнем ключей: поле плана, заведённое
        для файла подрядчику, в его ответ не попадает."""
        series_id = fx.series_with_years(db_session, SERIES_NAME, FULL_YEARS)
        ids = [
            fx.contract_with_dates(db_session, factories, {"1": ["1200000.00"]}, signed_date=dt.date(2024, 5, 28)),
            fx.contract_with_dates(db_session, factories, {"1": ["1200000.00"]}, signed_date=dt.date(2026, 7, 6)),
        ]

        data = cmp.build_comparison(
            db_session, ids, vat_mode=cmp.VAT_MODE_OWN, inflation_series_id=series_id, target_month=YearMonth(2026, 8)
        )

        assert set(data["inflation"]) == {
            "series_id",
            "series_name",
            "series_note",
            "series_updated_at",
            "target_month",
            "has_forecast",
            "used_years",
        }
