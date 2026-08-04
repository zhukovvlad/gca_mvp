"""Обвязка openpyxl для выгрузок §7.6.

Перенос из `backend/routers/export.py` источника (udp-tenders @ 98fb67667c44) —
именно обвязка: `_fill`, `_font`, `_align`, `_safe_str`, `_dev_font`, палитра,
секции с итогами. Домен УПД (счета, поставщики, классы материалов) не переносится
(§3 брифинга фазы 6).

Два отличия от источника, оба осознанные:

* **Деньги пишутся `Decimal`, а не `float`** — но у §3 здесь есть граница, и её
  надо знать. openpyxl принимает `Decimal`, однако **сам формат xlsx хранит числа
  как IEEE-754 double**, и «Decimal end-to-end» в файле не продолжается: точность
  ограничена форматом, а не кодом. Замер (`test_reports_api.py::
  test_xlsx_number_precision_boundary`):

  ```
  12000.55             → 12000.55              точно
  1234567890.12        → 1234567890.12         точно
  1075.35475           → 1075.35475            точно
  12345678901234567.89 → 1.234567890123457e+16 ПОТЕРЯ
  ```

  То есть до ~15 значащих цифр значение доезжает без искажений, дальше теряет
  младшие разряды. Суммы договоров ГП на стенде укладываются в 12 значащих цифр,
  то есть запас есть, но утверждать «полная точность» было бы неправдой.

  Альтернатива — писать деньги текстом — отвергнута: она ломает главное, ради чего
  выгрузку и просят. Банк складывает колонку в Excel, а `СУММ` по текстовым ячейкам
  даёт ноль. Числа с известной границей полезнее строк без неё.
* **Отдаём `bytes`, а не `StreamingResponse` с файловым объектом.** Грабли фазы 4:
  Starlette итерирует такой объект построчно и не закрывает хендл. Для xlsx из
  памяти (`BytesIO`) правильный ответ — вернуть содержимое целиком (§7 брифинга).
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
#  Стили (перенос из источника)
# ---------------------------------------------------------------------------

def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(bold: bool = False, color: str = "000000", size: int = 10) -> Font:
    return Font(bold=bold, color=color, name="Calibri", size=size)


def align(h: str = "left", v: str = "center", wrap: bool = False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

C_HEADER_BG = "1F4E79"      # тёмно-синий — шапка документа
C_CLASS_BG = "2E75B6"       # синий — заголовок секции класса
C_COL_BG = "4472C4"         # светлее — строка заголовков колонок
C_CLASS_TOTAL_BG = "9DC3E6" # итог по классу
C_TOTAL_BG = "BDD7EE"       # общий итог
C_ODD = "FFFFFF"
C_EVEN = "EBF3FB"
C_RED_TEXT = "C00000"       # превышение норматива
C_GREEN_TEXT = "375623"     # экономия

FMT_MONEY = '#,##0.00'
FMT_MONEY_CUR = '#,##0.00 "₽"'
FMT_QTY = "#,##0.###"
FMT_PCT = "+0.0;-0.0;0.0"   # проценты уже в п.п., знак обязателен
FMT_DATE = "DD.MM.YYYY"


def safe_str(value: Any) -> Any:
    """Защита от formula injection (перенос из источника).

    Excel считает формулой всё, что начинается с `=`, `+`, `-` или `@`. Наименование
    работы приходит из чужого XLSX, то есть это недоверенный ввод: строка вида
    `=1+1` или `-500` в ячейке превратилась бы в вычисление, а `=cmd|...` в старых
    сборках Excel — в запуск команды. Апостроф перед значением делает его текстом.

    Числа и даты не трогаем: у них нет такого разбора.
    """
    if isinstance(value, str):
        stripped = value.lstrip()
        if stripped.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def deviation_font(value: Decimal | None, *, bold: bool = False, size: int = 10) -> Font:
    """Цвет шрифта по знаку отклонения (перенос `_dev_font` из источника).

    `None` — «нет норматива», и цвета у него нет: §10 требует отличать этот случай
    от нуля, а покрасить его как экономию значило бы стереть разницу.
    """
    if value is None or value == 0:
        return font(bold=bold, size=size)
    return font(bold=bold, color=C_RED_TEXT if value > 0 else C_GREEN_TEXT, size=size)


# ---------------------------------------------------------------------------
#  Помощники листа
# ---------------------------------------------------------------------------

def write_cell(
    ws: Worksheet,
    row: int,
    col: int,
    value: Any,
    *,
    cell_font: Font | None = None,
    cell_fill: PatternFill | None = None,
    number_format: str | None = None,
    horizontal: str = "left",
    wrap: bool = False,
    border: bool = True,
):
    cell = ws.cell(row=row, column=col, value=safe_str(value))
    cell.font = cell_font or font()
    if cell_fill is not None:
        cell.fill = cell_fill
    if number_format is not None:
        cell.number_format = number_format
    cell.alignment = align(h=horizontal, wrap=wrap)
    if border:
        cell.border = BORDER
    return cell


def write_banner(ws: Worksheet, row: int, text: str, width: int, *, bg: str, size: int = 11) -> int:
    """Строка на всю ширину таблицы: шапка документа или заголовок секции."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row=row, column=1, value=safe_str(text))
    cell.fill = fill(bg)
    cell.font = font(bold=True, color="FFFFFF", size=size)
    cell.alignment = align(h="left")
    ws.row_dimensions[row].height = 18
    return row + 1


def write_column_headers(ws: Worksheet, row: int, columns: list[tuple[str, int, str | None, str]]) -> int:
    for index, (label, _width, _fmt, horizontal) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=label)
        cell.fill = fill(C_COL_BG)
        cell.font = font(bold=True, color="FFFFFF", size=9)
        cell.border = BORDER
        cell.alignment = align(h=horizontal, wrap=True)
    ws.row_dimensions[row].height = 30
    return row + 1


def apply_column_widths(ws: Worksheet, columns: list[tuple[str, int, str | None, str]]) -> None:
    for index, (_label, width, _fmt, _horizontal) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def workbook_bytes(wb: Workbook) -> bytes:
    """Готовый xlsx как `bytes` (см. модульную документацию про StreamingResponse)."""
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
