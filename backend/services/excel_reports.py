"""Листы выгрузок §7.6: свод по договору и отчёт «для банка».

Макет «для банка» согласован с пользователем (§6.1; таблица — в `crud/reports.py`):
класс → работа, колонки с отклонением в деньгах, итоги по каждому классу и общий,
шапка с реквизитами выборки и блоком подписей.

Правила, общие для обоих листов:

* **Деньги пишутся `Decimal`**, а формат ячейки отвечает за вид (§3). `float` здесь
  был бы тем же дефектом, что `jsonable_encoder` в фазе 5, только в файле.
* **«Нет норматива» — пустая ячейка плюс явная пометка, а не ноль** (§10). Ноль в
  колонке отклонения означал бы «ровно по нормативу», то есть прямо противоположное.
* **Отклонение в итогах — по суммам**, а не среднее арифметическое процентов
  (обоснование — `crud/reports.py`).
* Наименование работы прогоняется через `safe_str`: оно приходит из чужого XLSX, а
  Excel считает формулой всё, что начинается с `=`, `+`, `-` или `@`.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook

from services.excel import (
    C_CLASS_BG,
    C_CLASS_TOTAL_BG,
    C_EVEN,
    C_HEADER_BG,
    C_ODD,
    C_TOTAL_BG,
    FMT_MONEY,
    FMT_PCT,
    FMT_QTY,
    align,
    apply_column_widths,
    deviation_font,
    fill,
    font,
    safe_str,
    workbook_bytes,
    write_banner,
    write_cell,
    write_column_headers,
)

#: (заголовок, ширина, формат, выравнивание). Порядок колонок — из согласованного
#: макета §6.1 и менять его нельзя не спросив: по нему сверяются вручную.
_COLUMNS: list[tuple[str, int, str | None, str]] = [
    ("Работа", 60, "@", "left"),
    ("Ед.", 8, "@", "center"),
    ("Объём", 14, FMT_QTY, "right"),
    ("Ставка", 16, FMT_MONEY, "right"),
    ("Норматив", 16, FMT_MONEY, "right"),
    ("Стоимость", 18, FMT_MONEY, "right"),
    ("Отклонение, %", 14, FMT_PCT, "right"),
    ("Отклонение, ₽", 18, FMT_MONEY, "right"),
]
_N_COLS = len(_COLUMNS)

#: Пометка вместо нуля там, где норматива нет (§10).
NO_STANDARD = "нет норматива"


def _write_row(ws, row_num: int, row: dict, *, even: bool) -> None:
    background = fill(C_EVEN if even else C_ODD)
    deviation = row["deviation_pct"]
    has_standard = row["standard_unit_rate"] is not None

    values: list[tuple[object, str | None, str]] = [
        (row["job_title"], "@", "left"),
        (row["unit_code"] or "—", "@", "center"),
        (row["volume"], FMT_QTY, "right"),
        (row["rate"], FMT_MONEY, "right"),
        (row["standard_unit_rate"] if has_standard else NO_STANDARD,
         FMT_MONEY if has_standard else "@", "right"),
        (row["amount"], FMT_MONEY, "right"),
        (deviation if has_standard else NO_STANDARD, FMT_PCT if has_standard else "@", "right"),
        (row["deviation_money"] if has_standard else None, FMT_MONEY, "right"),
    ]
    for index, (value, number_format, horizontal) in enumerate(values, start=1):
        # Цветом выделяются только колонки отклонения: подкрашивать всю строку
        # значило бы утверждать, что «дорогая работа» и «превышение» — одно и то же.
        cell_font = deviation_font(deviation) if index in (7, 8) else font()
        write_cell(
            ws,
            row_num,
            index,
            value,
            cell_font=cell_font,
            cell_fill=background,
            number_format=number_format,
            horizontal=horizontal,
            wrap=index == 1,
        )


def _write_totals(ws, row_num: int, label: str, totals: dict, *, bg: str) -> int:
    background = fill(bg)
    deviation = totals["deviation_pct"]

    write_cell(
        ws, row_num, 1, label, cell_font=font(bold=True), cell_fill=background, horizontal="left"
    )
    for column in (2, 3, 4, 5):
        # Объёмы и ставки разных работ в разных единицах — суммировать их нельзя,
        # и прочерк честнее пустоты: видно, что здесь ничего не пропущено.
        write_cell(ws, row_num, column, "—", cell_font=font(bold=True), cell_fill=background,
                   horizontal="center")
    write_cell(ws, row_num, 6, totals["amount"], cell_font=font(bold=True), cell_fill=background,
               number_format=FMT_MONEY, horizontal="right")
    # Пустое отклонение печатается пометкой, а не пустотой и не нулём: в итоговой
    # строке отчёта для банка ноль означал бы «сошлось с нормативом» (см.
    # `crud.reports._empty_report_totals`).
    comparable = deviation is not None
    write_cell(ws, row_num, 7, deviation if comparable else NO_STANDARD,
               cell_font=deviation_font(deviation, bold=True), cell_fill=background,
               number_format=FMT_PCT if comparable else "@", horizontal="right")
    write_cell(ws, row_num, 8, totals["deviation_money"] if comparable else NO_STANDARD,
               cell_font=deviation_font(deviation, bold=True), cell_fill=background,
               number_format=FMT_MONEY if comparable else "@", horizontal="right")
    return row_num + 1


def _write_signatures(ws, row_num: int) -> int:
    """Блок подписей — требование согласованного макета §6.1."""
    row_num += 1
    for role in ("Составил", "Утвердил"):
        cell = ws.cell(
            row=row_num,
            column=1,
            value=f"{role}: ____________________  /____________________/  «___» __________ 20___ г.",
        )
        cell.font = font(size=9)
        cell.alignment = align(h="left")
        row_num += 2
    return row_num


def _write_generated_at(ws, row_num: int, generated_at: dt.date) -> int:
    cell = ws.cell(row=row_num, column=1, value=f"Сформирован: {generated_at.strftime('%d.%m.%Y')}")
    cell.font = font(size=9)
    return row_num + 1


# ---------------------------------------------------------------------------
#  (а) Свод расценок по договору
# ---------------------------------------------------------------------------

def build_contract_summary(data: dict, *, generated_at: dt.date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Свод расценок"
    apply_column_widths(ws, _COLUMNS)

    header = data["header"]
    row = write_banner(
        ws, 1, f"СВОД РАСЦЕНОК ПО ДОГОВОРУ {header['contract_number']}", _N_COLS, bg=C_HEADER_BG
    )

    amendment = header["estimate_amendment_no"]
    facts = [
        f"Объект: {header['object_title']}",
        f"Подрядчик: {header['contractor_title']}",
        f"Класс объектов: {header['rate_class_title']}",
        f"Подписант: {header['signer'] or '—'}",
        f"Дата договора: {_ru_date(header['signed_date'])}",
        f"Смета: {'исходная' if amendment is None else f'доп. соглашение № {amendment}'}"
        + f", дата {_ru_date(header['estimate_date'])}",
    ]
    for fact in facts:
        cell = ws.cell(row=row, column=1, value=safe_str(fact))
        cell.font = font(size=9)
        row += 1
    row = _write_generated_at(ws, row, generated_at)
    row += 1

    row = write_column_headers(ws, row, _COLUMNS)
    for index, data_row in enumerate(data["rows"]):
        _write_row(ws, row, data_row, even=index % 2 == 1)
        row += 1

    row = _write_totals(ws, row, "ИТОГО ПО ДОГОВОРУ", data["totals"], bg=C_TOTAL_BG)
    row = _write_footnote(ws, row, data["totals"])
    _write_signatures(ws, row)

    ws.freeze_panes = "B1"
    return workbook_bytes(wb)


# ---------------------------------------------------------------------------
#  (б) Сравнение с нормативами «для банка»
# ---------------------------------------------------------------------------

def build_bank_comparison(data: dict, *, generated_at: dt.date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение с нормативами"
    apply_column_widths(ws, _COLUMNS)

    header = data["header"]
    row = write_banner(ws, 1, "СРАВНЕНИЕ РАСЦЕНОК С НОРМАТИВАМИ", _N_COLS, bg=C_HEADER_BG)

    period = _period_text(header["date_from"], header["date_to"])
    for fact in (
        f"Период: {period}",
        f"Договоров: {header['contracts']}   Объектов: {header['objects']}   "
        f"Классов: {header['classes']}",
    ):
        cell = ws.cell(row=row, column=1, value=safe_str(fact))
        cell.font = font(size=9)
        row += 1
    row = _write_generated_at(ws, row, generated_at)
    row += 1

    for section in data["sections"]:
        row = write_banner(ws, row, f"КЛАСС: {section['rate_class_title']}", _N_COLS, bg=C_CLASS_BG)
        row = write_column_headers(ws, row, _COLUMNS)
        for index, data_row in enumerate(section["rows"]):
            _write_row(ws, row, data_row, even=index % 2 == 1)
            row += 1
        row = _write_totals(
            ws,
            row,
            f"ИТОГО ПО КЛАССУ «{section['rate_class_title']}»",
            section["totals"],
            bg=C_CLASS_TOTAL_BG,
        )
        # Оба счётчика per-класс: у секции теперь есть свой счётчик объёма
        # (замечание ревью — общий скаляр терял классы целиком).
        row = _write_excluded_counters(ws, row, section["totals"])
        row += 1

    row = _write_totals(ws, row, "ВСЕГО ПО ВЫБОРКЕ", data["totals"], bg=C_TOTAL_BG)
    row = _write_footnote(ws, row, data["totals"])
    _write_signatures(ws, row)

    ws.freeze_panes = "B1"
    return workbook_bytes(wb)


def _write_excluded_counters(ws, row_num: int, totals: dict) -> int:
    """Оба счётчика исключённого — под каждым итогом (макет §6.1).

    Печатаются всегда, даже нулём: отсутствие строки читалось бы как «таких позиций
    не проверяли», а ноль говорит «проверили, их нет».

    Подписи образуют **разбиение** (см. `crud/reports.py`): строки таблицы + эти два
    счётчика = все расценённые позиции выборки, без пересечений. «С объёмом, но без
    норматива» — уточнение по замечанию ревью: позиция без объёма и без норматива
    считается один раз, в счётчике объёма, и прежняя подпись «без норматива: 0» для
    такого файла была бы ложью.
    """
    for note in (
        # Первая строка делает разбиение проверяемым: строка отчёта агрегирует
        # работу, и по числу видимых строк позиции не сосчитать (замечание ревью).
        f"Сравнимых позиций (в расчёте отклонения): "
        f"{totals['comparable_positions']}",
        f"Позиций с объёмом, но без норматива (в отклонение не вошли): "
        f"{totals['positions_without_standard']}",
        f"Позиций с ценой, но без объёма (в расчёт не вошли): "
        f"{totals['positions_without_volume']}",
    ):
        cell = ws.cell(row=row_num, column=1, value=note)
        cell.font = font(size=9)
        row_num += 1
    return row_num


def _write_footnote(ws, row_num: int, totals: dict) -> int:
    """Как считалось отклонение — прямо на листе.

    Банк проверяет цифры на калькуляторе, и без этой сноски средневзвешенное
    отклонение выглядит расхождением: сумма процентов по строкам, поделённая на их
    число, даёт другое значение.
    """
    row_num = _write_excluded_counters(ws, row_num, totals)
    # Общий счёт посчитан независимо (count(*) по VIEW): совпадение с суммой трёх
    # счётчиков выше — проверяемый инвариант файла, а не тавтология.
    total_note = ws.cell(
        row=row_num,
        column=1,
        value=(
            f"Всего расценённых позиций: {totals['positions_priced']} "
            "(равно сумме трёх счётчиков выше)"
        ),
    )
    total_note.font = font(size=9, bold=True)
    row_num += 1
    cell = ws.cell(
        row=row_num,
        column=1,
        value=(
            "* Отклонение в итогах — средневзвешенное по объёму: (стоимость по ставкам − "
            "стоимость по нормативам) / стоимость по нормативам. Среднее арифметическое "
            "процентов по строкам даёт другое число и завышает вклад мелких работ."
        ),
    )
    cell.font = font(size=8)
    cell.alignment = align(h="left", wrap=True)
    return row_num + 2


def _ru_date(iso_date: str | None) -> str:
    if not iso_date:
        return "—"
    return dt.date.fromisoformat(iso_date).strftime("%d.%m.%Y")


def _period_text(date_from: str | None, date_to: str | None) -> str:
    if date_from is None and date_to is None:
        return "все сметы"
    return f"{_ru_date(date_from) if date_from else '…'} — {_ru_date(date_to) if date_to else '…'}"
