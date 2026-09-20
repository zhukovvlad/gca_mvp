"""Билдер книги «Изменения КП» — лист на участника (спека
2026-09-16-tender-changes-export-design.md §2.3, §2.5, §2.9, §2.10; план,
Task 4).

Модуль печатает готовый `Sheet` (задача 2, `services.changes_export.build_sheet`)
и не пересчитывает НИЧЕГО: числа на листе и в `Sheet` совпадают потому, что
здесь только раскладка по ячейкам Excel, а не вторая арифметика — тем же
приёмом, что `services/excel_comparison.py` печатает агрегат `crud.comparison`.

**Четыре места, где легче всего ошибиться, и как каждое закрыто здесь.**

1. **Прочерк у денежных ветвей печатается ПО `Cell.kind`, а не по причине.**
   У `MONEY_ONLY_KINDS` `unit_price` несёт `Money(None, REASON_NO_PRICE, False)`
   — ту же причину, что у настоящей погашенной цены работы. `_write_stage_cell`
   проверяет ветвь ПЕРВОЙ и для денежных ветвей пишет прочерк в «№ строк» /
   «Объём» / «Цена за ед.» безусловно, не разбирая `unit_price.reason`.
2. **Денежная ячейка ОДНОЯРУСНА.** `_write_money` пишет `Decimal` с
   `FMT_MONEY` через `write_cell`, без `wrap=True` и без второй строки внутри
   ячейки; неполнота — заливкой ТОЙ ЖЕ ячейки, а у агрегатов ДОПОЛНИТЕЛЬНО
   называется соседней колонкой «Полнота» блока подытогов
   (`_block_completeness_text`), а не второй строкой внутри денежной ячейки.
3. **В блоке подытогов `TEXT_NO_AMOUNT` не появляется НИ РАЗУ.** Подытог и
   общий итог несут ТОТ ЖЕ тип `Money`, что и ячейка, но по инварианту
   задачи 2 их `reason` бывает только `REASON_UNKNOWN_VAT_BASE` (гашение по
   оси) либо `None` (в т.ч. ноль с `incomplete=True`, спека §2.5, §2.9).
   `_write_money` не выбирает текст сама — печатает то, что получила, и
   корректность здесь наследуется от инварианта задачи 2, а не проверяется
   второй раз этим модулем.
4. **Автофильтр не покрывает блок подытогов.** `ws.auto_filter.ref` считается
   ДО того, как в лист дописан блок подытогов, и заканчивается на последней
   строке данных; блок подытогов начинается минимум двумя строками ниже
   (пустая строка-разделитель плюс баннер блока).

Заголовков статей внутри области данных нет вовсе: статья — первые две
колонки КАЖДОЙ строки, а не отдельная строка-разделитель (спека §2.10).
"""
from __future__ import annotations

from collections.abc import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from services.changes_export import (
    MONEY_ONLY_KINDS,
    REASON_MIX_INCOMPLETE,
    REASON_NO_AMOUNT,
    REASON_NO_PRICE,
    REASON_UNKNOWN_VAT_BASE,
    Cell,
    Money,
    Sheet,
    SheetRow,
    sheet_names,
)
from services.excel import (
    BORDER,
    C_COL_BG,
    C_HEADER_BG,
    C_TOTAL_BG,
    FMT_MONEY,
    FMT_QTY,
    align,
    apply_column_widths,
    fill,
    font,
    safe_str,
    workbook_bytes,
    write_banner,
    write_cell,
    write_column_headers,
)
from services.stage_summary import TAX_GROSS, TAX_NET, TAX_NONE

__all__ = [
    "LEFT_COLUMNS",
    "STAGE_COLUMNS",
    "TAIL_COLUMNS",
    "TEXT_NO_AMOUNT",
    "TEXT_NO_PRICE",
    "TEXT_MIX_INCOMPLETE",
    "TEXT_UNKNOWN_VAT",
    "build_changes_export",
]

#: Слева, на каждый этап и справа (спека §2.3) — `4 + 4 × N + 7` колонок листа.
LEFT_COLUMNS: tuple[str, ...] = ("Статья", "Наименование статьи", "Наименование работы", "Ед.")
STAGE_COLUMNS: tuple[str, ...] = ("№ строк", "Объём", "Цена за ед.", "Сумма")
TAIL_COLUMNS: tuple[str, ...] = (
    "Δ сумма", "Δ %", "Δ работы", "Δ материалы", "Δ косвенные", "Полнота", "Что двигалось",
)

#: Подписи причин недоступности денег словами (спека §2.5, §2.9) — тот же
#: приём, что `_REASON_LABELS` у `services/excel_comparison.py`: печатается
#: текст, а не код причины.
TEXT_NO_AMOUNT = "суммы нет"
TEXT_NO_PRICE = "цены нет"
TEXT_MIX_INCOMPLETE = "состав неполон"
TEXT_UNKNOWN_VAT = "неизвестна база НДС"

_REASON_TEXT: dict[str, str] = {
    REASON_NO_AMOUNT: TEXT_NO_AMOUNT,
    REASON_NO_PRICE: TEXT_NO_PRICE,
    REASON_MIX_INCOMPLETE: TEXT_MIX_INCOMPLETE,
    REASON_UNKNOWN_VAT_BASE: TEXT_UNKNOWN_VAT,
}

_DASH = "—"

#: Заливка неполноты — тот же оттенок, что «частичная свёртка» у макета
#: гейта 1 (`docs/superpowers/specs/2026-09-16-tender-changes-export/gen_mockup.py`,
#: `td.partial { background: #fff4e0; }`).
_INCOMPLETE_FILL = fill("FFF4E0")

#: Подписи налогового состава и ценового уровня — обе печатаются на КАЖДОМ
#: листе (AGENTS.md §10, спека §2.9). Слов «валовые, как в файлах» и «нетто»
#: держаться дословно: они же употреблены в `gen_mockup.py`.
_TAX_BASIS_CAPTION = {
    TAX_GROSS: "Деньги: валовые, как в файлах (одна известная ставка НДС на всех этапах)",
    TAX_NET: "Деньги: нетто (ставки НДС этапов различаются)",
    TAX_NONE: "Деньги: сумм нет — ни на одном этапе не известна база НДС",
}
_PRICE_LEVEL_CAPTION = "Ценовой уровень: номинальный, без приведения"

#: Формат процента — свой, не `FMT_PCT` из `services/excel.py`: тот формат
#: рассчитан на уже готовые «пункты», а колонка `Δ %` несёт готовое число со
#: знаком (спека §2.6), которому нужен только суффикс «%».
_PCT_FORMAT = '+0.0"%";-0.0"%";0.0"%"'


def _columns_spec(n_stages: int) -> list[tuple[str, int, str | None, str]]:
    """Плоский список колонок листа: `4 + 4 × n_stages + 7` (спека §2.3)."""
    spec: list[tuple[str, int, str | None, str]] = [
        ("Статья", 10, "@", "left"),
        ("Наименование статьи", 26, "@", "left"),
        ("Наименование работы", 42, "@", "left"),
        ("Ед.", 8, "@", "center"),
    ]
    for _ in range(n_stages):
        spec.append(("№ строк", 20, "@", "left"))
        spec.append(("Объём", 12, FMT_QTY, "right"))
        spec.append(("Цена за ед.", 14, FMT_MONEY, "right"))
        spec.append(("Сумма", 16, FMT_MONEY, "right"))
    spec.append(("Δ сумма", 16, FMT_MONEY, "right"))
    spec.append(("Δ %", 10, _PCT_FORMAT, "right"))
    spec.append(("Δ работы", 14, FMT_MONEY, "right"))
    spec.append(("Δ материалы", 14, FMT_MONEY, "right"))
    spec.append(("Δ косвенные", 14, FMT_MONEY, "right"))
    spec.append(("Полнота", 30, "@", "left"))
    spec.append(("Что двигалось", 36, "@", "left"))
    return spec


def _write_money(ws: Worksheet, row: int, col: int, money: Money, *, horizontal: str = "right") -> None:
    """Одноярусная денежная ячейка (риск 2 в докстроке модуля): число с
    `FMT_MONEY`, либо текст причины словами — никогда оба сразу и никогда
    переносом строки внутри одной ячейки. Неполнота (`incomplete=True`, только
    у агрегатов — у `Cell.amount` строки она всегда `False`, задача 2) — это
    заливка ТОЙ ЖЕ ячейки, не вторая строка внутри неё."""
    if money.value is not None:
        cell_fill = _INCOMPLETE_FILL if money.incomplete else None
        write_cell(ws, row, col, money.value, number_format=FMT_MONEY, horizontal=horizontal, cell_fill=cell_fill)
    else:
        text = _REASON_TEXT.get(money.reason, money.reason or _DASH)
        write_cell(ws, row, col, text, horizontal=horizontal)


def _write_dash(ws: Worksheet, row: int, col: int, *, horizontal: str = "center") -> None:
    write_cell(ws, row, col, _DASH, horizontal=horizontal)


def _write_stage_cell(ws: Worksheet, row: int, first_col: int, kind: str, cell: Cell | None) -> None:
    """Четыре колонки одного этапа одной строки (спека §2.3). `cell is None`
    значит «строки нет на этом этапе» — прочерк во всех четырёх: вклад в
    статью нулевой, а не «неизвестно» (спека §2.5)."""
    if cell is None:
        for offset in range(4):
            _write_dash(ws, row, first_col + offset)
        return

    if kind in MONEY_ONLY_KINDS:
        # Риск 1: различитель — ВЕТВЬ (`kind`), а не причина `unit_price.reason`
        # (у `MONEY_ONLY_KINDS` она всегда `REASON_NO_PRICE`, хотя цены здесь
        # нет по контракту, а не по нехватке данных, спека §2.2).
        _write_dash(ws, row, first_col)
        _write_dash(ws, row, first_col + 1)
        _write_dash(ws, row, first_col + 2)
        _write_money(ws, row, first_col + 3, cell.amount)
        return

    numbers_text = ", ".join(cell.numbers) if cell.numbers else _DASH
    write_cell(ws, row, first_col, numbers_text, horizontal="left", wrap=True)
    if cell.quantity is None:
        _write_dash(ws, row, first_col + 1, horizontal="right")
    else:
        write_cell(ws, row, first_col + 1, cell.quantity, number_format=FMT_QTY, horizontal="right")
    _write_money(ws, row, first_col + 2, cell.unit_price)
    _write_money(ws, row, first_col + 3, cell.amount)


def _write_components(ws: Worksheet, row: int, first_col: int, sheet_row: SheetRow) -> None:
    """Три колонки Δ состава (спека §2.3). У `MONEY_ONLY_KINDS` — прочерк, А
    НЕ подпись состояния: тождество на них не распространяется вовсе (спека
    §2.2), и `components_reason` у них всегда `None` (задача 2)."""
    if sheet_row.kind in MONEY_ONLY_KINDS:
        for offset in range(3):
            _write_dash(ws, row, first_col + offset)
        return
    if sheet_row.delta_components is not None:
        comp = sheet_row.delta_components
        for offset, value in enumerate((comp.works, comp.materials, comp.indirect)):
            write_cell(ws, row, first_col + offset, value, number_format=FMT_MONEY, horizontal="right")
        return
    text = _REASON_TEXT.get(sheet_row.components_reason, sheet_row.components_reason or TEXT_MIX_INCOMPLETE)
    for offset in range(3):
        write_cell(ws, row, first_col + offset, text, horizontal="right")


def _write_data_row(ws: Worksheet, row: int, sheet_row: SheetRow, n_stages: int) -> None:
    code_text = sheet_row.article.code if sheet_row.article.code is not None else _DASH
    write_cell(ws, row, 1, code_text, horizontal="left")
    write_cell(ws, row, 2, sheet_row.article.title or "", horizontal="left", wrap=True)
    write_cell(ws, row, 3, sheet_row.work_title or "", horizontal="left", wrap=True)
    write_cell(ws, row, 4, sheet_row.unit or _DASH, horizontal="center")

    col = 5
    for idx in range(n_stages):
        _write_stage_cell(ws, row, col, sheet_row.kind, sheet_row.cells[idx])
        col += 4

    _write_money(ws, row, col, sheet_row.delta_amount)
    col += 1
    if sheet_row.delta_pct is None:
        _write_dash(ws, row, col, horizontal="right")
    else:
        write_cell(ws, row, col, sheet_row.delta_pct, number_format=_PCT_FORMAT, horizontal="right")
    col += 1
    _write_components(ws, row, col, sheet_row)
    col += 3
    write_cell(ws, row, col, sheet_row.completeness or "", horizontal="left", wrap=True)
    col += 1
    route_text = "; ".join(sheet_row.route) if sheet_row.route else "без изменений"
    write_cell(ws, row, col, route_text, horizontal="left", wrap=True)


def _block_completeness_text(cells: Sequence[Money], stages) -> str:
    """Соседняя колонка «Полнота» блока подытогов (риск 2 в докстроке модуля):
    подытог со значением и `incomplete=True` остаётся числом в своей ячейке
    (заливка), а ЧТО именно неполно — называет эта колонка, не вторая строка
    внутри денежной ячейки (спека §2.5)."""
    notes = [
        f"Э{stage.stage_no}: неполно"
        for stage, money in zip(stages, cells, strict=True)
        if money.value is not None and money.incomplete
    ]
    return "; ".join(notes)


def _write_aggregate_row(
    ws: Worksheet,
    row: int,
    *,
    code: str,
    title: str,
    rows_count: int,
    by_stage: Sequence[Money],
    delta: Money,
    stages,
) -> None:
    write_cell(ws, row, 1, code, horizontal="left", cell_font=font(bold=True))
    write_cell(ws, row, 2, title, horizontal="left", cell_font=font(bold=True))
    write_cell(ws, row, 3, rows_count, horizontal="right")
    col = 4
    for money in by_stage:
        _write_money(ws, row, col, money)
        col += 1
    _write_money(ws, row, col, delta)
    col += 1
    write_cell(ws, row, col, _block_completeness_text(by_stage, stages), horizontal="left", wrap=True)


def _write_caption(ws: Worksheet, row: int, text: str, *, wrap: bool = False) -> int:
    """Строка подписи без рамки — тот же приём, что у `services/excel_comparison.py`
    (`caption_cell`): подписи шапки не часть таблицы, границ у них нет."""
    cell = ws.cell(row=row, column=1, value=safe_str(text))
    cell.font = font(bold=True, size=9)
    cell.alignment = align(h="left", wrap=wrap)
    return row + 1


def _write_sheet(ws: Worksheet, sheet: Sheet, tender_header: Mapping[str, str | None]) -> None:
    n_stages = len(sheet.stages)
    columns_spec = _columns_spec(n_stages)
    total_cols = len(columns_spec)
    apply_column_widths(ws, columns_spec)

    tender_bits = [tender_header.get("tender_number"), tender_header.get("tender_title")]
    tender_label = " ".join(bit for bit in tender_bits if bit) or "Тендер"
    # Главный баннер листа — C_HEADER_BG (шапка ДОКУМЕНТА), тем же цветом, что
    # `excel_reports.py` и `excel_comparison.py`: раздел «Имена» задачи 4
    # перечисляет `C_COL_BG` и `C_TOTAL_BG` как проверенные grep'ом, а не как
    # белый список разрешённых цветов; третья книга проекта не должна отличаться
    # от первых двух без причины (решение оркестратора, ревью задачи 4).
    # `C_COL_BG` остаётся в импортах — он красит баннер-группу этапа ниже.
    row = write_banner(ws, 1, f"{tender_label} — Изменения КП: {sheet.title}", total_cols, bg=C_HEADER_BG)

    # Подписи налогового состава и ценового уровня — ОБЕ, на КАЖДОМ листе
    # (AGENTS.md §10, спека §2.9): без них читатель не знает, в чём измерены
    # складываемые числа.
    row = _write_caption(ws, row, _TAX_BASIS_CAPTION.get(sheet.tax_basis.basis, ""))
    row = _write_caption(ws, row, _PRICE_LEVEL_CAPTION)
    # `moves_note` печатается ОДИН РАЗ, в шапке (спека §2.8) — не построчно.
    row = _write_caption(ws, row, f"Что двигалось между этапами (по классификатору): {sheet.moves_note}", wrap=True)
    row += 1  # пустая строка-разделитель перед таблицей

    # Баннер-группа на этап — merge над четырьмя колонками этапа (спека §2.3).
    # LEFT_COLUMNS и TAIL_COLUMNS в этой строке не подписаны: у них нет
    # понятия «этап», подпись за них несёт строка меток колонок ниже.
    group_banner_row = row
    stage_col = len(LEFT_COLUMNS) + 1
    for stage in sheet.stages:
        label = f"Этап {stage.stage_no}"
        if stage.held_on is not None:
            label += f" · {stage.held_on.strftime('%d.%m.%Y')}"
        ws.merge_cells(start_row=group_banner_row, start_column=stage_col,
                        end_row=group_banner_row, end_column=stage_col + 3)
        cell = ws.cell(row=group_banner_row, column=stage_col, value=safe_str(label))
        cell.fill = fill(C_COL_BG)
        cell.font = font(bold=True, color="FFFFFF", size=9)
        cell.alignment = align(h="center")
        cell.border = BORDER
        stage_col += 4
    ws.row_dimensions[group_banner_row].height = 18
    row += 1

    header_row = row
    data_start = write_column_headers(ws, header_row, columns_spec)

    data_row = data_start
    for sheet_row in sheet.rows:
        _write_data_row(ws, data_row, sheet_row, n_stages)
        data_row += 1
    last_data_row = data_row - 1

    # Автофильтр — на область данных, БЕЗ единой строки блока подытогов
    # (риск 4 в докстроке модуля): диапазон считается ЗДЕСЬ, до того как в
    # лист дописан блок подытогов.
    if sheet.rows:
        last_col_letter = get_column_letter(total_cols)
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"

    # Окно закреплено по шапке и первым четырём колонкам (спека §2.10):
    # столбец E — первый НЕзакреплённый (`LEFT_COLUMNS` — четыре колонки).
    ws.freeze_panes = f"{get_column_letter(len(LEFT_COLUMNS) + 1)}{data_start}"

    block_row = last_data_row + 2  # пустая строка-разделитель перед блоком
    block_row = write_banner(ws, block_row, "ПОДЫТОГИ ПО СТАТЬЯМ", total_cols, bg=C_TOTAL_BG)

    subtotal_headers: list[tuple[str, int, str | None, str]] = [
        ("Статья", 10, "@", "left"),
        ("Наименование", 26, "@", "left"),
        ("Строк", 8, "0", "right"),
    ]
    subtotal_headers += [(f"Этап {stage.stage_no}", 16, FMT_MONEY, "right") for stage in sheet.stages]
    if sheet.stages:
        subtotal_headers.append((f"Δ Э{sheet.stages[0].stage_no}→Э{sheet.stages[-1].stage_no}", 16, FMT_MONEY, "right"))
    else:
        subtotal_headers.append(("Δ", 16, FMT_MONEY, "right"))
    subtotal_headers.append(("Полнота", 30, "@", "left"))
    block_row = write_column_headers(ws, block_row, subtotal_headers)

    for subtotal in sheet.subtotals:
        code = subtotal.article.code if subtotal.article.code is not None else _DASH
        title = subtotal.article.title or ""
        _write_aggregate_row(
            ws, block_row, code=code, title=title, rows_count=subtotal.rows,
            by_stage=subtotal.by_stage, delta=subtotal.delta, stages=sheet.stages,
        )
        block_row += 1

    _write_aggregate_row(
        ws, block_row, code=_DASH, title="ИТОГО", rows_count=len(sheet.rows),
        by_stage=sheet.grand_by_stage, delta=sheet.grand_delta, stages=sheet.stages,
    )


def build_changes_export(sheets: Sequence[Sheet], *, tender_header: Mapping[str, str | None]) -> bytes:
    """Собрать книгу «Изменения КП» — лист на каждого сравнимого участника
    (спека §2.1, §2.3, §2.10). Ничего не пересчитывает: каждый `Sheet` уже
    несёт готовую свёртку, Δ и подытоги (`services.changes_export.build_sheet`,
    задача 2). `tender_header` — заголовочные поля тендера (`tender_number`,
    `tender_title`), печатаемые в баннере каждого листа; отсутствующие ключи
    молча опускаются, а не роняют сборку — состав книги важнее заголовка."""
    names = sheet_names([sheet.title for sheet in sheets])
    wb = Workbook()
    for index, (name, sheet) in enumerate(zip(names, sheets, strict=True)):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = name
        _write_sheet(ws, sheet, tender_header)
    return workbook_bytes(wb)
