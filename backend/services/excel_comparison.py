"""Лист выгрузки «Сравнение договоров» (спека 2026-08-17 §2.7, §2.8; план,
задача 6).

**Один агрегат — два представления (спека §2.7).** Этот модуль не пересчитывает
НИЧЕГО: он читает ровно тот же словарь, что возвращает
`crud.comparison.build_comparison`, и печатает его поля как есть. Числа листа и
экрана при одном режиме НДС совпадают потому, что оба смотрят в один агрегат, а
не потому, что здесь продублирована арифметика (DoD 19).

**Разрез — статьи × договоры**, в отличие от отчёта «для банка»
(`services/excel_reports.py`, класс → работа). У сравнения на экране есть
переключатель корзины «Итого / ДГП / ДС» — у листа его нет и быть не может
(бумага без интерактивности), поэтому спека требует печатать ВСЕ ТРИ корзины
колонками разом (§2.8). У каждой корзины — своя сумма, свой ₽/м² и своё
отклонение от медианы (DoD 21): медиана считается ПО КОРЗИНЕ (§2.2), и одно
отклонение на три колонки было бы отклонением неизвестно чего.

**Налоговый состав — на самом листе** (`AGENTS.md` §10 v6.8, DoD 20): читатель
обязан видеть, в чём измерены складываемые числа, а у листа нет всплывающих
подсказок. `data["caption"]` печатается баннером до таблицы; в режиме «своя
ставка» (`VAT_MODE_OWN`) у каждого договора ДОПОЛНИТЕЛЬНО печатается состав его
смет (`columns[].composition_caption`, например «ДГП 20 % · ДС 22 %») — без
этой подписи лист позволил бы сложить валовые суммы разных ставок незаметно.

**Ценовой уровень — на самом листе, и подробнее, чем на экране** (спека
инфляции §2.10, §2.12; DoD 13, 32). При сосчитанном приведении под подписью
состава печатаются ряд, целевой месяц и по каждому использованному году пять
фактов: год, коэффициент, ИСТОЧНИК, признак прогноза и дата правки ряда.
Источник года на экране сравнения не показывается, а здесь обязателен: версий у
ряда нет, воспроизводимость лежит на файле, и вопрос «откуда 8,3 %» задают именно
к нему. Без блока `inflation` лист не меняется ни на одну ячейку.

**Три состояния ячейки — три разных написания** (§2.1.2, §2.1.3), той же
логикой, что на экране:

* статьи нет ни в одной смете договора (`state == ABSENT`) — прочерк без
  пояснения: пояснять нечего, там, где статьи нет, не может быть неполноты;
* статья есть, но ячейка погашена `incomplete_reasons` — прочерк СО СЛОВАМИ
  причины («без цены», «с ошибкой», «неизвестна база НДС», «ставка показа не
  определена»), а не голый прочерк и не число. Ноль здесь читался бы как
  «предусмотрено и стоит ноль», хотя сумма НЕИЗВЕСТНА — то же предостережение,
  что и на экране;
* иначе — число, в т.ч. настоящий ноль (`state == ZERO`, статья есть и
  расценена в ноль), который прочерку не эквивалентен.

**Деньги — `Decimal`-числа, не строки с числом внутри** (см. модульную
документацию `services/excel.py`): банк и любой другой читатель суммирует
колонку формулой `СУММ`, а по текстовым ячейкам она молча даёт ноль.
"""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment

from crud.comparison import ABSENT, BUCKET_AMENDMENTS, BUCKET_BASE, BUCKET_TOTAL, VAT_MODE_OWN
from services.excel import (
    C_CLASS_BG,
    C_EVEN,
    C_HEADER_BG,
    C_ODD,
    C_TOTAL_BG,
    FMT_MONEY,
    FMT_PCT,
    align,
    apply_column_widths,
    deviation_font,
    fill,
    font,
    safe_str,
    workbook_bytes,
    write_banner,
    write_cell,
    write_column_headers,
)

__all__ = ["build_comparison_sheet"]

#: Порядок и подписи трёх корзин на листе (спека §2.2, §2.8). ФИКСИРОВАН: строки
#: листа читаются как «ДГП · ДС · Итого» слева направо на каждом договоре — тем
#: же порядком, что в `crud.comparison._composition_caption`.
_BUCKET_ORDER = (BUCKET_BASE, BUCKET_AMENDMENTS, BUCKET_TOTAL)
_BUCKET_LABELS = {BUCKET_BASE: "ДГП", BUCKET_AMENDMENTS: "ДС", BUCKET_TOTAL: "Итого"}

_METRIC_SUM = "Сумма"
_METRIC_PER_SQM = "₽/м²"
_METRIC_DEVIATION = "Откл., %"
#: Три метрики на корзину (сумма, ₽/м², отклонение) × три корзины — DoD 21.
_COLS_PER_CONTRACT = 3 * len(_BUCKET_ORDER)

#: Причины неполноты ячейки словами (спека §2.1.3, §2.3.2) — тот же приём, что
#: у паспортного `incompletenessCaption`: причины называются, а не кодами.
#: `display_rate_undefined` — единственная причина, зависящая от режима показа
#: (§2.1.3 п.4), остальные три приходят прямо из счётчиков VIEW.
_REASON_LABELS = {
    "unpriced_rows": "без цены",
    "not_finite_rows": "с ошибкой",
    "vat_base_unknown": "неизвестна база НДС",
    "display_rate_undefined": "ставка показа не определена",
}

_DASH = "—"


def _reason_text(reasons) -> str:
    names = [_REASON_LABELS.get(code, code) for code in sorted(reasons)]
    return ", ".join(names)


def _sum_cell(bucket_cell: dict) -> tuple[object, str | None]:
    """Значение и переопределение формата колонки «Сумма» (§2.1.2, §2.1.3).

    `ABSENT` — прочерк без пояснения (статьи нет вовсе). `shown is None` при
    состоянии, отличном от `ABSENT`, — ячейка есть, но погашена причинами:
    прочерк СО словами причины, а не голое число и не голый прочерк. Иначе —
    число как есть, `Decimal`, включая настоящий ноль.
    """
    if bucket_cell["state"] == ABSENT:
        return _DASH, "@"
    shown = bucket_cell["shown"]
    if shown is None:
        reasons = bucket_cell["incomplete_reasons"]
        text = f"{_DASH} ({_reason_text(reasons)})" if reasons else _DASH
        return text, "@"
    return shown, None


def _per_sqm_cell(bucket_cell: dict) -> tuple[object, str | None]:
    """₽/м² — прочерк без слов причины: без ТЭП объекта либо без суммы (§2.4)."""
    value = bucket_cell["shown_per_sqm"]
    if value is None:
        return _DASH, "@"
    return value, None


def _deviation_cell(bucket_cell: dict) -> tuple[object, str | None]:
    """Отклонение — уже посчитано по нетто на агрегате (§2.5 правило 2),
    здесь только печатается тем же значением, что видит экран (DoD 10)."""
    value = bucket_cell["deviation_pct"]
    if value is None:
        return _DASH, "@"
    return value, None


def _columns_spec(n_contracts: int) -> list[tuple[str, int, str | None, str]]:
    spec: list[tuple[str, int, str | None, str]] = [("Статья", 50, "@", "left")]
    for _ in range(n_contracts):
        for bucket in _BUCKET_ORDER:
            label = _BUCKET_LABELS[bucket]
            spec.append((f"{label} · {_METRIC_SUM}", 16, FMT_MONEY, "right"))
            spec.append((f"{label} · {_METRIC_PER_SQM}", 14, FMT_MONEY, "right"))
            spec.append((f"{label} · {_METRIC_DEVIATION}", 12, FMT_PCT, "right"))
    return spec


def _write_contract_group_headers(ws, row_num: int, columns: list[dict], vat_mode: str) -> int:
    """Баннер-группа на договор: номер, объект и, в режиме «своя ставка»,
    состав его смет (DoD 20 — подпись состава ПО ДОГОВОРУ)."""
    for index, column in enumerate(columns):
        start_col = 2 + index * _COLS_PER_CONTRACT
        end_col = start_col + _COLS_PER_CONTRACT - 1
        label = f"{column['contract_number']} — {column['object_title']}"
        if vat_mode == VAT_MODE_OWN:
            label += f" ({column['composition_caption']})"
        ws.merge_cells(start_row=row_num, start_column=start_col, end_row=row_num, end_column=end_col)
        cell = ws.cell(row=row_num, column=start_col, value=safe_str(label))
        cell.fill = fill(C_CLASS_BG)
        cell.font = font(bold=True, color="FFFFFF", size=9)
        cell.alignment = align(h="left", wrap=True)
    ws.row_dimensions[row_num].height = 24
    return row_num + 1


def _write_data_row(
    ws, row_num: int, label: str, level: int, cells: list[dict], *, background_hex: str, bold: bool = False
) -> int:
    """Одна строка листа: подпись статьи (с отступом по уровню, §2.1) и по три
    метрики на каждую из трёх корзин, для каждого договора (DoD 21)."""
    background = fill(background_hex)
    label_cell = write_cell(
        ws, row_num, 1, label, cell_font=font(bold=bold), cell_fill=background, horizontal="left"
    )
    # Иерархия строк — отступом, а не шевроном: у листа нет интерактивности
    # (спека §2.1, «строки листа должны читаться иерархично на бумаге»).
    label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=max(level - 1, 0))

    col = 2
    for cell_entry in cells:
        for bucket in _BUCKET_ORDER:
            bucket_cell = cell_entry[bucket]

            sum_value, sum_fmt = _sum_cell(bucket_cell)
            write_cell(ws, row_num, col, sum_value, cell_font=font(bold=bold), cell_fill=background,
                       number_format=sum_fmt or FMT_MONEY, horizontal="right")
            col += 1

            per_sqm_value, per_sqm_fmt = _per_sqm_cell(bucket_cell)
            write_cell(ws, row_num, col, per_sqm_value, cell_font=font(bold=bold), cell_fill=background,
                       number_format=per_sqm_fmt or FMT_MONEY, horizontal="right")
            col += 1

            deviation_value, deviation_fmt = _deviation_cell(bucket_cell)
            # Цвет — ТОЛЬКО у отклонения: закрашивать сумму значило бы утверждать,
            # что «дорого» и «превышение медианы» — одно и то же (тот же довод,
            # что у `services/excel_reports.py::_write_row`).
            dev_font = deviation_font(bucket_cell["deviation_pct"], bold=bold)
            write_cell(ws, row_num, col, deviation_value, cell_font=dev_font, cell_fill=background,
                       number_format=deviation_fmt or FMT_PCT, horizontal="right")
            col += 1
    return row_num + 1


#: Месяцы в именительном: на листе месяц стоит ПОСЛЕ двоеточия, то есть подписью
#: («Цены приведены к: август 2026»), а не продолжением фразы. В подписи оси
#: (`crud.comparison`) он в дательном, потому что там фраза сплошная.
_MONTHS_NOMINATIVE = (
    "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)


def _month_label(target_month: str) -> str:
    year, month = target_month.split("-")
    return f"{_MONTHS_NOMINATIVE[int(month) - 1]} {year}"


def _write_inflation_block(ws, row: int, inflation: dict) -> int:
    """Ряд, целевой месяц и ПО КАЖДОМУ году — пять фактов (DoD 13, спека §2.10).

    **Печатается только на листе.** Экран показывает примечание ряда и уровни по
    годам, но НЕ источники: три источника рядом с тремя процентами превращают
    ориентирующую строку в таблицу. Источник нужен там, где задают вопрос «откуда
    8,3 %», — а задают его к артефакту защиты перед банком, то есть к файлу.
    Разделение простое: экран ориентирует, файл защищает (§2.12).

    Пять фактов на год — коэффициент, источник, признак прогноза, дата правки и
    сам год — потому что версионности у ряда нет (§2.10): воспроизводимость
    лежит на файле, и защищённые числа заморожены в нём. Год без источника или без
    даты правки эту роль не исполнил бы.
    """
    ws.cell(
        row=row, column=1,
        value=safe_str(
            f"Цены приведены к: {_month_label(inflation['target_month'])}"
            f" · ряд: «{inflation['series_name']}»"
        ),
    ).font = font(bold=True, size=9)
    row += 1

    for item in inflation["used_years"]:
        kind = "прогноз" if item["is_forecast"] else "факт"
        parts = [
            str(item["year"]),
            format(item["coefficient"], "f"),
            safe_str(item["source"]),
            kind,
        ]
        if item["updated_at"]:
            parts.append(f"правлен {_iso_date_to_ru(item['updated_at'])}")
        ws.cell(row=row, column=1, value=safe_str(" · ".join(parts))).font = font(size=9)
        row += 1

    if not inflation["used_years"]:
        # Пустой список — законное состояние (цель совпала с месяцем сметы, DoD 5),
        # и молчать о нём нельзя: читатель обязан понять, что приведение включено,
        # а годов не потребовалось, — иначе он решит, что ряд не доехал до листа.
        ws.cell(
            row=row, column=1,
            value="Коэффициенты за годы не потребовались: цены уже в целевом уровне.",
        ).font = font(size=9)
        row += 1
    return row


def _iso_date_to_ru(value: str) -> str:
    """`2026-01-12T…` → `12.01.2026`. Дата правки ряда печатается по-русски.

    Разбор строкой, а не `datetime.fromisoformat`: значение уже пришло в ответе
    строкой (`crud.common.iso`), и второе преобразование туда-обратно только
    добавило бы место, где формат может разойтись.
    """
    date_part = value.split("T")[0]
    year, month, day = date_part.split("-")
    return f"{day}.{month}.{year}"


def build_comparison_sheet(data: dict, *, generated_at: dt.date) -> bytes:
    """Лист «статьи × договоры» из готового агрегата `crud.comparison.build_comparison`.

    `data` печатается КАК ЕСТЬ (спека §2.7): числа листа и экрана при одном
    режиме НДС совпадают потому, что оба читают один и тот же словарь, а не
    потому, что здесь что-то пересчитано заново.
    """
    columns = data["columns"]
    n_contracts = len(columns)
    columns_spec = _columns_spec(n_contracts)
    total_cols = len(columns_spec)

    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение договоров"
    apply_column_widths(ws, columns_spec)

    row = write_banner(ws, 1, "СРАВНЕНИЕ ДОГОВОРОВ ПО СТАТЬЯМ", total_cols, bg=C_HEADER_BG)

    # Подпись налогового состава — ОБЯЗАНА быть на листе (AGENTS.md §10 v6.8,
    # DoD 20): без неё читатель не знает, в чём измерены суммы, которые он
    # складывает, а у листа нет тултипов, чтобы сказать это иначе.
    caption_cell = ws.cell(row=row, column=1, value=safe_str(data["caption"]))
    caption_cell.font = font(bold=True, size=9)
    row += 1

    numbers = ", ".join(column["contract_number"] for column in columns) or "нет договоров"
    selection_cell = ws.cell(
        row=row, column=1, value=safe_str(f"Выборка: {len(columns)} договор(ов): {numbers}")
    )
    selection_cell.font = font(size=9)
    row += 1

    ws.cell(row=row, column=1, value=f"Сформирован: {generated_at.strftime('%d.%m.%Y')}").font = font(size=9)
    row += 1

    # Блок приведения печатается ТОЛЬКО когда приведение сосчитано: без него лист
    # не меняется ни на одну ячейку, и это стережёт golden-снимок дофичевого листа
    # (DoD 1). `data.get`, а не `data[...]`: номинальный ответ ключа не несёт вовсе.
    inflation = data.get("inflation")
    if inflation is not None:
        row = _write_inflation_block(ws, row, inflation)

    row += 1

    if columns:
        row = _write_contract_group_headers(ws, row, columns, data["vat_mode"])
    row = write_column_headers(ws, row, columns_spec)

    for index, row_data in enumerate(data["rows"]):
        background_hex = C_EVEN if index % 2 == 1 else C_ODD
        row = _write_data_row(
            ws, row, row_data["title"], row_data["level"], row_data["cells"], background_hex=background_hex,
        )

    row = _write_data_row(
        ws, row, "ИТОГО ПО ДОГОВОРУ", 1, data["totals"], background_hex=C_TOTAL_BG, bold=True,
    )

    ws.freeze_panes = "B1"
    return workbook_bytes(wb)
