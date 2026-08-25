"""Независимый замер сводной таблицы — класс 3 спеки §7.

Единственный образец `samples/`, чей снимок «до» был отказом (ширина 12 не
поддерживалась) и чей снимок «после» — успешный разбор. Итоговый JSON после
`normalize_lots_json_structure` не несёт всего, что нужно проверить: пустой
блок «Расчетная стоимость» заменяется заглушкой (его позиции из JSON исчезают),
а отклонения подрядчиков от базы вычищаются, хотя физическая колонка «% от р/с»
на листе есть (спека §2.8). Поэтому замер идёт ДВУМЯ этапами:

* этап A — до постобработки, все пять физических блоков подрядчика (базовый +
  четыре предложения) поколоночно, сверкой JSON с сырым чтением листа по
  литеральным координатам замера плана;
* этап B — после `parse_estimate`, пять предпосылок brief-а задачи 6 как
  утверждения внутри прогона (docs/insights/false-test-premises.md), а не
  тихие допущения.

Запуск из backend/ (без аргумента дайджест выводится сравнением снимков
before/after; можно передать явно первыми 12+ символами sha256):

    PYTHONIOENCODING=utf-8 uv run python -m scripts.verify_summary_sheet [<sha256>]

Печатает только счётчики сверенных ячеек, дайджест (первые 12 символов) и
имена ключей/классов JSON. Ни путь к файлу, ни реквизиты контрагентов, ни
денежные суммы в вывод не попадают (AGENTS.md §9) — при расхождении сообщение
называет блок/позицию/ключ, но не печатает сами значения.
"""
from __future__ import annotations

import hashlib
import json
import math
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from parser import estimate, parse_estimate
from parser.build_merged_shape_map import merged_rows_in_first_column
from parser.constants import (
    JSON_KEY_BASELINE_PROPOSAL,
    JSON_KEY_COMMENT_CONTRACTOR,
    JSON_KEY_CONTRACTOR_INDEX,
    JSON_KEY_CONTRACTOR_ITEMS,
    JSON_KEY_CONTRACTOR_POSITIONS,
    JSON_KEY_CONTRACTOR_TITLE,
    JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
    JSON_KEY_INDIRECT_COSTS,
    JSON_KEY_LOTS,
    JSON_KEY_MATERIALS,
    JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
    JSON_KEY_PROPOSALS,
    JSON_KEY_SUGGESTED_QUANTITY,
    JSON_KEY_TOTAL,
    JSON_KEY_TOTAL_COST,
    JSON_KEY_UNIT_COST,
    JSON_KEY_VAT_RATE,
    JSON_KEY_WORKS,
    TABLE_PARSE_ADDITIONAL_WORKS_TITLE,
)
from parser.parse_contractor_row import MONEY_KEYS
from parser.postprocess import BASELINE_MISSING_TITLE
from parser.read_contractors import read_contractors
from parser.read_lots_and_boundaries import find_lot_starts, read_lots_and_boundaries
from parser.resolve_contractor import resolve_contractor
from parser.sheet import cell_text_is_blank, normalized_cell_text, row_is_empty

REPO_ROOT = Path(__file__).resolve().parents[2]
SAMPLES_DIR = REPO_ROOT / "samples"
SNAPSHOTS_DIR = SAMPLES_DIR / "_snapshots"

# --- Эталон физической раскладки: литералы замера плана (25.08.2026) -------
#
# docs/superpowers/plans/2026-08-25-parser-columns-by-header.md, задача 6:
# col_start пяти блоков 10/19/31/44/57, header_row 11, первая строка данных
# лота 13. Ширины блоков — спека §7: «блоки 8 + 11 с «% от р/с» + 12×3» —
# первый блок (базовый, «Расчетная стоимость») шириной 8, остальные четыре —
# одна ширина 11 и три ширины 12, ВСЕ с «% от р/с». Здесь НЕ зашито, какой из
# четырёх блоков имеет ширину 11: она читается у read_contractors (геометрия
# объединения, не зависящая от resolve_contractor) и по ней выбирается
# канонический эталон ключей — так эталон не строится на догадке о порядке
# файлов, которую здесь неоткуда проверить.
EXPECTED_COL_STARTS = (10, 19, 31, 44, 57)
EXPECTED_HEADER_ROW = 11
EXPECTED_FIRST_DATA_ROW = 13
EXPECTED_BASELINE_WIDTH = 8

_UC, _TC = JSON_KEY_UNIT_COST, JSON_KEY_TOTAL_COST
#: Восемь денежных ключей в физическом порядке (tests/unit/parser/sheet_builders.py:KEYS_8).
MONEY_OCTET: tuple[str, ...] = tuple(
    f"{group}.{part}"
    for group in (_UC, _TC)
    for part in (JSON_KEY_MATERIALS, JSON_KEY_WORKS, JSON_KEY_INDIRECT_COSTS, JSON_KEY_TOTAL)
)

#: Канонический порядок ключей по ширине блока — те же кортежи, что
#: `tests/unit/parser/sheet_builders.py` называет KEYS_8 / KEYS_TENDER_11 /
#: KEYS_12, измеренные на реальных файлах при планировании фичи, НЕЗАВИСИМО от
#: этого конкретного файла.
EXPECTED_KEYS_BY_WIDTH: dict[int, tuple[str, ...]] = {
    8: MONEY_OCTET,
    11: (
        JSON_KEY_SUGGESTED_QUANTITY,
        *MONEY_OCTET,
        JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
        JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
    ),
    12: (
        JSON_KEY_SUGGESTED_QUANTITY,
        *MONEY_OCTET,
        JSON_KEY_ORGANIZER_QUANTITY_TOTAL_COST,
        JSON_KEY_COMMENT_CONTRACTOR,
        JSON_KEY_DEVIATION_FROM_CALCULATED_COST,
    ),
}


def expected_money_cell(value: Any) -> Any:
    """Зеркало контракта money_to_json на стороне ЭТАЛОНА: число → десятичная
    строка, пусто/bool/нечисловые nan-inf → None, текст (включая Excel-ошибки
    вида '#DIV/0!') — исходная строка как есть. Независимость замера — в
    КООРДИНАТАХ (литералы листа против раскладки резолвера), а не в кодировке
    значений: её контракт один на проект (AGENTS.md §3)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(Decimal(value))
    if isinstance(value, float):
        return str(Decimal(str(value))) if math.isfinite(value) else None
    return value


def _find_sample_path(digest: str) -> Path:
    """Находит файл с данным sha256 в samples/ — тем же обходом, что
    snapshot_parse_samples.py (пропускает временные ~$ и сами снимки)."""
    for path in sorted(SAMPLES_DIR.rglob("*.xlsx")):
        if path.name.startswith("~$") or "_snapshots" in path.parts:
            continue
        if hashlib.sha256(path.read_bytes()).hexdigest() == digest:
            return path
    raise SystemExit(f"файл с дайджестом {digest[:12]} не найден в samples/")


def _find_class3_digest() -> str:
    """Дайджест единственного файла класса 3 (§7): «до» — отказ, «после» — ok.

    Определяется по манифестам и снимкам, а не зашит константой: та же
    классификация, что печатает compare_parse_snapshots.py, просто без его
    зависимостей.
    """
    before_dir = SNAPSHOTS_DIR / "before"
    after_dir = SNAPSHOTS_DIR / "after"
    before_manifest = json.loads((before_dir / "manifest.json").read_text(encoding="utf-8"))
    after_manifest = json.loads((after_dir / "manifest.json").read_text(encoding="utf-8"))
    common = set(before_manifest["digests"]) & set(after_manifest["digests"])

    candidates: list[str] = []
    for digest in sorted(common):
        before_payload = json.loads((before_dir / f"{digest}.json").read_text(encoding="utf-8"))
        after_payload = json.loads((after_dir / f"{digest}.json").read_text(encoding="utf-8"))
        if before_payload["status"] == "error" and after_payload["status"] == "ok":
            candidates.append(digest)

    if len(candidates) != 1:
        raise SystemExit(
            "ожидался ровно один файл класса 3 (§7: «до» — отказ, «после» — успех), "
            f"найдено {len(candidates)}; передайте дайджест аргументом явно"
        )
    return candidates[0]


def _position_rows(ws: Worksheet, first_row: int, last_column: int) -> list[int]:
    """Номера строк позиций одного лота, определённые тем же правилом конца
    блока и пропуска строк, что `get_lot_positions` (объединённая ячейка
    колонки A — конец блока позиций; пустая строка — пропуск; агрегатная
    строка допработ — не позиция), но ОТДЕЛЬНЫМ проходом по листу.

    Раскладка блока (`BlockLayout`, то, что разрешает `resolve_contractor`) в
    этом проходе не участвует вообще — только геометрия объединений и общие
    колонки A/B/D, поэтому список строк независим от того, правильно ли
    резолвер разметил колонки: даже при неверной раскладке блока физические
    строки позиций остаются теми же.
    """
    merged_first_column_rows = merged_rows_in_first_column(ws)
    rows: list[int] = []
    for row in range(first_row, ws.max_row + 1):
        if row in merged_first_column_rows:
            break
        if row_is_empty(ws, row, last_column):
            continue
        number_blank = cell_text_is_blank(ws.cell(row=row, column=1).value)
        chapter_blank = cell_text_is_blank(ws.cell(row=row, column=2).value)
        if number_blank and chapter_blank:
            title = normalized_cell_text(ws.cell(row=row, column=4).value)
            if title.casefold() == TABLE_PARSE_ADDITIONAL_WORKS_TITLE.casefold():
                continue
        rows.append(row)
    return rows


def _stage_a(ws: Worksheet) -> tuple[list[str], dict[str, Any]]:
    """Этап A — до постобработки, пять блоков поколоночно.

    Цепочка `read_contractors` → `find_lot_starts` → `_validate_column_headers`
    → `resolve_contractor` на каждый блок → `read_lots_and_boundaries` — та же,
    что `parse_worksheet`, но останавливается ДО `normalize_lots_json_structure`
    (см. докстроку модуля).
    """
    problems: list[str] = []
    counters: dict[str, Any] = {"cells_checked": 0, "positions_per_block": None}

    contractors = read_contractors(ws)
    block_count = 0 if not contractors else len(contractors) - 1
    if block_count != len(EXPECTED_COL_STARTS):
        problems.append(
            f"read_contractors нашёл {block_count} блоков подрядчика вместо "
            f"{len(EXPECTED_COL_STARTS)}"
        )
        return problems, counters
    blocks = contractors[1:]

    col_starts = [block["column_start"] for block in blocks]
    if col_starts != list(EXPECTED_COL_STARTS):
        problems.append(f"col_start блоков {col_starts} ≠ ожидаемых {list(EXPECTED_COL_STARTS)}")

    widths = [block["merged_shape"]["colspan"] for block in blocks]
    if widths[0] != EXPECTED_BASELINE_WIDTH:
        problems.append(f"ширина первого (базового) блока {widths[0]} ≠ ожидаемой {EXPECTED_BASELINE_WIDTH}")
    if sorted(widths[1:]) != [11, 12, 12, 12]:
        problems.append(f"ширины четырёх предложений {widths[1:]} ≠ ожидаемых [11, 12, 12, 12] (спека §7)")

    lot_starts = find_lot_starts(ws)
    if len(lot_starts) != 1:
        problems.append(f"find_lot_starts нашёл {len(lot_starts)} лотов вместо 1")
        return problems, counters
    if lot_starts[0]["start_row"] != EXPECTED_FIRST_DATA_ROW:
        problems.append(
            f"первая строка данных лота {lot_starts[0]['start_row']} ≠ ожидаемой "
            f"{EXPECTED_FIRST_DATA_ROW}"
        )

    header_row = estimate._validate_column_headers(ws, contractors, lot_starts)
    if header_row != EXPECTED_HEADER_ROW:
        problems.append(f"header_row {header_row} ≠ ожидаемого {EXPECTED_HEADER_ROW}")

    resolved = [resolve_contractor(ws, block, header_row) for block in blocks]

    unreliable = [False] * len(blocks)
    for i, (width, contractor) in enumerate(zip(widths, resolved, strict=True)):
        expected_keys = EXPECTED_KEYS_BY_WIDTH.get(width)
        if expected_keys is None or contractor.layout.column_keys != expected_keys:
            problems.append(
                f"блок {i + 1}: раскладка резолвера не совпала с эталоном ширины {width} "
                "(замер плана) — координаты этого блока не сверяются"
            )
            unreliable[i] = True

    lots = read_lots_and_boundaries(ws, header_row=header_row, contractors=resolved)
    if list(lots.lots.keys()) != ["lot_1"]:
        problems.append(f"read_lots_and_boundaries вернул лоты {list(lots.lots.keys())} вместо ['lot_1']")
        return problems, counters

    proposals = lots.lots["lot_1"]["proposals"]
    positions_by_block: list[dict[str, Any]] = []
    for i in range(len(blocks)):
        key = f"{JSON_KEY_CONTRACTOR_INDEX}{i + 1}"
        proposal = proposals.get(key)
        if proposal is None:
            problems.append(f"в сырых proposals нет ключа {key}")
            return problems, counters
        positions_by_block.append(proposal[JSON_KEY_CONTRACTOR_ITEMS][JSON_KEY_CONTRACTOR_POSITIONS])

    counts = [len(positions) for positions in positions_by_block]
    counters["positions_per_block"] = counts
    if len(set(counts)) != 1:
        problems.append(f"число позиций по блокам не совпадает: {counts}")
        return problems, counters
    common_count = counts[0]
    if common_count == 0:
        problems.append("в блоках нет ни одной позиции — сверка первых/последних строк невозможна")
        return problems, counters

    last_column_overall = col_starts[-1] + widths[-1] - 1
    rows = _position_rows(ws, EXPECTED_FIRST_DATA_ROW, last_column_overall)
    if len(rows) != common_count:
        problems.append(
            f"независимый обход листа нашёл {len(rows)} строк позиций, а разбор — "
            f"{common_count}; поколоночная сверка по этому файлу пропущена"
        )
        return problems, counters

    sample_indices = sorted(
        set(range(1, min(3, common_count) + 1))
        | set(range(max(1, common_count - 1), common_count + 1))
    )

    for i, block_positions in enumerate(positions_by_block):
        if unreliable[i]:
            continue
        expected_keys = EXPECTED_KEYS_BY_WIDTH[widths[i]]
        col_start = col_starts[i]
        for pos_index in sample_indices:
            row = rows[pos_index - 1]
            position = block_positions[str(pos_index)]
            for offset, key in enumerate(expected_keys):
                column = col_start + offset
                raw = ws.cell(row=row, column=column).value
                expected_value = expected_money_cell(raw) if key in MONEY_KEYS else raw
                if "." in key:
                    head, _, tail = key.partition(".")
                    actual_value = position.get(head, {}).get(tail)
                else:
                    actual_value = position.get(key)
                counters["cells_checked"] += 1
                if actual_value != expected_value:
                    problems.append(
                        f"блок {i + 1}, позиция {pos_index}, ключ «{key}»: значение из JSON "
                        "не совпало с независимым чтением ячейки листа"
                    )

    return problems, counters


def _stage_b(path: Path) -> tuple[list[str], dict[str, Any]]:
    """Этап B — после `parse_estimate`, пять предпосылок brief-а как
    утверждения внутри прогона (docs/insights/false-test-premises.md)."""
    problems: list[str] = []
    counters: dict[str, Any] = {"proposals_checked": 0, "positions_checked": 0}

    result = parse_estimate(str(path))
    lots = result.data.get(JSON_KEY_LOTS, {})
    if list(lots.keys()) != ["lot_1"]:
        problems.append(f"после разбора найдено лотов {list(lots.keys())} вместо ['lot_1']")
        return problems, counters
    lot = lots["lot_1"]

    proposals = lot.get(JSON_KEY_PROPOSALS, {})
    expected_keys = {f"{JSON_KEY_CONTRACTOR_INDEX}{i}" for i in range(1, 5)}
    actual_keys = set(proposals.keys())
    if actual_keys != expected_keys:
        problems.append(f"proposals после постобработки: {sorted(actual_keys)} ≠ {sorted(expected_keys)}")
    else:
        counters["proposals_checked"] = len(proposals)

    baseline = lot.get(JSON_KEY_BASELINE_PROPOSAL) or {}
    baseline_title = baseline.get(JSON_KEY_CONTRACTOR_TITLE)
    if baseline_title == BASELINE_MISSING_TITLE:
        print(
            "этап B: baseline_proposal['title'] == BASELINE_MISSING_TITLE "
            f"«{BASELINE_MISSING_TITLE}» — подтверждено"
        )
    else:
        problems.append(
            "baseline_proposal['title'] НЕ равен BASELINE_MISSING_TITLE; фактическое "
            "значение не печатается (могло бы быть текстом из файла) — разобрать вручную "
            "и записать находку в devlog до ослабления проверки"
        )

    for key in sorted(actual_keys):
        proposal = proposals[key]
        vat_rate = proposal.get(JSON_KEY_VAT_RATE)
        if vat_rate != "22":
            problems.append(f"{key}: vat_rate = {vat_rate!r} ≠ '22'")

        positions = proposal.get(JSON_KEY_CONTRACTOR_ITEMS, {}).get(JSON_KEY_CONTRACTOR_POSITIONS) or {}
        for pos_key, position in positions.items():
            counters["positions_checked"] += 1
            if JSON_KEY_DEVIATION_FROM_CALCULATED_COST in position:
                problems.append(
                    f"{key}, позиция {pos_key}: ключ {JSON_KEY_DEVIATION_FROM_CALCULATED_COST} "
                    "присутствует, хотя правило «нет валидной базы» обязано было его вычистить"
                )

    if not any("найдено 5" in warning for warning in result.warnings):
        problems.append("предупреждения не содержат «найдено 5» (ожидалось число блоков подрядчика с базовым)")

    return problems, counters


def main() -> int:
    try:
        digest = sys.argv[1] if len(sys.argv) > 1 else _find_class3_digest()
        path = _find_sample_path(digest)

        wb = openpyxl.load_workbook(str(path), data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            problems_a, counters_a = _stage_a(ws)
        finally:
            wb.close()

        problems_b, counters_b = _stage_b(path)
    except Exception as exc:  # noqa: BLE001 — сообщение исключения может нести
        # текст из файла (подписи и координаты в EstimateParseError, реквизиты
        # в других отказах), поэтому наружу уходит только имя типа исключения.
        print(f"замер упал с исключением {type(exc).__name__} — сообщение не печатается")
        return 1

    print(f"файл класса 3 (§7): {digest[:12]}")
    print(
        f"этап A: сверено ячеек — {counters_a['cells_checked']}; "
        f"позиций на блок — {counters_a['positions_per_block']}"
    )
    for problem in problems_a:
        print(f"этап A, РАСХОЖДЕНИЕ: {problem}")

    print(
        f"этап B: предложений сверено — {counters_b['proposals_checked']}; "
        f"позиций проверено на отсутствие {JSON_KEY_DEVIATION_FROM_CALCULATED_COST} — "
        f"{counters_b['positions_checked']}"
    )
    for problem in problems_b:
        print(f"этап B, РАСХОЖДЕНИЕ: {problem}")

    if problems_a or problems_b:
        return 1

    print("этап A и этап B: расхождений нет")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
